from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "新北市青年政策公開資料AI系統_建立流程與研究方法手冊_V1.0.docx"
NAVY = "17365D"
BLUE = "2E74B5"
CYAN = "21A6B5"
LIGHT = "E8EEF5"
PALE = "F4F7FA"
GRAY = "5B6573"
RED = "B42318"
GREEN = "207A52"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(table) -> None:
    set_repeat_header(table.rows[0])
    for row in table.rows:
        no_split(row)


def set_east_asia(run, font="Microsoft JhengHei") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    paragraph.add_run(" 頁")


def hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color, before, after in (
        ("Title", 26, NAVY, 0, 18),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, NAVY, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Small Note" not in doc.styles:
        note = doc.styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = doc.styles["Small Note"]
    note.font.name = "Calibri"
    note.font.size = Pt(8.5)
    note.font.color.rgb = RGBColor.from_string(GRAY)
    note._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    note.paragraph_format.space_after = Pt(3)


def set_section(section, landscape=False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def configure_header_footer(section, label="新北市青年政策公開資料 AI 系統｜V1.0") -> None:
    header = section.header
    p = header.paragraphs[0]
    p.text = label
    p.style = "Small Note"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "研究與行政查核用｜產製日期：2026-08-16"
    p.style = "Small Note"
    add_page_number(p)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = str(header)
        shade(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Cm(widths[idx])
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[idx].width = Cm(widths[idx])
        if len(table.rows) % 2 == 1:
            for cell in cells:
                shade(cell, PALE)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    set_east_asia(run)
                    if row_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(NAVY)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    set_repeat_table_header(table)
    return table


def add_callout(doc, title, body, color=BLUE):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(0.35)
    table.columns[1].width = Cm(15.0)
    shade(table.cell(0, 0), color)
    shade(table.cell(0, 1), PALE)
    p = table.cell(0, 1).paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    set_east_asia(r)
    r = p.add_run(body)
    set_east_asia(r)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=140, bottom=140)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_flow(doc, steps):
    table = doc.add_table(rows=1, cols=len(steps))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, step in enumerate(steps):
        cell = table.cell(0, idx)
        cell.text = step
        shade(cell, NAVY if idx % 2 == 0 else BLUE)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8)
                set_east_asia(run)
        set_cell_margins(cell, top=160, bottom=160, start=80, end=80)
    return table


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build():
    sources = json.loads((ROOT / "config" / "sources.json").read_text(encoding="utf-8"))
    manifest = json.loads((ROOT / "data" / "manifests" / "latest.json").read_text(encoding="utf-8"))
    indicators = json.loads((ROOT / "data" / "curated" / "policy-indicators-latest.json").read_text(encoding="utf-8"))
    doc = Document()
    setup_styles(doc)
    set_section(doc.sections[0])
    configure_header_footer(doc.sections[0])

    # Cover — editorial_cover template.
    for _ in range(3):
        doc.add_paragraph()
    bar = doc.add_table(rows=1, cols=1)
    shade(bar.cell(0, 0), NAVY)
    p = bar.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("建立流程與研究方法手冊")
    r.font.size = Pt(25)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    set_east_asia(r)
    set_cell_margins(bar.cell(0, 0), top=380, bottom=380)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("新北市青年政策公開資料 AI 系統")
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    set_east_asia(r)
    p = doc.add_paragraph("公開資料 × 可重現分析 × AWS Bedrock AgentCore Harness")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = "Subtitle"
    for _ in range(4):
        doc.add_paragraph()
    add_table(doc, ["文件屬性", "內容"], [
        ("版本／日期", "V1.0／2026-08-16"),
        ("適用對象", "新北市政府承辦人、研究人員、資料工程及資訊安全人員"),
        ("資料批次", manifest["runId"]),
        ("文件定位", "研究與系統建立之可追溯操作基準；非對外統計公報"),
        ("敏感度", "公開資料＋治理中繼資料；不含帳密、API key 或個人資料"),
    ], widths=[4, 11.5], font_size=9.5)
    p = doc.add_paragraph("設計基準：compact_reference_guide｜封面：editorial_cover", style="Small Note")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    page_break(doc)
    doc.add_heading("文件控制與使用原則", level=1)
    add_callout(doc, "先看結論", "本手冊把資料、定義、介接、研究、圖表、AI 問答與發布 Gate 放入同一證據鏈。只有通過查核的資料可進入正式指標；資料不足時必須停止推論。")
    add_table(doc, ["控制項", "本版規則"], [
        ("唯一來源登錄", "config/sources.json"),
        ("介接契約", "config/integration-contracts.json"),
        ("年齡定義", "config/age-groups.json；核心 18–35"),
        ("資料版本", "data/manifests/<runId>.json＋SHA-256"),
        ("發布模式", "私人儀表板；未核准不得外部發布或寫入"),
        ("AWS 狀態", "Harness 已部署；API bridge 程式備妥但未部署"),
    ], widths=[4.2, 11.3], font_size=9)
    doc.add_heading("文件內容地圖", level=2)
    for item in [
        "第 1–4 章：來源、縮寫、介接、欄位對照。",
        "第 5–7 章：題目要求、研究流程、青年年齡定義與限制。",
        "第 8–10 章：圖表、可運用方向、AWS Harness AI Agent。",
        "第 11 章與附錄：人工查核、現況數值、官方來源與重用說明。",
    ]:
        bullet(doc, item)
    add_callout(doc, "本次資料批次", f"runId＝{manifest['runId']}。主要數值與其口徑列於附錄 A；本頁只說明文件控制，避免摘要表跨頁造成誤讀。", color=CYAN)

    page_break(doc)
    doc.add_heading("1｜來源資料", level=1)
    doc.add_paragraph("命題指定平台均已納入來源清冊；另以政府資料開放平臺與戶政司單一年齡 API 補足 18–35 精確人口。來源登錄不等於資料已自動取得，狀態須分開判讀。")
    rows = []
    integration_type = {
        "MOI-POP1Y": "分頁 JSON",
        "NTPC-POP": "JSON",
        "OAS-POP": "人工匯出",
        "NTPC-UNEMP": "JSON",
        "NTPC-EMPSTR": "JSON",
        "NTPC-EDU15P": "JSON",
        "STAT-WAGE-A": "官方 XLSX（清冊）",
        "STAT-WAGE-LOC": "官方 XLSX",
        "STAT-WAGE-EDU": "官方 XLSX",
    }
    for item in sources["datasets"]:
        rows.append((item["alias"], item["topic"], item["population"], item["geographyRole"], integration_type.get(item["alias"], "登錄"), item["status"]))
    add_table(doc, ["代號", "主題", "母體", "地理", "介接", "狀態"], rows, widths=[2.7, 2.2, 4.0, 2.2, 2.8, 2.5], font_size=7.8)
    add_callout(doc, "狀態解讀", "ACQUIRED＝本次批次已取得並留存雜湊；EVIDENCE_ONLY／PENDING_HUMAN_REVIEW＝只能作來源或查核證據，不可宣稱已自動介接。", color=CYAN)

    page_break(doc)
    doc.add_heading("2｜簡寫對照表", level=1)
    abbreviations = [
        ("MOI-POP1Y", "內政部戶政司村里戶數、單一年齡人口", "核心青年精確加總"),
        ("NTPC-POP", "新北市現住人口之年齡分配", "年度人口趨勢"),
        ("OAS-POP", "新北市統計資料庫人口表", "人工匯出鏡像查核"),
        ("NTPC-UNEMP", "新北市失業率－年齡別", "25–29 失業率"),
        ("NTPC-EMPSTR", "新北市就業者年齡結構", "就業補充觀察"),
        ("NTPC-EDU15P", "十五歲以上人口教育程度結構", "教育環境趨勢"),
        ("STAT-WAGE-EDU", "年齡×教育程度薪資表 5", "教育分層薪資"),
        ("STAT-WAGE-LOC", "工作場所縣市×年齡薪資表 6", "新北工作場所薪資"),
        ("OCR-DOC", "掃描文件 OCR", "無結構附件補充"),
        ("PK", "Primary Key", "資料列唯一鍵"),
        ("SHA-256", "256 位元檔案雜湊", "完整性與版本指紋"),
        ("MAE", "平均絕對誤差", "線性外推回測誤差"),
    ]
    add_table(doc, ["代號", "完整名稱", "系統用途"], abbreviations, widths=[3.5, 7.2, 5.2], font_size=8.5)
    doc.add_heading("狀態代號", level=2)
    add_table(doc, ["代號", "意義", "發布效果"], [
        ("EXACT", "關鍵定義一致", "可直接比較"),
        ("CROSSWALK_REQUIRED", "須用已版本化無損對照", "對照驗證後可比較"),
        ("PARTIAL_OR_CONFLICT", "只部分重疊或不可拆分", "只能並列"),
        ("PENDING_HUMAN_REVIEW", "人工證據未完成", "阻擋正式使用"),
        ("SCENARIO_NOT_OFFICIAL_FORECAST", "模型情境外推", "不得稱官方預測"),
    ], widths=[5.2, 6.1, 4.6], font_size=8.5)

    page_break(doc)
    doc.add_heading("3｜介接方式", level=1)
    add_flow(doc, ["來源登錄", "擷取", "原始快照", "格式驗證", "欄位正規化", "定義比對", "人工 Gate", "呈現／問答"])
    doc.add_paragraph()
    add_table(doc, ["來源類型", "資料代號", "做法", "成功證據", "異常處理"], [
        ("分頁 JSON", "MOI-POP1Y", "民國年月＋page 逐頁", "頁數、原檔、列數、雜湊", "重試後停止產製"),
        ("JSON", "NTPC-*", "資料集 API URL", "原 JSON、列數、雜湊", "欄位／空值異常停止"),
        ("XLSX", "STAT-WAGE-*", "官方附件＋工作表 selector", "原 XLSX、sheet、列欄、雜湊", "格式不符停止"),
        ("人工匯出", "OAS-POP", "網站查詢後下載", "條件、截圖、原檔、雜湊、覆核人", "標記待人工查核"),
        ("OCR", "OCR-DOC", "Textract／Tesseract", "頁碼、信心、逐欄複核", "低信心不得入指標"),
    ], widths=[2.6, 2.7, 4.1, 4.0, 3.8], font_size=8)
    doc.add_heading("介接契約共通參數", level=2)
    for item in [
        "逾時 30 秒、最多 3 次重試、退避 1／3／9 秒。",
        "原始檔版本化保存；不覆蓋前次快照；SHA-256 為完整性證據。",
        "fail-closed：必要來源或欄位失敗時，不更新最新正式指標。",
        "所有回答至少帶資料集代號、期別、母體、地理角色與限制。",
    ]:
        bullet(doc, item)
    add_callout(doc, "OAS 特別規則", "目前未確認穩定公開 API，因此只採人工匯出證據鏈。未具備查詢條件、截圖、原檔、SHA-256 與覆核人時，OAS 數字不得發布。", color=RED)

    # Landscape field matrix.
    section = doc.add_section()
    set_section(section, landscape=True)
    configure_header_footer(section, "新北市青年政策公開資料 AI 系統｜欄位對照")
    doc.add_heading("4｜各資料欄位比較表", level=1)
    doc.add_paragraph("列為系統標準參數，欄為參考資料庫。✓＝可對應；△＝部分對應／分組跨界／需轉換；—＝無此欄位。共同打勾欄位可列為未來介接候選；period、source_id、source_url、retrieved_at、sha256、review_status 優先作技術與稽核鍵，但仍須檢查母體與地理角色。")
    matrix_rows = [
        ("period", "期別", "✓", "✓", "△", "✓", "✓", "✓", "✓", "✓"),
        ("geo", "行政區／縣市", "✓", "✓", "✓", "新北", "新北", "新北", "全國", "工作地"),
        ("age", "年齡", "✓單歲", "△5歲", "依匯出", "△來源帶", "△來源帶", "—", "✓25–29", "✓25–29"),
        ("sex", "性別", "✓", "✓", "依匯出", "✓", "✓", "✓", "—", "—"),
        ("population", "人口數", "✓", "✓", "依匯出", "—", "—", "—", "—", "—"),
        ("unemp_rate", "失業率", "—", "—", "—", "✓", "—", "—", "—", "—"),
        ("employment", "就業結構", "—", "—", "—", "—", "✓", "—", "—", "—"),
        ("education", "教育程度", "—", "—", "—", "—", "—", "✓", "✓", "—"),
        ("salary_mean", "薪資平均", "—", "—", "—", "—", "—", "—", "✓", "✓"),
        ("salary_median", "薪資中位", "—", "—", "—", "—", "—", "—", "✓", "✓"),
        ("population_scope", "母體", "戶籍", "戶籍", "依匯出", "勞動力", "就業者", "15+人口", "受僱員工", "本國籍全時"),
        ("geography_role", "地理角色", "居住", "居住", "居住", "居住", "居住", "居住", "全國", "工作場所"),
        ("unit", "單位", "人", "人", "人", "%", "%", "%", "萬元", "萬元"),
        ("source_url", "官方 URL", "✓", "✓", "✓", "✓", "✓", "✓", "✓", "✓"),
        ("retrieved_at", "擷取時間", "✓", "✓", "人工", "✓", "✓", "✓", "✓", "✓"),
        ("sha256", "雜湊", "✓", "✓", "人工", "✓", "✓", "✓", "✓", "✓"),
    ]
    add_table(doc, ["標準參數", "定義", "MOI-POP1Y", "NTPC-POP", "OAS-POP", "NTPC-UNEMP", "NTPC-EMPSTR", "NTPC-EDU15P", "WAGE-EDU", "WAGE-LOC"], matrix_rows, widths=[2.5, 2.8, 2.2, 2.1, 2.0, 2.4, 2.4, 2.5, 2.2, 2.2], font_size=7.1)
    section = doc.add_section()
    set_section(section, landscape=False)
    configure_header_footer(section)
    doc.add_heading("5｜題目要求項目與符合度", level=1)
    requirements = [
        ("命題指定公開資料", "三平台均登錄；另納入 data.gov.tw／戶政司", "部分完成", "OAS 人工證據待補"),
        ("青年 18–35", "核心分三組；單一年齡精確人口", "完成", "其他率值只在來源邊界吻合時用"),
        ("人口、就業、薪資", "均已取得；另補教育", "完成", "母體不得混用"),
        ("互動式儀表板", "可篩選、看圖、查來源、問問題", "完成／私人", "未核准不得公開"),
        ("AI 系統", "AgentCore Harness 已部署", "部分完成", "bridge 未部署，仍是本地證據回答"),
        ("預測與洞察", "線性情境＋滾動回測 MAE", "研究版完成", "非官方預測、非因果"),
        ("可追溯可查核", "runId、原檔、URL、雜湊、規則、測試", "完成", "OAS／OCR 仍需人工"),
        ("安全與成本", "最小權限範本與人工 Gate", "待裁示", "Owner、預算、保存期、輪替"),
    ]
    add_table(doc, ["要求", "本版實作", "狀態", "未完成／限制"], requirements, widths=[3.8, 5.8, 2.7, 4.0], font_size=8.2)
    add_callout(doc, "本版缺口", "主要缺口不是資料視覺化，而是 OAS 證據、AWS bridge 實際部署、Custom Browser 網域政策及成本／保存期裁示。這些均維持 Gate 阻擋，不影響本地研究版重算。", color=RED)

    page_break(doc)
    doc.add_heading("6｜研究方法及步驟", level=1)
    for text in [
        "界定政策問題與完成定義：確認人口、就業、教育、薪資的決策用途。",
        "來源盤點與登錄：只收官方頁面及可查核附件；建立縮寫與版本。",
        "定義字典：記錄母體、年齡、地理、時間、單位、統計量與公式。",
        "資料擷取：依契約下載，保留原始快照、SHA-256、列數與時間。",
        "欄位正規化：將來源欄位映射到標準參數，但不消除來源語意。",
        "相容性判定：EXACT、CROSSWALK_REQUIRED、PARTIAL_OR_CONFLICT 或待人工裁示。",
        "指標計算：核心青年只用單一年齡；率與平均數不做比例拆分。",
        "分析與不確定性：描述趨勢；外推以尾端線性迴歸與滾動回測 MAE 表示。",
        "交叉驗證：比較官方頁面、附件、manifest 與儀表板輸出；異常停止發布。",
        "人工 Gate 與版本發布：具名查核、記錄部署 ID、保留回復路徑。",
    ]:
        numbered(doc, text)
    doc.add_heading("主要計算規則", level=2)
    add_table(doc, ["指標", "公式／方法", "禁止事項"], [
        ("核心青年人口", "Σ age=18…35（男＋女）", "不得由 15–19、35–39 按年數比例切割"),
        ("分組人口", "18–24、25–29、30–35 各自單歲加總", "15–17、36–40 不進核心"),
        ("25–29 失業率", "直接使用來源 25–29×性別欄", "不得外推全部 18–35"),
        ("教育×薪資", "群體趨勢與分層薪資並列", "不得建立個體因果解讀"),
        ("短期情境", "尾端視窗 OLS＋rolling-origin MAE", "不得稱官方預測或政策成效"),
    ], widths=[3.8, 6.2, 5.8], font_size=8.5)

    doc.add_heading("7｜18–35 歲分類方式及限制", level=1)
    add_table(doc, ["組別代號", "年齡", "政策階段", "是否計入核心"], [
        ("AGE-OBS-15-17", "15–17", "青年銜接觀察組", "否"),
        ("AGE-CORE-18-24", "18–24", "就學、初入職場與轉銜階段", "是"),
        ("AGE-CORE-25-29", "25–29", "職涯建立階段", "是"),
        ("AGE-CORE-30-35", "30–35", "職涯、居住與家庭形成階段", "是"),
        ("AGE-EXT-36-40", "36–40", "方案擴充組", "否"),
    ], widths=[4.2, 2.4, 6.6, 2.8], font_size=9)
    doc.add_heading("使用限制", level=2)
    for item in [
        "18–24 若來源只有 15–24，無法精確拆出；保留來源分組並標示部分對應。",
        "30–35 若來源分成 30–34 與 35–39，35 歲無法從後者無損抽出；率與平均數尤其不可比例拆分。",
        "25–29 是目前人口、失業率及薪資可共同呼應最完整的年齡帶，但三者母體與地理角色仍不同。",
        "戶籍人口描述設籍人口；不等於實住人口、勞動力、就業者或受僱員工。",
        "分組是政策分析層，不代表法律定義或個人身分；對外使用須引述核定版本。",
    ]:
        bullet(doc, item)
    add_callout(doc, "衝突處理原則", "先保留原值與來源定義，再判斷是否存在可驗證、無損的 crosswalk；無法無損轉換時並列呈現並交具名裁示，不能以 AI 猜測或補值。", color=RED)

    doc.add_heading("8｜呈現圖表及方式", level=1)
    chart_rows = [
        ("KPI 卡", "核心青年總數、占比、資料期別", "人口", "卡片旁固定顯示戶籍口徑"),
        ("年齡軌／堆疊條", "15–17、三核心組、36–40", "人口", "核心與非核心用標籤及線型雙重區隔"),
        ("折線圖", "25–29 失業率男女趨勢", "就業", "縱軸%、資料點、來源年份"),
        ("群組長條", "25–29 各教育程度平均／中位薪資", "薪資", "全國受僱員工；平均與中位分開"),
        ("雙面板並列", "教育結構趨勢＋薪資教育分層", "教育×薪資", "明示群體資料、不可作因果"),
        ("實績＋虛線情境", "次年線性外推與 MAE 範圍", "研究情境", "虛線、範圍帶、非官方預測標籤"),
        ("來源狀態表", "ACQUIRED／待人工、雜湊前綴、列數", "治理", "每列可開官方頁與 manifest"),
        ("查核看板", "定義、資料、AWS、發布 Gate", "行政", "阻擋項置頂；顯示責任人與證據"),
    ]
    add_table(doc, ["圖表", "呈現內容", "主題", "必要標示"], chart_rows, widths=[3.5, 5.3, 2.7, 4.4], font_size=8.2)
    doc.add_heading("視覺與無障礙規則", level=2)
    for item in [
        "圖名回答『誰、哪裡、何時、什麼統計量』；單位不可只放在滑鼠提示。",
        "顏色不作唯一差異；搭配標籤、圖例、線型、符號或直接數值。",
        "預測與實績分段；示意範圍不是信賴區間時，不使用信賴區間名稱。",
        "資料不足或衝突時顯示缺口，不以 0 或平滑線填補。",
    ]:
        bullet(doc, item)

    page_break(doc)
    doc.add_heading("9｜可運用的方向", level=1)
    add_table(doc, ["方向", "可回答問題", "所需資料", "邊界／人工裁示"], [
        ("青年人口資源配置", "各年齡階段規模與行政區分布？", "MOI-POP1Y", "戶籍不等於實住；小區域須注意隱私"),
        ("就學至就業轉銜", "25–29 失業率趨勢是否偏離？", "NTPC-UNEMP＋人口", "率值母體不同，不計失業人數"),
        ("職涯與薪資觀察", "新北工作場所 25–29 薪資趨勢？", "STAT-WAGE-LOC", "不是居住新北青年薪資"),
        ("教育環境與薪資描述", "教育結構與教育分層薪資如何變動？", "NTPC-EDU15P＋WAGE-EDU", "不可推論教育造成薪資"),
        ("政策服務分眾", "三核心階段需要何種服務入口？", "人口＋行政服務資料", "需再納入方案使用／成效資料"),
        ("資料品質監測", "來源是否過期、改欄或中斷？", "manifest＋契約", "異常自動阻擋、人工確認"),
        ("承辦人研究助理", "來源、定義、限制與圖表如何查？", "curated＋來源清冊＋AI", "回答須帶引用，不代替行政裁量"),
        ("方案成效評估", "介入前後是否改善？", "尚需個案、暴露、對照與結果資料", "本資料包不足以作因果成效"),
    ], widths=[3.7, 5.0, 3.8, 4.0], font_size=8)
    add_callout(doc, "最適合的近期用途", "先作人口／失業／教育／薪資的政策情勢儀表板與承辦人證據問答；若要評估方案成效，必須另建服務使用、時間、對照與結果指標，不能只靠公開橫斷資料。", color=GREEN)

    page_break(doc)
    doc.add_heading("10｜AWS Harness AI Agent 建立流程", level=1)
    add_flow(doc, ["私人 Sites", "/api/ask", "API Gateway", "Lambda", "InvokeHarness", "AgentCore 工具", "CloudWatch"])
    doc.add_paragraph()
    doc.add_heading("目前與目標狀態", level=2)
    add_table(doc, ["元件", "目前狀態", "目標／驗收"], [
        ("AgentCore Harness", "已部署，us-west-2，DEFAULT V3", "可回答人口、薪資、定義衝突；工具最小化"),
        ("互動式儀表板", "已私人部署；本地證據回答", "bridge 完成後才標示 AWS_HARNESS"),
        ("API bridge", "Lambda＋SAM/CloudFormation 範本已完成，未部署", "API key 僅伺服器端；request ID 可回查"),
        ("Browser", "內建 Browser＋提示限制", "Custom Browser 預設阻擋，只放行官方網域"),
        ("Knowledge Base／S3", "尚未建立", "僅上傳核准 PUBLIC 文件；保留版本與加密"),
        ("成本／日誌", "待裁示", "Owner、CostCenter、月預算、告警、輪替與保存期"),
    ], widths=[3.8, 5.3, 6.7], font_size=8.3)
    future_heading = doc.add_heading("未來 AWS 數據處理方式", level=2)
    future_heading.paragraph_format.page_break_before = True
    for idx, item in enumerate([
        "Amazon EventBridge Scheduler 觸發擷取工作；Step Functions 編排來源下載、驗證與人工 Gate。",
        "S3 以 raw／curated／published 分層保存，開啟版本、加密、Block Public Access 與生命週期。",
        "AWS Glue Data Catalog／Athena 提供可查詢資料表；資料契約改版須保留 schema 版本。",
        "Lambda 或受控 Code Interpreter 執行確定性計算；Textract 只用於掃描件並保留信心與覆核。",
        "AgentCore Harness 只讀核准的 curated 證據；CloudWatch 保存 request ID、錯誤與版本，不記錄秘密。",
        "儀表板以 API Gateway 私密橋接；API key／憑證存 Secrets Manager 或伺服器端 secret，定期輪替。",
    ], start=1):
        doc.add_paragraph(f"{idx}. {item}")
    add_callout(doc, "上線 Gate", "Owner、CostCenter、預算告警、API key 輪替、日誌保存與 Custom Browser 政策未核定前，不部署 bridge，不改為公開站台。", color=RED)

    page_break(doc)
    doc.add_heading("11｜人工查核與裁示表", level=1)
    checklist = [
        ("HR-OAS-01", "OAS 條件、截圖、原檔、雜湊、覆核人", "阻擋 OAS 數值"),
        ("HR-OCR-01", "OCR 數值逐欄雙人覆核", "阻擋 OCR 數值"),
        ("HR-AGE-01", "18–35 單一年齡精確加總", "阻擋發布"),
        ("HR-GEO-01", "居住／工作場所／全國分開", "阻擋發布"),
        ("HR-RES-01", "教育與薪資未作個體因果", "阻擋發布"),
        ("HR-FCST-01", "目標年、訓練期、MAE、結構轉折", "年度查核"),
        ("HR-AWS-02", "Owner、CostCenter、預算、告警、到期日", "阻擋 bridge"),
        ("HR-API-01", "API Gateway→Lambda→Harness→CloudWatch E2E", "阻擋 AWS 即時標示"),
        ("HR-PUB-01", "每個數值可回指 URL、期別、版本", "阻擋發布"),
        ("HR-PUB-03", "內容、受眾、管道、發布人與時間核准", "阻擋發布"),
    ]
    add_table(doc, ["代號", "查核事項", "Gate", "結果", "查核人／日期", "證據"], [r + ("□通過 □退回 □N/A", "", "") for r in checklist], widths=[2.8, 6.0, 3.0, 3.7, 3.5, 3.2], font_size=7.5)
    doc.add_heading("本版待裁示", level=2)
    add_table(doc, ["代號", "事項", "目前作法", "裁示"], [
        ("DEC-OAS-01", "OAS 是否納入正式數值", "只作人工查核來源", ""),
        ("DEC-AWS-01", "成本責任與預算", "bridge 未部署", ""),
        ("DEC-API-01", "API key 輪替／日誌保存", "待核定", ""),
        ("DEC-BROWSER-01", "Custom Browser 網域政策", "待部署", ""),
        ("DEC-FCST-01", "情境預測是否可對外", "只在私人研究版呈現", ""),
    ], widths=[3.2, 5.2, 5.2, 2.5], font_size=8.5)

    doc.add_heading("附錄 A｜本次批次證據", level=1)
    artifact_rows = []
    for item in manifest["artifacts"]:
        status_label = "待人工查核" if item["status"] == "PENDING_HUMAN_REVIEW" else item["status"]
        artifact_rows.append((item["sourceId"], status_label, item.get("rowCount") or "—", item.get("byteCount") or "—", (item.get("sha256") or "—")[:12]))
    add_table(doc, ["來源", "狀態", "列數", "Bytes", "SHA-256 前綴"], artifact_rows, widths=[3.7, 4.0, 2.5, 3.1, 3.2], font_size=8.2)
    doc.add_paragraph(f"runId：{manifest['runId']}；擷取時間：{manifest['retrievedAt']}；failClosed：{manifest['failClosed']}。", style="Small Note")
    doc.add_heading("品質檢驗", level=2)
    for item in [
        "Python 管線測試：10 項通過。",
        "AWS bridge 單元測試：2 項通過。",
        "儀表板 lint、build、路由測試及私人站台視覺檢查：通過。",
        "OAS-POP、OCR-DOC：PENDING_HUMAN_REVIEW。",
    ]:
        bullet(doc, item)

    page_break(doc)
    doc.add_heading("附錄 B｜官方來源與重用說明", level=1)
    official = [
        ("政府資料開放平臺", "https://data.gov.tw/", "命題指定公開資料入口"),
        ("單一年齡人口資料集", "https://data.gov.tw/dataset/77132", "18–35 人口基準"),
        ("新北資料開放平臺", "https://data.ntpc.gov.tw/openapi/", "人口、就業、教育 JSON"),
        ("新北市統計資料庫", "https://oas.bas.ntpc.gov.tw/DgbasWeb/Page/Default.aspx", "人工匯出鏡像查核"),
        ("主計總處薪資發布頁", "https://www.stat.gov.tw/News_Content.aspx?n=4580&s=232642", "薪資表 5、表 6"),
        ("AWS InvokeHarness API", "https://docs.aws.amazon.com/bedrock-agentcore/latest/APIReference/API_InvokeHarness.html", "Harness bridge API"),
        ("AWS Harness Security", "https://docs.aws.amazon.com/bedrock-agentcore/latest/devguide/harness-security.html", "IAM 與安全基準"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, h in enumerate(("來源", "官方連結", "用途")):
        table.cell(0, idx).text = h
        shade(table.cell(0, idx), LIGHT)
    for name, url, use in official:
        cells = table.add_row().cells
        cells[0].text = name
        hyperlink(cells[1].paragraphs[0], "開啟官方頁面", url)
        cells[2].text = use
    for row in table.rows:
        no_split(row)
        for cell in row.cells:
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(8)
                    set_east_asia(run)
    set_repeat_table_header(table)
    doc.add_heading("資料包重用順序", level=2)
    for idx, item in enumerate([
        "讀取來源清冊與簡寫表，確認本次題目需要哪些資料。",
        "沿用介接契約取新批次；不可直接改動既有 raw 快照。",
        "依欄位比較表判斷可 join 的技術鍵與不可混用的語意欄。",
        "執行管線、測試與人工 SOP，再建立新版本資料包。",
        "只有通過 Gate 的 curated 結果可進入儀表板或 AI 問答。",
    ], start=1):
        p = doc.add_paragraph(f"{idx}. {item}")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            run.font.size = Pt(9.5)
    add_callout(doc, "交付內容", "可沿用研究基線；無 AWS 憑證、API key、Cookie 或個資。", color=GREEN)
    trailing = doc.paragraphs[-1]
    if not trailing.text:
        trailing._element.getparent().remove(trailing._element)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
