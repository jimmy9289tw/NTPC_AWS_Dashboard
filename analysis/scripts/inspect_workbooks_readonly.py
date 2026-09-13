from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


def compact_row(values):
    return [[idx, value] for idx, value in enumerate(values, start=1) if value not in (None, "")]


for raw_path in sys.argv[1:]:
    path = Path(raw_path)
    print(f"\n=== {path.name} ===")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("sheets=" + json.dumps(workbook.sheetnames, ensure_ascii=False))
    for worksheet in workbook.worksheets:
        print(f"-- {worksheet.title}: rows={worksheet.max_row}, cols={worksheet.max_column}")
        shown = 0
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            compact = compact_row(row)
            if compact:
                print(f"R{row_number}: " + json.dumps(compact, ensure_ascii=False, default=str))
                shown += 1
            if shown >= 45:
                break
