"""Create the governed 18--35 age-integration methodology report."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF4CE"
PALE_RED = "FDECEC"
MUTED = "5B6573"
WHITE = "FFFFFF"
BLACK = "000000"
CJK_FONT = "Microsoft JhengHei"
LATIN_FONT = "Calibri"


def set_run_font(run, *, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None,
                 latin: str = LATIN_FONT, east_asia: str = CJK_FONT) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_text(cell, text: Any, *, bold: bool = False, color: str = BLACK,
                  size: float = 8.5, align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run("" if text is None else str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    suffix = paragraph.add_run(" 頁")
    set_run_font(suffix, size=9, color=MUTED)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_p.paragraph_format.space_after = Pt(0)
        run = header_p.add_run("新北市青年政策公開資料AI系統｜年齡整合方法")
        set_run_font(run, size=8.5, color=MUTED, bold=True)

        footer = section.footer
        footer_p = footer.paragraphs[0]
        add_page_field(footer_p)


def add_para(doc: Document, text: str = "", *, size: float = 11, bold: bool = False,
             color: str = BLACK, italic: bool = False,
             align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
             before: float = 0, after: float = 6, keep_next: bool = False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.10
    paragraph.paragraph_format.keep_with_next = keep_next
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return paragraph


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_callout(doc: Document, label: str, text: str, *, fill: str = PALE_BLUE,
                label_color: str = DARK_BLUE) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.10
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), label_color)
    borders.append(left)
    p_pr.append(borders)
    label_run = paragraph.add_run(f"{label}｜")
    set_run_font(label_run, size=10.5, bold=True, color=label_color)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10.5, color=BLACK)


def add_table(doc: Document, headers: list[str], data: Iterable[Iterable[Any]],
              widths: list[int], *, font_size: float = 8.5,
              header_fill: str = LIGHT_GRAY,
              keep_rows_together: bool = False) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=INK, size=font_size)
        set_cell_fill(table.rows[0].cells[index], header_fill)
    set_repeat_header(table.rows[0])
    for row_index, row_data in enumerate(data, start=1):
        row = table.add_row()
        if keep_rows_together:
            row_properties = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            row_properties.append(cant_split)
        for col_index, item in enumerate(row_data):
            set_cell_text(row.cells[col_index], item, size=font_size)
            if row_index % 2 == 0:
                set_cell_fill(row.cells[col_index], "FAFBFC")
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)


def source_note(doc: Document, text: str) -> None:
    paragraph = add_para(doc, text, size=8.5, color=MUTED, before=4, after=4)
    paragraph.paragraph_format.keep_with_next = False


def format_number(value: Any, decimals: int = 1) -> str:
    if value in (None, ""):
        return "—"
    number = float(value)
    if number.is_integer():
        return f"{int(number):,}"
    return f"{number:,.{decimals}f}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-json", type=Path, required=True)
    parser.add_argument("--input-csv", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--table-helper", type=Path, required=True)
    args = parser.parse_args()

    sys.path.insert(0, str(args.table_helper.parent))
    global apply_table_geometry
    from table_geometry import apply_table_geometry

    payload = json.loads(args.input_json.read_text(encoding="utf-8"))
    with args.input_csv.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))

    def pick(**criteria: Any) -> dict[str, str]:
        for row in rows:
            if all(row.get(key) == str(expected) for key, expected in criteria.items()):
                return row
        raise KeyError(criteria)

    summary = payload["summary"]
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    section = doc.sections[0]

    # Memo masthead: serious technical report, no decorative border.
    add_para(doc, "研究方法與資料治理報告", size=10.5, bold=True, color=BLUE, after=4)
    add_para(doc, "18–35歲青年年齡資料整合方法", size=24, bold=True, color=INK, after=4, keep_next=True)
    add_para(doc, "25–29歲第一階段驗證、18–35歲擴充規則與完整資料整合表說明",
             size=13, color=MUTED, after=14)
    metadata = [
        ("政策母體", "18–35歲；子群18–24、25–29、30–35"),
        ("驗證基準", "25–29歲"),
        ("證據快照", f"本地管線 {payload['generatedFromRunId']}；另納入主計總處2025年官方表27、32、37、50"),
        ("文件版本", "V1.1｜產製日期 2026-08-16"),
        ("發布原則", "發布機關與統計生成方式分欄；官方調查值仍標示為估計值"),
    ]
    for label, detail in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        r1 = p.add_run(f"{label}：")
        set_run_font(r1, size=10, bold=True, color=INK)
        r2 = p.add_run(detail)
        set_run_font(r2, size=10, color=BLACK)

    doc.add_heading("決策摘要", level=1)
    add_callout(
        doc,
        "結論",
        f"第一階段以25–29歲驗證是可行的；目前可用官方來源涵蓋人口、就業、失業、教育與薪資。"
        f"但整體18–35歲目前只有戶籍人口可由官方單一年齡行政紀錄確定性生成：{summary['exact18To35Population']:,}人，"
        f"占新北市總人口{summary['exact18To35PopulationSharePct']:.4f}%。其餘指標必須取得18、19及35歲的單一年齡資料、"
        "分子分母或微資料後才能形成正式18–35統計。",
    )
    add_bullet(doc, "整合的核心不是把所有來源硬切成相同年齡，而是以共同目標年齡、來源原生值、轉換方法及證據等級共同管理。")
    add_bullet(doc, "第一階段只建議發布E0官方直接值、E1精確加總值及E2由官方分子分母重算值。")
    add_bullet(doc, "『官方發布』不等於『完整計數』：人力資源調查由政府發布，但仍是抽樣加權估計；CSV已把兩個維度分欄。")
    add_bullet(doc, "M3校準或小區域模型只能放在研究估計層，必須有不確定區間、方法版本、回測及人工核准。")
    add_bullet(doc, "30–39歲薪資即使可拆成30–34與35–39，仍不能直接得到30–35；35歲仍需單一年齡或微資料。")

    doc.add_heading("研究目的與適用範圍", level=1)
    add_para(doc, "本方法用於新北市青年政策公開資料AI系統，把不同機關、不同母體、不同年齡帶的官方資料整理成可追溯、可重現、可查核的分析層。政策核心青年固定為18–35歲；資料不足時保留來源原生年齡帶，不以生成式AI補值。")
    add_para(doc, "本報告回答四個實務問題：", bold=True, color=INK, keep_next=True)
    for text in (
        "每份資料原本涵蓋哪些年齡，以及能否直接對應18–35歲。",
        "25–29歲是否足以作為第一階段的資料管線與定義驗證基準。",
        "何種統計量可以精確重組，何種必須使用模型或阻擋。",
        "儀表板與AI問答在18–35歲範圍內可以安全呈現哪些欄位。",
    ):
        add_number(doc, text)

    doc.add_heading("統一年齡整合架構", level=1)
    add_callout(
        doc,
        "治理原則",
        "共同資料模型統一的是欄位與證據狀態，不是假設每個統計量都能被切成18–35。原始值永遠保留；調整值另欄保存。",
        fill=PALE_GOLD,
        label_color="7A5A00",
    )
    method_rows = [(item["code"], item["definition"], item["publication"]) for item in payload["methodCatalog"]]
    add_table(doc, ["方法代號", "定義", "發布規則"], method_rows, [1900, 4800, 2660], font_size=9)

    doc.add_heading("官方精確值與官方調查估計值的差異", level=2)
    add_callout(
        doc,
        "判讀方式",
        "『官方』回答誰發布；『精確／估計』回答數字如何產生。行政紀錄通常是對既定登記母體的完整計數；抽樣調查則只訪問樣本，再以權數推估母體。",
        fill=PALE_GOLD,
        label_color="7A5A00",
    )
    origin_rows = [(item["label"], item["meaning"]) for item in payload["valueOriginCatalog"]]
    add_table(doc, ["CSV標記", "意義與發布要求"], origin_rows, [2900, 6460], font_size=8.7)
    add_bullet(doc, "行政精確值仍不是無條件的『真實世界全體』：它只對來源定義的行政母體、統計截止日及登記狀態成立。")
    add_bullet(doc, "官方調查估計值有官方品質管理，但仍應揭露抽樣設計、權數、發布位數、標準誤或信賴區間。")
    add_bullet(doc, "目前CSV沒有任何已填值的PROJECT_MODEL_ESTIMATE；資料不足的18–35欄位保持空白。")

    doc.add_heading("不同統計量的正確重組公式", level=2)
    formula_items = [
        ("計數", "N(18–35)=Σ N(a)，a=18,…,35。必須有完整單一年齡或完全落在邊界內的可加總群組。"),
        ("比率", "R=Σ numerator / Σ denominator。不可直接平均各年齡組失業率、勞參率或教育占比。"),
        ("平均數", "Mean=Σ(wg×Mean_g)/Σwg；權重必須是同一母體的觀察數或官方權重。"),
        ("中位數", "必須從微資料或完整分布重算加權第50百分位；彙總中位數不能加權合成。"),
        ("校準拆分", "若只有30–39平均薪資M，可用輔助收入比r及受僱人數n校準30–34與35–39；結果是模型估計，不是官方統計。"),
    ]
    add_table(doc, ["統計量", "可查核處理"], formula_items, [1800, 7560], font_size=9.2)
    source_note(doc, "方法基礎：校準估計（Deville & Särndal, 1992）及小區域估計的一般原則；本計畫採保守發布門檻。")

    doc.add_heading("資料來源年齡範圍與18–35對應", level=1)
    inventory = [
        ("MOI-POP1Y", "單一年齡", "戶籍人口", "18–35可精確加總", "正式主來源"),
        ("NTPC-POP", "5歲組：15–19、20–24、25–29、30–34、35–39…", "戶籍人口", "18、19、35歲跨界", "分頁已完整；僅作原生趨勢查核" if summary.get("ntpcPopPaginationPassed") else "目前快照不完整，封鎖"),
        ("NTPC-UNEMP", "15–24、25–29、30–34、35–39…", "勞動力", "25–29精確；18–35缺分子分母", "來源原生值可呈現"),
        ("NTPC-EMPSTR", "15–24、25–29、30–34、35–39…", "就業者", "25–29精確；18–35缺單歲就業數", "來源原生值可呈現"),
        ("NTPC-EDU15P", "15歲以上", "15歲以上人口", "無法拆出18–35", "僅教育環境趨勢"),
        ("STAT-WAGE-EDU", "未滿25、25–29、30–39…", "全國受僱員工", "25–29精確；30–35跨界", "25–29可發布"),
        ("STAT-WAGE-LOC", "未滿25、25–29、30–39…", "工作場所全時受僱員工", "25–29精確；30–35跨界", "須標示工作地"),
        ("DGBAS-HR-T27/T32/T37", "15–24、25–29、30–34、35–39…", "民間人口／就業者／勞動力", "25–29精確；18–35需單歲邊界", "輔助驗證"),
        ("DGBAS-HR-T50", "15–19、20–24、25–29、30–34、35–39…", "全國就業者", "25–29精確；18、19、35仍跨界", "教育×年齡驗證"),
        ("MOI-EDU5Y", "五歲年齡組15歲以上", "戶籍人口", "25–29可直接統計；18／19／35仍有邊界", "已取得並完成粒度驗證"),
    ]
    add_table(doc, ["資料代號", "來源年齡帶", "統計母體", "對18–35處理", "狀態"], inventory,
              [1550, 2100, 1650, 2500, 1560], font_size=7.8)
    source_note(doc, "官方欄位依政府資料開放平臺、新北市資料開放平臺及行政院主計總處2025年人力資源調查表。")

    doc.add_heading("第一階段：25–29歲驗證結果", level=1)
    add_para(doc, "25–29歲是目前最適合的基準組，因為人口、失業率、就業結構、薪資及就業者教育程度均有來源原生或精確加總值，不需處理18、19或35歲邊界。此階段驗證資料管線與定義治理，不宣稱不同母體間數值應完全相等。")

    p_25 = pick(source_alias="MOI-POP1Y", measure_code="resident_population_count", target_age_band="25–29", sex="all")
    u_m = pick(source_alias="NTPC-UNEMP", measure_code="unemployment_rate_pct", source_age_band="25–29", sex="male")
    u_f = pick(source_alias="NTPC-UNEMP", measure_code="unemployment_rate_pct", source_age_band="25–29", sex="female")
    e_m = pick(source_alias="NTPC-EMPSTR", measure_code="employment_share_pct", source_age_band="25–29", sex="male")
    e_f = pick(source_alias="NTPC-EMPSTR", measure_code="employment_share_pct", source_age_band="25–29", sex="female")
    t27 = pick(source_alias="DGBAS-HR-T27", measure_code="civilian_population_count", source_age_band="25–29")
    t32 = pick(source_alias="DGBAS-HR-T32", measure_code="employed_person_count", source_age_band="25–29")
    emp_ratio = pick(source_alias="DGBAS-HR-T27+T32", measure_code="employment_population_ratio_pct")
    t37 = pick(source_alias="DGBAS-HR-T37", measure_code="unemployment_rate_pct", sex="all")
    edu_share = pick(source_alias="DGBAS-HR-T50", measure_code="college_plus_share_pct")
    wage_mean = pick(source_alias="STAT-WAGE-EDU", measure_code="annual_salary_mean_10k_twd", education_level="全體")
    wage_median = pick(source_alias="STAT-WAGE-EDU", measure_code="annual_salary_median_10k_twd", education_level="全體")
    wage_loc_mean = pick(source_alias="STAT-WAGE-LOC", measure_code="annual_salary_mean_10k_twd")
    wage_loc_median = pick(source_alias="STAT-WAGE-LOC", measure_code="annual_salary_median_10k_twd")

    validation_rows = [
        ("戶籍人口", p_25["source_period"], f"{format_number(p_25['adjusted_value'], 0)} 人", "MOI-POP1Y", "行政紀錄確定性計算", "單一年齡加總；無抽樣誤差"),
        ("失業率－男／女", u_m["source_period"], f"{format_number(u_m['source_value'])}%／{format_number(u_f['source_value'])}%", "NTPC-UNEMP", "官方抽樣調查估計", "來源原生25–29；方法連結待確認"),
        ("就業者年齡結構－男／女", e_m["source_period"], f"{format_number(e_m['source_value'], 2)}%／{format_number(e_f['source_value'], 2)}%", "NTPC-EMPSTR", "官方抽樣調查估計", "各性別分母不同"),
        ("15歲以上民間人口", t27["source_period"], f"{format_number(t27['source_value'], 0)} 千人", "DGBAS-HR-T27", "官方抽樣調查估計", "千人四捨五入"),
        ("就業者人數", t32["source_period"], f"{format_number(t32['source_value'], 0)} 千人", "DGBAS-HR-T32", "官方抽樣調查估計", "千人四捨五入"),
        ("就業人口比率", emp_ratio["source_period"], f"{format_number(emp_ratio['adjusted_value'], 2)}%", "T27+T32", "調查估計值再計算", "215÷246；承接抽樣與四捨五入誤差"),
        ("失業率－全體", t37["source_period"], f"{format_number(t37['source_value'])}%", "DGBAS-HR-T37", "官方抽樣調查估計", "用於跨年欄位驗證"),
        ("就業者大專及以上占比", edu_share["source_period"], f"{format_number(edu_share['adjusted_value'], 2)}%", "DGBAS-HR-T50", "調查估計值再計算", "994÷1,253；全國就業者"),
        ("全年總薪資平均／中位", wage_mean["source_period"], f"{format_number(wage_mean['source_value'])}／{format_number(wage_median['source_value'])} 萬元", "STAT-WAGE-EDU", "官方行政大數據統計", "所得稅連結保險檔；全國受僱員工"),
        ("新北工作場所薪資平均／中位", wage_loc_mean["source_period"], f"{format_number(wage_loc_mean['source_value'])}／{format_number(wage_loc_median['source_value'])} 萬元", "STAT-WAGE-LOC", "官方行政大數據統計", "工作場所口徑"),
    ]
    add_table(doc, ["指標", "期別", "25–29數值", "來源", "數值性質", "方法／限制"], validation_rows,
              [1550, 800, 1400, 1200, 1850, 2560], font_size=7.7)

    doc.add_heading("驗證判定", level=2)
    add_bullet(doc, "結構驗證通過：25–29可在多個主題以原生年齡帶使用，無需比例拆分。")
    add_bullet(doc, "定義驗證通過：CSV已分開保存地理角色、母體、性別、教育層級、期別及單位。")
    add_bullet(doc, "數值一致性不以『相等』為標準：2026戶籍人口、2025民間人口與2025就業者是不同母體；2024與2025失業率也不是同年比較。")
    add_bullet(doc, "第一階段不能直接外推成18–35：25–29只是管線與方法的基準測試組。")

    doc.add_heading("第二階段：擴充至18–35歲", level=1)
    add_para(doc, "擴充採欄位級管理。每個指標都要先確認母體、地理角色、年齡邊界、期別、單位與測量型態，再決定是精確重組、研究估計或阻擋。")
    availability = [
        ("resident_population_count", "戶籍人口數", "MOI-POP1Y單一年齡", "可", f"{summary['exact18To35Population']:,}人", "E1"),
        ("resident_population_share_pct", "占全市人口比率", "MOI-POP1Y分子／分母", "可", f"{summary['exact18To35PopulationSharePct']:.4f}%", "E2"),
        ("education_by_level_count/share", "教育程度人數／占比", "MOI ODRP053五歲年齡組", "25–29已取得", "18–35邊界仍留空", "E5"),
        ("labor_force_participation_rate", "勞動力參與率", "人力資源微資料／客製表", "待申請", "目前留空", "E5"),
        ("employment_population_ratio", "就業人口比率", "就業者與民間人口分子分母", "待申請", "目前留空", "E5"),
        ("unemployment_rate_pct", "失業率", "失業者與勞動力分子分母", "待申請", "目前留空", "E5"),
        ("annual_salary_mean", "全年總薪資平均數", "受僱人數權重＋微資料或校準", "研究估計可行", "未核准不發布", "E3/E5"),
        ("annual_salary_median", "全年總薪資中位數", "個體分布／加權分位數", "待微資料", "目前留空", "E5"),
    ]
    add_table(doc, ["標準欄位", "中文指標", "所需資料", "18–35狀態", "目前數值", "證據"], availability,
              [1900, 1500, 2250, 1200, 1550, 960], font_size=7.9)

    doc.add_heading("18、19及35歲邊界處理", level=2)
    add_table(doc, ["目標群", "常見來源帶", "問題", "正式處理"], [
        ("18–24", "15–24或15–19＋20–24", "混入15–17，或無法只取18–19", "取得單一年齡18、19或官方18–24表；否則來源原生呈現"),
        ("25–29", "25–29", "通常完全一致", "直接使用，作第一階段基準"),
        ("30–35", "30–34＋35–39或30–39", "35歲與36歲以上混合", "取得35歲單一年齡或微資料；不得按1/5或1/10拆分率、中位數"),
        ("18–35", "多組彙總", "不同母體、分母與期別", "逐指標重算；不能把不同母體相加或平均"),
    ], [1300, 1850, 2450, 3760], font_size=8.5)

    doc.add_heading("改善邊界問題的資料與統計方式", level=2)
    boundary_methods = [
        ("人口", "MOI-POP1Y單一年齡", "直接加總18至35歲", "正式精確值", "已完成"),
        ("教育", "戶政ODRP053：五歲年齡組×教育程度×人口數", "25–29直接加總；18／19／35須客製表或M3估計", "部分精確／研究估計", "已取得並驗證"),
        ("就業／失業", "人力資源調查微資料或官方客製表", "以個體年齡與調查權數計算加權分子分母；以複雜抽樣方法估標準誤與95%區間", "官方客製或研究估計", "待申請"),
        ("薪資平均數", "受僱人數權重或個體行政資料", "Σ(人數×平均薪資)÷Σ人數；只在同母體且權重完整時使用", "可重算／估計", "邊界資料待取得"),
        ("薪資中位數", "個體分布或足夠細的薪資級距", "以個體權重重算第50百分位；不能合成已發布中位數", "需微資料", "待申請"),
        ("替代研究法", "五歲組＋單歲人口邊際＋歷史型態", "校準、MRP或小區域估計，並輸出區間、回測及敏感度", "研究模型估計", "不得覆寫官方值"),
    ]
    add_table(doc, ["主題", "優先資料", "統計處理", "結果性質", "目前狀態"], boundary_methods,
              [1050, 2050, 3400, 1550, 1310], font_size=7.6)
    add_callout(doc, "正式發布優先序", "官方客製18–35表 > 可申請微資料重算 > 行政單一年齡交叉表 > 研究模型估計。若只能使用模型，必須與官方值分層並附不確定區間。", fill=PALE_GOLD, label_color="7A5A00")

    doc.add_heading("NTPC人口API分頁：問題與修正", level=1)
    add_para(doc, "新北市資料頁顯示至少2,250列，但原介接網址固定page=0&size=1000，因此只下載第一頁1,000列；後面的page=1與page=2沒有進入本地快照。這是『介接不完整』，不是官方缺資料。")
    add_callout(
        doc,
        "本版結果",
        (
            f"分頁修正已執行並通過：共取得{summary.get('ntpcPopRowCount', 0):,}列，年度範圍{summary.get('ntpcPopYearSpan', 'unknown')}。"
            if summary.get("ntpcPopPaginationPassed")
            else f"分頁連接器已修正，但本地快照仍只有{summary.get('ntpcPopRowCount', 0):,}列；重新取得並通過完整性檢查前維持封鎖。"
        ),
        fill=PALE_GOLD,
        label_color="7A5A00",
    )
    for text in (
        "依官方開發指引改用page與size參數，從page=0開始逐頁下載；每頁1,000列時應至少取得page=0、1、2。",
        "以『最後一頁筆數小於size』作停止條件，並設最大頁數避免API重複回傳同一頁造成無限迴圈。",
        "合併後檢查總列數、重複頁、最大年度、欄位結構與SHA-256；未達官方目錄最低列數即fail closed。",
        "即使分頁修好，NTPC-POP仍是五歲年齡組，只適合歷史趨勢與鏡像查核；18、19、35歲正式人口仍以戶政單一年齡資料解決。",
    ):
        add_number(doc, text)

    doc.add_heading("完整CSV數據整合表", level=1)
    add_para(doc, f"隨附CSV共有{summary['recordCount']}筆紀錄、39個欄位。每一列是一個來源×期別×年齡帶×性別×教育層級×指標的可稽核紀錄；原始值、調整值與數值性質分欄保存。")
    dictionary = [
        ("識別與階段", "record_id, phase, topic, source_alias, source_name"),
        ("定義維度", "source_period, source_geography, geography_role, population_scope, source_age_band, target_age_band, sex, education_level"),
        ("指標與數值", "measure_code, measure_name, source_value, source_unit, adjusted_value, adjusted_unit"),
        ("方法與證據", "adjustment_method_code, value_status, evidence_class, interval_overlap_percent, can_answer_18_35, publish_status"),
        ("官方／估計標記", "value_origin_class, value_origin_label_zh, official_published_value, is_statistical_estimate, calculation_or_estimation_method, uncertainty_precision_note, method_verification_status"),
        ("來源證據鏈", "source_url, local_snapshot, sha256"),
        ("治理欄位", "definition_conflict, human_review_id, human_review_required, notes"),
    ]
    add_table(doc, ["欄位群", "CSV欄位"], dictionary, [1900, 7460], font_size=8.5)
    add_callout(doc, "機器可讀規則", "空白adjusted_value不是遺漏，而是方法阻擋。AI問答必須同時讀取value_status與publish_status，不能只取source_value。", fill=PALE_GOLD, label_color="7A5A00")

    doc.add_heading("資料品質發現與風險", level=1)
    quality = [
        (
            "低" if summary.get("ntpcPopPaginationPassed") else "高",
            "NTPC-POP分頁完整性" if summary.get("ntpcPopPaginationPassed") else "NTPC-POP本地快照不完整",
            (
                f"已取得{summary.get('ntpcPopRowCount', 0):,}列，年度{summary.get('ntpcPopYearSpan', 'unknown')}"
                if summary.get("ntpcPopPaginationPassed")
                else f"官方目錄至少2,250列；介接僅{summary.get('ntpcPopRowCount', 0):,}列，年度{summary.get('ntpcPopYearSpan', 'unknown')}"
            ),
            "保留自動列數／重複頁／最大年度檢查；仍不作18–35邊界拆分" if summary.get("ntpcPopPaginationPassed") else "不得用作最新人口；修正完整分頁",
        ),
        ("高", "年齡邊界缺口", "18、19、35歲常落在跨界帶", "取得單一年齡／微資料；否則封鎖18–35正式值"),
        ("高", "母體混用", "戶籍人口、民間人口、就業者、受僱員工不同", "圖表、CSV及問答強制顯示population_scope"),
        ("高", "地理角色混用", "居住地與工作場所薪資不可合併", "分卡、分色並在標題顯示地理角色"),
        ("中", "統計期不一致", "人口2026、就業2025、薪資與新北開放資料2024", "採最近可得期並列，禁止無註記橫向差值"),
        ("中", "抽樣及四捨五入", "人力資源表以千人與抽樣估計發布", "保留單位、來源及抽樣註記；取得誤差後設門檻"),
        ("中", "教育來源年齡不符", "NTPC-EDU15P為15歲以上整體；ODRP053為五歲組", "25–29可用ODRP053；18–35須客製表或核准M3估計"),
    ]
    add_table(doc, ["嚴重度", "發現", "證據", "處置"], quality, [850, 2100, 3000, 3410], font_size=8.3)

    doc.add_heading("人工查核表", level=1)
    severity_zh = {"HIGH": "高", "MEDIUM": "中", "LOW": "低"}
    checklist_rows = [
        (
            item["id"],
            severity_zh.get(item["severity"], item["severity"]),
            item["item"],
            item["decision"],
            item["recommended"],
        )
        for item in payload["humanReviewChecklist"]
    ]
    add_table(doc, ["代號", "等級", "查核項目", "需要裁示", "建議"], checklist_rows,
              [1000, 780, 1750, 3200, 2630], font_size=7.7,
              header_fill=PALE_GOLD, keep_rows_together=True)
    add_callout(doc, "建議裁示順序", "先核定HR-AGE-02雙層政策，再處理HR-AGE-01資料完整性與HR-AGE-04單一年齡資料；其餘顯示與發布規則接續辦理。")

    doc.add_heading("AWS數據處理與AI問答落地", level=1)
    add_para(doc, "未來在AWS上應把『來源取得、定義比較、年齡轉換、人工核准、AI問答』拆成獨立且可稽核的步驟。Bedrock Agent不得直接從raw層挑數值回答。")
    aws_steps = [
        ("S3 raw", "依source_alias/run_id保存原檔；啟用版本、雜湊與不可變保存政策。"),
        ("Glue Data Catalog", "登錄來源原生schema、年齡欄位、母體、地理角色、單位與期別。"),
        ("Lambda／Glue ETL", "執行M0–M2確定性轉換；M3模型另建研究工作並輸出信賴區間。"),
        ("Step Functions", "串接取得、驗證、轉換、品質檢查與人工核准Gate；任一衝突即fail closed。"),
        ("S3 curated／Athena", "發布本CSV所示共同schema；以publish_status控制可查詢資料列。"),
        ("Bedrock AgentCore", "只讀核准curated view；回答必須回傳來源、期別、母體、地理角色、方法與證據等級。"),
        ("CloudWatch／CloudTrail", "保存每次管線執行、查詢、人工裁示、模型版本與輸出雜湊。"),
    ]
    add_table(doc, ["AWS元件", "職責與控制"], aws_steps, [2050, 7310], font_size=8.8)

    doc.add_heading("建議新增Gate與驗收條件", level=1)
    gates = [
        ("G4-AGE-A", "來源年齡盤點", "所有資料皆有source_age_band、population_scope、geography_role與period。"),
        ("G4-AGE-B", "25–29基準驗證", "人口、就業、失業、教育、薪資資料均可讀；值與來源表抽查一致。"),
        ("G4-AGE-C", "18–35精確重組", "計數使用單一年齡；率有分子分母；平均數有同母體權重；中位數由分布重算。"),
        ("G4-AGE-D", "模型估計", "保留方法版本、訓練資料、回測、不確定區間與人工核准；不得覆寫官方值。"),
        ("G4-AGE-E", "發布測試", "BLOCKED/PENDING列不進正式API；問答必須揭露來源與限制。"),
    ]
    add_table(doc, ["Gate", "目的", "驗收條件"], gates, [1400, 1900, 6060], font_size=8.8)

    doc.add_heading("限制與後續工作", level=1)
    for text in (
        "本報告的新增主計總處表為2025年官方下載快照；人力資源調查值屬抽樣估計且以千人四捨五入。",
        "目前未取得足以精確切出18、19、35歲的勞動力與薪資微資料，因此CSV刻意保留空值並標示BLOCKED/PENDING。",
        "若後續採M3校準拆分，應至少進行3至5年比率穩定性、重聚合回原值、替代模型敏感度及抽樣誤差傳遞。",
        "25–29驗證通過只代表管線與定義治理可運作，不代表25–29可代表全部18–35青年。",
    ):
        add_bullet(doc, text)

    doc.add_heading("參考資料", level=1)
    references = [
        "政府資料開放平臺｜村里戶數、單一年齡人口：https://data.gov.tw/dataset/77132",
        "新北市資料開放平臺｜現住人口之年齡分配：https://data.ntpc.gov.tw/datasets/8308ab58-62d1-424e-8314-24b65b7ab492",
        "新北市資料開放平臺｜失業率－年齡別：https://data.ntpc.gov.tw/datasets/c29c80d4-bef1-452c-8d9a-659e72f07831",
        "新北市資料開放平臺｜就業者年齡結構：https://data.ntpc.gov.tw/datasets/c285509a-7fb2-434f-8542-0b4986c337a8",
        "行政院主計總處｜薪資中位數及分布統計：https://www.stat.gov.tw/News_Content.aspx?n=4580&s=232642",
        "行政院主計總處｜114年人力資源調查統計（表27、32、37、50）：https://www.stat.gov.tw/News_Content.aspx?n=4001&s=236078",
        "政府資料開放平臺｜15歲以上現住人口按性別、年齡、婚姻狀況及教育程度分：https://data.gov.tw/dataset/117988",
        "政府資料開放平臺｜就業者之教育程度與年齡：https://data.gov.tw/dataset/34112",
        "新北市資料開放平臺｜開發指引（page、size分頁）：https://data.ntpc.gov.tw/applications",
        "行政院主計總處｜人力資源調查編製方法（分層二階段抽樣、比率估計與校正）：https://www.stat.gov.tw/public/Attachment/6325104830T64V6LTY.pdf",
        "Deville, J.-C. & Särndal, C.-E. (1992). Calibration Estimators in Survey Sampling. Journal of the American Statistical Association.",
        "Statistics Canada｜Small area estimation methodology：https://www150.statcan.gc.ca/n1/pub/12-001-x/2020001/article/00001/04-eng.htm",
    ]
    for ref in references:
        add_bullet(doc, ref)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(json.dumps({"output": str(args.output), "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
