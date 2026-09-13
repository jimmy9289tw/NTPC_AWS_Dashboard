from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    args = parser.parse_args()

    for year_dir in sorted(path for path in args.root.iterdir() if path.is_dir()):
        print(f"\n=== ROC {year_dir.name} ===")
        for table in (27, 28, 32, 36, 37):
            path = year_dir / f"table{table}.xlsx"
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            print(
                f"table{table}: sheet={sheet.title!r} rows={sheet.max_row} cols={sheet.max_column}"
            )
            for row_number in (2, 5, 9, 10, 11, 12, 15):
                columns = [2, *range(13, 27)]
                values = [sheet.cell(row_number, column).value for column in columns]
                if any(value not in (None, "") for value in values):
                    compact = [
                        f"c{column}={value!r}"
                        for column, value in zip(columns, values)
                        if value not in (None, "")
                    ]
                    print(f"  r{row_number}: " + " | ".join(compact))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
