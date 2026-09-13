from __future__ import annotations

import argparse
import csv
import io
import json
import math
import os
from contextlib import redirect_stdout
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from process_marital_status_v23 import select_pclm


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PACKAGE_ROOT = (
    PROJECT_ROOT
    / "deliverables"
    / "NTPC_Youth_18-35_Government_Open_Data_Package_V2.3_20260825"
)
DEFAULT_TABLE_41 = PACKAGE_ROOT / "02_失業率" / "raw" / "DGBAS-NTPC-UR-AGE-T41-11412.xlsx"
DEFAULT_TABLE_42 = PACKAGE_ROOT / "02_失業率" / "raw" / "DGBAS-NTPC-LABOR-AGE-T42-114H2.xlsx"
DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "deliverables"
    / "NTPC_Labor_18-35_Sprague_PCLM_Pilot_V1.0_20260825"
)

PERIOD = "114年下半年平均（2025-07至2025-12）"
GEOGRAPHY = "新北市"
UNIT = "千人"

# Table 42, total-sex rows. The official workbook is bilingual but its Chinese
# strings are mojibake in the downloaded file; row positions and English table
# structure are stable in the archived source.
FIVE_YEAR_ROWS = {
    15: 19,
    20: 20,
    25: 22,
    30: 23,
    35: 24,
    40: 25,
    45: 27,
    50: 28,
    55: 29,
    60: 30,
    65: 31,  # 65+
}

# Table 41, 114 H2 row 48. The first published age rate is 15-24; the
# remaining age rates align with the five-year bands through 65+.
RATE_COLUMNS = {
    "15-24": 9,
    "25-29": 10,
    "30-34": 11,
    "35-39": 12,
    "40-44": 13,
    "45-49": 14,
    "50-54": 15,
    "55-59": 16,
    "60-64": 17,
    "65+": 18,
}

TARGET_GROUPS = {
    "18-24": (18, 24),
    "25-29": (25, 29),
    "30-35": (30, 35),
    "18-35": (18, 35),
}


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        raise ValueError(f"No rows to write: {path}")
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def extract_source(table_41: Path, table_42: Path):
    wb42 = openpyxl.load_workbook(table_42, read_only=True, data_only=True)
    ws42 = wb42["42"]

    raw = {}
    for age_start, row_number in FIVE_YEAR_ROWS.items():
        raw[age_start] = {
            "age_group": "65+" if age_start == 65 else f"{age_start}-{age_start + 4}",
            "P_reported": float(ws42.cell(row_number, 2).value),
            "LF_reported": float(ws42.cell(row_number, 3).value),
            "E_reported": float(ws42.cell(row_number, 4).value),
            "U_reported": float(ws42.cell(row_number, 5).value),
            "NLF_reported": float(ws42.cell(row_number, 6).value),
            "table42_row": row_number,
        }

    aggregate_15_24 = {
        "age_group": "15-24",
        "P_reported": float(ws42.cell(18, 2).value),
        "LF_reported": float(ws42.cell(18, 3).value),
        "E_reported": float(ws42.cell(18, 4).value),
        "U_reported": float(ws42.cell(18, 5).value),
        "NLF_reported": float(ws42.cell(18, 6).value),
        "table42_row": 18,
    }

    wb41 = openpyxl.load_workbook(table_41, read_only=True, data_only=True)
    ws41 = wb41["41"]
    rates = {label: float(ws41.cell(48, column).value) for label, column in RATE_COLUMNS.items()}
    return raw, aggregate_15_24, rates


def rounding_interval(reported: float, half_width: float = 0.5) -> tuple[float, float]:
    return reported - half_width, reported + half_width


def reconcile_labor_components(
    lf_reported: float,
    e_reported: float,
    u_reported: float,
    rate_pct_reported: float,
) -> tuple[float, float, float, dict]:
    """Find a latent LF/E/U point inside all published rounding intervals."""
    lf_low, lf_high = rounding_interval(lf_reported)
    e_low, e_high = rounding_interval(e_reported)
    u_low_reported, u_high_reported = rounding_interval(u_reported)
    rate_low = max(rate_pct_reported - 0.05, 0.0)
    rate_high = rate_pct_reported + 0.05

    best = None
    for labor_force in np.linspace(lf_low, lf_high, 4001):
        u_low = max(
            u_low_reported,
            labor_force - e_high,
            labor_force * rate_low / 100.0,
        )
        u_high = min(
            u_high_reported,
            labor_force - e_low,
            labor_force * rate_high / 100.0,
        )
        if u_low > u_high + 1e-12:
            continue
        unemployment = float(np.clip(labor_force * rate_pct_reported / 100.0, u_low, u_high))
        employment = labor_force - unemployment
        implied_rate = 100.0 * unemployment / labor_force
        objective = (
            (labor_force - lf_reported) ** 2
            + (employment - e_reported) ** 2
            + (unemployment - u_reported) ** 2
            + ((implied_rate - rate_pct_reported) / 0.1) ** 2
        )
        if best is None or objective < best[0]:
            best = (objective, labor_force, employment, unemployment, implied_rate)

    if best is None:
        raise ValueError(
            "No LF/E/U latent point satisfies the joint count and rate rounding intervals: "
            f"LF={lf_reported}, E={e_reported}, U={u_reported}, rate={rate_pct_reported}"
        )

    _, labor_force, employment, unemployment, implied_rate = best
    return labor_force, employment, unemployment, {
        "joint_rounding_feasible": True,
        "implied_rate_pct": implied_rate,
        "rate_error_pp": implied_rate - rate_pct_reported,
        "lf_distance_from_reported_thousand": labor_force - lf_reported,
        "e_distance_from_reported_thousand": employment - e_reported,
        "u_distance_from_reported_thousand": unemployment - u_reported,
    }


def reconcile_population(
    labor_force: float,
    p_reported: float,
    nlf_reported: float,
) -> tuple[float, float, dict]:
    p_low, p_high = rounding_interval(p_reported)
    nlf_low, nlf_high = rounding_interval(nlf_reported)
    feasible_low = max(p_low, labor_force + nlf_low)
    feasible_high = min(p_high, labor_force + nlf_high)
    if feasible_low > feasible_high + 1e-12:
        raise ValueError(
            "No P/NLF latent point satisfies rounding intervals: "
            f"P={p_reported}, LF={labor_force}, NLF={nlf_reported}"
        )
    civilian_population = float(np.clip(p_reported, feasible_low, feasible_high))
    not_in_labor_force = civilian_population - labor_force
    return civilian_population, not_in_labor_force, {
        "joint_rounding_feasible": True,
        "p_distance_from_reported_thousand": civilian_population - p_reported,
        "nlf_distance_from_reported_thousand": not_in_labor_force - nlf_reported,
    }


def reconcile_15_24_subgroups(raw: dict, aggregate: dict, official_rate: float):
    total_lf, total_e, total_u, labor_diag = reconcile_labor_components(
        aggregate["LF_reported"],
        aggregate["E_reported"],
        aggregate["U_reported"],
        official_rate,
    )
    total_p, total_nlf, pop_diag = reconcile_population(
        total_lf,
        aggregate["P_reported"],
        aggregate["NLF_reported"],
    )

    first, second = raw[15], raw[20]
    x1_low, x1_high = rounding_interval(first["LF_reported"])
    x2_low, x2_high = rounding_interval(second["LF_reported"])
    x1_low = max(x1_low, total_lf - x2_high)
    x1_high = min(x1_high, total_lf - x2_low)

    best = None
    for lf1 in np.linspace(x1_low, x1_high, 4001):
        lf2 = total_lf - lf1
        u1_low = max(
            first["U_reported"] - 0.5,
            lf1 - (first["E_reported"] + 0.5),
        )
        u1_high = min(
            first["U_reported"] + 0.5,
            lf1 - (first["E_reported"] - 0.5),
        )
        u2_low = max(
            second["U_reported"] - 0.5,
            lf2 - (second["E_reported"] + 0.5),
        )
        u2_high = min(
            second["U_reported"] + 0.5,
            lf2 - (second["E_reported"] - 0.5),
        )
        u1_low = max(u1_low, total_u - u2_high)
        u1_high = min(u1_high, total_u - u2_low)
        if u1_low > u1_high + 1e-12:
            continue
        target_share = first["U_reported"] / (
            first["U_reported"] + second["U_reported"]
        )
        u1 = float(np.clip(total_u * target_share, u1_low, u1_high))
        u2 = total_u - u1
        e1, e2 = lf1 - u1, lf2 - u2
        objective = sum(
            value**2
            for value in (
                lf1 - first["LF_reported"],
                lf2 - second["LF_reported"],
                e1 - first["E_reported"],
                e2 - second["E_reported"],
                u1 - first["U_reported"],
                u2 - second["U_reported"],
            )
        )
        if best is None or objective < best[0]:
            best = (objective, lf1, lf2, e1, e2, u1, u2)
    if best is None:
        raise ValueError("No feasible 15-19 / 20-24 allocation inside published rounding intervals")

    _, lf1, lf2, e1, e2, u1, u2 = best

    # Allocate P while preserving the official 15-24 total and both P/NLF
    # rounding intervals.
    p1_low = max(first["P_reported"] - 0.5, lf1 + first["NLF_reported"] - 0.5)
    p1_high = min(first["P_reported"] + 0.5, lf1 + first["NLF_reported"] + 0.5)
    p2_low = max(second["P_reported"] - 0.5, lf2 + second["NLF_reported"] - 0.5)
    p2_high = min(second["P_reported"] + 0.5, lf2 + second["NLF_reported"] + 0.5)
    p1_low = max(p1_low, total_p - p2_high)
    p1_high = min(p1_high, total_p - p2_low)
    if p1_low > p1_high + 1e-12:
        raise ValueError("No feasible 15-19 / 20-24 P allocation inside rounding intervals")
    p1 = float(np.clip(first["P_reported"], p1_low, p1_high))
    p2 = total_p - p1

    reconciled = {
        15: {"P": p1, "LF": lf1, "E": e1, "U": u1, "NLF": p1 - lf1},
        20: {"P": p2, "LF": lf2, "E": e2, "U": u2, "NLF": p2 - lf2},
    }
    diagnostic = {
        "aggregate_labor": labor_diag,
        "aggregate_population": pop_diag,
        "aggregate_P": total_p,
        "aggregate_LF": total_lf,
        "aggregate_E": total_e,
        "aggregate_U": total_u,
        "aggregate_NLF": total_nlf,
    }
    return reconciled, diagnostic


def build_reconciled_bands(raw: dict, aggregate_15_24: dict, rates: dict):
    reconciled, diag_15_24 = reconcile_15_24_subgroups(
        raw, aggregate_15_24, rates["15-24"]
    )
    diagnostics = {"15-24": diag_15_24}
    for age_start in range(25, 70, 5):
        source = raw[age_start]
        label = source["age_group"]
        labor_force, employment, unemployment, labor_diag = reconcile_labor_components(
            source["LF_reported"],
            source["E_reported"],
            source["U_reported"],
            rates[label],
        )
        civilian_population, not_in_labor_force, pop_diag = reconcile_population(
            labor_force,
            source["P_reported"],
            source["NLF_reported"],
        )
        reconciled[age_start] = {
            "P": civilian_population,
            "LF": labor_force,
            "E": employment,
            "U": unemployment,
            "NLF": not_in_labor_force,
        }
        diagnostics[label] = {"labor": labor_diag, "population": pop_diag}
    return reconciled, diagnostics


def sprague_matrix(group_count: int) -> np.ndarray:
    """DemoTools-compatible Sprague matrix with an open final age group."""
    if group_count < 6:
        raise ValueError("Sprague requires at least six grouped ages including the open group")
    g1g2 = np.array(
        [
            [0.3616, -0.2768, 0.1488, -0.0336, 0.0000],
            [0.2640, -0.0960, 0.0400, -0.0080, 0.0000],
            [0.1840, 0.0400, -0.0320, 0.0080, 0.0000],
            [0.1200, 0.1360, -0.0720, 0.0160, 0.0000],
            [0.0704, 0.1968, -0.0848, 0.0176, 0.0000],
            [0.0336, 0.2272, -0.0752, 0.0144, 0.0000],
            [0.0080, 0.2320, -0.0480, 0.0080, 0.0000],
            [-0.0080, 0.2160, -0.0080, 0.0000, 0.0000],
            [-0.0160, 0.1840, 0.0400, -0.0080, 0.0000],
            [-0.0176, 0.1408, 0.0912, -0.0144, 0.0000],
        ],
        dtype=float,
    )
    g3 = np.array(
        [
            [-0.0128, 0.0848, 0.1504, -0.0240, 0.0016],
            [-0.0016, 0.0144, 0.2224, -0.0416, 0.0064],
            [0.0064, -0.0336, 0.2544, -0.0336, 0.0064],
            [0.0064, -0.0416, 0.2224, 0.0144, -0.0016],
            [0.0016, -0.0240, 0.1504, 0.0848, -0.0128],
        ],
        dtype=float,
    )
    g4g5 = np.array(
        [
            [0.0000, -0.0144, 0.0912, 0.1408, -0.0176],
            [0.0000, -0.0080, 0.0400, 0.1840, -0.0160],
            [0.0000, 0.0000, -0.0080, 0.2160, -0.0080],
            [0.0000, 0.0080, -0.0480, 0.2320, 0.0080],
            [0.0000, 0.0144, -0.0752, 0.2272, 0.0336],
            [0.0000, 0.0176, -0.0848, 0.1968, 0.0704],
            [0.0000, 0.0160, -0.0720, 0.1360, 0.1200],
            [0.0000, 0.0080, -0.0320, 0.0400, 0.1840],
            [0.0000, -0.0080, 0.0400, -0.0960, 0.2640],
            [0.0000, -0.0336, 0.1488, -0.2768, 0.3616],
        ],
        dtype=float,
    )

    row_count = group_count * 5 - 4
    matrix = np.zeros((row_count, group_count), dtype=float)
    matrix[0:10, 0:5] = g1g2
    middle_panels = group_count - 5
    for panel in range(middle_panels):
        row_start = 10 + panel * 5
        matrix[row_start : row_start + 5, panel : panel + 5] = g3
    matrix[row_count - 11 : row_count - 1, group_count - 5 : group_count] = g4g5
    matrix[-1, -1] = 1.0
    return matrix


def calibrate_closed_bands(
    ages: np.ndarray,
    seed: np.ndarray,
    targets: dict[int, float],
) -> tuple[np.ndarray, dict]:
    calibrated = seed.astype(float).copy()
    negative_before = int(np.sum(calibrated < 0))
    min_before = float(np.min(calibrated))
    calibrated = np.maximum(calibrated, 0.0)
    max_error = 0.0
    for start in range(15, 65, 5):
        mask = (ages >= start) & (ages <= start + 4)
        current = float(calibrated[mask].sum())
        target = float(targets[start])
        if current <= 0:
            calibrated[mask] = target / int(mask.sum())
        else:
            calibrated[mask] *= target / current
        max_error = max(max_error, abs(float(calibrated[mask].sum()) - target))
    return calibrated, {
        "negative_single_ages_before_calibration": negative_before,
        "minimum_before_calibration": min_before,
        "maximum_post_calibration_reaggregation_error_thousand": max_error,
    }


def graduate_sprague(reconciled: dict[int, dict]):
    group_starts = list(range(15, 70, 5))
    matrix = sprague_matrix(len(group_starts))
    all_ages = np.arange(15, 66)
    closed_mask = all_ages <= 64
    ages = all_ages[closed_mask]
    components = {}
    diagnostics = {}
    for component in ("E", "U", "NLF"):
        grouped = np.array([reconciled[start][component] for start in group_starts])
        raw_single = matrix @ grouped
        calibrated, diag = calibrate_closed_bands(
            ages,
            raw_single[closed_mask],
            {start: reconciled[start][component] for start in range(15, 65, 5)},
        )
        components[component] = calibrated
        diagnostics[component] = diag
    return ages, components, diagnostics


def graduate_pclm(reconciled: dict[int, dict]):
    ages = np.arange(15, 65)
    groups = [(start, start + 4) for start in range(15, 65, 5)]
    components = {}
    diagnostics = {}
    for component in ("E", "U", "NLF"):
        grouped = np.array([reconciled[start][component] for start in range(15, 65, 5)])
        seed, diag = select_pclm(grouped, ages, groups)
        calibrated, calibration_diag = calibrate_closed_bands(
            ages,
            seed,
            {start: reconciled[start][component] for start in range(15, 65, 5)},
        )
        components[component] = calibrated
        diagnostics[component] = {**diag, **calibration_diag}
    return ages, components, diagnostics


def component_frame(method: str, ages: np.ndarray, components: dict[str, np.ndarray]):
    frame = pd.DataFrame(
        {
            "method": method,
            "age": ages.astype(int),
            "employed_thousand": components["E"],
            "unemployed_thousand": components["U"],
            "not_in_labor_force_thousand": components["NLF"],
        }
    )
    frame["labor_force_thousand"] = frame["employed_thousand"] + frame["unemployed_thousand"]
    frame["civilian_population_thousand"] = frame["labor_force_thousand"] + frame["not_in_labor_force_thousand"]
    frame["unemployment_rate_pct"] = 100 * frame["unemployed_thousand"] / frame["labor_force_thousand"]
    frame["labor_force_participation_rate_pct"] = 100 * frame["labor_force_thousand"] / frame["civilian_population_thousand"]
    frame["employment_to_population_rate_pct"] = 100 * frame["employed_thousand"] / frame["civilian_population_thousand"]
    frame["employment_share_of_labor_force_pct"] = 100 * frame["employed_thousand"] / frame["labor_force_thousand"]
    return frame


def summarize_targets(frame: pd.DataFrame):
    rows = []
    method = str(frame["method"].iloc[0])
    for label, (start, end) in TARGET_GROUPS.items():
        subset = frame[(frame["age"] >= start) & (frame["age"] <= end)]
        e = float(subset["employed_thousand"].sum())
        u = float(subset["unemployed_thousand"].sum())
        nlf = float(subset["not_in_labor_force_thousand"].sum())
        lf = e + u
        p = lf + nlf
        exact_band = label == "25-29"
        rows.append(
            {
                "period": PERIOD,
                "geography": GEOGRAPHY,
                "age_group": label,
                "method": method,
                "value_origin_class": (
                    "OFFICIAL_SURVEY_ESTIMATE_AGE_BAND_EXACT"
                    if exact_band
                    else "MODEL_ESTIMATE_FROM_OFFICIAL_SURVEY_ESTIMATES"
                ),
                "age_boundary_status": (
                    "OFFICIAL_5Y_BAND_NO_BOUNDARY_ESTIMATION"
                    if exact_band
                    else "ESTIMATED_BOUNDARY_AGES"
                ),
                "civilian_population_thousand": p,
                "labor_force_thousand": lf,
                "employed_thousand": e,
                "unemployed_thousand": u,
                "not_in_labor_force_thousand": nlf,
                "unemployment_rate_pct": 100 * u / lf,
                "labor_force_participation_rate_pct": 100 * lf / p,
                "employment_to_population_rate_pct": 100 * e / p,
                "employment_share_of_labor_force_pct": 100 * e / lf,
                "unit_for_counts": UNIT,
                "source_time_grain": "HALF_YEAR_AVERAGE",
                "official_or_model_note": (
                    "25-29為官方發布年齡帶；其一致化小數仍屬官方千人四捨五入調和值。"
                    if exact_band
                    else "含年齡邊界拆分；不得標為官方發布精確值。"
                ),
            }
        )
    return rows


def build_source_reconciliation_rows(raw: dict, reconciled: dict, rates: dict):
    rows = []
    for start in range(15, 70, 5):
        source = raw[start]
        final = reconciled[start]
        label = source["age_group"]
        official_rate = rates.get(label)
        rows.append(
            {
                "period": PERIOD,
                "age_group": label,
                "table42_row": source["table42_row"],
                "P_reported_thousand": source["P_reported"],
                "LF_reported_thousand": source["LF_reported"],
                "E_reported_thousand": source["E_reported"],
                "U_reported_thousand": source["U_reported"],
                "NLF_reported_thousand": source["NLF_reported"],
                "official_unemployment_rate_pct": official_rate if official_rate is not None else "",
                "reconciled_P_thousand": final["P"],
                "reconciled_LF_thousand": final["LF"],
                "reconciled_E_thousand": final["E"],
                "reconciled_U_thousand": final["U"],
                "reconciled_NLF_thousand": final["NLF"],
                "identity_LF_minus_E_minus_U": final["LF"] - final["E"] - final["U"],
                "identity_P_minus_LF_minus_NLF": final["P"] - final["LF"] - final["NLF"],
                "rate_check_note": (
                    "No official subgroup rate; reconciled under official 15-24 rate"
                    if start in (15, 20)
                    else "Within joint published rounding intervals"
                ),
            }
        )
    return rows


def build_sensitivity_rows(pclm_rows: list[dict], sprague_rows: list[dict]):
    pclm = {row["age_group"]: row for row in pclm_rows}
    sprague = {row["age_group"]: row for row in sprague_rows}
    metrics = [
        "civilian_population_thousand",
        "labor_force_thousand",
        "employed_thousand",
        "unemployed_thousand",
        "unemployment_rate_pct",
        "labor_force_participation_rate_pct",
        "employment_to_population_rate_pct",
    ]
    rows = []
    for age_group in TARGET_GROUPS:
        for metric in metrics:
            p_value = float(pclm[age_group][metric])
            s_value = float(sprague[age_group][metric])
            is_rate = metric.endswith("_pct")
            rows.append(
                {
                    "age_group": age_group,
                    "metric": metric,
                    "PCLM_value": p_value,
                    "Sprague_value": s_value,
                    "Sprague_minus_PCLM": s_value - p_value,
                    "absolute_difference": abs(s_value - p_value),
                    "relative_difference_pct_of_PCLM": (
                        abs(s_value - p_value) / abs(p_value) * 100 if p_value else ""
                    ),
                    "difference_unit": "percentage_point" if is_rate else UNIT,
                }
            )
    return rows


def build_validation_rows(
    frames: dict[str, pd.DataFrame],
    reconciled: dict[int, dict],
    diagnostics: dict[str, dict],
    target_rows: dict[str, list[dict]],
):
    checks = []
    for method, frame in frames.items():
        identity_lf = float(
            np.max(
                np.abs(
                    frame["labor_force_thousand"]
                    - frame["employed_thousand"]
                    - frame["unemployed_thousand"]
                )
            )
        )
        identity_p = float(
            np.max(
                np.abs(
                    frame["civilian_population_thousand"]
                    - frame["labor_force_thousand"]
                    - frame["not_in_labor_force_thousand"]
                )
            )
        )
        minimum_component = float(
            frame[
                [
                    "employed_thousand",
                    "unemployed_thousand",
                    "not_in_labor_force_thousand",
                ]
            ].min().min()
        )
        reaggregation_error = 0.0
        for start in range(15, 65, 5):
            subset = frame[(frame["age"] >= start) & (frame["age"] <= start + 4)]
            for source_component, column in (
                ("E", "employed_thousand"),
                ("U", "unemployed_thousand"),
                ("NLF", "not_in_labor_force_thousand"),
            ):
                reaggregation_error = max(
                    reaggregation_error,
                    abs(float(subset[column].sum()) - reconciled[start][source_component]),
                )
        exact_25 = next(row for row in target_rows[method] if row["age_group"] == "25-29")
        source_25 = reconciled[25]
        exact_25_error = max(
            abs(float(exact_25["employed_thousand"]) - source_25["E"]),
            abs(float(exact_25["unemployed_thousand"]) - source_25["U"]),
            abs(float(exact_25["not_in_labor_force_thousand"]) - source_25["NLF"]),
        )
        for check_id, check, observed, threshold in (
            ("VAL-NONNEG", "All E/U/NLF single-age values are non-negative", minimum_component, ">= 0"),
            ("VAL-LF-ID", "LF = E + U at every single age", identity_lf, "<= 1e-9 thousand"),
            ("VAL-P-ID", "P = LF + NLF at every single age", identity_p, "<= 1e-9 thousand"),
            ("VAL-REAGG", "Single ages reaggregate to reconciled five-year margins", reaggregation_error, "<= 1e-8 thousand"),
            ("VAL-25-29", "Official 25-29 age band is exactly recovered", exact_25_error, "<= 1e-8 thousand"),
        ):
            if check_id == "VAL-NONNEG":
                passed = observed >= -1e-12
            elif check_id in ("VAL-LF-ID", "VAL-P-ID"):
                passed = observed <= 1e-9
            else:
                passed = observed <= 1e-8
            checks.append(
                {
                    "method": method,
                    "check_id": check_id,
                    "check": check,
                    "observed": observed,
                    "threshold": threshold,
                    "status": "PASS" if passed else "FAIL",
                }
            )

        for component, diag in diagnostics[method].items():
            if method == "PCLM":
                checks.append(
                    {
                        "method": method,
                        "check_id": f"VAL-PCLM-CONV-{component}",
                        "check": f"PCLM converged for {component}",
                        "observed": bool(diag.get("converged")),
                        "threshold": "True",
                        "status": "PASS" if diag.get("converged") else "FAIL",
                    }
                )
            checks.append(
                {
                    "method": method,
                    "check_id": f"VAL-RAW-NEG-{component}",
                    "check": f"Raw {method} negative single-age count for {component}",
                    "observed": diag["negative_single_ages_before_calibration"],
                    "threshold": "diagnostic; zero preferred",
                    "status": (
                        "PASS"
                        if diag["negative_single_ages_before_calibration"] == 0
                        else "WARN"
                    ),
                }
            )
    return checks


def flatten_diagnostics(diagnostics: dict[str, dict]):
    rows = []
    for method, components in diagnostics.items():
        for component, values in components.items():
            row = {"method": method, "component": component}
            for key, value in values.items():
                if isinstance(value, (np.floating, np.integer)):
                    value = value.item()
                row[key] = value
            rows.append(row)
    all_fields = []
    for row in rows:
        for field in row:
            if field not in all_fields:
                all_fields.append(field)
    return [{field: row.get(field, "") for field in all_fields} for row in rows]


def build_notebook(output_dir: Path) -> Path:
    notebook_path = output_dir / "labor_age_graduation_audit.ipynb"
    cells = [
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": [
                "# 新北市就業／失業 18–35 歲 Sprague 與 PCLM 試算\n",
                "\n",
                "此筆記讀取同資料夾已產製的可稽核 CSV。所有人數單位為千人；期間為 114 年下半年平均。\n",
            ],
        },
        {
            "cell_type": "code",
            "execution_count": None,
            "metadata": {},
            "outputs": [],
            "source": [
                "from pathlib import Path\n",
                "import pandas as pd\n",
                "BASE = Path.cwd()\n",
                "results = pd.read_csv(BASE / 'target_age_results.csv', encoding='utf-8-sig')\n",
                "validation = pd.read_csv(BASE / 'validation_checks.csv', encoding='utf-8-sig')\n",
                "sensitivity = pd.read_csv(BASE / 'method_sensitivity.csv', encoding='utf-8-sig')\n",
                "print('result rows=', len(results), 'validation rows=', len(validation))\n",
            ],
        },
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": ["## Results\n"],
        },
        {
            "cell_type": "code",
            "execution_count": None,
            "metadata": {},
            "outputs": [],
            "source": [
                "cols = ['age_group','method','civilian_population_thousand','labor_force_thousand','employed_thousand','unemployed_thousand','unemployment_rate_pct','labor_force_participation_rate_pct','employment_to_population_rate_pct']\n",
                "print(results[cols].round(4).to_string(index=False))\n",
            ],
        },
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": ["## Validation\n"],
        },
        {
            "cell_type": "code",
            "execution_count": None,
            "metadata": {},
            "outputs": [],
            "source": [
                "print(validation[['method','check_id','status','observed','threshold']].to_string(index=False))\n",
                "assert not (validation['status'] == 'FAIL').any()\n",
            ],
        },
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": ["## Method sensitivity\n"],
        },
        {
            "cell_type": "code",
            "execution_count": None,
            "metadata": {},
            "outputs": [],
            "source": [
                "focus = sensitivity[sensitivity['metric'].isin(['unemployment_rate_pct','employment_to_population_rate_pct'])]\n",
                "print(focus[['age_group','metric','PCLM_value','Sprague_value','Sprague_minus_PCLM']].round(6).to_string(index=False))\n",
            ],
        },
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": [
                "## Takeaways\n",
                "\n",
                "- 25–29 歲為官方原生五歲年齡帶，不需估計邊界。\n",
                "- 18–24 與 30–35 含邊界年齡拆分，應標示為模型估計。\n",
                "- PCLM 作為主估計，Sprague 作為敏感度基準；兩者差異不是抽樣信賴區間。\n",
            ],
        },
    ]

    namespace = {"__name__": "__notebook__"}
    execution_count = 0
    previous_cwd = Path.cwd()
    try:
        os.chdir(output_dir)
        for cell in cells:
            if cell["cell_type"] != "code":
                continue
            execution_count += 1
            stream = io.StringIO()
            with redirect_stdout(stream):
                exec("".join(cell["source"]), namespace)
            cell["execution_count"] = execution_count
            cell["outputs"] = [
                {
                    "name": "stdout",
                    "output_type": "stream",
                    "text": stream.getvalue().splitlines(keepends=True),
                }
            ]
    finally:
        os.chdir(previous_cwd)

    notebook = {
        "cells": cells,
        "metadata": {
            "kernelspec": {
                "display_name": "Python 3",
                "language": "python",
                "name": "python3",
            },
            "language_info": {"name": "python", "version": "3.12"},
            "execution": {
                "engine": "top-to-bottom Python exec using the project runtime",
                "status": "executed",
            },
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    notebook_path.write_text(json.dumps(notebook, ensure_ascii=False, indent=2), encoding="utf-8")

    # Raw-JSON validation because nbformat is not installed in the bundled runtime.
    reloaded = json.loads(notebook_path.read_text(encoding="utf-8"))
    assert reloaded["nbformat"] == 4
    assert all(
        cell["execution_count"] is not None
        for cell in reloaded["cells"]
        if cell["cell_type"] == "code"
    )
    return notebook_path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--table-41", type=Path, default=DEFAULT_TABLE_41)
    parser.add_argument("--table-42", type=Path, default=DEFAULT_TABLE_42)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    raw, aggregate_15_24, rates = extract_source(args.table_41, args.table_42)
    reconciled, reconciliation_diagnostics = build_reconciled_bands(
        raw, aggregate_15_24, rates
    )

    sprague_ages, sprague_components, sprague_diag = graduate_sprague(reconciled)
    pclm_ages, pclm_components, pclm_diag = graduate_pclm(reconciled)
    frames = {
        "PCLM": component_frame("PCLM", pclm_ages, pclm_components),
        "SPRAGUE_NN_CALIBRATED": component_frame(
            "SPRAGUE_NN_CALIBRATED", sprague_ages, sprague_components
        ),
    }
    diagnostics = {
        "PCLM": pclm_diag,
        "SPRAGUE_NN_CALIBRATED": sprague_diag,
    }
    target_rows = {method: summarize_targets(frame) for method, frame in frames.items()}
    combined_targets = target_rows["PCLM"] + target_rows["SPRAGUE_NN_CALIBRATED"]
    sensitivity_rows = build_sensitivity_rows(
        target_rows["PCLM"], target_rows["SPRAGUE_NN_CALIBRATED"]
    )
    validation_rows = build_validation_rows(
        frames, reconciled, diagnostics, target_rows
    )
    source_rows = build_source_reconciliation_rows(raw, reconciled, rates)

    single_age = pd.concat(frames.values(), ignore_index=True)
    single_age.insert(0, "period", PERIOD)
    single_age.insert(1, "geography", GEOGRAPHY)
    single_age.to_csv(
        args.output_dir / "single_age_estimates.csv",
        index=False,
        encoding="utf-8-sig",
    )
    write_csv(args.output_dir / "target_age_results.csv", combined_targets)
    write_csv(args.output_dir / "method_sensitivity.csv", sensitivity_rows)
    write_csv(args.output_dir / "validation_checks.csv", validation_rows)
    write_csv(args.output_dir / "source_rounding_reconciliation.csv", source_rows)
    write_csv(args.output_dir / "model_diagnostics.csv", flatten_diagnostics(diagnostics))

    source_metadata = {
        "period": PERIOD,
        "geography": GEOGRAPHY,
        "unit": UNIT,
        "table41": str(args.table_41),
        "table42": str(args.table_42),
        "table41_url": "https://ws.dgbas.gov.tw/001/Upload/463/relfile/11040/235759/table41.xlsx",
        "table42_url": "https://ws.dgbas.gov.tw/001/Upload/463/relfile/11040/235759/table42.xlsx",
        "source_class": "OFFICIAL_SURVEY_ESTIMATE",
        "rounding_reconciliation": reconciliation_diagnostics,
        "method_references": {
            "Sprague": "https://timriffe.github.io/DemoTools/articles/graduation_with_demotools.html",
            "PCLM": "Rizzi, Gampe and Eilers (2015), Efficient estimation of smooth distributions from coarsely grouped data",
        },
    }
    (args.output_dir / "source_and_run_metadata.json").write_text(
        json.dumps(source_metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    notebook_path = build_notebook(args.output_dir)
    fail_count = sum(row["status"] == "FAIL" for row in validation_rows)
    summary = {
        "output_dir": str(args.output_dir),
        "target_rows": len(combined_targets),
        "validation_failures": fail_count,
        "validation_warnings": sum(row["status"] == "WARN" for row in validation_rows),
        "notebook": str(notebook_path),
        "pclm_selected_lambda": {
            component: diag["lambda"] for component, diag in pclm_diag.items()
        },
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 1 if fail_count else 0


if __name__ == "__main__":
    raise SystemExit(main())
