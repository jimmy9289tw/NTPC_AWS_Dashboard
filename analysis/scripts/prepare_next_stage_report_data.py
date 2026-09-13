from __future__ import annotations

import csv
import io
import json
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


PROJECT = Path(r"D:\github\work\019fe05e-e1a3-71f2-840e-c026180d9474\NtpcYouthAI")
PACKAGE = PROJECT / "deliverables" / "NTPC_Youth_18-35_Government_Open_Data_Package_V2.3_20260825"
POP_JSON = PACKAGE / "01_年齡人口結構" / "raw" / "MOI-POP1Y-11412.json"
MARITAL_CSV = PACKAGE / "07_婚姻狀態" / "processed" / "114" / "MOI-MARITAL-NTPC-114-TARGET-AGE.csv"
EDU_ZIP = PACKAGE / "04_教育程度" / "raw" / "MOI-EDU5Y-114.zip"
WAGE_XLSX = PACKAGE / "05_薪資平均數" / "raw" / "DGBAS-STAT-WAGE-LOC-current.xlsx"
OUT = PROJECT / "tmp" / "next_stage_report_data.json"

AGE_BANDS = {
    "18-24": range(18, 25),
    "25-29": range(25, 30),
    "30-35": range(30, 36),
    "18-35": range(18, 36),
}

EDU_MAP = {
    "博畢": "研究所",
    "碩畢": "研究所",
    "大畢": "大學",
    "專畢": "專科",
    "高中畢": "高中職",
    "國中畢": "國中及以下",
    "國小畢以下": "國中及以下",
}


def read_population() -> dict:
    payload = json.loads(POP_JSON.read_text(encoding="utf-8"))
    rows = payload["responseData"]
    single = defaultdict(int)
    for row in rows:
        if not str(row.get("site_id", "")).startswith("新北市"):
            continue
        if not str(row.get("district_code", "")).startswith("65"):
            continue
        for age in range(0, 100):
            single[("男", age)] += int(row.get(f"people_age_{age:03d}_m", 0) or 0)
            single[("女", age)] += int(row.get(f"people_age_{age:03d}_f", 0) or 0)
    out = []
    for band, ages in AGE_BANDS.items():
        for sex in ("男", "女", "合計"):
            sexes = ("男", "女") if sex == "合計" else (sex,)
            out.append({
                "period": "11412",
                "age_band": band,
                "sex": sex,
                "population": sum(single[(sx, age)] for sx in sexes for age in ages),
                "value_origin_class": "OFFICIAL_ADMIN_EXACT",
                "method_code": "M0_DIRECT_SINGLE_AGE_SUM",
            })
    return {"single_age": {f"{sex}_{age}": value for (sex, age), value in single.items()}, "target_bands": out}


def read_marriage() -> list[dict]:
    rows = []
    with MARITAL_CSV.open("r", encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            if row["sex"] != "合計" or row["classification_view"] != "DERIVED_BINARY":
                continue
            rows.append({
                "age_band": row["target_age_band"],
                "status": row["marital_status_zh"],
                "population": float(row["population_count"]),
                "share_pct": float(row["share_pct"]),
                "sensitivity_low": float(row["sensitivity_low"]),
                "sensitivity_high": float(row["sensitivity_high"]),
                "value_origin_class": row["value_origin_class"],
                "method_code": row["method_code"],
            })
    return rows


def read_education() -> dict:
    aggregate = defaultdict(int)
    raw_categories = set()
    rows_read = 0
    rows_ntpc = 0
    with zipfile.ZipFile(EDU_ZIP) as zf:
        name = zf.namelist()[0]
        with zf.open(name) as raw:
            text = io.TextIOWrapper(raw, encoding="utf-8-sig", newline="")
            for row in csv.DictReader(text):
                rows_read += 1
                year = (row.get("statistic_yyy") or "").lstrip("\ufeff").strip()
                if year == "統計年" or year != "114":
                    continue
                district = (row.get("district_code") or "").strip()
                site = (row.get("site_id") or "").strip()
                if not district.startswith("65") or not site.startswith("新北市"):
                    continue
                rows_ntpc += 1
                sex = (row.get("sex") or "").strip()
                age = (row.get("age") or "").strip().replace("~", "-").replace("歲", "")
                edu_raw = (row.get("edu") or "").strip()
                raw_categories.add(edu_raw)
                edu = EDU_MAP.get(edu_raw, "其他/未註記")
                try:
                    population = int((row.get("population") or "0").strip())
                except ValueError:
                    continue
                aggregate[(sex, age, edu)] += population
    rows = []
    age_order = ["15-19", "20-24", "25-29", "30-34", "35-39"]
    edu_order = ["研究所", "大學", "專科", "高中職", "國中及以下", "其他/未註記"]
    for age in age_order:
        for sex in ("男", "女", "合計"):
            sexes = ("男", "女") if sex == "合計" else (sex,)
            values = {edu: sum(aggregate[(sx, age, edu)] for sx in sexes) for edu in edu_order}
            total = sum(values.values())
            for edu in edu_order:
                if values[edu] == 0 and edu == "其他/未註記":
                    continue
                rows.append({
                    "source_age_band": age,
                    "sex": sex,
                    "education": edu,
                    "population": values[edu],
                    "share_pct": values[edu] / total * 100 if total else None,
                    "value_origin_class": "OFFICIAL_ADMIN_EXACT",
                    "method_code": "M0_DIRECT_5Y_AGGREGATION",
                })
    return {
        "rows_read": rows_read,
        "rows_ntpc": rows_ntpc,
        "raw_categories": sorted(raw_categories),
        "official_5y": rows,
    }


def read_wage() -> list[dict]:
    wb = load_workbook(WAGE_XLSX, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        if not ws.title.startswith("表6("):
            continue
        year = ws.title.removeprefix("表6(").removesuffix("年)")
        ntpc = None
        for row in ws.iter_rows(values_only=True):
            if str(row[0]).strip() == "新北市":
                ntpc = row
                break
        if ntpc is None:
            continue
        for measure, offset in (("平均數", 0), ("中位數", 8)):
            for age_band, col in (("未滿25", 3 + offset), ("25-29", 4 + offset), ("30-39", 5 + offset)):
                rows.append({
                    "year": year,
                    "geography": "新北市",
                    "geography_role": "WORKPLACE",
                    "universe": "本國籍全時受僱員工",
                    "measure": measure,
                    "source_age_band": age_band,
                    "annual_salary_10k_ntd": float(ntpc[col]),
                    "value_origin_class": "OFFICIAL_SOURCE_NATIVE",
                })
    return sorted(rows, key=lambda r: (r["year"], r["measure"], r["source_age_band"]))


def main() -> None:
    payload = {
        "population": read_population(),
        "marriage": read_marriage(),
        "education": read_education(),
        "wage": read_wage(),
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "output": str(OUT),
        "population_rows": len(payload["population"]["target_bands"]),
        "marriage_rows": len(payload["marriage"]),
        "education_rows_read": payload["education"]["rows_read"],
        "education_rows_ntpc": payload["education"]["rows_ntpc"],
        "wage_rows": len(payload["wage"]),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
