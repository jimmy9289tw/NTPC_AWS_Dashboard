from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
PACKAGE = PROJECT / "deliverables" / "NTPC_Youth_Complete_Data_Package_G5_V3.2_20260826"
OUTPUT = PACKAGE / "06_跨母體儀表板規格" / "每張圖七項資料證據揭露與AWS_Kiro交接規格_V1.0.docx"

FONT = "Calibri"
EAST_ASIA_FONT = "Microsoft JhengHei"
NAVY = "16324F"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GRAY = "5B6573"
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
PALE_GREEN = "E2F0D9"
PALE_AMBER = "FFF2CC"
PALE_RED = "FCE8E6"
WHITE = "FFFFFF"
CONTENT_DXA = 9360


def font_run(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EAST_ASIA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def format_cell(cell, *, bold=False, color=NAVY, size=8.7, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            font_run(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        table.rows[0].cells[index].text = text
        shade(table.rows[0].cells[index], BLUE)
        format_cell(table.rows[0].cells[index], bold=True, color=WHITE, size=font_size)
    set_repeat_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, text in enumerate(values):
            row.cells[index].text = str(text)
            if row_index % 2:
                shade(row.cells[index], LIGHT_GRAY)
            format_cell(row.cells[index], size=font_size)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, label, text, fill):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.1
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    r = p.add_run(f"{label}｜")
    font_run(r, size=10.5, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    font_run(r, size=10.5, color=NAVY)


def create_numbering(doc, *, bullet=False):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    level.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_list(doc, items, num_id):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.1
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        n = OxmlElement("w:numId")
        n.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(n)
        p._p.get_or_add_pPr().append(num_pr)
        r = p.add_run(item)
        font_run(r, size=11, color=NAVY)


def add_page_field(paragraph):
    paragraph.add_run("第 ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    paragraph.add_run(" 頁")


def set_document_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("新北市青年資料研究｜儀表板證據揭露規格 V1.0")
    font_run(r, size=8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer)
    for run in footer.runs:
        font_run(run, size=8.5, color=GRAY)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("研究方法與系統規格")
    font_run(r, size=11, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("每張圖七項資料證據揭露")
    font_run(r, size=24, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("AWS／Kiro資料沿襲、畫面契約與驗收規格")
    font_run(r, size=13, color=DARK_BLUE)
    add_table(doc, ["文件代號", "資料版本", "規格狀態", "更新日期"], [["DOC-G5-UI-701", "G5 V3.2", "已描述／尚未全面實作", "2026-08-27"]], [1700, 1700, 3260, 2700], 9)


def build():
    doc = Document()
    set_document_styles(doc)
    bullet_id = create_numbering(doc, bullet=True)
    number_id = create_numbering(doc, bullet=False)
    add_title_block(doc)
    add_callout(doc, "核心要求", "每張資料圖都必須固定顯示七項證據。全頁頁尾、tooltip或預設收合區只能補充，不能取代圖表本身的固定揭露。", PALE_GREEN)
    doc.add_heading("1. 文件目的與實作邊界", level=1)
    doc.add_paragraph("本文件把政府行政查核所需的資料證據，轉成AWS資料服務、Kiro規格及前端圖表可以共同遵循的契約。它不改變三母體、年齡換算或現行統計值，只補強每張圖的來源、方法、更新與可用性揭露。")
    add_callout(doc, "現況聲明", "三份發布CSV已具有所需欄位，但目前儀表板精簡JSON與圖表元件尚未完整傳遞並固定顯示，因此狀態為DESCRIBED_NOT_IMPLEMENTED。", PALE_AMBER)
    add_list(doc, [
        "人工讀者以本Word理解定義、規則、範例與驗收方式。",
        "Kiro以chart-evidence-contract.json及chart-evidence-requirements.md產生程式與測試。",
        "AWS資料層以三份長格式CSV、來源主檔及資料列來源關聯作唯一真實來源。",
        "文件存在不代表功能上線；必須通過資料匯出、API、元件及八類圖表驗收。",
    ], bullet_id)

    doc.add_heading("2. 七項固定資料證據定義", level=1)
    add_table(doc, ["固定顯示項目", "行政與研究定義", "公開畫面規則"], [
        ["母體定義", "指出指標的分母與涵蓋對象：戶籍人口、民間人口及勞動市場、受僱員工薪資三者不得混用。", "使用完整中文名稱；同一分析圖不得混入兩種母體。"],
        ["資料年度", "數值所描述的年度或期間，與網站更新日期不同。", "單年顯示民國年與時間基準；趨勢顯示起訖年及缺年。"],
        ["官方精確／官方調查／模型估計", "區分行政全數登記、抽樣調查估計及由官方資料再模型換算。", "顯示value_origin_label_zh；混合時標示混合資料身分。"],
        ["計算方法", "公式、直接加總、PCLM、Sprague、IPF、錨定校準或未估計原因。", "優先顯示formula_zh及method_name_zh；不得將模型值寫成直接加總。"],
        ["更新日期", "本圖使用紀錄最近一次成功擷取或發布快照日期。", "顯示日期；不得拿更新日期取代資料年度。"],
        ["資料來源", "本圖每一筆紀錄實際依據的原子官方來源。", "來源去重、顯示中文名稱並保留可開啟官方URL。"],
        ["可用性狀態", "由publish_status與qa_status共同判定能否查詢與發布。", "顯示可用、部分可用、尚無可發布值或停止發布；缺值不是0。"],
    ], [1800, 4240, 3320], 8.5)

    doc.add_heading("3. 三份CSV欄位對照", level=1)
    doc.add_paragraph("三份長格式CSV使用相同欄位契約；Kiro不得另造第二套名稱。")
    add_table(doc, ["七項證據", "CSV欄位", "資料沿襲補充"], [
        ["母體定義", "universe_code、universe_name_zh", "圖表查詢先依universe_code分流。"],
        ["資料年度", "roc_year、source_period、period_basis", "同時保留資料年度與年／月／年末存量基準。"],
        ["數值身分", "value_origin_class、value_origin_label_zh", "公開用label；class供API與測試。"],
        ["計算方法", "formula_zh、method_code、method_name_zh、model_version", "模型值另保留uncertainty_low／high。"],
        ["更新日期", "retrieved_at、source_snapshot", "retrieved_at為本圖證據更新日期候選。"],
        ["資料來源", "source_alias、source_name、source_url", "record_id再連04來源主檔與05資料列來源關聯。"],
        ["可用性狀態", "publish_status、qa_status", "兩欄均符合查詢條件才可顯示數值。"],
    ], [1650, 4550, 3160], 8.5)
    add_callout(doc, "完整性查核", "截至本次封裝，三份CSV共37,412列；上述母體、年度、數值身分、方法、擷取日期、來源、發布狀態及QA狀態欄位均有值。缺口位於前端傳遞與呈現，不在發布CSV。", LIGHT_BLUE)

    doc.add_heading("4. 現況與目標差距", level=1)
    add_table(doc, ["層級", "目前已有", "尚待補強", "完成判定"], [
        ["發布CSV", "七項證據所需欄位完整，並有record_id與來源雜湊。", "維持schema與發布政策自動測試。", "欄位完整、主鍵唯一、狀態合法。"],
        ["前端JSON／API", "目前保留部分origin、method及全域generatedAt／sources。", "逐紀錄保留母體、公式、擷取日期、來源、publish與QA。", "chartEvidence物件schema驗證通過。"],
        ["圖表元件", "部分圖卡或展開表已有年度、方法或資料身分。", "建立固定可見ChartEvidenceStrip，不能只放details或頁尾。", "八類圖表均含同一共用元件。"],
        ["測試", "現有資料QA與部分前端測試。", "新增56項必要顯示、混合身分、HOLD、來源與視覺回歸。", "T-107、T-108全部通過後才完成。"],
    ], [1500, 2780, 2920, 2160], 8.3)

    doc.add_heading("5. 固定顯示元件規格", level=1)
    doc.add_heading("5.1 建議畫面結構", level=2)
    add_table(doc, ["第一列", "第二列", "互動補充"], [["母體定義｜資料年度｜數值身分｜可用性狀態", "計算方法｜更新日期｜資料來源", "可展開逐年度、逐序列、來源快照與不確定性，但七項摘要仍須固定可見。"]], [3050, 3050, 3260], 8.6)
    doc.add_heading("5.2 多年度與多序列規則", level=2)
    add_list(doc, [
        "資料年度：趨勢圖顯示110–114年；若114缺值，另寫114年尚無可發布值，不縮寫成110–113年造成誤解。",
        "數值身分：若同圖同時含官方抽樣與模型換算，顯示混合資料身分，並在明細列出年度／序列對應。",
        "計算方法：同一方法去重；多方法顯示方法名稱摘要，展開區保留每筆公式與模型版本。",
        "更新日期：摘要採本圖證據紀錄max(retrieved_at)，展開區保留各來源日期；不代表官方資料年度。",
        "資料來源：以record_id連接05資料列來源關聯，再由source_id連04來源主檔；來源名稱與URL去重。",
        "母體衝突：偵測到多個universe_code時fail closed，拆圖後才能發布。",
    ], bullet_id)

    doc.add_heading("6. 可用性狀態判定", level=1)
    add_table(doc, ["優先序", "條件", "公開顯示", "系統行為"], [
        ["1", "QA不在允許清單，且非刻意保留的不可用列", "停止發布", "不得查詢；保留前版並告警。"],
        ["2", "全部為HOLD_DATA_GAP或DATA_GAP_NOT_ESTIMATED", "尚無可發布值", "顯示原因、最近可用年度與來源，不顯示0。"],
        ["3", "部分年度／序列可查詢，部分HOLD或缺漏", "部分可用", "繪製可用點，缺口保留空白並註記。"],
        ["4", "全部publish_status可查詢且qa_status合格", "可用", "正常顯示並保留來源與方法。"],
    ], [900, 3920, 1700, 2840], 8.5)
    add_callout(doc, "薪資範例", "110–113年平均數可用；非25–29歲中位數為HOLD_DATA_GAP。平均與中位數趨勢圖應顯示部分可用，而不是把中位數空白轉成0。", PALE_AMBER)

    doc.add_heading("7. 八類圖表套用矩陣", level=1)
    add_table(doc, ["圖表", "母體", "年度／身分與方法要點", "可用性與來源要點"], [
        ["新北市29區行政區戶籍青年地圖", "戶籍人口母體", "選取年度；人口為官方行政精確值，單一年齡直接加總。", "可用；人口官方來源與最新擷取日。"],
        ["戶籍人口五年趨勢", "戶籍人口母體", "110–114年；官方行政精確值；單一年齡加總。", "逐年完整性與人口來源。"],
        ["男女比例", "戶籍人口母體", "選取年度；官方行政精確值；男／女除以同群合計。", "分母完整才可用；人口來源。"],
        ["教育程度結構", "戶籍人口母體", "選取年度；官方行政邊際推估；PCLM種子＋行政區IPF。", "模型標示與不確定性；教育、人口來源。"],
        ["婚姻狀態結構", "戶籍人口母體", "選取年度；官方行政邊際推估；PCLM種子＋行政區IPF。", "模型標示與不確定性；婚姻、人口來源。"],
        ["勞動力中就業與失業", "民間人口及勞動市場母體", "選取年度；官方抽樣或模型換算；E／LF與U／LF。", "UR＋ESLF=100%；人力資源調查來源。"],
        ["勞參率與就業人口比率趨勢", "民間人口及勞動市場母體", "110–114年；PCLM主模型、Sprague敏感度；由計數重算。", "缺年與模型身分逐年顯示；勞動來源。"],
        ["平均數與中位數趨勢", "受僱員工薪資母體", "110–114年；官方調查／錨定模型；中位數不可由平均數合成。", "114與非原生中位數HOLD；薪資與輔助來源。"],
    ], [1880, 1740, 3240, 2500], 8.1)

    doc.add_heading("8. AWS與Kiro資料流程", level=1)
    add_list(doc, [
        "AWS ingestion保存官方raw快照、來源期間、SHA-256及retrieved_at。",
        "轉換作業產生三份長格式資料，保留record_id、母體、身分、公式、方法、發布與QA狀態。",
        "Machine QA及人工Gate通過後，published區才提供儀表板與ChatBot查詢。",
        "export_g5_dashboard_json.py或API依record_id合併來源主檔與資料列來源關聯，輸出chartEvidence。",
        "Kiro產生ChartEvidenceStrip與型別，所有圖表wrapper強制要求chartEvidence參數。",
        "前端依可用性規則繪圖；不可用值保持null，記錄run_id、catalog_version及錯誤狀態。",
        "CloudWatch監控缺欄位、母體混用、失效URL、HOLD誤顯示、過期更新日及元件漏掛。",
    ], number_id)
    add_callout(doc, "Fail-closed原則", "只要母體衝突、來源沿襲缺失、QA不合格或七項必要欄位無法形成，該圖不得以正式狀態發布；維持上一個已核准版本。", PALE_RED)

    doc.add_heading("9. API／JSON回應契約摘要", level=1)
    add_table(doc, ["物件", "必要欄位", "說明"], [
        ["chartEvidence", "universeDefinition、dataPeriod、valueIdentity、calculationMethod、updatedAt、sources、availability", "每張圖必要，不能為null；sources至少一筆，資料缺口圖除外。"],
        ["sources[]", "sourceId、nameZh、officialUrl、retrievedAt、role", "由來源主檔與資料列來源關聯形成，不接受臨時人工貼URL。"],
        ["availability", "code、labelZh、queryable、reasonZh、missingPeriods", "同時反映publish_status與qa_status。"],
        ["lineage", "recordIds、catalogVersion、runId、bundleSha256", "公開畫面可隱藏代號，但稽核API與日誌必須保留。"],
    ], [1700, 4100, 3560], 8.5)
    doc.add_paragraph("機器完整契約：06_跨母體儀表板規格/chart-evidence-contract.json。Kiro執行需求：12_AWS_Kiro_實作檔案/KIRO規格/specs/ntpc-youth-g5/chart-evidence-requirements.md。")

    doc.add_heading("10. 驗收與人工查核表", level=1)
    add_table(doc, ["查核ID", "必要查核", "預期結果", "責任／證據"], [
        ["G5-UI-701", "八類圖表是否都固定顯示七項", "8×7＝56項全部通過", "前端測試＋人工截圖"],
        ["G5-UI-702", "公開文字是否只有中文或阿拉伯數字", "無內部代號外露", "視覺回歸＋文字掃描"],
        ["G5-UI-703", "趨勢圖的混合身分與缺年", "摘要、逐年明細與空白點一致", "fixture測試"],
        ["G5-UI-704", "HOLD或QA失敗是否顯示為0", "不得出現；顯示不可用原因", "publication-policy測試"],
        ["G5-UI-705", "每張圖來源是否可回到record_id", "來源連結及bundle SHA-256可追溯", "lineage join測試"],
        ["G5-UI-706", "資料年度與更新日期是否混淆", "兩者分列且語意正確", "人工文案查核"],
        ["G5-UI-707", "同圖是否混用母體", "偵測即fail closed", "API與元件單元測試"],
        ["G5-UI-708", "來源連結、鍵盤與讀屏", "連結有效且可操作", "無障礙與URL探測"],
    ], [1250, 3100, 2720, 2290], 8.3)
    add_list(doc, [
        "T-107：完成共用元件與八類圖表套用。",
        "T-108：完成資料沿襲欄位、56項顯示、HOLD、來源連結及視覺回歸測試。",
        "兩項均未通過前，文件與發布說明維持已描述／尚未全面實作。",
    ], bullet_id)

    doc.add_heading("11. 版本、變更與使用限制", level=1)
    add_table(doc, ["版本", "日期", "變更內容", "資料數值影響"], [["V1.0", "2026-08-27", "建立七項證據、聚合、可用性、八圖矩陣、AWS／Kiro流程與驗收契約。", "無；三份CSV維持G5 V3.2。"]], [1200, 1400, 4900, 1860], 8.6)
    add_callout(doc, "使用限制", "本文件是實作規格與查核依據，不是已部署證明。正式上線仍須完成Kiro任務、AWS dev／staging Gate、人工核准及production驗收。", PALE_AMBER)
    doc.add_heading("11.1 主要檔案索引", level=2)
    add_list(doc, [
        "三份長格式CSV：09_發布資料包/01、02、03開頭之CSV。",
        "來源主檔與關聯：09_發布資料包/04_來源主檔.csv、05_資料列來源關聯.csv。",
        "機器契約：06_跨母體儀表板規格/chart-evidence-contract.json。",
        "Kiro需求：12_AWS_Kiro_實作檔案/KIRO規格/specs/ntpc-youth-g5/chart-evidence-requirements.md。",
        "發布政策：11_介接設定與追溯紀錄/config/publication-policy.json。",
        "目前匯出程式：scripts/export_g5_dashboard_json.py；後續須依本契約補欄位。",
    ], bullet_id)

    doc.add_heading("11.2 ZIP解壓後的Kiro讀取順序", level=2)
    handoff_number_id = create_numbering(doc, bullet=False)
    add_list(doc, [
        "讀取根目錄TOC.txt、README.txt及PACKAGE_RELEASE_NOTES_20260827.txt，確認版本與尚未完成項目。",
        "讀取本Word與chart-evidence-contract.json，建立七項證據的共同語意。",
        "讀取KIRO規格的requirements.md、design.md、tasks.md及chart-evidence-requirements.md，不得只依單一檔案產碼。",
        "讀取三份長格式CSV、04來源主檔及05資料列來源關聯，建立record_id到原子來源的join。",
        "修改匯出程式與API schema，再建立ChartEvidenceStrip；先跑fixture及HOLD測試，後跑八類圖表視覺驗收。",
        "完成dev與staging驗收、人工裁示及rollback演練後，才可更新狀態並部署production。",
    ], handoff_number_id)
    doc.add_heading("11.3 禁止事項與部署前裁示", level=2)
    add_callout(doc, "禁止事項", "不得把HOLD或null改成0、不得跨母體合成、不得以全頁來源清單取代每圖來源、不得把更新日期當資料年度、不得將文件完成宣稱為功能上線。", PALE_RED)
    add_table(doc, ["裁示項目", "目前狀態", "部署前需填入或確認"], [
        ["功能狀態", "已描述／尚未全面實作", "T-107、T-108通過後才能改為已完成。"],
        ["AWS環境", "尚未正式部署", "帳號、區域、Owner、預算、告警信箱及TransformWorkflowArn。"],
        ["登入與權限", "JWT／白名單程式與本地測試完成", "正式issuer、audience、Gateway與Policy ENFORCE。"],
        ["發布核准", "本地QA通過", "dev／staging測試、人工核准與production rollback演練。"],
    ], [1850, 3000, 4510], 8.5)

    doc.core_properties.title = "每張圖七項資料證據揭露與AWS／Kiro交接規格"
    doc.core_properties.subject = "新北市青年18–35歲儀表板資料沿襲與查核"
    doc.core_properties.keywords = "AWS,Kiro,資料沿襲,儀表板,母體,資料來源,可用性"
    doc.core_properties.comments = "G5 V3.2；七項證據為已描述、尚未全面實作。"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
