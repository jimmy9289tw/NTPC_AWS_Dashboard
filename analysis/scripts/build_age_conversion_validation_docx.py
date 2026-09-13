from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REGISTRY = ROOT / "docs" / "age-range-conversion-registry-v1.0.csv"
METHODS = ROOT / "docs" / "age-conversion-method-registry-v1.0.csv"
OUTPUT = ROOT / "deliverables" / "各資料表18-35歲換算方法與驗證算例_V1.0.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5F6B76"
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
CAUTION_FILL = "FFF4CE"
RISK_FILL = "FDEBEC"
POSITIVE_FILL = "EAF3EC"
WHITE = "FFFFFF"
GRID = "B8C2CC"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}


WHY = {
    "ARC-001": "原表是五歲年齡組且又按教育程度交叉；18、19、35位於兩個邊界組內。微資料可直接依權數重估，若只能取得分組人數，才以PCLM拆分各教育類別，再用IPF維持人口與教育邊際。",
    "ARC-002": "勞參率、就業人口比率與失業率都是比率，必須用同一批18–35歲樣本的分子與分母重算；人力資源調查又是複雜抽樣，因此不能忽略權數或平均公開年齡組率。",
    "ARC-003": "目前擷取只有25–29、30–34、35–39；18–24完全缺漏，且35歲與36–39混在一起。只要補到單齡／微資料即可用加總，否則沒有唯一的18–35答案。",
    "ARC-004": "就業人口比率的分母是民間人口。各年齡組分母不同，故必須先合併就業人數與民間人口，再相除；比率的算術平均沒有統計意義。",
    "ARC-005": "就業人數具可加總性，但目前缺18–24且35與36–39混合；以微資料加權或官方單齡表補齊後可直接相加。",
    "ARC-006": "失業率的分母是勞動力人口，不是總人口。只有25–29失業率不能推到18–35；需要各年齡失業人數與勞動力人口或微資料。",
    "ARC-007": "教育程度占比必須由各教育類別就業人數除以總就業人數。直接平均不同年齡組百分比會忽略組別人數差異。",
    "ARC-008": "平均數是線性統計量；取得個體權數，或取得完全落在18–35內的組平均數與人數，即可依權數合併。",
    "ARC-009": "中位數是分布位置，不具可加總性。只有個體薪資與權數，或足以重建累積分布的細格資料，才能重新取得18–35中位數。",
    "ARC-010": "教育交叉表有20–34完整組，但18、19、35必須拆邊界。官方單齡交叉表最佳；研究備援以PCLM建立非負單齡種子，再用IPF同時符合單齡人口與教育邊際。",
    "ARC-011": "來源本身已提供單一年齡行政人數，完全不需要模型。直接加總可避免邊界假設，也是18–35人口的最高證據等級。",
    "ARC-012": "15歲以上是開放區間且只有一個教育占比；同一個總占比可由無數種18–35與36歲以上組合產生，因此18–35在統計上不可識別。",
    "ARC-013": "就業結構占比的分母通常是全體就業者；要切出18、19、35，必須有單齡就業人數或可信的組內分布，不能只拆百分比。",
    "ARC-014": "NTPC表為五歲組；專案已有MOI單齡真值，故核心人口直接改接MOI。Sprague／PCLM只保留作來源中斷時的模型備援與回測，不可冒充行政精確值。",
    "ARC-015": "公開欄位是分組失業率，沒有失業人數與勞動力人口；即使年齡組看似相鄰，也不能平均率或直接拆率。",
    "ARC-016": "平均薪資可依受僱人數加權，但30–39跨越35歲，且18–24缺漏。須取得18–24、30–35的薪資總額與人數，或直接以微資料重估。",
    "ARC-017": "25–29與30–39兩個中位數不足以決定合併後中位數；中位數必須由18–35個體薪資分布重新排序及累積權重。",
    "ARC-018": "新北工作場所表目前只有25–29平均薪資；工作場所與居住地不可混用，且沒有其他年齡組的人數與平均數，不能外推18–35。",
    "ARC-019": "新北工作場所表目前只有25–29中位薪資；單一中位數不包含分布資訊，不能外推18–35。",
}


EXAMPLES = {
    "ARC-001": (
        "示範算例／不可作政策值",
        "假設微資料限制18–35後，加權就業人數為大專70千人、高中25千人、其他5千人，則大專占比=70/(70+25+5)=70%。驗證時再限制25–29，回算結果必須與官方25–29表在四捨五入容許差內；若改採PCLM+IPF，拆分後合併回原五歲組也必須完全重現原表。",
    ),
    "ARC-002": (
        "示範算例／微資料取得後執行",
        "假設18–35加權民間人口100千人、勞動力72千人、就業68千人、失業4千人：勞參率=72/100=72.00%，就業人口比率=68/100=68.00%，失業率=4/72=5.56%。先檢查72=68+4，再以25–29子樣本回算官方表；正式結果另報SE、95% CI與RSE。",
    ),
    "ARC-003": (
        "來源實值＋邊界示範",
        "來源原生25–29=246千人、30–34=273千人、35–39=277千人。273+277=550千人會錯把36–39納入，且漏掉18–29。若補件後得到N18=40、N19=42、N20–24=210、N35=54千人，示範合計為40+42+210+246+273+54=865千人；40、42、210、54均為教學假設，不能發布。",
    ),
    "ARC-004": (
        "官方表回算",
        "25–29官方四捨五入人數為就業215千人、民間人口246千人，因此215/246×100=87.398%，與整合表一致。18–35亦須以ΣE/ΣP計算；若ΣE=780、ΣP=900千人，示範比率為86.667%，而不是各組比率的簡單平均。",
    ),
    "ARC-005": (
        "來源實值＋邊界示範",
        "來源原生25–29就業215千人、30–34為251千人、35–39為253千人。若補得E18=30、E19=32、E20–24=170、E35=49千人，示範18–35就業人數=30+32+170+215+251+49=747千人。49必須是單齡35，不能用253千人代替。",
    ),
    "ARC-006": (
        "官方原生值＋率的示範",
        "目前25–29官方原生失業率為6.4%，可用於欄位驗證但不是18–35。若18–35加總失業人數30千人、勞動力700千人，示範失業率=30/700×100=4.286%；只有各組率而沒有U與LF時，此步驟必須被系統阻擋。",
    ),
    "ARC-007": (
        "官方四捨五入人數回算",
        "表50之25–29就業者總數1,253千人、大專及以上994千人，994/1,253×100=79.3296%，與整合表一致。正式18–35驗證相同，但分子與分母必須先由18–35微資料／單齡表加總；不能平均各年齡組教育占比。",
    ),
    "ARC-008": (
        "示範加權平均",
        "假設完全位於18–35內的三組為：18–24人數100、平均45；25–29人數200、平均60；30–35人數150、平均75（萬元/年）。合併平均=(100×45+200×60+150×75)/(100+200+150)=61.667萬元/年。若最後一組其實是30–39，則此算例失效。",
    ),
    "ARC-009": (
        "示範加權中位數",
        "假設18–35薪資排序為40、50、60、80萬元，權數為1、2、4、1；總權數8，累積權數為1、3、7、8，首次達到8×50%=4的位置是60萬元，因此加權中位數=60萬元。平均四個組中位數不能得到相同答案。",
    ),
    "ARC-010": (
        "IPF可重算示範／非官方值",
        "假設PCLM產生的2×2種子表為[[40,60],[30,70]]，指定年齡列邊際[120,80]、教育欄邊際[90,110]。交替列、欄縮放至收斂後為[[59.2014,60.7986],[30.7986,49.2014]]；列和=120、80，欄和=90、110。此例驗證IPF同時符合邊際，但不能證明種子分布正確，故仍要做替代種子與遮蔽回測。",
    ),
    "ARC-011": (
        "官方行政實值驗證",
        "MOI新北市115年7月快照：18歲34,724人、19歲34,391人、20–24歲186,990人、25–29歲244,695人、30–34歲274,842人、35歲56,572人。加總=34,724+34,391+186,990+244,695+274,842+56,572=832,214人；與逐歲18–35加總完全一致，差額0。",
    ),
    "ARC-012": (
        "不可識別證明／阻擋驗證",
        "假設15歲以上人口10,000人、大專以上4,213.5人，總占比同為42.135%。情境A：18–35有4,000人且60%大專以上（2,400人），36歲以上為1,813.5/6,000=30.225%；情境B：18–35仍4,000人但30%大專以上（1,200人），36歲以上為3,013.5/6,000=50.225%。兩者總占比相同，18–35卻差30個百分點，證明不能由15+總占比換算。",
    ),
    "ARC-013": (
        "示範分子加總／不可作政策值",
        "假設單齡／細組就業人數為E18=20、E19=22、E20–24=120、E25–29=210、E30–34=250、E35=48千人，則18–35就業人數=670千人；若全體就業者2,000千人，18–35占比=670/2,000=33.5%。只持有15–24與35–39占比時，20、22、48無法被唯一識別。",
    ),
    "ARC-014": (
        "遮蔽回測設計＋官方基準",
        "先把MOI單齡真值合併成NTPC五歲組，再用Sprague／PCLM拆回，目標真值固定為832,214人。每一折記錄估計值Y_hat、誤差Y_hat−832,214與相對誤差；合併回原五歲組必須一致。合成 sanity test 可用線性單齡序列98、99、…、115檢查係數實作；正式發布仍採MOI精確值而非模型值。",
    ),
    "ARC-015": (
        "率的反例／不可作政策值",
        "假設四段失業人數為10、8、6、1千人，勞動力為100、200、250、60千人，各組率為10%、4%、2.4%、1.667%。正確合併率=(10+8+6+1)/(100+200+250+60)=4.098%；四率簡單平均=4.517%，差0.419個百分點。現行公開表缺U與LF，因此系統應阻擋18–35。",
    ),
    "ARC-016": (
        "可識別與不可識別雙重示範",
        "若取得恰好18–24、25–29、30–35的人數100、200、150及平均45、61.1、72萬元，合併平均=(4,500+12,220+10,800)/450=61.156萬元。反之，來源30–39平均72.4可由30–35平均65、36–39平均83.5組成，也可由30–35平均75、36–39平均68.5組成（假設人數600與400）；所以72.4不能唯一推出30–35。",
    ),
    "ARC-017": (
        "中位數不可合併證明",
        "兩組都可能有中位數50，但分布A=[40,50,50,50]、分布B=[50,50,100,100]；合併權數或人數改變時，總中位數會受完整排序影響。現有25–29中位數52.1與30–39中位數58.1不足以決定18–35中位數，故驗證結果應為BLOCKED，而非產生估計值。",
    ),
    "ARC-018": (
        "官方原生值＋外推阻擋",
        "新北工作場所25–29平均薪資為59.9萬元/年，只能驗證該原生年齡帶。即使假設18–24平均45、30–35平均72，若沒有各組受僱人數，(45+59.9+72)/3=58.97也不是合法合併平均；驗證應確認系統要求人數／權數並阻擋外推。",
    ),
    "ARC-019": (
        "官方原生值＋分布阻擋",
        "新北工作場所25–29中位薪資為52.4萬元/年，只能原生呈現。新增18–24與30–35中位數仍不足以合併；測試輸入數個組中位數時，系統預期回覆「需18–35個體薪資分布與權數」，不得回傳平均後的中位數。",
    ),
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, **kwargs: int) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin in ["top", "start", "bottom", "end"]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(kwargs.get(margin, CELL_MARGIN_DXA[margin])))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None, italic: bool | None = None, name: str = "Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size: float = 11, color: str = "000000") -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color)


def add_label_paragraph(doc: Document, label: str, text: str, *, after: int = 6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(label + "：")
    set_run_font(r1, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2)
    return p


def add_status_paragraph(doc: Document, text: str, fill: str):
    """Render a source status as an in-flow paragraph, avoiding Word table overlap."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Pt(7)
    p.paragraph_format.right_indent = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    run = p.add_run(text)
    set_run_font(run, bold=True, color=INK, size=10)
    return p


def add_formula(doc: Document, formula: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Pt(7)
    p.paragraph_format.right_indent = Pt(7)
    p.paragraph_format.line_spacing = 1.1
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT_FILL)
    p_pr.append(shd)
    run = p.add_run(formula)
    set_run_font(run, name="Consolas", size=9.5, color=INK)


def add_validation_box(doc: Document, kind: str, text: str, status: str) -> None:
    fill = POSITIVE_FILL if "官方" in kind else (RISK_FILL if status == "否" else CAUTION_FILL)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Pt(7)
    p.paragraph_format.right_indent = Pt(7)
    p.paragraph_format.line_spacing = 1.2
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    lead = p.add_run("驗證算例｜" + kind)
    set_run_font(lead, bold=True, color=INK)
    lead.add_break()
    run = p.add_run(text)
    set_run_font(run, size=10.5)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r = OxmlElement("w:r")
    r.append(fld_char1)
    r.append(instr_text)
    r.append(fld_char2)
    paragraph._p.append(r)
    tail = paragraph.add_run(" 頁")
    set_run_font(tail, size=9, color=MUTED)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_document(doc: Document) -> None:
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("新北市青年政策公開資料 AI 系統｜年齡換算研究方法")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    add_page_number(fp)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("研究方法與驗證手冊")
    set_run_font(r, size=12, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("各資料表如何換算為 18–35 歲")
    set_run_font(r, size=28, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(42)
    r = p.add_run("逐來源方法選擇、公式、不確定性與驗證算例")
    set_run_font(r, size=14, color=DARK_BLUE)

    meta = doc.add_table(rows=4, cols=2)
    rows = [
        ("適用範圍", "新北市青年政策公開資料 AI 系統；核心青年18–35歲"),
        ("方法版本", "V1.0｜逐來源登錄19筆；方法登錄9種"),
        ("證據基準", "青年18-35歲資料整合表_V1.3.csv"),
        ("編製日期", date.today().isoformat()),
    ]
    for i, (label, value) in enumerate(rows):
        meta.cell(i, 0).text = label
        meta.cell(i, 1).text = value
        shade(meta.cell(i, 0), HEADER_FILL)
        for run in meta.cell(i, 0).paragraphs[0].runs:
            set_run_font(run, bold=True, color=INK)
        for run in meta.cell(i, 1).paragraphs[0].runs:
            set_run_font(run)
    set_repeat_table_header(meta.rows[0])
    set_table_geometry(meta, [2100, 7260])
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("行政研究文件｜估計值不得冒充官方精確值")
    set_run_font(r, size=10.5, bold=True, color="9B1C1C")
    doc.add_page_break()


def add_overview(doc: Document, registry: list[dict[str, str]], methods: list[dict[str, str]]) -> None:
    doc.add_heading("1. 使用目的與判讀原則", level=1)
    doc.add_paragraph(
        "本手冊回答兩個問題：每一個政府資料表應用哪一種方法換算到18–35歲，以及如何證明算法有正確執行。"
        "同一個來源若同時包含平均數與中位數，會分成不同登錄，因為兩者的數學性質不同。"
    )
    add_validation_box(
        doc,
        "核心原則",
        "只有來源年齡組可由既有欄位無損組成18–35時，才可稱為精確重組。若只有20–34總數，Y(18–35)=Y(20–34)+Y18+Y19+Y35；缺Y18、Y19、Y35時沒有唯一解，必須補單齡／微資料或標示模型估計。",
        "條件式",
    )
    for text in [
        "行政單齡加總：標示「官方行政精確值」，不另造模型信賴區間。",
        "調查微資料加權：標示「官方調查估計值」，報SE、95% CI、RSE與有效樣本數。",
        "Sprague、PCLM、IPF：標示「模型估計值」，報回測誤差、敏感度與預測／bootstrap區間。",
        "無法識別：標示BLOCKED，不以0、前期值、25–29或年齡寬度比例代替。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("2. 方法選擇與不確定性", level=1)
    table = doc.add_table(rows=1, cols=4)
    headers = ["代號", "方法", "適用資料", "不確定性輸出"]
    for j, value in enumerate(headers):
        table.cell(0, j).text = value
        shade(table.cell(0, j), HEADER_FILL)
        for run in table.cell(0, j).paragraphs[0].runs:
            set_run_font(run, bold=True, color=INK, size=9.5)
    set_repeat_table_header(table.rows[0])
    for item in methods:
        cells = table.add_row().cells
        values = [item["method_code"], item["method_name"], item["minimum_inputs"], item["required_uncertainty_output"]]
        for j, value in enumerate(values):
            cells[j].text = value
            for p in cells[j].paragraphs:
                set_paragraph_font(p, size=9)
                p.paragraph_format.space_after = Pt(2)
    set_table_geometry(table, [1500, 1900, 2800, 3160])
    p = doc.add_paragraph("方法來源：")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    set_run_font(p.runs[0], size=9, color=MUTED)
    p.add_run("完整計算步驟見 age-conversion-method-registry-v1.0.csv；本手冊後續逐表套用。")
    set_paragraph_font(p, size=9, color=MUTED)

    doc.add_heading("2.1 驗證類型", level=2)
    for label, text in [
        ("官方實值驗證", "直接以官方單齡或分子分母回算，應與原表一致。"),
        ("遮蔽回測", "有單齡真值時先合併成粗年齡組，再拆回並計算bias、MAE、RMSE與區間。"),
        ("示範算例", "用清楚標示的假設值驗證公式與系統阻擋規則；不得作政策值。"),
        ("不可識別證明", "建立兩組都符合現有公開值、但18–35答案不同的情境，證明不能換算。"),
    ]:
        add_label_paragraph(doc, label, text, after=3)

    doc.add_heading("3. 各資料表換算方法與驗證", level=1)
    summary = doc.add_table(rows=1, cols=4)
    for j, value in enumerate(["登錄", "來源代號", "指標", "18–35判定"]):
        summary.cell(0, j).text = value
        shade(summary.cell(0, j), HEADER_FILL)
        for run in summary.cell(0, j).paragraphs[0].runs:
            set_run_font(run, bold=True, size=9.5)
    set_repeat_table_header(summary.rows[0])
    for item in registry:
        cells = summary.add_row().cells
        for j, value in enumerate([item["registry_id"], item["source_alias"], item["indicator_type"], item["can_publish_18_35"]]):
            cells[j].text = value
            for p in cells[j].paragraphs:
                set_paragraph_font(p, size=8.7)
                p.paragraph_format.space_after = Pt(1)
        shade(cells[3], POSITIVE_FILL if item["can_publish_18_35"] == "是" else (RISK_FILL if item["can_publish_18_35"] == "否" else CAUTION_FILL))
    set_table_geometry(summary, [1100, 2400, 3600, 2260])


def add_source_sections(doc: Document, registry: list[dict[str, str]]) -> None:
    for idx, item in enumerate(registry, start=1):
        source_section = doc.add_section(WD_SECTION.NEW_PAGE)
        source_section.page_width = Inches(8.5)
        source_section.page_height = Inches(11)
        source_section.top_margin = Inches(1)
        source_section.bottom_margin = Inches(1)
        source_section.left_margin = Inches(1)
        source_section.right_margin = Inches(1)
        source_section.header_distance = Inches(0.492)
        source_section.footer_distance = Inches(0.492)
        source_section.different_first_page_header_footer = False
        source_section.header.is_linked_to_previous = True
        source_section.footer.is_linked_to_previous = True
        heading = doc.add_heading(f"3.{idx} {item['registry_id']}｜{item['source_alias']}", level=2)
        heading.paragraph_format.page_break_before = False
        heading.paragraph_format.keep_with_next = True
        rows = [
            ("資料表", item["source_name"]),
            ("指標型態", item["indicator_type"]),
            ("原始年齡", item["source_age_band"]),
            ("目標年齡", item["target_age_band"]),
            ("首選方法", item["preferred_method_code"]),
        ]
        for label, value in rows:
            meta_p = add_label_paragraph(doc, label, value, after=2)
            meta_p.paragraph_format.keep_with_next = True
        status_p = add_status_paragraph(
            doc,
            "發布判定：" + item["can_publish_18_35"] + "｜" + item["estimate_class"],
            POSITIVE_FILL if item["can_publish_18_35"] == "是" else (RISK_FILL if item["can_publish_18_35"] == "否" else CAUTION_FILL),
        )
        status_p.paragraph_format.keep_with_next = False

        add_label_paragraph(doc, "為何採用", WHY[item["registry_id"]])
        add_label_paragraph(doc, "必要輸入", item["required_auxiliary_data"])
        add_label_paragraph(doc, "計算公式／規則", "")
        add_formula(doc, item["formal_formula_or_rule"])
        add_label_paragraph(doc, "不確定性與查核", item["validation_required"])
        kind, example = EXAMPLES[item["registry_id"]]
        add_validation_box(doc, kind, example, item["can_publish_18_35"])
        if item["blocked_reason"]:
            add_label_paragraph(doc, "阻擋條件", item["blocked_reason"], after=4)
        add_label_paragraph(doc, "正式結論", (
            "可產製18–35；輸出仍須附來源版本與查核紀錄。" if item["can_publish_18_35"] == "是"
            else "目前不可產製正式18–35；先補齊必要資料並通過上述驗證。" if item["can_publish_18_35"] == "否"
            else "只有在取得補充資料、完成算法與驗證後才可發布；目前來源原生年齡帶可另行呈現。"
        ), after=10)


def add_sprague_pclm_ipf_details(doc: Document) -> None:
    doc.add_heading("4. Sprague、PCLM、IPF的共同驗證程序", level=1)
    doc.add_heading("4.1 Sprague", level=2)
    doc.add_paragraph("把相鄰五歲人數組排成向量y，依聯合國手冊的係數矩陣S計算單齡估計x_hat=S y。係數與端點處理必須版本化；若有負值、回合併不一致或回測失敗，改用PCLM。")
    add_formula(doc, "x_hat = S y；Y_hat(18–35)=Σ x_hat(a), a=18,…,35")
    doc.add_paragraph("Sprague沒有內建SE。以歷史單齡真值做遮蔽回測，令r_h=log(Y_hat_h/Y_h)，95%經驗預測區間為[Y_hat·exp(−q.975), Y_hat·exp(−q.025)]。")

    doc.add_heading("4.2 PCLM", level=2)
    doc.add_paragraph("分組計數y經組成矩陣C連到單齡期望數γ，logγ以B-spline基底B平滑；λ由凍結的交叉驗證／回測規則選擇。")
    add_formula(doc, "y_g~Poisson(μ_g)；μ=Cγ；logγ=Bθ；θ_hat=argmax{ℓ(y;θ)−(λ/2)||D^dθ||²}")
    doc.add_paragraph("輸出λ、deviance、bias、MAE、RMSE與95%經驗／bootstrap區間。行政總數的bootstrap只代表模型敏感度，不可稱為行政資料抽樣CI。")

    doc.add_heading("4.3 IPF", level=2)
    doc.add_paragraph("從非負種子表開始交替做列、欄縮放，直到所有可信邊際在ε內一致。IPF負責校準，不會憑空創造未觀測資訊。")
    add_formula(doc, "列：m_ij←m_ij·r_i/Σ_jm_ij；欄：m_ij←m_ij·c_j/Σ_im_ij")
    doc.add_paragraph("以不同合理種子、PCLM平滑參數與邊際版本做敏感度分析，報迭代次數、最大邊際誤差與結果範圍。IPF收斂不等於估計真實。")

    doc.add_heading("5. 人工查核與發布清單", level=1)
    checks = [
        "來源期別、地理、母體、性別與單位一致。",
        "原始年齡是單齡、封閉分組或開放區間，已正確登錄。",
        "指標類型（人數／率／平均數／中位數）與方法一致。",
        "18、19、35已取得或已標示模型估計與區間。",
        "率由分子分母重算；平均數有權數；中位數由個體分布重算。",
        "PCLM／Sprague合併回原年齡組一致；IPF邊際誤差在ε內。",
        "25–29官方表回算與18／19／35遮蔽回測均通過。",
        "輸出明確標示官方行政精確值、官方調查估計值或模型估計值。",
        "AWS紀錄輸入版本、程式版本、參數、輸出雜湊、區間與查核人。",
        "儀表板遇到BLOCKED時顯示資料不足，不補0、不沿用25–29。",
    ]
    for check in checks:
        doc.add_paragraph("□ " + check)


def add_sources(doc: Document, registry: list[dict[str, str]]) -> None:
    doc.add_heading("6. 來源與方法文獻", level=1)
    seen: set[str] = set()
    for item in registry:
        url = item["source_url"]
        key = item["source_alias"] + url
        if key in seen:
            continue
        seen.add(key)
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item["source_alias"] + "｜" + item["source_name"])
        set_run_font(r)
        if url:
            p.add_run("：")
            add_hyperlink(p, "官方來源", url)
        else:
            p.add_run("：微資料／官方客製表申請待確認")
        set_paragraph_font(p)

    literature = [
        ("United Nations, Methods for Population Projections by Sex and Age（Sprague multipliers）", "https://www.un.org/development/desa/pd/sites/www.un.org/development.desa.pd/files/files/documents/2020/Jan/un_1956_manual_iii_-_methods_for_population_projections_by_sex_and_age_0.pdf"),
        ("United Nations, World Population Prospects 2024 Methodology（PCLM）", "https://population.un.org/wpp/Publications/Files/WPP2024_Methodology_Advance_Unedited.pdf"),
        ("Eilers (2007), Composite Link Model and Penalized Likelihood", "https://doi.org/10.1177/1471082X0700700302"),
        ("Deming & Stephan (1940), Iterative adjustment with known margins", "https://doi.org/10.1214/aoms/1177731829"),
        ("Deville & Särndal (1992), Calibration Estimators in Survey Sampling", "https://doi.org/10.1080/01621459.1992.10475217"),
        ("行政院主計總處，人力資源調查統計編製方法概述", "https://www.stat.gov.tw/public/Data/9226181433OM26MHO7.pdf"),
    ]
    doc.add_heading("6.1 方法文獻", level=2)
    for label, url in literature:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)

    doc.add_heading("7. 文件版本與可追溯性", level=1)
    doc.add_paragraph(
        "本手冊由 age-range-conversion-registry-v1.0.csv、age-conversion-method-registry-v1.0.csv 與青年18-35歲資料整合表_V1.3.csv產生。"
        "未來任一來源欄位或年齡定義改版，應先更新登錄表與測試算例，再由Kiro修改程式，經本地測試、AWS預檢及人工Gate後部署。"
    )


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "各資料表如何換算為18–35歲：方法與驗證算例"
    props.subject = "新北市青年政策公開資料AI系統年齡整合方法"
    props.author = "新北市青年政策公開資料AI系統研究文件"
    props.keywords = "18-35, Sprague, PCLM, IPF, AWS, Kiro, 新北市青年"
    props.comments = "含逐來源算法、示範算例與發布限制；估計值不得冒充官方精確值。"


def build() -> Path:
    registry = read_csv(REGISTRY)
    methods = read_csv(METHODS)
    if {item["registry_id"] for item in registry} != set(EXAMPLES) or set(EXAMPLES) != set(WHY):
        raise ValueError("WHY／EXAMPLES與來源登錄表不一致")

    doc = Document()
    configure_styles(doc)
    configure_document(doc)
    set_core_properties(doc)
    add_cover(doc)
    add_overview(doc, registry, methods)
    add_source_sections(doc, registry)
    add_sprague_pclm_ipf_details(doc)
    add_sources(doc, registry)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"registry={len(registry)} methods={len(methods)} examples={len(EXAMPLES)}")
    return OUTPUT


if __name__ == "__main__":
    build()
