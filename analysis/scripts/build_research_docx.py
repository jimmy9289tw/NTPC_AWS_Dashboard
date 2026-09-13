from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft JhengHei"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "64748B"
LIGHT = "F2F4F7"
CALLOUT = "F4F6F9"
RISK = "9B1C1C"
GOLD = "7A5A00"


def set_run_font(run, size=11, bold=None, italic=None, color=INK):
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size, color=INK, bold=False):
    style.font.name = FONT_LATIN
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = doc.styles["Title"]
    set_style_font(title, 23, color="000000", bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, 13, color=MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        set_style_font(style, size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def set_cell_margins(table, top=80, bottom=80, start=120, end=120):
    tbl_pr = table._tbl.tblPr
    mar = tbl_pr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(mar)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths: Sequence[int], indent=120, borders=True):
    assert sum(widths) == 9360, widths
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    if borders:
        border_root = tbl_pr.find(qn("w:tblBorders"))
        if border_root is None:
            border_root = OxmlElement("w:tblBorders")
            tbl_pr.append(border_root)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = border_root.find(qn(f"w:{edge}"))
            if el is None:
                el = OxmlElement(f"w:{edge}")
                border_root.append(el)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "CBD5E1")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    set_cell_margins(table)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def apply_cell_text(cell, text, bold=False, color=INK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold, color=color)


def add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int], font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        shade_cell(table.rows[0].cells[i], LIGHT)
        apply_cell_text(table.rows[0].cells[i], header, bold=True, color=DARK_BLUE, size=font_size)
    set_repeat_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            apply_cell_text(row.cells[i], value, size=font_size, align=align)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)
    return table


def add_table_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=9, italic=True, color=MUTED)


def add_numbering(doc: Document, kind="bullet"):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_para(doc, text, bold_lead=None, italic=False, color=INK):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, italic=italic, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic, color=color)
    return p


def add_callout(doc, label, text, color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    r1 = p.add_run(label + "｜")
    set_run_font(r1, size=10.5, bold=True, color=color)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    set_table_geometry(table, [9360], borders=False)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_field(paragraph):
    paragraph.add_run("\t第 ")
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    r2 = paragraph.add_run(" 頁")
    set_run_font(r2, size=8.5, color=MUTED)


def configure_page(doc: Document, doc_code: str, running_title: str):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run(running_title)
    set_run_font(hr, size=8.5, color=MUTED)

    fp = section.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    fr = fp.add_run(f"{doc_code}｜研究草案｜V1.0")
    set_run_font(fr, size=8.5, color=MUTED)
    add_page_field(fp)


def add_masthead(doc, kicker, title, subtitle, metadata, status_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(kicker.upper())
    set_run_font(r, size=9.5, bold=True, color=BLUE)

    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph(style="Subtitle")
    p.add_run(subtitle)

    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(label + "：")
        set_run_font(r1, size=10, bold=True, color=INK)
        r2 = p.add_run(value)
        set_run_font(r2, size=10, color=INK)

    rule = doc.add_paragraph()
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    rule.paragraph_format.space_after = Pt(12)
    add_callout(doc, "文件狀態", status_text)


def new_doc(code, running_title, title, subtitle, status):
    doc = Document()
    configure_styles(doc)
    configure_page(doc, code, running_title)
    doc.core_properties.author = "新北市青年政策公開資料研究原型"
    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    add_masthead(
        doc,
        code,
        title,
        subtitle,
        [("編製日期", "2026-08-16"), ("版本", "V1.0"), ("適用階段", "研究原型／行政查核前")],
        status,
    )
    return doc, add_numbering(doc, "bullet"), add_numbering(doc, "decimal")


def build_research_method():
    doc, bullets, numbers = new_doc(
        "DOC-001-RM",
        "新北市青年政策公開資料 AI 研究方法",
        "新北市青年政策公開資料 AI 系統研究方法",
        "公開資料、青年定義、AWS Harness 與互動式儀表板之可追溯研究設計",
        "方法已完成本地測試與 AWS Harness T02 驗收；正式發布仍受成本、Custom Browser、KB 文件上傳及發布核准 Gate 約束。",
    )

    doc.add_heading("一、研究目的與適用範圍", level=1)
    add_para(doc, "本研究建立一套可供新北市青年政策承辦人查詢、比較及查核公開資料的 AI 輔助方法。系統不替代行政裁量，也不把來源名稱相同視為定義相同；任何正式指標均須能回到來源、版本、母體、地理角色、時間與計算規則。")
    add_list_item(doc, "研究主題：人口、就業、薪資；後續主題須另立資料契約與測試。", bullets)
    add_list_item(doc, "核心地理範圍：新北市；居住地、工作場所與全國基準分開標記。", bullets)
    add_list_item(doc, "初版使用情境：研究草稿、資料盤點、定義衝突檢查、可視化與人工查核。", bullets)
    add_list_item(doc, "不適用情境：個案資格判定、法規解釋、未核准對外發布、含個資或非公開資料分析。", bullets)

    doc.add_heading("二、研究問題", level=1)
    for text in [
        "各官方平台是否提供可對應 18–35 歲核心青年之人口、就業或薪資資料？",
        "資料的母體、年齡邊界、地理角色、時間、單位及公式是否一致，若不一致應如何阻擋錯誤合併？",
        "如何使每一次 AI 回答保留來源、輸入、規則版本、工具呼叫、計算程式及人工裁示證據？",
        "如何以 AWS AgentCore Harness 與私人互動式儀表板實現最小權限、成本上限、可觀測性與回復？",
    ]:
        add_list_item(doc, text, numbers)

    doc.add_heading("三、核定青年年齡定義", level=1)
    add_table(
        doc,
        ["代號", "年齡", "政策解讀", "核心總數"],
        [
            ["AGE-OBS-15-17", "15–17", "青年銜接觀察組", "不計入"],
            ["AGE-CORE-18-24", "18–24", "就學、初入職場與轉銜", "計入"],
            ["AGE-CORE-25-29", "25–29", "職涯建立", "計入"],
            ["AGE-CORE-30-35", "30–35", "職涯、居住與家庭形成", "計入"],
            ["AGE-EXT-36-40", "36–40", "方案擴充", "不計入"],
        ],
        [2200, 1100, 3900, 2160],
    )
    add_table_note(doc, "核心青年總數只能由 18–24、25–29、30–35 三組或 18–35 單一年齡精確加總。")

    doc.add_heading("四、資料來源與縮寫代號", level=1)
    add_table(
        doc,
        ["代號", "平台／資料", "主題", "粒度", "目前用途"],
        [
            ["MOI-POP1Y", "政府資料開放平臺／戶政司單一年齡人口", "人口", "單一年齡", "核心總數"],
            ["OAS-POP", "新北市統計資料庫", "人口", "五歲年齡帶", "鏡像查核"],
            ["NTPC-POP", "新北市資料開放平臺", "人口", "五歲年齡帶", "年度趨勢"],
            ["NTPC-UNEMP", "新北市資料開放平臺", "失業", "來源年齡帶", "25–29 精確"],
            ["NTPC-EMPSTR", "新北市資料開放平臺", "就業結構", "來源年齡帶", "25–29 精確"],
            ["STAT-WAGE-A", "中華民國統計資訊網", "薪資", "全國年齡帶", "基準"],
            ["STAT-WAGE-LOC", "中華民國統計資訊網", "薪資", "縣市×年齡", "25–29 精確"],
        ],
        [1500, 3200, 1100, 1500, 2060],
        font_size=8.8,
    )

    doc.add_heading("五、定義衝突判定方法", level=1)
    add_para(doc, "比較順序固定為：政策目的、統計母體、年齡、地理角色、時間語意、單位、分類版本、計算公式。先比定義，再算數字；不得以數值相近反推定義一致。")
    add_table(
        doc,
        ["分類", "判定條件", "允許處理", "發布"],
        [
            ["EXACT", "定義與邊界完全一致", "可依資料版本計算", "可進下一 Gate"],
            ["CROSSWALK_REQUIRED", "存在可驗證且無損對照", "保存對照表版本", "對照查核後"],
            ["PARTIAL_OR_CONFLICT", "部分重疊或不可拆分", "並列、不合併", "不可作核心指標"],
            ["PENDING_HUMAN_\nDECISION", "影響母體、地理角色、單位或正式解讀", "列人工查核表", "阻擋"],
        ],
        [2100, 3100, 2100, 2060],
        font_size=8.8,
    )
    add_callout(doc, "硬性規則", "來源若只有 15–24、30–34、35–39 等跨界分組，保留原分組；不得依年數比例拆分人數，也不得拆分率、百分比或平均數。", color=RISK)

    doc.add_heading("六、可重現資料處理流程", level=1)
    for text in [
        "登錄來源：保存資料集代號、官方網址、資料期間、擷取時間、授權及版本識別。",
        "原始層保存：下載檔或 API 回應保留原檔、SHA-256 與擷取紀錄，不在原地覆寫。",
        "標準化：欄位改名、型態、單位、地理角色與年齡映射分開執行並保存程式碼。",
        "品質檢查：檢查缺漏、重複、異常值、時間連續性、修訂、母體與年齡邊界。",
        "計算與輸出：確定性加總由程式執行；每張表與回答附來源及限制。",
        "人工裁示：衝突或發布項目由具名決策人核准，保存理由、範圍、期限與證據。",
    ]:
        add_list_item(doc, text, numbers)

    doc.add_heading("七、AWS 與互動式儀表板實作", level=1)
    add_para(doc, "AWS 初版以 Amazon Bedrock AgentCore Harness 為受控執行面；模型採 global.amazon.nova-2-lite-v1:0。允許工具僅為 Browser、Code Interpreter 與 request_human_approval，自訂函數用於任何登入、上傳、發布或外部寫入前停止並要求人工核准。")
    add_table(
        doc,
        ["層次", "目前狀態", "正式上線要求"],
        [
            ["Harness", "版本 3 READY；IAM；Memory off", "持續監控版本、成本與錯誤"],
            ["Browser", "AWS 內建 Browser；公有網路 MVP", "Custom Browser 網域企業政策"],
            ["Code Interpreter", "僅確定性計算與品質檢查", "保留程式、輸入與輸出摘要"],
            ["Knowledge Base", "未建立；未上傳本機文件", "先核准文件清冊、S3、KB、Gateway"],
            ["儀表板", "私人站台；本地證據問答", "IAM API bridge 後才啟用雲端即時問答"],
        ],
        [1900, 3300, 4160],
        font_size=9,
    )

    doc.add_heading("八、驗證、效能與安全證據", level=1)
    add_list_item(doc, "Harness schema validate：通過。", bullets)
    add_list_item(doc, "治理單元測試：7/7 通過；preflight 通過。", bullets)
    add_list_item(doc, "儀表板：3/3 路由測試及正式建置通過。", bullets)
    add_list_item(doc, "AWS T02：年齡邊界衝突驗收通過；830 input、1,131 output、延遲 6.69 秒、工具呼叫 0。", bullets)
    add_list_item(doc, "正式執行依賴 npm audit --omit=dev：0；開發工具鏈仍有 20 項待受控升級。", bullets)

    doc.add_heading("九、研究限制與人工查核", level=1)
    add_table(
        doc,
        ["代號", "待裁示事項", "阻擋範圍"],
        [
            ["HR-AWS-02", "Owner、CostCenter、月預算、告警收件人、到期日", "正式上線"],
            ["HR-AWS-04", "Custom Browser 官方網域企業政策", "Browser 正式上線"],
            ["HR-AWS-05", "可上傳 AWS 的研究文件清冊與資料分級", "KB 建立"],
            ["HR-AWS-07", "錄影 7 天、稽核中繼資料 90 天", "保存政策"],
            ["HR-API-01", "私人儀表板之 IAM API bridge 與受眾", "雲端即時問答"],
            ["HR-SEC-02", "開發工具鏈相容性升級", "正式維護版"],
        ],
        [1700, 5000, 2660],
        font_size=9,
    )

    doc.add_heading("十、參考資料與追溯", level=1)
    refs = [
        "新北市政府青年局 AI 黑客松命題文件。",
        "成功大學 AI 研究啟動會議簡報（研究架構與啟動會議表達方式之參考；本文件為重新研究設計）。",
        "藥品缺藥計畫啟動會議簡報（資料盤點、欄位對照與查核方法之參考；未沿用其研究結論）。",
        "Bedrock AgentCore 工作坊操作手冊、2026 新北市 AI 智慧城市黑客松競賽實戰工作坊。",
        "AWS Well-Architected Framework 與 Security、Reliability、Performance Efficiency、Cost Optimization pillars。",
        "官方資料入口：data.gov.tw、data.ntpc.gov.tw、oas.bas.ntpc.gov.tw、stat.gov.tw、ris.gov.tw。",
    ]
    for ref in refs:
        add_list_item(doc, ref, bullets)

    path = OUT / "研究方法_新北市青年政策公開資料AI系統_V1.0.docx"
    doc.save(path)
    return path


def build_doc_002():
    doc, bullets, numbers = new_doc(
        "DOC-002",
        "資料定義與政府公開資料源研究會議文件",
        "DOC-002｜資料定義與政府公開資料源研究會議文件",
        "青年參數、來源資料庫、欄位互應與衝突處理之研究版",
        "本文件為本案重新產出的研究會議參考資料；吸收啟動會議文件的結構化方法，但不複製原文或他案結論。",
    )

    doc.add_heading("一、會議目的", level=1)
    add_para(doc, "供研究、資料與政策承辦人共同確認：各分析參數由哪些官方資料庫支持、哪些欄位可介接、哪些定義不能合併，以及衝突發生時應由誰裁示。")
    add_callout(doc, "建議結論", "人口核心總數以 MOI-POP1Y 單一年齡資料為基準；OAS-POP 與 NTPC-POP 作鏡像與趨勢查核。就業與薪資資料按來源年齡帶並列，除 25–29 外不強行轉成核定三組。")

    doc.add_heading("二、會議前已核定定義", level=1)
    for text in [
        "核心青年為 18–35 歲，拆為 18–24、25–29、30–35。",
        "15–17 為銜接觀察組；36–40 為方案擴充組，均不進核心總數。",
        "新北市為主要範圍，但居住地與工作場所必須分欄。",
        "政府資料開放平臺、新北市統計資料庫、新北市資料開放平臺及中華民國統計資訊網均納入。",
        "未經人工核准不得對外發布或執行外部寫入。",
    ]:
        add_list_item(doc, text, bullets)

    doc.add_heading("三、參數 × 資料庫互應矩陣", level=1)
    add_table(
        doc,
        ["分析參數", "MOI-POP1Y", "OAS-POP", "NTPC 系列", "STAT-WAGE-LOC", "處理判定"],
        [
            ["單一年齡", "有", "無", "多為無", "無", "人口核心採 MOI"],
            ["18–35 核心總數", "可精確", "不可精確", "不可精確", "不適用", "只用 MOI"],
            ["25–29", "可精確", "視分組", "部分可精確", "可精確", "可跨主題並列"],
            ["30–35", "可精確", "跨界", "跨界", "30–39 跨界", "不拆分"],
            ["性別", "有", "依表", "依資料集", "依表", "保存原欄位"],
            ["居住地", "有", "有", "多為有", "無", "不得與工作地混用"],
            ["工作場所", "無", "無", "依資料集", "有", "獨立 geography_role"],
            ["薪資平均／中位", "無", "無", "部分", "有", "不得以人口為分母"],
            ["來源版本／期間", "有", "有", "有", "有", "必填稽核欄位"],
            ["官方 URL", "有", "有", "有", "有", "回答須回指"],
        ],
        [1700, 1200, 1200, 1450, 1700, 2110],
        font_size=7.8,
    )
    add_table_note(doc, "「可精確」只表示年齡邊界可對應，不代表母體、地理角色、期間或單位相同。")

    doc.add_heading("四、統一資料欄位建議", level=1)
    add_table(
        doc,
        ["欄位", "中文名稱", "用途／規則"],
        [
            ["dataset_id", "資料集代號", "固定縮寫；連回來源清冊"],
            ["source_url", "官方網址", "必須為官方來源或核准 KB URI"],
            ["retrieved_at", "擷取時間", "ISO 8601；含時區"],
            ["reference_period", "資料期間", "不得與擷取時間混用"],
            ["population_scope", "統計母體", "戶籍人口、就業者、受僱員工等"],
            ["geography_role", "地理角色", "residence、workplace、national benchmark"],
            ["age_lower／age_upper", "年齡邊界", "整數閉區間；跨界不得比例拆分"],
            ["measure_type", "測量型態", "count、rate、percentage、mean、median"],
            ["unit", "單位", "人、%、萬元等；保存換算"],
            ["mapping_status", "對應狀態", "EXACT、CROSSWALK_REQUIRED、PARTIAL_OR_CONFLICT"],
            ["publication_gate", "發布 Gate", "PASS、BLOCKED、PENDING_HUMAN_DECISION"],
        ],
        [2400, 2200, 4760],
        font_size=8.6,
    )

    doc.add_heading("五、共同打勾欄位與系統介接", level=1)
    add_para(doc, "多來源共同具備的欄位可列為未來系統介接候選，但「共同有欄位」不等於「欄位定義一致」。介接採兩層：第一層保存原始欄位與來源；第二層經定義檢查後才建立標準欄位。")
    for text in [
        "第一優先：資料集 ID、官方 URL、資料期間、擷取時間、地理名稱、統計值、單位。",
        "第二優先：年齡上下界、性別、母體、地理角色、測量型態、分類版本。",
        "不得自動介接：僅名稱相同、跨越 18 或 35 邊界、分母未知、平均與中位數混用、居住地與工作地混用。",
    ]:
        add_list_item(doc, text, bullets)

    doc.add_heading("六、衝突處理決策樹", level=1)
    for text in [
        "確認政策目的與統計母體是否相同；不同即並列。",
        "確認年齡邊界是否完全落在核定組；跨界即 AGE_BOUNDARY_CONFLICT。",
        "確認地理角色；residence 與 workplace 不得合併為同一政策漏斗。",
        "確認期間、單位、測量型態與公式；不能無損對照即 PARTIAL_OR_CONFLICT。",
        "若影響正式指標或對外解讀，建立 PENDING_HUMAN_DECISION 並阻擋發布。",
    ]:
        add_list_item(doc, text, numbers)

    doc.add_heading("七、目前資料品質發現", level=1)
    add_table(
        doc,
        ["代號", "發現", "處理"],
        [
            ["DQ-003", "多數來源年齡帶跨越核定邊界", "除 25–29 外不進正式核心總數"],
            ["DQ-004", "薪資為工作地、人口為居住地", "並列並加 geography_role 註記"],
            ["DQ-007", "戶籍人口、民間人口、就業者、全時受僱員工母體不同", "禁止串成單一漏斗"],
            ["DQ-008", "人口 2026-07；其他多至 2024", "分別標示資料期"],
            ["DQ-009", "2013 年三行政區年齡加總短少", "來源確認前排除相關 9 筆趨勢"],
        ],
        [1500, 4300, 3560],
        font_size=8.8,
    )

    doc.add_heading("八、建議會議裁示", level=1)
    add_table(
        doc,
        ["裁示代號", "建議議題", "建議選項"],
        [
            ["DEC-DATA-01", "正式人口基準", "MOI-POP1Y 為主；OAS／NTPC 僅查核"],
            ["DEC-DATA-02", "跨界年齡", "並列原分組；不估算、不拆分"],
            ["DEC-DATA-03", "人口與薪資", "分面呈現；禁止同一母體解讀"],
            ["DEC-DATA-04", "資料修訂", "保存版本與雜湊；不得覆寫原檔"],
            ["DEC-DATA-05", "發布責任", "具名承辦／主管一次性核准"],
        ],
        [1800, 3900, 3660],
        font_size=9,
    )

    doc.add_heading("九、會議查核紀錄欄", level=1)
    add_table(
        doc,
        ["項目", "填寫"],
        [
            ["會議日期／地點", "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"],
            ["主席／出席單位", "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"],
            ["核定事項", "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"],
            ["退回補件", "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"],
            ["決策人／日期", "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"],
            ["證據連結／版本", "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"],
        ],
        [2600, 6760],
    )

    path = OUT / "DOC-002_資料定義與政府公開資料源研究會議文件_V1.0.docx"
    doc.save(path)
    return path


def build_doc_003():
    doc, bullets, numbers = new_doc(
        "DOC-003",
        "AWS Harness 與互動式儀表板研究會議文件",
        "DOC-003｜AWS Harness 與互動式儀表板研究會議文件",
        "受控 AI Agent、部署驗收、Well-Architected 與人工 Gate 之研究版",
        "AWS Harness V3 與私人儀表板已可運行；雲端即時問答、KB 與 Custom Browser 仍待核准後續建置。",
    )

    doc.add_heading("一、會議目的與決策摘要", level=1)
    add_para(doc, "確認初版 AI Agent 如何在 AWS 上以最小權限運行、回答問題、保存稽核證據，並以互動式儀表板提供人口、就業、薪資、來源與風險查詢。")
    add_callout(doc, "目前成果", "AgentCore Harness ntpc_youth_ai_v1 版本 3 已就緒；Nova 2 Lite T02 通過。私人儀表板版本 4 已部署，可提問並使用本地已驗證證據回答。")

    doc.add_heading("二、初版架構", level=1)
    for text in [
        "使用者在私人儀表板查詢人口、薪資、來源及定義衝突。",
        "儀表板目前先由本地已驗證證據回覆，回應標示 LOCAL_EVIDENCE_FALLBACK。",
        "AWS Harness 提供受控研究 Agent；模型、工具、迭代、逾時與字元上限皆版本化。",
        "Browser 只讀核准官方公開來源；Code Interpreter 執行確定性計算與品質檢查。",
        "request_human_approval 攔截登入、上傳、發布、建立、修改、刪除或外部寫入。",
        "未來由 IAM 保護的 API bridge 串接私人儀表板與 Harness；KB 只收核准文件。",
    ]:
        add_list_item(doc, text, numbers)

    doc.add_heading("三、AWS 實際部署紀錄", level=1)
    add_table(
        doc,
        ["欄位", "值"],
        [
            ["帳號／角色", "787498376572／WSParticipantRole/Participant"],
            ["區域", "us-west-2"],
            ["Harness", "ntpc_youth_ai_v1；ID ntpc_youth_ai_v1-jIOgOhhfaT"],
            ["Endpoint", "DEFAULT；版本 3；READY"],
            ["模型", "global.amazon.nova-2-lite-v1:0"],
            ["服務角色", "AmazonBedrockAgentCoreHarnessDefaultServiceRole-3r0zx"],
            ["傳入驗證／網路", "IAM／PUBLIC MVP"],
            ["Memory", "Disabled"],
            ["限制", "12 iterations；15 分鐘；8192 字元；idle 5 分；lifetime 1 小時"],
        ],
        [2700, 6660],
        font_size=9,
    )

    doc.add_heading("四、工具邊界與人工核准", level=1)
    add_table(
        doc,
        ["工具", "允許用途", "禁止／停止條件"],
        [
            ["aws_browser_v1", "官方公開資料查詢、讀取、下載", "登入、表單送出、個人 Google、非核准網域"],
            ["aws_codeinterpreter_v1", "型態、缺漏、重複、雜湊、加總、重現", "IAM 繞過、任意外部連線、未保存程式"],
            ["request_human_approval", "產生具冪等鍵之核准請求", "呼叫後不得自行繼續"],
            ["Shell／file operations", "未開放", "不得以其他工具繞過"],
        ],
        [2200, 3600, 3560],
        font_size=8.8,
    )

    doc.add_heading("五、模型選擇與錯誤處理紀錄", level=1)
    add_para(doc, "部署過程保留錯誤、原因、修正與不採取事項，避免以擴張權限掩蓋模型或帳號限制。")
    add_table(
        doc,
        ["事件", "原因", "處理", "結果"],
        [
            ["ValidationException", "Sonnet 4.6 不接受同時設定 temperature 與 Top P", "保留 temperature 0.1；移除 Top P", "參數錯誤排除"],
            ["AccessDeniedException", "工作坊帳號無 Marketplace ViewSubscriptions／Subscribe", "不擴張 IAM；改選有存取權模型", "避免權限擴張"],
            ["模型切換", "需兼顧中文、工具與工作坊可用性", "採 Amazon Nova 2 Lite V1", "Harness V3 READY"],
        ],
        [2200, 3000, 2600, 1560],
        font_size=8.3,
    )

    doc.add_heading("六、AWS T02 驗收結果", level=1)
    add_table(
        doc,
        ["驗收欄位", "紀錄"],
        [
            ["測試情境", "來源僅有 15–24、25–29、30–34、35–39，判斷能否計算 18–35 核心青年"],
            ["預期", "AGE_BOUNDARY_CONFLICT；不得比例拆分；不可發布；列人工查核表"],
            ["實際", "符合預期；未呼叫 Browser"],
            ["Session ID", "695e5a30-c8a9-4d31-84d0-fb66045b103d"],
            ["用量／延遲", "830 input；1,131 output；總計 1,961；6,690 ms"],
        ],
        [2500, 6860],
        font_size=9,
    )

    doc.add_heading("七、互動式儀表板功能", level=1)
    for text in [
        "可提問：核心青年人口、年齡定義、薪資、來源平台及定義衝突。",
        "可篩選：性別及全部／核心年齡範圍。",
        "可追溯：來源縮寫連回官方網址；顯示資料期、母體、地理角色與限制。",
        "可查核：顯示資料品質風險與人工裁示事項。",
        "可降級：AWS 尚未由 IAM API bridge 串接時只用本地證據，不偽裝成雲端回答。",
        "私人發布：需登入；未改為公開站台。",
    ]:
        add_list_item(doc, text, bullets)

    doc.add_heading("八、Well-Architected 檢視", level=1)
    add_table(
        doc,
        ["面向", "目前控制", "後續缺口"],
        [
            ["Security", "IAM、工具 allowlist、Memory off、人工核准", "Custom Browser、KB 文件分級、API bridge"],
            ["Reliability", "版本化 endpoint、錯誤紀錄、回復至已驗證版", "T01–T06 全套自動驗收、告警"],
            ["Performance", "Nova 2 Lite；T02 6.69 秒；限制 12 iterations", "建立 P95 基準、冷啟動與 Browser 延遲"],
            ["Cost", "Lite 模型、迭代／字元／生命週期上限", "月預算、告警收件人、到期日"],
            ["Operations", "CloudWatch 入口、runbook、session ID", "日誌交付、保存期與事件演練"],
        ],
        [1600, 3800, 3960],
        font_size=8.6,
    )

    doc.add_heading("九、安全與供應鏈稽核", level=1)
    add_list_item(doc, "正式執行依賴 npm audit --omit=dev：0 low、0 moderate、0 high、0 critical。", bullets)
    add_list_item(doc, "完整開發／建置工具鏈：1 low、4 moderate、15 high、0 critical；不得用 --force 自動改版。", bullets)
    add_list_item(doc, "正式維護版應在獨立變更單更新 vinext、Vite、Cloudflare 工具及相關傳遞依賴，重跑完整建置與部署。", bullets)

    doc.add_heading("十、待裁示事項", level=1)
    add_table(
        doc,
        ["代號", "裁示內容", "未裁示影響"],
        [
            ["DEC-AWS-01", "Owner、CostCenter、月預算、告警收件人、到期日", "不得正式上線"],
            ["DEC-AWS-02", "是否核准建立 S3、KB、Gateway 與上傳文件清冊", "僅能公開資料／本地證據"],
            ["DEC-AWS-03", "Custom Browser 允許網域與錄影保存期", "Browser 僅 MVP"],
            ["DEC-API-01", "儀表板雲端問答受眾、驗證與 API bridge", "維持本地證據模式"],
            ["DEC-SEC-01", "開發工具鏈升級維護窗口", "保留非執行期風險"],
        ],
        [1800, 4800, 2760],
        font_size=8.8,
    )

    doc.add_heading("十一、建議下一 Gate", level=1)
    for text in [
        "先核定成本責任、保存期、KB 文件清冊與儀表板受眾。",
        "建立 Custom Browser 與官方網域政策，執行 T04、T05。",
        "建立 S3、Managed Knowledge Base、Gateway，執行 T06 引用驗收。",
        "設計 IAM API bridge，將 Harness 回答以來源、限制、trace ID 回傳私人儀表板。",
        "完成 T01–T06、成本告警、日誌保存、回復演練後，再送正式發布核准。",
    ]:
        add_list_item(doc, text, numbers)

    path = OUT / "DOC-003_AWS-Harness與互動式儀表板研究會議文件_V1.0.docx"
    doc.save(path)
    return path


def main():
    paths = [build_research_method(), build_doc_002(), build_doc_003()]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
