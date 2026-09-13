from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import re
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import numpy as np
import openpyxl

from process_marital_status_v23 import ipf, select_pclm


def resolve_project_root() -> Path:
    """Find the repository root from either scripts/ or a packaged script copy."""
    for candidate in Path(__file__).resolve().parents:
        if (candidate / "config" / "g5-refresh-policy.json").is_file() and (candidate / "scripts").is_dir():
            return candidate
    raise RuntimeError("Cannot locate NtpcYouthAI project root")


PROJECT = resolve_project_root()
DEFAULT_PACKAGE = (
    PROJECT
    / "deliverables"
    / "NTPC_Youth_Complete_Data_Package_G5_V3.2_20260826"
)
LABOR_RESULTS = (
    PROJECT
    / "deliverables"
    / "NTPC_Labor_18-35_Sprague_PCLM_Pilot_V1.1_20260825"
    / "annual_110_114"
    / "annual_target_age_results.csv"
)
LABOR_SOURCE_ROOT = LABOR_RESULTS.parents[1]
LABOR_ANNUAL_RAW_ROOT = LABOR_SOURCE_ROOT / "raw" / "annual"
LABOR_H2_RECONCILIATION = LABOR_SOURCE_ROOT / "half_year_114H2" / "source_rounding_reconciliation.csv"
LABOR_ANNUAL_URLS = {
    110: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=207903",
    111: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=231112",
    112: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=234726",
    113: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=234885",
    114: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=236078",
}
LABOR_H2_URL = "https://www.stat.gov.tw/News_Content.aspx?n=4003&s=235759"

LABOR_LINEAGE = {
    "CIVILIAN_POPULATION": (
        "DGBAS-HR-T27",
        "人力資源調查年報表27：15歲以上民間人口之教育程度與年齡",
        (27,),
    ),
    "LABOR_FORCE": (
        "DGBAS-HR-T32+DGBAS-HR-T36|QA:DGBAS-HR-T28",
        "年報表32就業者＋表36失業者重建勞動力；表28勞動力作四捨五入一致性查核",
        (32, 36, 28),
    ),
    "EMPLOYED": (
        "DGBAS-HR-T32",
        "人力資源調查年報表32：就業者之教育程度與年齡",
        (32,),
    ),
    "UNEMPLOYED": (
        "DGBAS-HR-T36",
        "人力資源調查年報表36：失業者之教育程度與年齡",
        (36,),
    ),
    "NOT_IN_LABOR_FORCE": (
        "DGBAS-HR-T27+DGBAS-HR-T32+DGBAS-HR-T36",
        "年報表27民間人口減表32就業者與表36失業者",
        (27, 32, 36),
    ),
    "UNEMPLOYMENT_RATE": (
        "DGBAS-HR-T32+DGBAS-HR-T36",
        "年報表36失業者除以表32就業者與表36失業者之和",
        (32, 36),
    ),
    "LABOR_FORCE_PARTICIPATION_RATE": (
        "DGBAS-HR-T27+DGBAS-HR-T32+DGBAS-HR-T36",
        "年報表32就業者與表36失業者之和除以表27民間人口",
        (27, 32, 36),
    ),
    "EMPLOYMENT_TO_POPULATION_RATE": (
        "DGBAS-HR-T27+DGBAS-HR-T32",
        "年報表32就業者除以表27民間人口",
        (27, 32),
    ),
    "EMPLOYMENT_SHARE_OF_LABOR_FORCE": (
        "DGBAS-HR-T32+DGBAS-HR-T36",
        "年報表32就業者除以表32就業者與表36失業者之和",
        (32, 36),
    ),
}
WAGE_WORKBOOK = (
    PROJECT
    / "data"
    / "raw"
    / "WAGE-AUX"
    / "20260826"
    / "STAT-WAGE-LOC-TABLE6-current.xlsx"
)
WAGE_AUX_DIR = PROJECT / "data" / "raw" / "WAGE-AUX" / "20260826"

YEARS = (110, 111, 112, 113, 114)
SOURCE_GROUPS = ((15, 19), (20, 24), (25, 29), (30, 34), (35, 39))
TARGET_BANDS = {
    "18-24": tuple(range(18, 25)),
    "25-29": tuple(range(25, 30)),
    "30-35": tuple(range(30, 36)),
    "18-35": tuple(range(18, 36)),
}
SEXES = ("男", "女")

SCHEMA_FIELDS = [
    "record_id",
    "roc_year",
    "gregorian_year",
    "source_period",
    "period_basis",
    "universe_code",
    "universe_name_zh",
    "geography_code",
    "geography_name_zh",
    "geography_level",
    "geography_role_zh",
    "age_band",
    "sex_code",
    "sex_name_zh",
    "category_dimension_code",
    "category_dimension_name_zh",
    "category_code",
    "category_name_zh",
    "metric_code",
    "metric_name_zh",
    "value",
    "unit",
    "numerator",
    "denominator",
    "formula_zh",
    "value_origin_class",
    "value_origin_label_zh",
    "method_code",
    "method_name_zh",
    "model_version",
    "uncertainty_low",
    "uncertainty_high",
    "uncertainty_type",
    "source_alias",
    "source_name",
    "source_url",
    "source_snapshot",
    "source_sha256",
    "retrieved_at",
    "publish_status",
    "qa_status",
    "note_zh",
]

ORIGIN_LABELS = {
    "OFFICIAL_ADMIN_EXACT": "官方行政精確值",
    "MODEL_ESTIMATE_FROM_OFFICIAL_MARGINS": "官方行政邊際推估值",
    "OFFICIAL_SURVEY_ESTIMATE_AGE_BAND_EXACT": "官方抽樣調查年齡帶估計值",
    "MODEL_ESTIMATE_FROM_OFFICIAL_SURVEY_ESTIMATES": "由官方抽樣調查估計值再模型換算",
    "OFFICIAL_SURVEY_STATISTIC": "官方調查統計值",
    "MODEL_ESTIMATE_CALIBRATED_OFFICIAL_WAGE_AUX": "官方薪資錨定之輔助資料模型估計值",
    "MODEL_ESTIMATE_CALIBRATED_OFFICIAL_WAGE_DISTRIBUTION": "官方薪資錨定之分布模型估計值",
    "DATA_GAP_NOT_ESTIMATED": "資料缺口／未估計",
}

METHOD_LABELS = {
    "M0_DIRECT_SINGLE_AGE_SUM": "單一年齡直接加總",
    "M0_DIRECT_5Y_AGGREGATION": "官方五歲年齡帶直接加總",
    "PCLM_CITY_SEED_IPF_GEO": "全市PCLM種子加行政區IPF邊際校準",
    "PCLM": "PCLM年齡拆分",
    "M0_OFFICIAL_AGE_BAND": "官方原生年齡帶",
    "WAGE_PCLM_RATIO_CALIBRATION": "薪資官方錨定＋PCLM輔助年齡輪廓比率校準",
    "WAGE_LOGNORMAL_SHAPE_TRANSFER_PCLM_MIXTURE": "官方薪資錨定＋對數常態形狀轉移＋PCLM權重混合",
    "REQUIRES_MICRODATA_DISTRIBUTION": "需個體薪資分布或可重建之級距頻數",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def file_observed_at(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).isoformat()


def write_csv(path: Path, rows: Iterable[dict]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=SCHEMA_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in SCHEMA_FIELDS})
            count += 1
    return count


def normalize_float(value: float | int | str | None, digits: int = 6) -> str:
    if value is None or value == "":
        return ""
    number = float(value)
    if not math.isfinite(number):
        return ""
    return f"{number:.{digits}f}"


def find_population_snapshots(years: tuple[int, ...]) -> dict[int, Path]:
    candidates: dict[int, list[Path]] = defaultdict(list)
    for path in (PROJECT / "data" / "raw" / "MOI-POP1Y").glob("*/MOI-POP1Y.json"):
        with path.open("rb") as handle:
            prefix = handle.read(300).decode("utf-8", errors="ignore")
        match = re.search(r'"period":"?(\d{5})', prefix)
        if not match:
            match = re.search(r'"statistic_yyymm":"?(\d{5})', prefix)
        if match and match.group(1).endswith("12"):
            candidates[int(match.group(1)[:3])].append(path)
    selected = {}
    for year in years:
        matches = candidates.get(year, [])
        if not matches:
            raise FileNotFoundError(f"Missing MOI-POP1Y December snapshot for ROC {year}")
        selected[year] = max(matches, key=lambda item: item.stat().st_mtime)
    return selected


def district_code_from_row(row: dict) -> str:
    code = str(row.get("district_code", ""))
    return code[:8] if len(code) >= 8 else code


def read_population(path: Path, expected_period: str):
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    if str(payload.get("period")) != expected_period:
        raise ValueError(f"Population period mismatch in {path}: {payload.get('period')}")
    rows = payload.get("responseData", [])
    expected_size = int(payload.get("totalDataSize", len(rows)))
    if len(rows) != expected_size:
        raise ValueError(f"Population pagination incomplete: {len(rows)} != {expected_size}")

    totals: dict[tuple[str, str, str, int], int] = defaultdict(int)
    district_names: dict[str, str] = {"65000000": "新北市"}
    all_age_totals: dict[tuple[str, str, str], int] = defaultdict(int)
    for row in rows:
        site = str(row.get("site_id", ""))
        code = str(row.get("district_code", ""))
        if not site.startswith("新北市") or not code.startswith("65"):
            continue
        district_code = district_code_from_row(row)
        district_names[district_code] = site
        geographies = ((district_code, site), ("65000000", "新北市"))
        for geo_code, geo_name in geographies:
            for sex, suffix in (("男", "m"), ("女", "f")):
                all_age_totals[(geo_code, geo_name, sex)] += int(row.get(f"people_total_{suffix}", 0) or 0)
                for age in range(0, 100):
                    totals[(geo_code, geo_name, sex, age)] += int(
                        row.get(f"people_age_{age:03d}_{suffix}", 0) or 0
                    )
    district_names = dict(sorted(district_names.items(), key=lambda item: item[0]))
    return totals, all_age_totals, district_names, len(rows)


def parse_int(value: str | int | None) -> int:
    try:
        return int(str(value or "0").replace(",", "").strip())
    except ValueError:
        return 0


def education_category(raw: str) -> tuple[str, str]:
    value = raw.strip()
    if value.startswith(("博", "碩")):
        return "GRADUATE", "研究所"
    if value.startswith("大"):
        return "UNIVERSITY", "大學"
    if value.startswith("專"):
        return "COLLEGE", "專科"
    if value.startswith(("高中", "高職")):
        return "SENIOR_HIGH", "高中職"
    if value.startswith(("國中", "國小")):
        return "JUNIOR_HIGH_OR_BELOW", "國中及以下"
    return "OTHER", "其他或未註記"


def marital_category(raw: str) -> tuple[str, str]:
    value = raw.strip()
    if value.startswith("未婚"):
        return "NEVER_MARRIED", "未婚"
    if value.startswith("有偶"):
        return "MARRIED_WITH_SPOUSE", "有偶"
    if value.startswith(("離婚", "終止結婚")):
        return "DIVORCED", "離婚或終止結婚"
    if value.startswith("喪偶"):
        return "WIDOWED", "喪偶"
    return "OTHER", "其他或未註記"


def source_group(label: str) -> tuple[int, int] | None:
    match = re.fullmatch(r"(\d+)~(\d+)歲", label.strip())
    if not match:
        return None
    group = int(match.group(1)), int(match.group(2))
    return group if group in SOURCE_GROUPS else None


def read_categorical_zip(
    path: Path,
    year: int,
    *,
    dimension: str,
):
    if dimension not in {"education", "marital"}:
        raise ValueError(dimension)
    aggregate: dict[tuple[str, str, str, tuple[int, int], str], int] = defaultdict(int)
    labels: dict[str, str] = {}
    rows_read = 0
    rows_ntpc = 0
    with zipfile.ZipFile(path) as archive:
        names = [name for name in archive.namelist() if name.lower().endswith(".csv")]
        if len(names) != 1:
            raise ValueError(f"Expected one CSV in {path}; got {names}")
        with archive.open(names[0]) as raw:
            text = io.TextIOWrapper(raw, encoding="utf-8-sig", newline="")
            for row in csv.DictReader(text):
                rows_read += 1
                source_year = (row.get("statistic_yyy") or "").lstrip("\ufeff").strip()
                if source_year != str(year):
                    continue
                site = (row.get("site_id") or "").strip()
                district_code = (row.get("district_code") or "").strip()
                if not site.startswith("新北市") or not district_code.startswith("65"):
                    continue
                age_group = source_group(row.get("age") or "")
                sex = (row.get("sex") or "").strip()
                if age_group is None or sex not in SEXES:
                    continue
                rows_ntpc += 1
                if dimension == "education":
                    category_code, category_name = education_category(row.get("edu") or "")
                else:
                    category_code, category_name = marital_category(row.get("marital_status") or "")
                if category_code == "OTHER" and parse_int(row.get("population")) == 0:
                    continue
                labels[category_code] = category_name
                value = parse_int(row.get("population"))
                geo_code = district_code[:8]
                for code, name in ((geo_code, site), ("65000000", "新北市")):
                    aggregate[(code, name, sex, age_group, category_code)] += value
    return aggregate, labels, {"rows_read": rows_read, "rows_ntpc": rows_ntpc}


def fit_city_seed(
    aggregate: dict,
    categories: list[str],
    sex: str,
) -> tuple[np.ndarray, list[dict]]:
    ages = np.arange(15, 40)
    seed = np.zeros((len(ages), len(categories)), dtype=float)
    diagnostics = []
    for ci, category in enumerate(categories):
        y = np.array(
            [aggregate.get(("65000000", "新北市", sex, group, category), 0) for group in SOURCE_GROUPS],
            dtype=float,
        )
        gamma, diagnostic = select_pclm(y, ages, list(SOURCE_GROUPS))
        seed[:, ci] = gamma
        diagnostics.append({"sex": sex, "category": category, **diagnostic})
    return seed, diagnostics


def graduate_categories(
    aggregate: dict,
    labels: dict[str, str],
    population: dict,
    district_names: dict[str, str],
):
    categories = sorted(labels)
    ages = np.arange(15, 40)
    seeds = {}
    pclm_diagnostics = []
    for sex in SEXES:
        seeds[sex], diagnostics = fit_city_seed(aggregate, categories, sex)
        pclm_diagnostics.extend(diagnostics)

    estimates: dict[tuple[str, str, str, int, str], float] = {}
    baselines: dict[tuple[str, str, str, int, str], float] = {}
    ipf_diagnostics = []
    # Estimate each district first.  Model-based boundary allocations are not
    # additive when the city and districts are fitted independently, so the
    # published city result is reconciled as the sum of the 29 district
    # results below.  This preserves map drill-down accounting identities.
    for geo_code, geo_name in district_names.items():
        if geo_code == "65000000":
            continue
        for sex in SEXES:
            for start, end in SOURCE_GROUPS:
                age_indices = np.where((ages >= start) & (ages <= end))[0]
                row_targets = np.array(
                    [population.get((geo_code, geo_name, sex, int(ages[index])), 0) for index in age_indices],
                    dtype=float,
                )
                col_targets = np.array(
                    [aggregate.get((geo_code, geo_name, sex, (start, end), category), 0) for category in categories],
                    dtype=float,
                )
                if col_targets.sum() <= 0:
                    continue
                fitted, fitted_diag = ipf(seeds[sex][age_indices, :], row_targets, col_targets)
                proportional_seed = np.outer(np.maximum(row_targets, 1e-12), np.maximum(col_targets, 1e-12))
                baseline, _ = ipf(proportional_seed, row_targets, col_targets)
                for local_index, global_index in enumerate(age_indices):
                    age = int(ages[global_index])
                    for category_index, category in enumerate(categories):
                        estimates[(geo_code, geo_name, sex, age, category)] = float(
                            fitted[local_index, category_index]
                        )
                        baselines[(geo_code, geo_name, sex, age, category)] = float(
                            baseline[local_index, category_index]
                        )
                ipf_diagnostics.append(
                    {
                        "geography_code": geo_code,
                        "geography_name_zh": geo_name,
                        "sex": sex,
                        "source_age_band": f"{start}-{end}",
                        **fitted_diag,
                    }
                )

    for sex in SEXES:
        for age in ages:
            for category in categories:
                estimates[("65000000", "新北市", sex, int(age), category)] = sum(
                    estimates.get((geo_code, geo_name, sex, int(age), category), 0.0)
                    for geo_code, geo_name in district_names.items()
                    if geo_code != "65000000"
                )
                baselines[("65000000", "新北市", sex, int(age), category)] = sum(
                    baselines.get((geo_code, geo_name, sex, int(age), category), 0.0)
                    for geo_code, geo_name in district_names.items()
                    if geo_code != "65000000"
                )
    return estimates, baselines, pclm_diagnostics, ipf_diagnostics


def base_row(
    *,
    year: int,
    universe_code: str,
    universe_name: str,
    geography_code: str,
    geography_name: str,
    geography_level: str,
    geography_role: str,
    age_band: str,
    sex: str,
    source_period: str,
    period_basis: str,
):
    sex_code = {"合計": "ALL", "男": "M", "女": "F"}[sex]
    return {
        "roc_year": year,
        "gregorian_year": year + 1911,
        "source_period": source_period,
        "period_basis": period_basis,
        "universe_code": universe_code,
        "universe_name_zh": universe_name,
        "geography_code": geography_code,
        "geography_name_zh": geography_name,
        "geography_level": geography_level,
        "geography_role_zh": geography_role,
        "age_band": age_band,
        "sex_code": sex_code,
        "sex_name_zh": sex,
    }


def finalize_row(row: dict) -> dict:
    stable = "|".join(
        str(row.get(field, ""))
        for field in (
            "roc_year",
            "universe_code",
            "geography_code",
            "age_band",
            "sex_code",
            "category_dimension_code",
            "category_code",
            "metric_code",
        )
    )
    row["record_id"] = hashlib.sha256(stable.encode("utf-8")).hexdigest()[:24]
    row["value_origin_label_zh"] = ORIGIN_LABELS.get(
        row.get("value_origin_class", ""), row.get("value_origin_class", "")
    )
    row["method_name_zh"] = METHOD_LABELS.get(row.get("method_code", ""), row.get("method_code", ""))
    return row


def population_rows(
    *,
    year: int,
    population: dict,
    all_age_totals: dict,
    district_names: dict[str, str],
    source_path: Path,
):
    rows = []
    source_hash = sha256(source_path)
    for geo_code, geo_name in district_names.items():
        level = "CITY" if geo_code == "65000000" else "DISTRICT"
        for age_band, ages in TARGET_BANDS.items():
            counts = {
                sex: sum(population.get((geo_code, geo_name, sex, age), 0) for age in ages)
                for sex in SEXES
            }
            counts["合計"] = counts["男"] + counts["女"]
            for sex in ("合計", "男", "女"):
                base = base_row(
                    year=year,
                    universe_code="REGISTERED_POPULATION",
                    universe_name="戶籍人口母體",
                    geography_code=geo_code,
                    geography_name=geo_name,
                    geography_level=level,
                    geography_role="戶籍登記地",
                    age_band=age_band,
                    sex=sex,
                    source_period=f"{year}12",
                    period_basis="12月31日年末存量",
                )
                denominator = (
                    all_age_totals.get((geo_code, geo_name, "男"), 0)
                    + all_age_totals.get((geo_code, geo_name, "女"), 0)
                    if sex == "合計"
                    else all_age_totals.get((geo_code, geo_name, sex), 0)
                )
                count = counts[sex]
                common = {
                    **base,
                    "category_dimension_code": "NONE",
                    "category_dimension_name_zh": "無",
                    "category_code": "ALL",
                    "category_name_zh": "合計",
                    "value_origin_class": "OFFICIAL_ADMIN_EXACT",
                    "method_code": "M0_DIRECT_SINGLE_AGE_SUM",
                    "model_version": "G5-V3.0",
                    "uncertainty_type": "NONE_ADMIN_COUNT",
                    "source_alias": "MOI-POP1Y",
                    "source_name": "村里戶數、單一年齡人口",
                    "source_url": "https://data.gov.tw/dataset/77132",
                    "source_snapshot": str(source_path.relative_to(PROJECT)).replace("\\", "/"),
                    "source_sha256": source_hash,
                    "retrieved_at": file_observed_at(source_path),
                    "publish_status": "PUBLISH_EXACT",
                    "qa_status": "PASS",
                    "note_zh": "戶籍人口行政紀錄；不等同民間人口或受僱員工。",
                }
                rows.append(
                    finalize_row(
                        {
                            **common,
                            "metric_code": "REGISTERED_POPULATION_COUNT",
                            "metric_name_zh": "戶籍人口數",
                            "value": count,
                            "unit": "人",
                            "numerator": count,
                            "denominator": "",
                            "formula_zh": f"加總{age_band}歲各單一年齡戶籍人口",
                            "uncertainty_low": count,
                            "uncertainty_high": count,
                        }
                    )
                )
                rows.append(
                    finalize_row(
                        {
                            **common,
                            "metric_code": "REGISTERED_POPULATION_SHARE_PCT",
                            "metric_name_zh": "占該地區戶籍人口比率",
                            "value": normalize_float(count / denominator * 100 if denominator else None),
                            "unit": "%",
                            "numerator": count,
                            "denominator": denominator,
                            "formula_zh": "目標年齡戶籍人口數÷該地區同一性別全年齡戶籍人口數×100%",
                            "uncertainty_low": normalize_float(count / denominator * 100 if denominator else None),
                            "uncertainty_high": normalize_float(count / denominator * 100 if denominator else None),
                        }
                    )
                )
                if sex in SEXES:
                    age_total = counts["合計"]
                    rows.append(
                        finalize_row(
                            {
                                **common,
                                "metric_code": "SEX_SHARE_WITHIN_AGE_BAND_PCT",
                                "metric_name_zh": "目標年齡性別占比",
                                "value": normalize_float(count / age_total * 100 if age_total else None),
                                "unit": "%",
                                "numerator": count,
                                "denominator": age_total,
                                "formula_zh": "該性別目標年齡人口數÷目標年齡男女合計人口數×100%",
                                "uncertainty_low": normalize_float(count / age_total * 100 if age_total else None),
                                "uncertainty_high": normalize_float(count / age_total * 100 if age_total else None),
                            }
                        )
                    )
    return rows


def categorical_rows(
    *,
    year: int,
    dimension: str,
    labels: dict[str, str],
    estimates: dict,
    baselines: dict,
    district_names: dict[str, str],
    source_path: Path,
    population_path: Path,
):
    rows = []
    source_hash = sha256(source_path)
    pop_hash = sha256(population_path)
    if dimension == "education":
        dim_code, dim_name = "EDUCATION", "教育程度"
        count_metric, count_name = "EDUCATION_POPULATION_COUNT", "教育程度人口數"
        share_metric, share_name = "EDUCATION_SHARE_PCT", "教育程度占比"
        source_alias = "MOI-EDU5Y+MOI-POP1Y"
        source_name = "15歲以上現住人口按性別、年齡、婚姻及教育程度分＋單一年齡人口"
        source_url = "https://data.gov.tw/dataset/117988"
    else:
        dim_code, dim_name = "MARITAL_STATUS", "婚姻狀態"
        count_metric, count_name = "MARITAL_STATUS_POPULATION_COUNT", "婚姻狀態人口數"
        share_metric, share_name = "MARITAL_STATUS_SHARE_PCT", "婚姻狀態占比"
        source_alias = "MOI-MARITAL5Y+MOI-POP1Y"
        source_name = "15歲以上現住人口按性別、年齡及婚姻狀況分＋單一年齡人口"
        source_url = "https://data.gov.tw/dataset/117986"

    for geo_code, geo_name in district_names.items():
        level = "CITY" if geo_code == "65000000" else "DISTRICT"
        for age_band, ages in TARGET_BANDS.items():
            for sex in ("合計", "男", "女"):
                sexes = SEXES if sex == "合計" else (sex,)
                category_values = {
                    category: sum(
                        estimates.get((geo_code, geo_name, sx, age, category), 0.0)
                        for sx in sexes
                        for age in ages
                    )
                    for category in labels
                }
                category_baselines = {
                    category: sum(
                        baselines.get((geo_code, geo_name, sx, age, category), 0.0)
                        for sx in sexes
                        for age in ages
                    )
                    for category in labels
                }
                denominator = sum(category_values.values())
                for category, category_name in sorted(labels.items()):
                    value = category_values[category]
                    baseline = category_baselines[category]
                    exact = age_band == "25-29"
                    origin = "OFFICIAL_ADMIN_EXACT" if exact else "MODEL_ESTIMATE_FROM_OFFICIAL_MARGINS"
                    method = "M0_DIRECT_5Y_AGGREGATION" if exact else "PCLM_CITY_SEED_IPF_GEO"
                    low = value if exact else min(value, baseline)
                    high = value if exact else max(value, baseline)
                    base = base_row(
                        year=year,
                        universe_code="REGISTERED_POPULATION",
                        universe_name="戶籍人口母體",
                        geography_code=geo_code,
                        geography_name=geo_name,
                        geography_level=level,
                        geography_role="戶籍登記地",
                        age_band=age_band,
                        sex=sex,
                        source_period=str(year),
                        period_basis="12月31日年末存量",
                    )
                    common = {
                        **base,
                        "category_dimension_code": dim_code,
                        "category_dimension_name_zh": dim_name,
                        "category_code": category,
                        "category_name_zh": category_name,
                        "value_origin_class": origin,
                        "method_code": method,
                        "model_version": "G5-V3.0-PCLM-IPF-1",
                        "uncertainty_low": normalize_float(low, 3),
                        "uncertainty_high": normalize_float(high, 3),
                        "uncertainty_type": "NONE_ADMIN_COUNT" if exact else "METHOD_SENSITIVITY_ENVELOPE_NOT_CI",
                        "source_alias": source_alias,
                        "source_name": source_name,
                        "source_url": source_url,
                        "source_snapshot": str(source_path.relative_to(PROJECT)).replace("\\", "/"),
                        "source_sha256": f"{source_hash}+{pop_hash}",
                        "retrieved_at": file_observed_at(source_path),
                        "publish_status": "PUBLISH_EXACT" if exact else "PUBLISH_MODEL_WITH_METHOD_LABEL_AFTER_REVIEW",
                        "qa_status": "PASS",
                        "note_zh": (
                            "25–29歲為官方原生五歲帶；其餘含18、19或35歲邊界，採PCLM種子與IPF邊際校準。"
                        ),
                    }
                    rows.append(
                        finalize_row(
                            {
                                **common,
                                "metric_code": count_metric,
                                "metric_name_zh": count_name,
                                "value": normalize_float(value, 3),
                                "unit": "人",
                                "numerator": normalize_float(value, 3),
                                "denominator": "",
                                "formula_zh": "加總目標單一年齡之類別估計人口；IPF保持官方五歲組類別總數與單歲人口邊際",
                            }
                        )
                    )
                    rows.append(
                        finalize_row(
                            {
                                **common,
                                "metric_code": share_metric,
                                "metric_name_zh": share_name,
                                "value": normalize_float(value / denominator * 100 if denominator else None),
                                "unit": "%",
                                "numerator": normalize_float(value, 3),
                                "denominator": normalize_float(denominator, 3),
                                "formula_zh": "該類別目標年齡人口數÷同地區同性別目標年齡全部類別人口數×100%",
                                "uncertainty_low": normalize_float(low / denominator * 100 if denominator else None),
                                "uncertainty_high": normalize_float(high / denominator * 100 if denominator else None),
                            }
                        )
                    )
    return rows


def build_registered_csv(package: Path, years: tuple[int, ...]) -> tuple[Path, dict]:
    output = package / "09_發布資料包" / "01_戶籍人口母體_長格式.csv"
    pop_paths = find_population_snapshots(years)
    all_rows = []
    diagnostics = {"years": {}, "sources": []}
    for year in years:
        pop_path = pop_paths[year]
        population, all_age_totals, district_names, pop_row_count = read_population(pop_path, f"{year}12")
        all_rows.extend(
            population_rows(
                year=year,
                population=population,
                all_age_totals=all_age_totals,
                district_names=district_names,
                source_path=pop_path,
            )
        )

        education_path = (
            PROJECT / "data" / "raw" / "MOI-EDU-SINGLEAGE" / str(year) / f"opendata{year}Y051.zip"
        )
        marital_path = (
            PROJECT / "data" / "raw" / "MOI-MARITAL5Y" / str(year) / f"opendata{year}Y031.zip"
        )
        education_agg, education_labels, education_ingest = read_categorical_zip(
            education_path, year, dimension="education"
        )
        education_est, education_base, education_pclm, education_ipf = graduate_categories(
            education_agg, education_labels, population, district_names
        )
        all_rows.extend(
            categorical_rows(
                year=year,
                dimension="education",
                labels=education_labels,
                estimates=education_est,
                baselines=education_base,
                district_names=district_names,
                source_path=education_path,
                population_path=pop_path,
            )
        )

        marital_agg, marital_labels, marital_ingest = read_categorical_zip(
            marital_path, year, dimension="marital"
        )
        marital_est, marital_base, marital_pclm, marital_ipf = graduate_categories(
            marital_agg, marital_labels, population, district_names
        )
        all_rows.extend(
            categorical_rows(
                year=year,
                dimension="marital",
                labels=marital_labels,
                estimates=marital_est,
                baselines=marital_base,
                district_names=district_names,
                source_path=marital_path,
                population_path=pop_path,
            )
        )
        diagnostics["years"][str(year)] = {
            "population_api_rows": pop_row_count,
            "district_count_including_city": len(district_names),
            "education_ingest": education_ingest,
            "education_categories": education_labels,
            "education_pclm_converged": all(bool(item["converged"]) for item in education_pclm),
            "education_max_ipf_row_error": max(item["row_error"] for item in education_ipf),
            "education_max_ipf_col_error": max(item["col_error"] for item in education_ipf),
            "marital_ingest": marital_ingest,
            "marital_categories": marital_labels,
            "marital_pclm_converged": all(bool(item["converged"]) for item in marital_pclm),
            "marital_max_ipf_row_error": max(item["row_error"] for item in marital_ipf),
            "marital_max_ipf_col_error": max(item["col_error"] for item in marital_ipf),
        }
        for source_path, alias in (
            (pop_path, "MOI-POP1Y"),
            (education_path, "MOI-EDU5Y"),
            (marital_path, "MOI-MARITAL5Y"),
        ):
            diagnostics["sources"].append(
                {
                    "year": year,
                    "source_alias": alias,
                    "path": str(source_path.relative_to(PROJECT)).replace("\\", "/"),
                    "sha256": sha256(source_path),
                    "bytes": source_path.stat().st_size,
                }
            )
    row_count = write_csv(output, all_rows)
    diagnostics["published_rows"] = row_count
    return output, diagnostics


def build_labor_csv(package: Path, years: tuple[int, ...]) -> tuple[Path, dict]:
    output = package / "09_發布資料包" / "02_民間人口與勞動市場母體_長格式.csv"
    grouped = defaultdict(dict)
    with LABOR_RESULTS.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            year = int(row["roc_year"])
            if year not in years:
                continue
            grouped[(year, row["age_group"])][row["method"]] = row

    metric_specs = (
        ("civilian_population_thousand", "CIVILIAN_POPULATION", "民間人口", "千人", "P", ""),
        ("labor_force_thousand", "LABOR_FORCE", "勞動力人口", "千人", "LF", ""),
        ("employed_thousand", "EMPLOYED", "就業人口", "千人", "E", ""),
        ("unemployed_thousand", "UNEMPLOYED", "失業人口", "千人", "U", ""),
        ("not_in_labor_force_thousand", "NOT_IN_LABOR_FORCE", "非勞動力人口", "千人", "NLF", "P−LF"),
        ("unemployment_rate_pct", "UNEMPLOYMENT_RATE", "失業率", "%", "UR", "U÷LF×100%"),
        ("labor_force_participation_rate_pct", "LABOR_FORCE_PARTICIPATION_RATE", "勞動力參與率", "%", "LFPR", "LF÷P×100%"),
        ("employment_to_population_rate_pct", "EMPLOYMENT_TO_POPULATION_RATE", "就業人口比率", "%", "EPR", "E÷P×100%"),
        ("employment_share_of_labor_force_pct", "EMPLOYMENT_SHARE_OF_LABOR_FORCE", "勞動力中就業占比", "%", "ESLF", "E÷LF×100%"),
    )
    rows = []
    source_diagnostics = []
    for (year, age_band), methods in sorted(grouped.items()):
        primary = methods.get("PCLM")
        alternate = methods.get("SPRAGUE_NN_CALIBRATED")
        if not primary:
            raise ValueError(f"Missing PCLM labor result: {year} {age_band}")
        exact = age_band == "25-29"
        for field, metric_code, metric_name, unit, symbol, formula in metric_specs:
            alias, source_name, table_numbers = LABOR_LINEAGE[metric_code]
            source_paths = [LABOR_ANNUAL_RAW_ROOT / str(year) / f"table{number}.xlsx" for number in table_numbers]
            if year == 114:
                if "|QA:" in alias:
                    primary_alias, qa_alias = alias.split("|QA:", 1)
                    alias = f"{primary_alias}+AUX:DGBAS-NTPC-H2-T41+T42|QA:{qa_alias}"
                else:
                    alias += "+AUX:DGBAS-NTPC-H2-T41+T42"
                source_name += "；114年另以同年下半年表41／42之E、U、NLF年齡占比拆分15–24歲"
                source_paths.append(LABOR_H2_RECONCILIATION)
            source_paths.append(LABOR_RESULTS)
            missing_sources = [path for path in source_paths if not path.exists()]
            if missing_sources:
                raise FileNotFoundError(f"Missing labor lineage source(s): {missing_sources}")
            source_snapshot = "|".join(
                str(path.relative_to(PROJECT)).replace("\\", "/") for path in source_paths
            )
            source_hash = combined_sha256(source_paths)
            retrieved_at = max(file_observed_at(path) for path in source_paths)
            value = float(primary[field])
            alternate_value = float(alternate[field]) if alternate else value
            base = base_row(
                year=year,
                universe_code="CIVILIAN_LABOR_MARKET",
                universe_name="民間人口及勞動市場母體",
                geography_code="65000000",
                geography_name="新北市",
                geography_level="CITY",
                geography_role="人力資源調查地區別",
                age_band=age_band,
                sex="合計",
                source_period=str(year),
                period_basis="全年12個月平均",
            )
            if unit == "%":
                if metric_code == "UNEMPLOYMENT_RATE":
                    numerator, denominator = primary["unemployed_thousand"], primary["labor_force_thousand"]
                elif metric_code == "LABOR_FORCE_PARTICIPATION_RATE":
                    numerator, denominator = primary["labor_force_thousand"], primary["civilian_population_thousand"]
                elif metric_code == "EMPLOYMENT_TO_POPULATION_RATE":
                    numerator, denominator = primary["employed_thousand"], primary["civilian_population_thousand"]
                else:
                    numerator, denominator = primary["employed_thousand"], primary["labor_force_thousand"]
            else:
                numerator, denominator = value, ""
            rows.append(
                finalize_row(
                    {
                        **base,
                        "category_dimension_code": "NONE",
                        "category_dimension_name_zh": "無",
                        "category_code": "ALL",
                        "category_name_zh": "合計",
                        "metric_code": metric_code,
                        "metric_name_zh": metric_name,
                        "value": normalize_float(value),
                        "unit": unit,
                        "numerator": normalize_float(numerator),
                        "denominator": normalize_float(denominator) if denominator != "" else "",
                        "formula_zh": formula or symbol,
                        "value_origin_class": (
                            "OFFICIAL_SURVEY_ESTIMATE_AGE_BAND_EXACT"
                            if exact
                            else "MODEL_ESTIMATE_FROM_OFFICIAL_SURVEY_ESTIMATES"
                        ),
                        "method_code": "M0_OFFICIAL_AGE_BAND" if exact else "PCLM",
                        "model_version": "G5-V3.0-LABOR-PCLM-1",
                        "uncertainty_low": normalize_float(min(value, alternate_value)),
                        "uncertainty_high": normalize_float(max(value, alternate_value)),
                        "uncertainty_type": "SOURCE_ROUNDING_ONLY" if exact else "PCLM_SPRAGUE_METHOD_ENVELOPE_NOT_CI",
                        "source_alias": alias,
                        "source_name": source_name,
                        "source_url": LABOR_ANNUAL_URLS[year],
                        "source_snapshot": source_snapshot,
                        "source_sha256": source_hash,
                        "retrieved_at": retrieved_at,
                        "publish_status": "PUBLISH_SOURCE_NATIVE" if exact else "PUBLISH_MODEL_WITH_METHOD_LABEL_AFTER_REVIEW",
                        "qa_status": "PASS",
                        "note_zh": (
                            "P、E、U為官方抽樣調查估計；LF=E+U，NLF=P−LF；率以相符母體分子與分母重算。"
                            + (f" 114年15–24歲先使用同年下半年輔助占比拆分；輔助來源：{LABOR_H2_URL}。" if year == 114 else "")
                        ),
                    }
                )
            )
            source_diagnostics.append(
                {
                    "year": year,
                    "metric_code": metric_code,
                    "source_alias": alias,
                    "source_url": LABOR_ANNUAL_URLS[year],
                    "source_paths": [str(path.relative_to(PROJECT)).replace("\\", "/") for path in source_paths],
                    "source_sha256": source_hash,
                    "h2_auxiliary_url": LABOR_H2_URL if year == 114 else None,
                }
            )
    count = write_csv(output, rows)
    diagnostics = {
        "derived_source": str(LABOR_RESULTS.relative_to(PROJECT)).replace("\\", "/"),
        "derived_source_sha256": sha256(LABOR_RESULTS),
        "official_annual_urls": {str(year): LABOR_ANNUAL_URLS[year] for year in years},
        "source_lineage": source_diagnostics,
        "published_rows": count,
        "years": sorted({int(row["roc_year"]) for row in rows}),
        "identities": ["LF=E+U", "NLF=P-LF", "UR+ESLF=100%"],
    }
    return output, diagnostics


def workbook_year(sheet_name: str) -> int | None:
    match = re.search(r"(10\d|11\d)", sheet_name)
    return int(match.group(1)) if match else None


def combined_sha256(paths: Iterable[Path]) -> str:
    digest = hashlib.sha256()
    for path in sorted(paths, key=lambda item: str(item)):
        digest.update(path.name.encode("utf-8"))
        digest.update(sha256(path).encode("ascii"))
    return digest.hexdigest()


def read_bli_ntpc_age_counts(path: Path) -> tuple[dict[tuple[int, int], float], float]:
    column_map = {
        (15, 19): "年齡15-19歲人數",
        (20, 24): "年齡20-24歲人數",
        (25, 29): "年齡25-29歲人數",
        (30, 34): "年齡30-34歲人數",
        (35, 39): "年齡35-39歲人數",
    }
    grouped = {group: 0.0 for group in SOURCE_GROUPS}
    under_15 = 0.0
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if row.get("地區別") != "新北市":
                continue
            under_15 += float(row.get("年齡未滿15歲人數") or 0)
            for group, column in column_map.items():
                grouped[group] += float(row.get(column) or 0)
    if sum(grouped.values()) <= 0:
        raise ValueError(f"No New Taipei age counts in {path}")
    return grouped, under_15


def read_pension_age_wages(path: Path) -> tuple[dict[tuple[int, int], float], float]:
    grouped: dict[tuple[int, int], float] = {}
    under_15 = 0.0
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            label = (row.get("年齡級距") or "").strip()
            value = float(row.get("平均提繳工資金額") or 0)
            if "未滿15" in label:
                under_15 = value
                continue
            bounds = [int(item) for item in re.findall(r"\d+", label)]
            if len(bounds) == 2 and tuple(bounds) in SOURCE_GROUPS:
                grouped[tuple(bounds)] = value
    missing = [group for group in SOURCE_GROUPS if group not in grouped]
    if missing:
        raise ValueError(f"Missing pension wage age groups in {path}: {missing}")
    return grouped, under_15


def wage_auxiliary_profiles(
    count_path: Path,
    wage_path: Path,
) -> tuple[dict[int, float], dict[int, float], dict[int, float], dict[int, float], dict]:
    grouped_counts, under_15_count = read_bli_ntpc_age_counts(count_path)
    grouped_wages, under_15_wage = read_pension_age_wages(wage_path)
    ages = np.arange(15, 40)
    count_margins = np.array([grouped_counts[group] for group in SOURCE_GROUPS], dtype=float)
    wage_mass_margins = np.array(
        [grouped_counts[group] * grouped_wages[group] for group in SOURCE_GROUPS],
        dtype=float,
    )
    pclm_counts, count_diag = select_pclm(count_margins, ages, list(SOURCE_GROUPS))
    pclm_mass, mass_diag = select_pclm(wage_mass_margins, ages, list(SOURCE_GROUPS))

    # PCLM is a smoothing model and may miss a published five-year margin by a
    # small amount.  Recalibrate within every source group so the auxiliary
    # single-age profile reproduces both official count and wage-mass margins.
    for index, (start, end) in enumerate(SOURCE_GROUPS):
        mask = (ages >= start) & (ages <= end)
        count_sum = float(pclm_counts[mask].sum())
        mass_sum = float(pclm_mass[mask].sum())
        if count_sum <= 0 or mass_sum <= 0:
            raise ValueError(f"Non-positive PCLM group before calibration: {(start, end)}")
        pclm_counts[mask] *= count_margins[index] / count_sum
        pclm_mass[mask] *= wage_mass_margins[index] / mass_sum

    uniform_counts: dict[int, float] = {}
    uniform_mass: dict[int, float] = {}
    for group in SOURCE_GROUPS:
        for age in range(group[0], group[1] + 1):
            uniform_counts[age] = grouped_counts[group] / 5.0
            uniform_mass[age] = grouped_counts[group] * grouped_wages[group] / 5.0

    counts = {int(age): float(pclm_counts[index]) for index, age in enumerate(ages)}
    mass = {int(age): float(pclm_mass[index]) for index, age in enumerate(ages)}
    counts[-1] = under_15_count
    mass[-1] = under_15_count * under_15_wage
    uniform_counts[-1] = under_15_count
    uniform_mass[-1] = under_15_count * under_15_wage
    diagnostics = {
        "count_pclm": count_diag,
        "wage_mass_pclm": mass_diag,
        "post_calibration_max_count_margin_error": max(
            abs(sum(counts[age] for age in range(start, end + 1)) - grouped_counts[(start, end)])
            for start, end in SOURCE_GROUPS
        ),
        "post_calibration_max_wage_mass_margin_error": max(
            abs(
                sum(mass[age] for age in range(start, end + 1))
                - grouped_counts[(start, end)] * grouped_wages[(start, end)]
            )
            for start, end in SOURCE_GROUPS
        ),
        "under_15_count": under_15_count,
        "under_15_auxiliary_monthly_wage": under_15_wage,
    }
    return counts, mass, uniform_counts, uniform_mass, diagnostics


def profile_mean(counts: dict[int, float], mass: dict[int, float], ages: Iterable[int]) -> float:
    selected = list(ages)
    denominator = sum(counts.get(age, 0.0) for age in selected)
    if denominator <= 0:
        raise ValueError(f"Non-positive auxiliary denominator for ages {selected}")
    return sum(mass.get(age, 0.0) for age in selected) / denominator


def calibrated_wage_means(
    official_means: dict[str, float],
    counts: dict[int, float],
    mass: dict[int, float],
) -> dict[str, float]:
    under_25_ages = [-1, *range(15, 25)]
    target_18_24 = range(18, 25)
    broad_30_39 = range(30, 40)
    target_30_35 = range(30, 36)
    estimates = {
        "18-24": official_means["UNDER_25"]
        * profile_mean(counts, mass, target_18_24)
        / profile_mean(counts, mass, under_25_ages),
        "25-29": official_means["25-29"],
        "30-35": official_means["30-39"]
        * profile_mean(counts, mass, target_30_35)
        / profile_mean(counts, mass, broad_30_39),
    }
    target_counts = {
        "18-24": sum(counts.get(age, 0.0) for age in target_18_24),
        "25-29": sum(counts.get(age, 0.0) for age in range(25, 30)),
        "30-35": sum(counts.get(age, 0.0) for age in target_30_35),
    }
    total_weight = sum(target_counts.values())
    estimates["18-35"] = sum(
        estimates[band] * target_counts[band] for band in ("18-24", "25-29", "30-35")
    ) / total_weight
    return estimates


def wage_target_counts(counts: dict[int, float]) -> dict[str, float]:
    """Return the worker-count weights used for target-age wage aggregation."""
    target_ages = {
        "18-24": range(18, 25),
        "25-29": range(25, 30),
        "30-35": range(30, 36),
    }
    return {
        band: sum(counts.get(age, 0.0) for age in ages)
        for band, ages in target_ages.items()
    }


def lognormal_cdf(value: float, mean: float, median: float) -> float:
    """CDF of a lognormal distribution parameterized by its mean and median."""
    if value <= 0:
        return 0.0
    if mean <= 0 or median <= 0 or mean < median:
        raise ValueError(f"Invalid lognormal anchors: mean={mean}, median={median}")
    sigma_squared = 2.0 * math.log(mean / median)
    if sigma_squared <= 1e-12:
        return 0.0 if value < median else 1.0
    sigma = math.sqrt(sigma_squared)
    z_score = (math.log(value) - math.log(median)) / sigma
    return 0.5 * (1.0 + math.erf(z_score / math.sqrt(2.0)))


def lognormal_mixture_median(
    means: dict[str, float],
    medians: dict[str, float],
    weights: dict[str, float],
) -> float:
    """Solve F(x)=0.5 for a weighted mixture of subgroup lognormal CDFs."""
    bands = ("18-24", "25-29", "30-35")
    total_weight = sum(weights[band] for band in bands)
    if total_weight <= 0:
        raise ValueError("Non-positive worker weight for 18-35 wage median")

    def mixture_cdf(value: float) -> float:
        return sum(
            weights[band] * lognormal_cdf(value, means[band], medians[band])
            for band in bands
        ) / total_weight

    lower = min(medians[band] for band in bands) * 0.1
    upper = max(means[band] for band in bands) * 4.0
    while mixture_cdf(upper) < 0.5:
        upper *= 2.0
    for _ in range(100):
        midpoint = (lower + upper) / 2.0
        if mixture_cdf(midpoint) < 0.5:
            lower = midpoint
        else:
            upper = midpoint
    return (lower + upper) / 2.0


def calibrated_wage_medians(
    official_means: dict[str, float],
    official_medians: dict[str, float],
    target_means: dict[str, float],
    counts: dict[int, float],
) -> dict[str, float]:
    """Estimate target medians without linearly averaging grouped medians.

    For a target subgroup contained in a published broad age band, preserve the
    broad band's median/mean ratio and apply it to the already calibrated target
    mean.  The 18-35 median is then the 50th percentile of a worker-weighted
    mixture of the three subgroup lognormal distributions.
    """
    estimates = {
        "18-24": target_means["18-24"]
        * official_medians["UNDER_25"]
        / official_means["UNDER_25"],
        "25-29": official_medians["25-29"],
        "30-35": target_means["30-35"]
        * official_medians["30-39"]
        / official_means["30-39"],
    }
    estimates["18-35"] = lognormal_mixture_median(
        target_means,
        estimates,
        wage_target_counts(counts),
    )
    return estimates


def held_out_wage_median_validation(
    official_means: dict[str, float],
    official_medians: dict[str, float],
) -> dict[str, float]:
    """Predict the native 25-29 median without using it as a model input."""
    adjacent_shape_ratio = math.sqrt(
        (official_medians["UNDER_25"] / official_means["UNDER_25"])
        * (official_medians["30-39"] / official_means["30-39"])
    )
    predicted = official_means["25-29"] * adjacent_shape_ratio
    observed = official_medians["25-29"]
    return {
        "predicted_median_25_29": predicted,
        "official_median_25_29": observed,
        "absolute_percentage_error_pct": abs(predicted - observed) / observed * 100.0,
        "validation_role": "HELD_OUT_SANITY_CHECK_NOT_PROOF",
    }


def build_wage_csv(package: Path, years: tuple[int, ...]) -> tuple[Path, dict]:
    output = package / "09_發布資料包" / "03_受僱員工薪資母體_長格式.csv"
    workbook = openpyxl.load_workbook(WAGE_WORKBOOK, read_only=True, data_only=True)
    rows = []
    available_years = []
    year_diagnostics = {}
    for sheet_name in workbook.sheetnames:
        year = workbook_year(sheet_name)
        if year is None or year not in years:
            continue
        available_years.append(year)
        sheet = workbook[sheet_name]
        # Official Table 6 order is total followed by New Taipei City at row 9.
        # Mean: under 25=col 4, 25-29=col 5, 30-39=col 6.
        # Median: under 25=col 12, 25-29=col 13, 30-39=col 14.
        official_means = {
            "UNDER_25": sheet.cell(9, 4).value,
            "25-29": sheet.cell(9, 5).value,
            "30-39": sheet.cell(9, 6).value,
        }
        official_medians = {
            "UNDER_25": sheet.cell(9, 12).value,
            "25-29": sheet.cell(9, 13).value,
            "30-39": sheet.cell(9, 14).value,
        }
        if not all(isinstance(value, (int, float)) for value in official_means.values()):
            raise ValueError(f"Unexpected wage mean values at {sheet_name}: {official_means!r}")
        if not all(isinstance(value, (int, float)) for value in official_medians.values()):
            raise ValueError(f"Unexpected wage median values at {sheet_name}: {official_medians!r}")

        count_path = WAGE_AUX_DIR / f"BLI-AGE-REGION-SEX-{year}.csv"
        pension_path = WAGE_AUX_DIR / f"BLI-PENSION-AGE-WAGE-{year}.csv"
        auxiliary_paths = [WAGE_WORKBOOK, count_path, pension_path]
        if any(not path.exists() for path in auxiliary_paths):
            missing = [str(path) for path in auxiliary_paths if not path.exists()]
            raise FileNotFoundError(f"Missing wage auxiliary sources for {year}: {missing}")

        pclm_counts, pclm_mass, uniform_counts, uniform_mass, profile_diag = wage_auxiliary_profiles(
            count_path, pension_path
        )
        estimates = calibrated_wage_means(official_means, pclm_counts, pclm_mass)
        baseline_estimates = calibrated_wage_means(official_means, uniform_counts, uniform_mass)
        median_estimates = calibrated_wage_medians(
            official_means, official_medians, estimates, pclm_counts
        )
        baseline_median_estimates = calibrated_wage_medians(
            official_means, official_medians, baseline_estimates, uniform_counts
        )
        median_validation = held_out_wage_median_validation(official_means, official_medians)
        source_hash = combined_sha256(auxiliary_paths)
        source_snapshot = " + ".join(
            str(path.relative_to(PROJECT)).replace("\\", "/") for path in auxiliary_paths
        )
        retrieved_at = max(file_observed_at(path) for path in auxiliary_paths)

        for age_band in ("18-24", "25-29", "30-35", "18-35"):
            raw_value = estimates[age_band]
            exact = age_band == "25-29"
            base = base_row(
                year=year,
                universe_code="EMPLOYEE_WAGE",
                universe_name="受僱員工薪資母體",
                geography_code="65000000",
                geography_name="新北市",
                geography_level="CITY",
                geography_role="工作場所所在地",
                age_band=age_band,
                sex="合計",
                source_period=str(year),
                period_basis="全年統計",
            )
            rows.append(
                finalize_row(
                    {
                        **base,
                        "category_dimension_code": "STATISTIC",
                        "category_dimension_name_zh": "統計量",
                        "category_code": "MEAN",
                        "category_name_zh": "平均數",
                        "metric_code": "ANNUAL_TOTAL_SALARY_MEAN",
                        "metric_name_zh": "全年總薪資平均數",
                        "value": normalize_float(raw_value, 1),
                        "unit": "萬元/年",
                        "numerator": "",
                        "denominator": "",
                        "formula_zh": (
                            "官方表6之新北市25–29歲原生平均數"
                            if exact
                            else "官方表6相鄰原生年齡帶平均數×（目標年齡輔助平均提繳工資÷原生年齡帶輔助平均提繳工資）；18–35再依新北市勞保年齡人數加權"
                        ),
                        "value_origin_class": (
                            "OFFICIAL_SURVEY_STATISTIC"
                            if exact
                            else "MODEL_ESTIMATE_CALIBRATED_OFFICIAL_WAGE_AUX"
                        ),
                        "method_code": "M0_OFFICIAL_AGE_BAND" if exact else "WAGE_PCLM_RATIO_CALIBRATION",
                        "model_version": "G5-V3.2-WAGE-PCLM-CAL-1",
                        "uncertainty_low": (
                            ""
                            if exact
                            else normalize_float(min(raw_value, baseline_estimates[age_band]), 1)
                        ),
                        "uncertainty_high": (
                            ""
                            if exact
                            else normalize_float(max(raw_value, baseline_estimates[age_band]), 1)
                        ),
                        "uncertainty_type": (
                            "OFFICIAL_TABLE_NO_CI_PUBLISHED"
                            if exact
                            else "PCLM_VS_UNIFORM_AUXILIARY_METHOD_ENVELOPE_NOT_CI"
                        ),
                        "source_alias": (
                            "STAT-WAGE-LOC"
                            if exact
                            else "STAT-WAGE-LOC+BLI-AGE-REGION-SEX+BLI-PENSION-AGE-WAGE"
                        ),
                        "source_name": (
                            "表6 工業及服務業全年總薪資統計－按工作場所所在地區及年齡別分"
                            if exact
                            else "主計總處表6薪資＋勞保新北市年齡人數＋勞退年齡平均提繳工資"
                        ),
                        "source_url": "https://www.stat.gov.tw/News_Content.aspx?n=4580&s=232642",
                        "source_snapshot": (
                            str(WAGE_WORKBOOK.relative_to(PROJECT)).replace("\\", "/")
                            if exact
                            else source_snapshot
                        ),
                        "source_sha256": sha256(WAGE_WORKBOOK) if exact else source_hash,
                        "retrieved_at": file_observed_at(WAGE_WORKBOOK) if exact else retrieved_at,
                        "publish_status": (
                            "PUBLISH_SOURCE_NATIVE"
                            if exact
                            else "PUBLISH_MODEL_WITH_METHOD_LABEL_AFTER_REVIEW"
                        ),
                        "qa_status": "PASS_WITH_SCOPE_NOTE",
                        "note_zh": (
                            "母體為工作場所位於新北市之本國籍全時受僱員工；不是戶籍人口或全部青年。"
                            if exact
                            else "估計值以官方表6為薪資水準錨點；勞保人數與勞退提繳工資只作年齡輪廓及權數，不是實領薪資。不得標示為官方精確值。"
                        ),
                    }
                )
            )

        # Keep the release table rectangular: 4 years x 4 age bands x 2
        # statistics = 32 rows.  The native 25--29 median remains an official
        # survey statistic; the other target bands are explicitly labeled
        # distribution-model estimates and retain a sensitivity envelope.
        for age_band in ("18-24", "25-29", "30-35", "18-35"):
            native_median = age_band == "25-29"
            median_base = base_row(
                year=year,
                universe_code="EMPLOYEE_WAGE",
                universe_name="受僱員工薪資母體",
                geography_code="65000000",
                geography_name="新北市",
                geography_level="CITY",
                geography_role="工作場所所在地",
                age_band=age_band,
                sex="合計",
                source_period=str(year),
                period_basis="全年統計",
            )
            rows.append(
                finalize_row(
                    {
                        **median_base,
                        "category_dimension_code": "STATISTIC",
                        "category_dimension_name_zh": "統計量",
                        "category_code": "MEDIAN",
                        "category_name_zh": "中位數",
                        "metric_code": "ANNUAL_TOTAL_SALARY_MEDIAN",
                        "metric_name_zh": "全年總薪資中位數",
                        "value": normalize_float(median_estimates[age_band], 1),
                        "unit": "萬元/年",
                        "numerator": "",
                        "denominator": "",
                        "formula_zh": (
                            "官方表6之新北市、25–29歲原生中位數"
                            if native_median
                            else (
                                "先以目標年齡平均數×（官方相鄰寬年齡帶中位數÷官方相鄰寬年齡帶平均數）校準子群中位數；"
                                "18–35歲再以各子群平均數與中位數校準對數常態分布，依PCLM勞保年齡人數權重混合三組分布，數值解F(x)=0.5"
                            )
                        ),
                        "value_origin_class": (
                            "OFFICIAL_SURVEY_STATISTIC"
                            if native_median
                            else "MODEL_ESTIMATE_CALIBRATED_OFFICIAL_WAGE_DISTRIBUTION"
                        ),
                        "method_code": (
                            "M0_OFFICIAL_AGE_BAND"
                            if native_median
                            else "WAGE_LOGNORMAL_SHAPE_TRANSFER_PCLM_MIXTURE"
                        ),
                        "model_version": "G5-V3.3-WAGE-MEDIAN-CAL-1",
                        "uncertainty_low": (
                            ""
                            if native_median
                            else normalize_float(
                                min(median_estimates[age_band], baseline_median_estimates[age_band]), 1
                            )
                        ),
                        "uncertainty_high": (
                            ""
                            if native_median
                            else normalize_float(
                                max(median_estimates[age_band], baseline_median_estimates[age_band]), 1
                            )
                        ),
                        "uncertainty_type": (
                            "OFFICIAL_TABLE_NO_CI_PUBLISHED"
                            if native_median
                            else "PCLM_VS_UNIFORM_AUXILIARY_METHOD_ENVELOPE_NOT_CI"
                        ),
                        "source_alias": (
                            "STAT-WAGE-LOC"
                            if native_median
                            else "STAT-WAGE-LOC+BLI-AGE-REGION-SEX+BLI-PENSION-AGE-WAGE"
                        ),
                        "source_name": (
                            "表6 工業及服務業全年總薪資統計－按工作場所所在地區及年齡別分"
                            if native_median
                            else "主計總處表6薪資＋勞保新北市年齡人數＋勞退年齡平均提繳工資"
                        ),
                        "source_url": "https://www.stat.gov.tw/News_Content.aspx?n=4580&s=232642",
                        "source_snapshot": (
                            str(WAGE_WORKBOOK.relative_to(PROJECT)).replace("\\", "/")
                            if native_median
                            else source_snapshot
                        ),
                        "source_sha256": sha256(WAGE_WORKBOOK) if native_median else source_hash,
                        "retrieved_at": file_observed_at(WAGE_WORKBOOK) if native_median else retrieved_at,
                        "publish_status": (
                            "PUBLISH_SOURCE_NATIVE"
                            if native_median
                            else "PUBLISH_MODEL_WITH_METHOD_LABEL_AFTER_REVIEW"
                        ),
                        "qa_status": "PASS_WITH_SCOPE_NOTE",
                        "note_zh": (
                            "官方原生25–29歲中位數；母體為工作場所位於新北市之本國籍全時受僱員工。"
                            if native_median
                            else "模型以官方相鄰寬年齡帶平均數與中位數為分布錨點，並以PCLM勞保年齡人數作18–35混合權重；low/high為PCLM與組內均勻替代法的敏感度範圍，不是信賴區間。不得標示為官方值。"
                        ),
                    }
                )
            )
        year_diagnostics[str(year)] = {
            "official_means": official_means,
            "official_medians": official_medians,
            "pclm_estimated_means": estimates,
            "uniform_sensitivity_means": baseline_estimates,
            "pclm_estimated_medians": median_estimates,
            "uniform_sensitivity_medians": baseline_median_estimates,
            "median_held_out_validation": median_validation,
            "profile_diagnostics": profile_diag,
            "source_paths": [str(path.relative_to(PROJECT)).replace("\\", "/") for path in auxiliary_paths],
        }

    count = write_csv(
        output,
        sorted(rows, key=lambda row: (int(row["roc_year"]), row["age_band"], row["metric_code"])),
    )
    diagnostics = {
        "source": str(WAGE_WORKBOOK.relative_to(PROJECT)).replace("\\", "/"),
        "source_sha256": sha256(WAGE_WORKBOOK),
        "published_rows": count,
        "available_years": sorted(available_years),
        "requested_years": list(years),
        "missing_years": sorted(set(years) - set(available_years)),
        "publishable_age_bands_by_metric": {
            "ANNUAL_TOTAL_SALARY_MEAN": ["18-24", "25-29", "30-35", "18-35"],
            "ANNUAL_TOTAL_SALARY_MEDIAN": ["18-24", "25-29", "30-35", "18-35"],
        },
        "rectangular_grid": {
            "formula": "4年×1地理×4年齡×1性別×2統計量",
            "expected_rows": len(available_years) * 1 * 4 * 1 * 2,
            "numeric_rows": len(available_years) * 4 * 2,
            "structural_gap_rows": 0,
            "blank_value_semantics": "NO_BLANKS_IN_AVAILABLE_YEARS",
        },
        "sex_dimension": "ALL_ONLY_BECAUSE_OFFICIAL_TABLE_6_HAS_NO_NT_PC_X_AGE_X_SEX_CROSS_TAB",
        "not_published": {
            "male_female": "官方表1有性別、表6有縣市×年齡，但沒有新北市×年齡×性別三維交叉值。",
        },
        "median_method": {
            "subgroup_formula": "target_mean * (official_broad_median / official_broad_mean)",
            "combined_18_35_formula": "weighted_lognormal_mixture_quantile(F=0.5)",
            "weight_source": "PCLM single-age profile calibrated to New Taipei labor-insurance five-year counts",
            "sensitivity": "PCLM versus within-band uniform auxiliary profiles; envelope is not a confidence interval",
            "validation": "25-29 official median held out and predicted from adjacent broad-band shape ratios",
        },
        "year_diagnostics": year_diagnostics,
    }
    return output, diagnostics


def validate_published_csv(path: Path) -> dict:
    keys = set()
    rows = 0
    errors = []
    years = set()
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != SCHEMA_FIELDS:
            errors.append("schema_mismatch")
        for row in reader:
            rows += 1
            years.add(int(row["roc_year"]))
            if row["record_id"] in keys:
                errors.append(f"duplicate_record_id:{row['record_id']}")
            keys.add(row["record_id"])
            if row["publish_status"].startswith("PUBLISH") and row["value"] == "":
                errors.append(f"published_blank_value:{row['record_id']}")
            if row["unit"] == "%" and row["value"]:
                value = float(row["value"])
                if not -1e-6 <= value <= 100.000001:
                    errors.append(f"percentage_out_of_range:{row['record_id']}:{value}")
    return {
        "path": str(path),
        "sha256": sha256(path),
        "rows": rows,
        "years": sorted(years),
        "status": "PASS" if not errors else "FAIL",
        "errors": errors[:100],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--package", type=Path, default=DEFAULT_PACKAGE)
    parser.add_argument("--years", default=",".join(str(year) for year in YEARS))
    args = parser.parse_args()
    years = tuple(sorted({int(value.strip()) for value in args.years.split(",") if value.strip()}))
    package = args.package.resolve()
    qa_dir = package / "08_人工查核與異常處理" / "machine_qa"
    qa_dir.mkdir(parents=True, exist_ok=True)

    registered_path, registered_diag = build_registered_csv(package, years)
    labor_path, labor_diag = build_labor_csv(package, years)
    wage_path, wage_diag = build_wage_csv(package, years)

    diagnostics = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "display_year_rule": "current_roc_year_minus_5_through_current_roc_year_minus_1",
        "execution_years": list(years),
        "registered_population": registered_diag,
        "civilian_labor_market": labor_diag,
        "employee_wage": wage_diag,
    }
    (qa_dir / "data_build_diagnostics.json").write_text(
        json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    validations = [validate_published_csv(path) for path in (registered_path, labor_path, wage_path)]
    catalog = {
        "catalog_version": "G5-V3.3",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "default_view": "同年度檢視",
        "display_year_formula": "民國執行年-5 至 民國執行年-1",
        "current_display_years": list(years),
        "files": [
            {
                "universe_name_zh": name,
                "path": str(path.relative_to(package)).replace("\\", "/"),
                "sha256": validation["sha256"],
                "rows": validation["rows"],
                "years": validation["years"],
                "qa_status": validation["status"],
            }
            for name, path, validation in zip(
                ("戶籍人口母體", "民間人口及勞動市場母體", "受僱員工薪資母體"),
                (registered_path, labor_path, wage_path),
                validations,
            )
        ],
    }
    (package / "09_發布資料包" / "data_catalog.json").write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (qa_dir / "published_csv_validation.json").write_text(
        json.dumps(validations, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    failed = [item for item in validations if item["status"] != "PASS"]
    print(json.dumps({"catalog": catalog, "failed": failed}, ensure_ascii=False, indent=2))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
