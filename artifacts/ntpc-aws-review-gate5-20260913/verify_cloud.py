"""Read-only deployment checks; optional labelled report drill and ETL invocation."""
import argparse
import base64
import hashlib
import io
import json
from pathlib import Path
from datetime import datetime, timezone
import urllib.request
import urllib.error
import boto3
from botocore.config import Config
import deploy as d

ROOT = Path(__file__).resolve().parent
def http(url):
    try:
        with urllib.request.urlopen(urllib.request.Request(url, headers={'User-Agent':'NTPC-deployment-verification'}), timeout=25) as r:
            body = r.read()
            return {'status': r.status, 'bytes': len(body), 'sha256': hashlib.sha256(body).hexdigest()}, body
    except urllib.error.HTTPError as e:
        return {'status': e.code}, e.read()

def invoke(name, event, timeout=60):
    r = boto3.client('lambda', region_name=d.REGION, config=Config(read_timeout=timeout, connect_timeout=7, retries={'max_attempts':0})).invoke(FunctionName=name, Payload=json.dumps(event).encode())
    with r['Payload'] as stream:
        value = json.loads(stream.read())
    if r.get('FunctionError'):
        raise RuntimeError(json.dumps({'function':name,'error_type':value.get('errorType'),'message':value.get('errorMessage')},ensure_ascii=False))
    return value

def main():
    p=argparse.ArgumentParser(); p.add_argument('--invoke-monthly',action='store_true'); p.add_argument('--report-drill',action='store_true'); args=p.parse_args()
    assert d.client('sts').get_caller_identity()['Account']==d.ACCOUNT
    receipt=json.loads((ROOT/'deployment.json').read_text(encoding='utf-8'))
    baseline=json.loads((ROOT/'cloud-baseline.json').read_text(encoding='utf-8'))
    result={'at':datetime.now(timezone.utc).isoformat(),'account':d.ACCOUNT,'checks':{},'pending':[]}
    c=result['checks']; lam=d.client('lambda'); s3=d.client('s3')
    c['lambdas']={}
    for suffix in ('control','report','api'):
        name=d.NAME+'-'+suffix; f=lam.get_function_configuration(FunctionName=name)
        c['lambdas'][suffix]={'state':f['State'],'update_status':f['LastUpdateStatus'],'concurrency':lam.get_function_concurrency(FunctionName=name).get('ReservedConcurrentExecutions')}
        assert f['State']=='Active' and f['LastUpdateStatus']=='Successful'
    c['site_code_unchanged']=lam.get_function_configuration(FunctionName='ntpc-youth-web-gate4')['CodeSha256']==baseline['site_code_sha256']
    assert c['site_code_unchanged']
    c['private_bucket']=all(s3.get_public_access_block(Bucket=d.BUCKET)['PublicAccessBlockConfiguration'].values()); assert c['private_bucket']
    c['effective_versions']={}
    for label,row in baseline['pointers'].items():
        obj=s3.get_object(Bucket=d.BUCKET,Key=row['key']); raw=obj['Body'].read(); obj['Body'].close()
        c['effective_versions'][label]={'sha256':hashlib.sha256(raw).hexdigest(),'unchanged_from_predeployment':hashlib.sha256(raw).hexdigest()==row['sha256']}
    c['controller_status']={k:invoke(d.NAME+'-control',{'action':'status','pipeline':k}) for k in ('MONTHLY','ANNUAL')}
    month=lam.get_function_configuration(FunctionName=d.S['monthly_function'])
    annual=d.client('codebuild').batch_get_projects(names=[d.S['annual_project']])['projects'][0]
    c['etl_hooks']={'monthly':month['Environment']['Variables'].get('REVIEW_CONTROL_FUNCTION')==receipt['controller'], 'annual':next(x['value'] for x in annual['environment']['environmentVariables'] if x['name']=='REVIEW_CONTROL_FUNCTION')==receipt['controller']}
    for pipeline,role in [('monthly',month['Role']),('annual',annual['serviceRole'])]:
        policy=d.client('iam').get_role_policy(RoleName=role.rsplit('/',1)[-1],PolicyName='CurrentWriteRequiresReviewController')['PolicyDocument']
        c['etl_hooks'][pipeline+'_direct_write_denied']=any(x['Effect']=='Deny' and 's3:PutObject' in x['Action'] for x in policy['Statement'])
    assert all(c['etl_hooks'].values())
    table=d.client('dynamodb').describe_table(TableName=d.NAME)['Table']
    c['cases_table']=table['TableStatus']; assert c['cases_table']=='ACTIVE'
    c['point_in_time_recovery']=d.client('dynamodb').describe_continuous_backups(TableName=d.NAME)['ContinuousBackupsDescription']['PointInTimeRecoveryDescription']['PointInTimeRecoveryStatus']
    c['cloudfront']=d.client('cloudfront').get_distribution(Id=d.S['distribution_id'])['Distribution']['Status']
    api=next(x for x in d.client('apigatewayv2').get_apis()['Items'] if x['Name']==d.NAME)
    routes=d.client('apigatewayv2').get_routes(ApiId=api['ApiId'])['Items']
    c['protected_api_routes']=all(x.get('AuthorizationType')=='JWT' and 'ntpc-review/access' in x.get('AuthorizationScopes',[]) for x in routes if '/api/' in x['RouteKey'] or x['RouteKey']=='$default'); assert c['protected_api_routes']
    c['http']={}
    for label,url in [('public',d.S['website']+'/'),('review',d.S['website']+'/review/'),('config',d.S['website']+'/review/config'),('unauthenticated_cases',d.S['website']+'/api/review/cases'),('direct_origin',api['ApiEndpoint']+'/review/')]:
        check,body=http(url); c['http'][label]=check
        if label=='config' and check['status']==200:
            config=json.loads(body); c['oauth_scope_present']='aws.cognito.signin.user.admin' in config['scope']
    subs=d.client('sns').list_subscriptions_by_topic(TopicArn=receipt['topic_arn'])['Subscriptions']
    confirmed=any(x.get('Endpoint')==d.S['notification_email'] and x.get('SubscriptionArn','').startswith('arn:') for x in subs)
    c['email_subscription']='CONFIRMED' if confirmed else 'PENDING_CONFIRMATION'
    user=d.client('cognito-idp').admin_get_user(UserPoolId=receipt['user_pool_id'],Username=d.S['admin_email'])
    c['admin_user_status']=user['UserStatus']
    if not confirmed:result['pending'].append('SNS email subscription confirmation by mailbox owner')
    if user['UserStatus']=='FORCE_CHANGE_PASSWORD':result['pending'].append('First administrator login and password setup')
    if args.invoke_monthly:
        print('INVOKING_MONTHLY_ETL',flush=True)
        r=invoke(d.S['monthly_function'],{},timeout=930)
        (ROOT/'monthly-live-result.json').write_text(json.dumps(r,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
        c['monthly_invocation']={'response_saved':'monthly-live-result.json','status':r.get('status',r.get('statusCode'))}
    if args.report_drill:
        drillpath=ROOT/'report-drill.json'
        if not drillpath.exists():
            from review_core import encoded
            now=datetime.now(timezone.utc).isoformat(); case_id='DRILL-'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')
            case={'case_id':case_id,'revision':1,'pipeline':'DEPLOYMENT_TEST','status':'CLOSED_NO_CHANGE','created_at':now,'updated_at':now,'baseline':{'pointer':None},'candidate':None,
                'issues':[{'severity':'TEST','rule':'DEPLOYMENT_REPORT_TEST','scope':'部署驗證；非真實資料異常','expected':'自動產生三張工作表並通知指定管理者','actual':'這是已結束的部署測試案件，不涉及任何正式資料發布。','evidence':'Production SQS → report Lambda → private S3; notification after subscription confirmation','resolution':'TEST_ONLY'}],
                'history':[{'revision':1,'at':now,'event':'DEPLOYMENT_TEST','actor':'DEPLOYMENT_VERIFICATION','reason':'驗證雲端自動產表與通知；未執行管理者簽核'}]}
            key=d.PREFIX+'cases/'+case_id+'/revision-000001-test.json'
            d.put_immutable(s3,key,encoded(case))
            boto3.resource('dynamodb',region_name=d.REGION).Table(d.NAME).put_item(Item={'pk':'CASE#'+case_id,'case_id':case_id,'pipeline':'DEPLOYMENT_TEST','revision':1,'status':'CLOSED_NO_CHANGE','updated_at':now,'snapshot_key':key,'issue_count':1},ConditionExpression='attribute_not_exists(pk)')
            queue=d.client('sqs').get_queue_url(QueueName=d.NAME+'-reports')['QueueUrl']
            message=d.client('sqs').send_message(QueueUrl=queue,MessageBody=json.dumps({'snapshot_key':key}))
            drill={'case_id':case_id,'snapshot_key':key,'message_id':message['MessageId'],'report_key':d.PREFIX+'reports/'+case_id+'/r1/查核表.xlsx','delivery_key':d.PREFIX+'reports/'+case_id+'/r1/delivery.json'}
            drillpath.write_text(json.dumps(drill,ensure_ascii=False,indent=2),encoding='utf-8')
        else:drill=json.loads(drillpath.read_text(encoding='utf-8'))
        c['report_drill']={'case_id':drill['case_id'],'sqs_enqueued':True}
        found=s3.list_objects_v2(Bucket=d.BUCKET,Prefix=drill['report_key'],MaxKeys=1).get('Contents',[])
        if any(x['Key']==drill['report_key'] for x in found):
            obj=s3.get_object(Bucket=d.BUCKET,Key=drill['report_key']); raw=obj['Body'].read();obj['Body'].close()
            (ROOT/'雲端產表_部署測試.xlsx').write_bytes(raw)
            from openpyxl import load_workbook
            w=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
            c['report_drill']['cloud_generated_sheets']=w.sheetnames;assert len(w.sheetnames)==3;w.close()
        found=s3.list_objects_v2(Bucket=d.BUCKET,Prefix=drill['delivery_key'],MaxKeys=1).get('Contents',[])
        c['report_drill']['sns_accepted']=any(x['Key']==drill['delivery_key'] for x in found)
    result['status']='DEPLOYED_USER_ACTIVATION_PENDING' if result['pending'] else 'DEPLOYED_SERVICE_CHECKED'
    (ROOT/'cloud-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
