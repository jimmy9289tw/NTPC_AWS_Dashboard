import copy
import io
import json
from pathlib import Path
import unittest
import zipfile
import os
from unittest.mock import patch, MagicMock
from xml.etree import ElementTree as ET
from review_core import Engine, ReviewError
from report_worker import workbook


class FakePort:
    def __init__(self):
        self.cases = {}; self.current_value = {'pointer': {'manifest_key': 'old'}, 'etag': 'old-etag'}
        self.queue_failed = False; self.publish_count = 0; self.tampered = False
        self.fail_after_commit = False; self.tick = 0
    def now(self):
        self.tick += 1
        return '2026-09-13T00:00:' + str(self.tick).zfill(2) + 'Z'
    def current(self, pipeline): return copy.deepcopy(self.current_value)
    def case(self, case_id): return copy.deepcopy(self.cases.get(case_id))
    def active(self, pipeline):
        return next((copy.deepcopy(x) for x in self.cases.values() if x['pipeline'] == pipeline and x['status'] not in ('PUBLISHED','CLOSED_NO_CHANGE')), None)
    def event_case(self, pipeline, event_id):
        return next((copy.deepcopy(x) for x in self.cases.values() if x['pipeline'] == pipeline and event_id in x.get('event_ids', [])), None)
    def save(self, c):
        if self.fail_after_commit and c['status'] == 'PUBLISHED':
            self.fail_after_commit = False; raise RuntimeError('timeout after S3 commit')
        old = self.cases.get(c['case_id'])
        if old and c['revision'] != old['revision'] + 1: raise ReviewError('stale write')
        self.cases[c['case_id']] = copy.deepcopy(c)
    def enqueue(self, c):
        if self.queue_failed: raise RuntimeError('queue unavailable')
    def log(self, *args): pass
    def verify_candidate(self, r):
        if self.tampered or r.get('qa') == 'FAIL': raise ReviewError('candidate invalid')
        return {k: r[k] for k in ('pipeline','run_id','manifest_key','manifest_sha256','pointer','baseline_etag')}
    def publish(self, pipeline, pointer, etag):
        if etag != self.current_value['etag']: raise ReviewError('CAS conflict')
        self.publish_count += 1
        self.current_value = {'pointer': copy.deepcopy(pointer), 'etag': 'new-etag'}


def candidate(suffix='a'):
    return {'pipeline':'MONTHLY','run_id':'retest-'+suffix,'manifest_key':'candidate-'+suffix,
            'manifest_sha256':suffix*64,'pointer':{'manifest_key':'candidate-'+suffix},'baseline_etag':'old-etag'}


class WorkflowTests(unittest.TestCase):
    def setUp(self): self.p = FakePort(); self.e = Engine(self.p)
    def anomaly(self):
        return self.e.anomaly({'pipeline':'MONTHLY','event_id':'failed-run-1','issue':{
            'rule':'GEO-01','severity':'BLOCKER','actual':'僅取得28區','expected':'應取得29區','evidence':'qa.json'}})
    def decision(self,c,action='APPROVE'):
        return {'case_id':c['case_id'],'revision':c['revision'],'decision':action,
                'reason':'已核對官方修訂與完整復驗證據','manifest_sha256':(c.get('candidate') or {}).get('manifest_sha256')}
    def ready(self):
        self.anomaly(); self.e.stage(candidate()); return self.p.active('MONTHLY')
    def test_clean_batch_can_publish(self):
        self.assertEqual(self.e.stage(candidate())['status'],'PUBLISHED'); self.assertEqual(self.p.publish_count,1)
    def test_failure_keeps_old_version(self):
        self.anomaly(); self.assertEqual(self.p.current_value['etag'],'old-etag')
    def test_duplicate_anomaly_same_case(self):
        a=self.anomaly(); b=self.anomaly(); self.assertEqual(a,b); self.assertEqual(len(b['issues']),1)
    def test_retest_does_not_bypass_approval(self):
        c=self.ready(); self.assertEqual(c['status'],'READY_REVIEW'); self.assertEqual(self.p.publish_count,0)
    def test_blocker_without_retest_cannot_approve(self):
        c=self.anomaly()
        with self.assertRaises(ReviewError):self.e.decision(self.decision(c),'named-admin')
    def test_acknowledgement_is_not_publication(self):
        c=self.anomaly(); self.e.decision(self.decision(c,'ACKNOWLEDGE'),'named-admin'); self.assertEqual(self.p.publish_count,0)
    def test_acknowledgement_preserves_ready_candidate(self):
        c=self.ready(); r=self.e.decision(self.decision(c,'ACKNOWLEDGE'),'named-admin'); self.assertEqual(r['status'],'READY_REVIEW')
    def test_signed_retest_publishes(self):
        c=self.ready(); r=self.e.decision(self.decision(c),'named-admin'); self.assertEqual(r['status'],'PUBLISHED')
    def test_reject_keeps_gate_closed_for_next_batch(self):
        c=self.ready(); self.e.decision(self.decision(c,'REJECT'),'named-admin'); self.assertEqual(self.e.stage(candidate('b'))['status'],'REVIEW_REQUIRED')
    def test_stale_revision_rejected(self):
        c=self.ready(); req=self.decision(c); req['revision']-=1
        with self.assertRaises(ReviewError):self.e.decision(req,'named-admin')
    def test_wrong_hash_rejected(self):
        c=self.ready(); req=self.decision(c); req['manifest_sha256']='x'*64
        with self.assertRaises(ReviewError):self.e.decision(req,'named-admin')
    def test_tampered_files_rejected(self):
        c=self.ready(); self.p.tampered=True
        with self.assertRaises(ReviewError):self.e.decision(self.decision(c),'named-admin')
    def test_system_cannot_impersonate_admin(self):
        c=self.ready()
        with self.assertRaises(ReviewError):self.e.decision(self.decision(c),'SYSTEM')
    def test_new_candidate_changes_revision(self):
        a=self.ready(); self.e.stage(candidate('b')); b=self.p.active('MONTHLY'); self.assertGreater(b['revision'],a['revision'])
        with self.assertRaises(ReviewError):self.e.decision(self.decision(a),'named-admin')
    def test_baseline_conflict_prevents_clean_and_reviewed_publish(self):
        c=self.ready(); self.p.current_value['etag']='other-batch'
        with self.assertRaises(ReviewError):self.e.decision(self.decision(c),'named-admin')
        with self.assertRaises(ReviewError):self.e.stage(candidate())
    def test_queue_failure_does_not_open_gate(self):
        self.p.queue_failed=True; c=self.anomaly(); self.assertEqual(self.p.active('MONTHLY')['case_id'],c['case_id']); self.assertEqual(self.p.publish_count,0)
    def test_publish_receipt_recovers_after_timeout(self):
        c=self.ready(); self.p.fail_after_commit=True
        with self.assertRaises(RuntimeError):self.e.decision(self.decision(c),'named-admin')
        self.assertEqual(self.p.active('MONTHLY')['status'],'PUBLISHING')
        self.e.resume(self.p.active('MONTHLY')); self.assertEqual(self.p.publish_count,1)
    def test_delayed_duplicate_does_not_reopen_closed_case(self):
        c=self.ready(); self.e.decision(self.decision(c),'named-admin'); result=self.anomaly(); self.assertEqual(result['status'],'PUBLISHED')
    def test_clean_retry_does_not_duplicate_commit(self):
        self.e.stage(candidate()); self.e.stage(candidate()); self.assertEqual(self.p.publish_count,1)
    def test_failed_qa_never_staged(self):
        with self.assertRaises(ReviewError):self.e.stage({**candidate(),'qa':'FAIL'})
        self.assertEqual(self.p.publish_count,0)
    def test_workbook_includes_all_issues_and_no_formulas(self):
        c=self.anomaly(); c['issues']=[{**c['issues'][0],'actual':'=HYPERLINK("https://invalid.example","test")'} for _ in range(57)]
        data=workbook(c); self.assertEqual(data,workbook(c))
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            self.assertIsNone(z.testzip()); ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            for name in z.namelist():
                if name.endswith('.xml'):ET.fromstring(z.read(name))
            issues=ET.fromstring(z.read('xl/worksheets/sheet2.xml'))
            self.assertEqual(len(issues.findall('s:sheetData/s:row',ns)),58)
            self.assertEqual(len(issues.findall('.//s:f',ns)),0)
            text=''.join(issues.itertext()); self.assertIn('=HYPERLINK',text)
            self.assertEqual(len(ET.fromstring(z.read('xl/workbook.xml')).findall('s:sheets/s:sheet',ns)),3)
    def test_real_integration_patch_preserves_baselines(self):
        import prepare_integration
        prepare_integration.build()
        root=Path(__file__).parent
        monthly=(root/'prepared/monthly/etl/monthly.py').read_text(encoding='utf-8')
        annual=(root/'prepared/annual/run_annual.py').read_text(encoding='utf-8')
        self.assertIn('review_pending',monthly); self.assertIn('review_run_id',monthly)
        self.assertNotIn('Key=pointer_key,Body=json.dumps(pointer)',annual)
        self.assertIn('annual_stage(',annual); self.assertIn('annual_failure(report)',annual)

    def test_xlsx_loads_in_independent_reader(self):
        from openpyxl import load_workbook
        c=self.anomaly(); w=load_workbook(io.BytesIO(workbook(c)),read_only=True,data_only=False)
        self.assertEqual(w.sheetnames,['管理者查核表','問題明細','處置與簽核紀錄'])
        self.assertEqual(w['問題明細']['F2'].value,'僅取得28區'); w.close()

    def test_api_rejects_origin_bypass(self):
        import review_api
        with patch.dict(os.environ,{'ORIGIN_SECRET':'test-origin','ALLOW_IPS':'127.0.0.1'}):
            r=review_api.handler({'rawPath':'/api/review/cases'},None)
        self.assertEqual(r['statusCode'],403)

    def test_api_rejects_missing_named_identity(self):
        import review_api
        env={'ORIGIN_SECRET':'test-origin','ALLOW_IPS':'127.0.0.1','ADMIN_SUB':'admin-sub','CLIENT_ID':'client'}
        event={'rawPath':'/api/review/cases','headers':{'x-ntpc-origin':'test-origin','x-ntpc-viewer-ip':'127.0.0.1'},'requestContext':{'http':{'method':'GET'}}}
        with patch.dict(os.environ,env):r=review_api.handler(event,None)
        self.assertEqual(r['statusCode'],403)

    def test_controller_rechecks_cognito_subject(self):
        import aws_control
        fake=MagicMock(); fake.get_user.return_value={'Username':'other','UserAttributes':[{'Name':'sub','Value':'other-sub'}]}
        with patch.object(aws_control,'Port',return_value=self.p),patch.object(aws_control,'client',return_value=fake),patch.dict(os.environ,{'ADMIN_SUB':'admin-sub'}):
            with self.assertRaises(ReviewError):aws_control.handler({'action':'decision','access_token':'test-token'},None)
        self.assertEqual(self.p.publish_count,0)

    def test_first_report_is_written_before_unconfirmed_email_retry(self):
        import report_worker
        c=self.anomaly(); s3=MagicMock(); sns=MagicMock()
        s3.get_object.return_value={'Body':io.BytesIO(json.dumps(c).encode())}
        s3.list_objects_v2.return_value={}
        sns.list_subscriptions_by_topic.return_value={'Subscriptions':[{'Endpoint':'admin@example.com','SubscriptionArn':'PendingConfirmation'}]}
        env={'DATA_BUCKET':'test','DATA_ACCOUNT':'123456789012','REVIEW_PREFIX':'review/','TOPIC_ARN':'topic','NOTIFICATION_EMAIL':'admin@example.com','REVIEW_URL':'https://example.invalid/review/'}
        with patch.object(report_worker.boto3,'client',side_effect=lambda name:s3 if name=='s3' else sns),patch.dict(os.environ,env):
            with self.assertRaisesRegex(RuntimeError,'EMAIL_SUBSCRIPTION_PENDING'):
                report_worker.process({'body':json.dumps({'snapshot_key':'review/cases/test.json'})})
        self.assertTrue(s3.put_object.called); self.assertTrue(s3.put_object.call_args.kwargs['Key'].endswith('查核表.xlsx'))
        sns.publish.assert_not_called()


if __name__=='__main__':
    unittest.main(verbosity=2)
