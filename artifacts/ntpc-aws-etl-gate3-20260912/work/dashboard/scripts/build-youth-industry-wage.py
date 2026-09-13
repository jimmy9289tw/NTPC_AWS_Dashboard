from __future__ import annotations

import csv
import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl


DASHBOARD = Path(__file__).resolve().parents[1]
PROJECT = DASHBOARD
SOURCE_DIR = PROJECT / "data" / "raw" / "DGBAS-WAGE-INDUSTRY-AGE" / "20260907"
TABLE2 = SOURCE_DIR / "DGBAS-WAGE-TABLE2-INDUSTRY-AGE.xlsx"
TABLE6 = SOURCE_DIR / "DGBAS-WAGE-TABLE6-COUNTY-AGE.xlsx"
AGE_COUNTS = SOURCE_DIR / "BLI-AGE-REGION-SEX-113.csv"
G5_DATA = DASHBOARD / "app" / "data" / "g5-dashboard-data.json"
MOL_CONTEXT = DASHBOARD / "app" / "data" / "mol-salary-context.json"
OUTPUT = DASHBOARD / "app" / "data" / "youth-industry-wage.json"

SOURCE_GROUPS = ((15, 19), (20, 24), (25, 29), (30, 34), (35, 39))
TARGET_AGES = {
    "18-24": range(18, 25),
    "25-29": range(25, 30),
    "30-35": range(30, 36),
}
TABLE2_COLUMNS = {"18-24": 4, "25-29": 5, "30-35": 6}
TABLE2_SOURCE_AGE = {"18-24": "未滿25歲", "25-29": "25–29歲", "30-35": "30–39歲"}

# Table 2 has 17 detailed industries. Agriculture and public administration
# appear in the MOL context but are outside this DGBAS salary table's detailed
# industry universe, so the published matrix retains them as explicit gaps.
TABLE2_INDUSTRIES = (
    (10, "礦業及土石採取業"),
    (11, "製造業"),
    (12, "電力及燃氣供應業"),
    (13, "用水供應及污染整治業"),
    (14, "營建工程業"),
    (16, "批發及零售業"),
    (17, "運輸及倉儲業"),
    (18, "住宿及餐飲業"),
    (19, "出版影音及資通訊業"),
    (20, "金融及保險業"),
    (21, "不動產業"),
    (22, "專業、科學及技術服務業"),
    (23, "支援服務業"),
    (24, "教育業"),
    (25, "醫療保健及社會工作服務業"),
    (26, "藝術、娛樂及休閒服務業"),
    (27, "其他服務業"),
)
UNAVAILABLE_INDUSTRIES = ("公共行政及國防；強制性社會安全", "農、林、漁、牧業")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def load_age_counts(path: Path) -> dict[tuple[int, int], float]:
    columns = {
        (15, 19): "年齡15-19歲人數",
        (20, 24): "年齡20-24歲人數",
        (25, 29): "年齡25-29歲人數",
        (30, 34): "年齡30-34歲人數",
        (35, 39): "年齡35-39歲人數",
    }
    totals = {group: 0.0 for group in SOURCE_GROUPS}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if row.get("地區別") != "新北市":
                continue
            for group, column in columns.items():
                totals[group] += float(row.get(column) or 0)
    if any(value <= 0 for value in totals.values()):
        raise ValueError(f"Missing New Taipei age counts: {totals}")
    return totals


def bspline_basis(ages: np.ndarray, left: float, right: float, step: int = 5, degree: int = 3) -> np.ndarray:
    internal = list(np.arange(left + step, right, step, dtype=float))
    knots = np.array([left] * (degree + 1) + internal + [right] * (degree + 1), dtype=float)
    n_basis = len(knots) - degree - 1
    x = ages.astype(float) + 0.5
    basis = np.zeros((len(x), n_basis), dtype=float)
    for index in range(n_basis):
        basis[:, index] = ((x >= knots[index]) & (x < knots[index + 1])).astype(float)
    for order in range(1, degree + 1):
        next_basis = np.zeros_like(basis)
        for index in range(n_basis):
            left_denominator = knots[index + order] - knots[index]
            right_denominator = knots[index + order + 1] - knots[index + 1] if index + 1 < n_basis else 0.0
            if left_denominator > 0:
                next_basis[:, index] += (x - knots[index]) / left_denominator * basis[:, index]
            if right_denominator > 0 and index + 1 < n_basis:
                next_basis[:, index] += (knots[index + order + 1] - x) / right_denominator * basis[:, index + 1]
        basis = next_basis
    return basis


def poisson_deviance(observed: np.ndarray, expected: np.ndarray) -> float:
    expected = np.maximum(expected, 1e-12)
    terms = expected.copy()
    positive = observed > 0
    terms[positive] = observed[positive] * np.log(observed[positive] / expected[positive]) - (observed[positive] - expected[positive])
    return float(2.0 * np.sum(terms))


def fit_pclm(observed: np.ndarray, ages: np.ndarray, groups: list[tuple[int, int]], penalty_lambda: float):
    basis = bspline_basis(ages, float(ages.min()), float(ages.max() + 1))
    composition = np.zeros((len(groups), len(ages)), dtype=float)
    for group_index, (start, end) in enumerate(groups):
        composition[group_index, (ages >= start) & (ages <= end)] = 1.0
    second_difference = np.diff(np.eye(basis.shape[1]), n=2, axis=0)
    penalty = second_difference.T @ second_difference
    group_basis = (composition @ basis) / np.maximum(composition.sum(axis=1, keepdims=True), 1.0)
    target = np.log((observed + 0.5) / np.maximum(composition.sum(axis=1), 1.0))
    theta = np.linalg.solve(group_basis.T @ group_basis + 1e-4 * np.eye(basis.shape[1]), group_basis.T @ target)

    def objective(candidate: np.ndarray) -> float:
        fitted = np.exp(np.clip(basis @ candidate, -30, 30))
        expected = np.maximum(composition @ fitted, 1e-12)
        log_likelihood = float(np.sum(observed * np.log(expected) - expected))
        return log_likelihood - 0.5 * penalty_lambda * float(candidate @ penalty @ candidate)

    converged = False
    for iteration in range(1, 151):
        fitted = np.exp(np.clip(basis @ theta, -30, 30))
        expected = np.maximum(composition @ fitted, 1e-12)
        jacobian = composition @ (fitted[:, None] * basis)
        information = jacobian.T @ (jacobian / expected[:, None])
        score = jacobian.T @ ((observed - expected) / expected) - penalty_lambda * penalty @ theta
        system = information + penalty_lambda * penalty + 1e-8 * np.eye(basis.shape[1])
        delta = np.linalg.solve(system, score)
        old_objective = objective(theta)
        step = 1.0
        while step >= 1e-5 and objective(theta + step * delta) < old_objective:
            step *= 0.5
        theta = theta + step * delta
        if float(np.max(np.abs(step * delta))) < 1e-8:
            converged = True
            break
    fitted = np.exp(np.clip(basis @ theta, -30, 30))
    expected = np.maximum(composition @ fitted, 1e-12)
    jacobian = composition @ (fitted[:, None] * basis)
    information = jacobian.T @ (jacobian / expected[:, None])
    system = information + penalty_lambda * penalty + 1e-8 * np.eye(basis.shape[1])
    effective_df = float(np.trace(np.linalg.solve(system, information)))
    deviance = poisson_deviance(observed, expected)
    return fitted, {
        "lambda": penalty_lambda,
        "deviance": deviance,
        "edf": effective_df,
        "gcv": deviance / max(len(groups) - effective_df, 0.25) ** 2,
        "iterations": iteration,
        "converged": converged,
        "max_group_reaggregation_error": float(np.max(np.abs(expected - observed))),
    }


def select_pclm(observed: np.ndarray, ages: np.ndarray, groups: list[tuple[int, int]]):
    candidates = []
    for penalty_lambda in (0.1, 1.0, 10.0, 100.0, 1000.0, 10000.0):
        fitted, diagnostics = fit_pclm(observed, ages, groups, penalty_lambda)
        holdout_deviance = 0.0
        for holdout in range(1, len(groups) - 1):
            training_groups = [group for index, group in enumerate(groups) if index != holdout]
            training_observed = np.array([value for index, value in enumerate(observed) if index != holdout], dtype=float)
            heldout_fitted, _ = fit_pclm(training_observed, ages, training_groups, penalty_lambda)
            start, end = groups[holdout]
            prediction = float(heldout_fitted[(ages >= start) & (ages <= end)].sum())
            holdout_deviance += poisson_deviance(
                np.array([observed[holdout]], dtype=float),
                np.array([prediction], dtype=float),
            )
        diagnostics["interior_leave_one_group_out_deviance"] = holdout_deviance
        candidates.append((holdout_deviance, fitted, diagnostics))
    _, fitted, diagnostics = min(candidates, key=lambda item: item[0])
    return fitted, diagnostics


def pclm_age_profile(grouped_counts: dict[tuple[int, int], float]) -> tuple[dict[int, float], dict]:
    ages = np.arange(15, 40)
    margins = np.array([grouped_counts[group] for group in SOURCE_GROUPS], dtype=float)
    fitted, diagnostics = select_pclm(margins, ages, list(SOURCE_GROUPS))
    for index, (start, end) in enumerate(SOURCE_GROUPS):
        mask = (ages >= start) & (ages <= end)
        fitted[mask] *= margins[index] / float(fitted[mask].sum())
    profile = {int(age): float(fitted[index]) for index, age in enumerate(ages)}
    diagnostics["post_calibration_max_margin_error"] = max(
        abs(sum(profile[age] for age in range(start, end + 1)) - grouped_counts[(start, end)])
        for start, end in SOURCE_GROUPS
    )
    return profile, diagnostics


def uniform_age_profile(grouped_counts: dict[tuple[int, int], float]) -> dict[int, float]:
    return {
        age: grouped_counts[(start, end)] / (end - start + 1)
        for start, end in SOURCE_GROUPS
        for age in range(start, end + 1)
    }


def target_weights(profile: dict[int, float]) -> dict[str, float]:
    return {
        band: sum(profile[age] for age in ages)
        for band, ages in TARGET_AGES.items()
    }


def weighted_mean(values: dict[str, float], weights: dict[str, float]) -> float:
    denominator = sum(weights[key] for key in values)
    if denominator <= 0:
        raise ValueError("Non-positive denominator")
    return sum(values[key] * weights[key] for key in values) / denominator


def normalize_ratios(raw: dict[str, float], industry_weights: dict[str, float]) -> dict[str, float]:
    center = weighted_mean(raw, industry_weights)
    return {industry: ratio / center for industry, ratio in raw.items()}


def round_or_none(value: float | None, digits: int = 1) -> float | None:
    return None if value is None else round(value, digits)


def main() -> None:
    required = [TABLE2, TABLE6, AGE_COUNTS, G5_DATA, MOL_CONTEXT]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Missing inputs: {missing}")

    g5 = load_json(G5_DATA)
    mol = load_json(MOL_CONTEXT)
    anchors = {
        row["ageBand"]: {
            "value": float(row["metrics"]["全年總薪資平均數"]),
            "low": float(row["meta"]["全年總薪資平均數"]["low"] or row["metrics"]["全年總薪資平均數"]),
            "high": float(row["meta"]["全年總薪資平均數"]["high"] or row["metrics"]["全年總薪資平均數"]),
            "origin": row["meta"]["全年總薪資平均數"]["origin"],
        }
        for row in g5["wage"]
        if row["year"] == 113
    }
    expected_bands = {"18-24", "25-29", "30-35", "18-35"}
    if set(anchors) != expected_bands:
        raise ValueError(f"Unexpected wage anchors: {sorted(anchors)}")

    workbook2 = openpyxl.load_workbook(TABLE2, read_only=True, data_only=True)
    table2 = workbook2[workbook2.sheetnames[0]]
    workbook6 = openpyxl.load_workbook(TABLE6, read_only=True, data_only=True)
    table6 = workbook6[workbook6.sheetnames[0]]
    if not math.isclose(float(table6.cell(9, 5).value), anchors["25-29"]["value"], abs_tol=0.05):
        raise ValueError("Table 6 native 25-29 mean does not match the dashboard anchor")

    mol_industries = {row["name"]: row for row in mol["industries"]}
    available_names = [industry for _, industry in TABLE2_INDUSTRIES]
    if any(name not in mol_industries for name in available_names):
        raise ValueError("DGBAS-to-MOL industry mapping is incomplete")
    industry_weights = {name: float(mol_industries[name]["headcount"]) for name in available_names}

    grouped_counts = load_age_counts(AGE_COUNTS)
    pclm_profile, pclm_diagnostics = pclm_age_profile(grouped_counts)
    uniform_profile = uniform_age_profile(grouped_counts)
    pclm_weights = target_weights(pclm_profile)
    uniform_weights = target_weights(uniform_profile)

    national_means: dict[str, dict[str, float]] = {}
    for band, column in TABLE2_COLUMNS.items():
        national_means[band] = {
            industry: float(table2.cell(row_number, column).value)
            for row_number, industry in TABLE2_INDUSTRIES
        }
        national_means[band]["__TOTAL__"] = float(table2.cell(8, column).value)

    raw_national_ratios: dict[str, dict[str, float]] = {}
    for band in ("18-24", "25-29", "30-35"):
        total = national_means[band]["__TOTAL__"]
        raw_national_ratios[band] = {
            industry: national_means[band][industry] / total
            for industry in available_names
        }
    raw_national_ratios["18-35"] = {
        industry: (
            sum(national_means[band][industry] * pclm_weights[band] for band in ("18-24", "25-29", "30-35"))
            / sum(pclm_weights.values())
        ) / (
            sum(national_means[band]["__TOTAL__"] * pclm_weights[band] for band in ("18-24", "25-29", "30-35"))
            / sum(pclm_weights.values())
        )
        for industry in available_names
    }
    uniform_composite_ratios = {
        industry: (
            sum(national_means[band][industry] * uniform_weights[band] for band in ("18-24", "25-29", "30-35"))
            / sum(uniform_weights.values())
        ) / (
            sum(national_means[band]["__TOTAL__"] * uniform_weights[band] for band in ("18-24", "25-29", "30-35"))
            / sum(uniform_weights.values())
        )
        for industry in available_names
    }

    national_ratios = {
        band: normalize_ratios(raw_national_ratios[band], industry_weights)
        for band in expected_bands
    }
    uniform_composite_ratios = normalize_ratios(uniform_composite_ratios, industry_weights)
    local_contribution_ratios = normalize_ratios(
        {
            industry: float(mol_industries[industry]["averageContributionWage"])
            / float(mol["newTaipei"]["averageContributionWage"])
            for industry in available_names
        },
        industry_weights,
    )

    rows = []
    for industry in available_names:
        estimates = {}
        for band in ("18-24", "25-29", "30-35", "18-35"):
            anchor = anchors[band]
            point_ratio = national_ratios[band][industry]
            point = anchor["value"] * point_ratio
            alternative_ratios = [point_ratio, local_contribution_ratios[industry]]
            if band == "18-35":
                alternative_ratios.append(uniform_composite_ratios[industry])
            candidates = [
                anchor_value * ratio
                for anchor_value in (anchor["low"], anchor["value"], anchor["high"])
                for ratio in alternative_ratios
            ]
            estimates[band] = {
                "value": round_or_none(point),
                "low": round_or_none(min(candidates)),
                "high": round_or_none(max(candidates)),
                "sourceAgeBand": (
                    "18–24、25–29、30–35歲組合"
                    if band == "18-35"
                    else TABLE2_SOURCE_AGE[band]
                ),
                "nationalAgeIndustryIndex": round(point_ratio, 6),
                "localAllAgeContributionIndex": round(local_contribution_ratios[industry], 6),
                "anchor": anchor["value"],
                "anchorOrigin": anchor["origin"],
                "status": "MODEL_ESTIMATE",
            }
        rows.append({
            "industry": industry,
            "supportingAllAgeHeadcount": int(industry_weights[industry]),
            "estimates": estimates,
            "availability": "MODEL_ESTIMATE_AVAILABLE",
        })

    rows.sort(key=lambda row: row["estimates"]["18-35"]["value"], reverse=True)
    rows.extend({
        "industry": industry,
        "supportingAllAgeHeadcount": int(mol_industries[industry]["headcount"]),
        "estimates": {
            band: {
                "value": None,
                "low": None,
                "high": None,
                "sourceAgeBand": None,
                "nationalAgeIndustryIndex": None,
                "localAllAgeContributionIndex": None,
                "anchor": anchors[band]["value"],
                "anchorOrigin": anchors[band]["origin"],
                "status": "SOURCE_UNIVERSE_EXCLUDED",
            }
            for band in ("18-24", "25-29", "30-35", "18-35")
        },
        "availability": "NOT_IN_DGBAS_TABLE2_INDUSTRY_UNIVERSE",
    } for industry in UNAVAILABLE_INDUSTRIES)

    reaggregation = {}
    for band in ("18-24", "25-29", "30-35", "18-35"):
        modeled = {row["industry"]: row["estimates"][band]["value"] for row in rows if row["availability"] == "MODEL_ESTIMATE_AVAILABLE"}
        reaggregated = weighted_mean(modeled, industry_weights)
        reaggregation[band] = {
            "anchor": anchors[band]["value"],
            "reaggregated": round(reaggregated, 6),
            "absoluteError": round(abs(reaggregated - anchors[band]["value"]), 6),
        }

    payload = {
        "meta": {
            "version": "G6-WAGE-INDUSTRY-1.0",
            "rocYear": 113,
            "ageBands": ["18-35", "18-24", "25-29", "30-35"],
            "unit": "萬元/年",
            "statistic": "全年總薪資平均數",
            "population": "工作場所位於新北市之本國籍全時受僱員工；以全國各業受僱員工年齡別薪資差異作結構輔助",
            "geographyBasis": "工作場所所在地；不是戶籍地或居住地",
            "identity": "模型估計值",
            "formula": "青年行業薪資估計＝新北市目標年齡全年總薪資平均數錨點×校準後全國同行業同年齡薪資指數",
            "calibration": "各年齡組的行業薪資指數以目前官方端點可取得之新北市114年勞退提繳者全年齡行業人數為權重中心化，使17個可用行業加權回加至113年同年齡新北市薪資錨點",
            "uncertainty": "上下限取新北年齡薪資錨點既有敏感度、全國同年齡行業薪資指數、以及新北全年齡行業提繳工資指數的組合包絡；18–35另納入PCLM與均勻拆分的年齡權重差。這是方法敏感度，不是信賴區間或顯著性檢定",
            "limitations": [
                "官方未發布新北市×青年年齡×行業的直接交叉薪資，本表不得標示為官方原生值",
                "18–24使用全國未滿25歲行業差異作輔助；30–35使用全國30–39歲行業差異作輔助，邊界差異已納入方法限制",
                "勞退提繳人數僅作行業校準權重；提繳工資是級距化申報值，不是全年總薪資",
                "目前勞退公開端點僅提供114年，故113年薪資模型以114年全年齡行業人數作校準權重；這是跨一年的結構輔助，不能解讀為113年青年行業人數",
                "農林漁牧業及公共行政未列入主計總處表2的詳細行業薪資範圍，保留缺值，不以勞退資料補造",
            ],
            "sources": [
                {
                    "alias": "DGBAS-WAGE-T2",
                    "name": "表2 各業受僱員工全年總薪資統計－按年齡別分",
                    "url": "https://www.stat.gov.tw/News_Content.aspx?n=4580&s=232642",
                    "snapshot": str(TABLE2.relative_to(PROJECT)).replace("\\", "/"),
                    "sha256": sha256(TABLE2),
                },
                {
                    "alias": "DGBAS-WAGE-T6",
                    "name": "表6 工業及服務業全年總薪資統計－按工作場所縣市及年齡別分",
                    "url": "https://www.stat.gov.tw/News_Content.aspx?n=4580&s=232642",
                    "snapshot": str(TABLE6.relative_to(PROJECT)).replace("\\", "/"),
                    "sha256": sha256(TABLE6),
                },
                {
                    "alias": "BLI-AGE-REGION-SEX",
                    "name": "勞工保險人數－按年齡、地區及性別",
                    "url": "https://data.gov.tw/dataset/162821",
                    "snapshot": str(AGE_COUNTS.relative_to(PROJECT)).replace("\\", "/"),
                    "sha256": sha256(AGE_COUNTS),
                },
                {
                    "alias": "BLI-PENSION-REGION-INDUSTRY-SIZE",
                    "name": "勞工退休金提繳統計年報－按地區、行業及規模別",
                    "url": mol["meta"]["sourceUrl"],
                    "snapshot": str(MOL_CONTEXT.relative_to(PROJECT)).replace("\\", "/"),
                    "sha256": sha256(MOL_CONTEXT),
                },
            ],
            "generatedAt": datetime.now(timezone.utc).isoformat(),
        },
        "rows": rows,
        "validation": {
            "availableIndustryCount": len(available_names),
            "displayedIndustryCount": len(rows),
            "unavailableIndustryCount": len(UNAVAILABLE_INDUSTRIES),
            "reaggregation": reaggregation,
            "pclmSelectedLambda": pclm_diagnostics["lambda"],
            "pclmPostCalibrationMaxMarginError": round(pclm_diagnostics["post_calibration_max_margin_error"], 9),
            "native2529AnchorMatchesTable6": True,
            "methodEnvelopeIsConfidenceInterval": False,
        },
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(json.dumps({
        "output": str(OUTPUT),
        "rows": len(rows),
        "available": len(available_names),
        "maxReaggregationError": max(item["absoluteError"] for item in reaggregation.values()),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
