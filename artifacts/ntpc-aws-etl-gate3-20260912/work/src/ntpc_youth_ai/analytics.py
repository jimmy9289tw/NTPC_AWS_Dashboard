"""Curate policy indicators and transparent low-complexity forecasts."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
from statistics import mean
from typing import Any, Sequence

from openpyxl import load_workbook

from .pipeline import acquired_path


TAIPEI = timezone(timedelta(hours=8))
AGE_GROUPS = (
    (15, 17, "AGE-OBS-15-17", False),
    (18, 24, "AGE-CORE-18-24", True),
    (25, 29, "AGE-CORE-25-29", True),
    (30, 35, "AGE-CORE-30-35", True),
    (36, 40, "AGE-EXT-36-40", False),
)


def _json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _rows(path: Path) -> list[dict[str, Any]]:
    payload = _json(path)
    if isinstance(payload, list):
        return payload
    return payload["responseData"]


def _number(value: Any) -> float:
    if value in (None, "", "－", "-"):
        raise ValueError("missing numeric value")
    return float(str(value).replace(",", "").strip())


def _gregorian_roc_period(period: str) -> str:
    if len(period) != 5 or not period.isdigit():
        return period
    return f"{int(period[:3]) + 1911:04d}-{period[3:]}"


def curate_population(path: Path) -> dict[str, Any]:
    payload = _json(path)
    period = str(payload["period"])
    rows = [row for row in payload["responseData"] if str(row.get("site_id", "")).startswith("新北市")]
    if not rows:
        raise ValueError("MOI-POP1Y contains no New Taipei rows")
    groups: list[dict[str, Any]] = []
    for minimum, maximum, code, core in AGE_GROUPS:
        male = sum(int(row.get(f"people_age_{age:03d}_m", 0) or 0) for row in rows for age in range(minimum, maximum + 1))
        female = sum(int(row.get(f"people_age_{age:03d}_f", 0) or 0) for row in rows for age in range(minimum, maximum + 1))
        groups.append({"code": code, "minAge": minimum, "maxAge": maximum, "core": core, "male": male, "female": female, "total": male + female})
    total_population = sum(int(row.get("people_total", 0) or 0) for row in rows)
    core_total = sum(group["total"] for group in groups if group["core"])
    return {
        "period": _gregorian_roc_period(period),
        "sourcePeriod": period,
        "geography": "新北市",
        "geographyRole": "residence",
        "population": "戶籍人口",
        "totalPopulation": total_population,
        "coreYouth18To35": core_total,
        "coreSharePercent": round(core_total / total_population * 100, 4),
        "groups": groups,
    }


def linear_forecast(points: Sequence[tuple[int, float]], *, window: int = 8) -> dict[str, Any]:
    if len(points) < 4:
        raise ValueError("at least four observations are required")
    ordered = sorted(points)[-window:]

    def fit(sample: Sequence[tuple[int, float]]) -> tuple[float, float]:
        xs = [float(x) for x, _ in sample]
        ys = [float(y) for _, y in sample]
        xbar, ybar = mean(xs), mean(ys)
        denominator = sum((x - xbar) ** 2 for x in xs)
        slope = sum((x - xbar) * (y - ybar) for x, y in zip(xs, ys)) / denominator if denominator else 0.0
        return ybar - slope * xbar, slope

    backtest_errors: list[float] = []
    start = max(3, len(ordered) - 4)
    for index in range(start, len(ordered)):
        train = ordered[:index]
        intercept, slope = fit(train)
        predicted = intercept + slope * ordered[index][0]
        backtest_errors.append(abs(predicted - ordered[index][1]))
    intercept, slope = fit(ordered)
    target_year = ordered[-1][0] + 1
    forecast = intercept + slope * target_year
    mae = mean(backtest_errors) if backtest_errors else 0.0
    return {
        "method": "ordinary_least_squares_trailing_window",
        "trainingYears": [ordered[0][0], ordered[-1][0]],
        "observationCount": len(ordered),
        "targetYear": target_year,
        "forecast": round(forecast, 2),
        "backtestMAE": round(mae, 2),
        "indicativeRange": [round(forecast - mae, 2), round(forecast + mae, 2)],
        "status": "SCENARIO_NOT_OFFICIAL_FORECAST",
        "caveat": "線性外推只作需求觀察情境；不得視為政策成效或官方預測。",
    }


def curate_education(path: Path) -> dict[str, Any]:
    rows = sorted(_rows(path), key=lambda row: int(row["field1"]))
    series = [
        {"year": int(row["field1"]), "male": _number(row["itemvalue2"]), "female": _number(row["itemvalue3"])}
        for row in rows
    ]
    return {
        "measure": "15歲以上人口教育程度結構－大專及以上",
        "geography": "新北市",
        "geographyRole": "residence",
        "population": "15歲以上人口",
        "unit": "%",
        "series": series,
        "latest": series[-1],
        "forecast": {
            "male": linear_forecast([(row["year"], row["male"]) for row in series]),
            "female": linear_forecast([(row["year"], row["female"]) for row in series]),
        },
        "alignmentStatus": "PARTIAL_OR_CONFLICT",
        "limitation": "來源非18–35青年專屬資料，只能作教育環境趨勢。",
    }


def curate_unemployment(path: Path) -> dict[str, Any]:
    rows = sorted(_rows(path), key=lambda row: int(row["field1"]))
    series = [
        {"year": int(row["field1"]), "male": _number(row["item value4"]), "female": _number(row["item value5"])}
        for row in rows
    ]
    return {
        "measure": "25–29歲失業率",
        "geography": "新北市",
        "geographyRole": "residence",
        "population": "勞動力統計母體",
        "unit": "%",
        "series": series,
        "latest": series[-1],
        "forecast": {
            "male": linear_forecast([(row["year"], row["male"]) for row in series]),
            "female": linear_forecast([(row["year"], row["female"]) for row in series]),
        },
        "alignmentStatus": "EXACT_FOR_25_29_ONLY",
        "limitation": "18–24與30–35無法由來源年齡帶精確重建。",
    }


def _sheet_year(sheet_name: str) -> int:
    roc = int(sheet_name.split("(", 1)[1].split("年", 1)[0])
    return roc + 1911


def curate_salary_education(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    series: list[dict[str, Any]] = []
    levels = ["全體", "國中及以下", "高級中等", "專科及大學", "研究所"]
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        target = None
        for row in sheet.iter_rows(min_row=1, max_col=11, values_only=True):
            if str(row[0]).replace(" ", "") == "25-29歲":
                target = row
                break
        if target:
            series.append({
                "year": _sheet_year(sheet_name),
                "ageBand": "25–29",
                "mean": {level: _number(target[index + 1]) for index, level in enumerate(levels)},
                "median": {level: _number(target[index + 6]) for index, level in enumerate(levels)},
            })
    series.sort(key=lambda item: item["year"])
    return {
        "measure": "25–29歲全年總薪資－按教育程度",
        "geography": "全國",
        "geographyRole": "national",
        "population": "工業及服務業受僱員工",
        "unit": "萬元",
        "series": series,
        "latest": series[-1],
        "alignmentStatus": "PARTIAL_OR_CONFLICT",
        "limitation": "全國受僱員工資料；教育分層差異是描述性關聯，不是因果效果。",
    }


def curate_salary_location(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    series: list[dict[str, Any]] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        target = None
        for row in sheet.iter_rows(min_row=1, max_col=13, values_only=True):
            if str(row[0]).strip() == "新北市":
                target = row
                break
        if target:
            series.append({"year": _sheet_year(sheet_name), "mean": _number(target[4]), "median": _number(target[12])})
    series.sort(key=lambda item: item["year"])
    return {
        "measure": "新北市工作場所25–29歲全年總薪資",
        "geography": "新北市",
        "geographyRole": "workplace",
        "population": "本國籍全時受僱員工",
        "unit": "萬元",
        "series": series,
        "latest": series[-1],
        "forecast": {
            "mean": linear_forecast([(row["year"], row["mean"]) for row in series], window=6),
            "median": linear_forecast([(row["year"], row["median"]) for row in series], window=6),
        },
        "alignmentStatus": "EXACT_FOR_25_29_WORKPLACE_ONLY",
        "limitation": "工作場所口徑不可視為居住在新北市青年的薪資。",
    }


def build_policy_indicators(root: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    output = {
        "schemaVersion": "1.0.0",
        "generatedAt": datetime.now(TAIPEI).isoformat(),
        "runId": manifest["runId"],
        "population": curate_population(acquired_path(root, manifest, "MOI-POP1Y")),
        "education": curate_education(acquired_path(root, manifest, "NTPC-EDU15P")),
        "unemployment": curate_unemployment(acquired_path(root, manifest, "NTPC-UNEMP")),
        "salaryByEducation": curate_salary_education(acquired_path(root, manifest, "STAT-WAGE-EDU")),
        "salaryByWorkplace": curate_salary_location(acquired_path(root, manifest, "STAT-WAGE-LOC")),
        "methodRules": {
            "coreYouth": "18–24 + 25–29 + 30–35；僅單一年齡精確加總",
            "crossDomain": "教育與薪資僅並列描述，不進行個體連結或因果推論",
            "forecast": "尾端視窗線性迴歸＋滾動回測MAE；標記情境而非官方預測",
        },
    }
    curated_dir = root / "data" / "curated"
    curated_dir.mkdir(parents=True, exist_ok=True)
    serialized = json.dumps(output, ensure_ascii=False, indent=2) + "\n"
    (curated_dir / f"policy-indicators-{manifest['runId']}.json").write_text(serialized, encoding="utf-8")
    (curated_dir / "policy-indicators-latest.json").write_text(serialized, encoding="utf-8")
    return output
