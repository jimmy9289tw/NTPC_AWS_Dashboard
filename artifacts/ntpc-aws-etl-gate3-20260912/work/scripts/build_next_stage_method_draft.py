from __future__ import annotations

import json
import math
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT = Path(r"D:\github\work\019fe05e-e1a3-71f2-840e-c026180d9474\NtpcYouthAI")
DATA_PATH = PROJECT / "tmp" / "next_stage_report_data.json"
OUT_DIR = PROJECT / "deliverables" / "NTPC_Youth_Demographic_Education_Wage_Method_Draft_V0.1_20260825"
ASSET_DIR = OUT_DIR / "assets"
OUTPUT = OUT_DIR / "新北市青年18-35歲人口婚姻教育與薪資計算及視覺化方法草稿_V0.1.docx"

NAVY = "173A5E"
TEAL = "168A8A"
BLUE = "2E6F9E"
SKY = "E8F2F7"
PALE_TEAL = "E7F4F2"
AMBER = "D9912B"
PALE_AMBER = "FFF3DF"
RED = "A33A3A"
PALE_RED = "FBE9E7"
GREEN = "2F7D54"
PALE_GREEN = "EAF4EC"
GRAY = "5A6773"
LIGHT_GRAY = "EEF1F4"
DARK = "1F2933"
WHITE = "FFFFFF"

FONT_CJK = "Microsoft JhengHei"
FONT_LATIN = "Aptos"
MONO = "Consolas"


SOURCES = {
    "S01": ("內政部戶政司｜各村里單一年齡人口", "https://data.gov.tw/dataset/77132"),
    "S02": ("內政部戶政司｜婚姻狀態（含同婚）", "https://data.gov.tw/dataset/117986"),
    "S03": ("內政部戶政司｜年齡×婚姻×教育交叉表", "https://data.gov.tw/dataset/117988"),
    "S04": ("主計總處｜工作場所縣市×年齡全年總薪資（表6）", "https://www.stat.gov.tw/News_Content.aspx?n=4580&s=232642"),
    "S05": ("主計總處｜114年全年薪資與生產力統計", "https://www.stat.gov.tw/News_Content.aspx?n=2724&s=235863"),
    "S06": ("主計總處｜113年全年總薪資中位數與分布", "https://www.stat.gov.tw/News_Content.aspx?n=2716&s=235514"),
    "S07": ("主計總處｜薪資中位數統計編製方法", "https://www.stat.gov.tw/public/Data/61111105359N4GQ42NS.pdf"),
    "S08": ("主計總處｜普（抽）查第一類資料申請單", "https://www.stat.gov.tw/public/data/dgbas04/bc3/orderform/orderform1.pdf"),
    "S09": ("Rizzi、Gampe、Eilers（2015）PCLM", "https://doi.org/10.1093/aje/kwv020"),
    "S10": ("Deming、Stephan（1940）IPF", "https://doi.org/10.1214/aoms/1177731829"),
    "S11": ("Statistics Canada｜Sprague multipliers 應用", "https://www150.statcan.gc.ca/n1/pub/91-528-x/91-528-x2003001-eng.pdf"),
    "S12": ("主計總處｜114年人力運用調查報告", "https://www.stat.gov.tw/News_Content.aspx?n=3105&s=236084"),
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "1")


def set_repeat_on_each_page(row) -> None:
    set_repeat_table_header(row)


def set_cell_text(cell, text: str, bold=False, color=DARK, size=8.3, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = FONT_LATIN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_table_borders(table, color="CBD3DA", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_col_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths=None, header_fill=NAVY,
              font_size=8.0, status_col=None) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_on_each_page(hdr)
    for idx, text in enumerate(headers):
        set_cell_shading(hdr.cells[idx], header_fill)
        set_cell_text(hdr.cells[idx], text, bold=True, color=WHITE, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ridx, values in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(values):
            fill = WHITE if ridx % 2 == 0 else "F7F9FB"
            if status_col is not None and cidx == status_col:
                value_text = str(value)
                if "HOLD" in value_text or "停止" in value_text:
                    fill = PALE_RED
                elif "模型" in value_text or "條件" in value_text:
                    fill = PALE_AMBER
                elif "精確" in value_text or "可發布" in value_text:
                    fill = PALE_GREEN
            set_cell_shading(cells[cidx], fill)
            set_cell_text(cells[cidx], value, size=font_size)
    if widths:
        set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    clr = OxmlElement("w:color")
    clr.set(qn("w:val"), color)
    r_pr.append(clr)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CJK)
    r_pr.append(fonts)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_source_ref(paragraph, code: str) -> None:
    label, url = SOURCES[code]
    paragraph.add_run(" ")
    add_hyperlink(paragraph, f"[{code}]", url)


def set_run_font(run, size=9.5, bold=None, color=DARK, name=FONT_LATIN) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK if name != MONO else MONO)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_body(doc: Document, text: str, bold_prefix: str | None = None, refs: list[str] | None = None,
             color=DARK, size=9.5, after=4) -> object:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=size, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=size, color=color)
    else:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color)
    for ref in refs or []:
        add_source_ref(p, ref)
    return p


def add_bullet(doc: Document, text: str, level=0, refs: list[str] | None = None) -> object:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, size=9.2)
    for ref in refs or []:
        add_source_ref(p, ref)
    return p


def add_heading(doc: Document, text: str, level=1, kicker=None) -> object:
    if kicker:
        p0 = doc.add_paragraph()
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(kicker.upper())
        set_run_font(r0, size=7.8, bold=True, color=TEAL)
        r0.font.all_caps = True
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_callout(doc: Document, title: str, body: str, fill=PALE_TEAL, accent=TEAL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=160, start=180, bottom=160, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=9.2, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    r2 = p2.add_run(body)
    set_run_font(r2, size=9.0, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_formula(doc: Document, label: str, formula: str, definitions: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F6F9")
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(label)
    set_run_font(r, size=8.4, bold=True, color=TEAL)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(2 if definitions else 0)
    r2 = p2.add_run(formula)
    set_run_font(r2, size=9.0, color=NAVY, name=MONO)
    if definitions:
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(definitions)
        set_run_font(r3, size=8.2, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=7.6, color=GRAY)
    r.italic = True


def add_picture(doc: Document, path: Path, width_cm: float, alt: str, caption: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Cm(width_cm))
    shape._inline.docPr.set("descr", alt)
    if caption:
        add_caption(doc, caption)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=7.5, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 頁")
    set_run_font(run2, size=7.5, color=GRAY)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color in (("Title", 28, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 12.5, TEAL), ("Heading 3", 10.5, NAVY)):
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(9.2)

    if "Small Note" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(7.5)
        style.font.color.rgb = RGBColor.from_string(GRAY)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p.add_run("NTPC YOUTH AI  /  METHOD NOTE 0.1")
    set_run_font(r1, size=7.2, bold=True, color=TEAL)
    p2 = header.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("方法草稿｜不得直接標示為官方統計")
    set_run_font(r2, size=7.0, color=GRAY)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


FONT_REGULAR_PATH = r"C:\Windows\Fonts\msjh.ttc"
FONT_BOLD_PATH = r"C:\Windows\Fonts\msjhbd.ttc"


def pil_font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_BOLD_PATH if bold else FONT_REGULAR_PATH, size=size)


def draw_centered(draw: ImageDraw.ImageDraw, box, text: str, font, fill, spacing=6) -> None:
    left, top, right, bottom = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=spacing)
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(((left + right - width) / 2, (top + bottom - height) / 2), text, font=font, fill=fill, align="center", spacing=spacing)


def draw_arrow(draw: ImageDraw.ImageDraw, start, end, color="#5A6773", width=5) -> None:
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    for delta in (math.radians(150), math.radians(-150)):
        point = (end[0] + length * math.cos(angle + delta), end[1] + length * math.sin(angle + delta))
        draw.line([end, point], fill=color, width=width)


def draw_flow(path: Path, title: str, boxes: list[tuple[str, str]], arrows: list[tuple[int, int]], note: str) -> None:
    width, height = 1800, 620
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 38), title, font=pil_font(42, True), fill="#173A5E")
    n = len(boxes)
    box_w, box_h = 330, 230
    side_margin = box_w / 2 + 55
    centers = [side_margin + i * ((width - 2 * side_margin) / max(n - 1, 1)) for i in range(n)]
    y0 = 190
    for i, ((head, body), cx) in enumerate(zip(boxes, centers)):
        rect = (int(cx - box_w / 2), y0, int(cx + box_w / 2), y0 + box_h)
        fill = "#E8F2F7" if i in (0, n - 1) else "#E7F4F2"
        draw.rounded_rectangle(rect, radius=14, fill=fill, outline="#168A8A", width=4)
        draw_centered(draw, (rect[0] + 12, rect[1] + 28, rect[2] - 12, rect[1] + 95), head, pil_font(27, True), "#173A5E")
        draw_centered(draw, (rect[0] + 18, rect[1] + 100, rect[2] - 18, rect[3] - 20), body, pil_font(22), "#384857")
    for start, end in arrows:
        if start == end:
            continue
        y = y0 + box_h / 2
        if start < end:
            p1 = (centers[start] + box_w / 2 + 10, y)
            p2 = (centers[end] - box_w / 2 - 10, y)
        else:
            p1 = (centers[start] - box_w / 2 - 10, y)
            p2 = (centers[end] + box_w / 2 + 10, y)
        draw_arrow(draw, p1, p2)
    wrapped = "\n".join(textwrap.wrap(note, width=70))
    draw.multiline_text((55, 500), wrapped, font=pil_font(22), fill="#5A6773", spacing=6)
    image.save(path)


def make_figures(data: dict) -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    figures = {}

    pop = data["population"]["target_bands"]
    bands = ["18-24", "25-29", "30-35"]
    male = [next(x["population"] for x in pop if x["age_band"] == b and x["sex"] == "男") for b in bands]
    female = [next(x["population"] for x in pop if x["age_band"] == b and x["sex"] == "女") for b in bands]
    image = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "114年12月底新北市青年戶籍人口｜三個核心年齡組", font=pil_font(42, True), fill="#173A5E")
    left, top, right, bottom = 150, 170, 1720, 730
    max_v = max(male + female) * 1.12
    for tick in range(0, 181000, 30000):
        y = bottom - (tick / max_v) * (bottom - top)
        draw.line((left, y, right, y), fill="#E4E9ED", width=2)
        draw.text((45, y - 13), f"{tick/1000:.0f}千", font=pil_font(20), fill="#5A6773")
    group_space = (right - left) / len(bands)
    bar_w = 100
    for i, band in enumerate(bands):
        cx = left + group_space * (i + 0.5)
        for value, offset, color, label in ((male[i], -62, "#2E6F9E", "男"), (female[i], 62, "#D06B86", "女")):
            bar_h = value / max_v * (bottom - top)
            rect = (cx + offset - bar_w / 2, bottom - bar_h, cx + offset + bar_w / 2, bottom)
            draw.rectangle(rect, fill=color)
            text = f"{value/1000:.1f}千"
            bbox = draw.textbbox((0, 0), text, font=pil_font(20, True))
            draw.text((cx + offset - (bbox[2]-bbox[0])/2, bottom - bar_h - 32), text, font=pil_font(20, True), fill=color)
        bbox = draw.textbbox((0, 0), band, font=pil_font(25, True))
        draw.text((cx - (bbox[2]-bbox[0])/2, bottom + 25), band, font=pil_font(25, True), fill="#384857")
    draw.rectangle((1330, 105, 1360, 130), fill="#2E6F9E")
    draw.text((1375, 100), "男", font=pil_font(22), fill="#384857")
    draw.rectangle((1455, 105, 1485, 130), fill="#D06B86")
    draw.text((1500, 100), "女", font=pil_font(22), fill="#384857")
    p = ASSET_DIR / "population_age_sex_11412.png"
    image.save(p)
    figures["population"] = p

    marr = data["marriage"]
    bands4 = ["18-24", "25-29", "30-35", "18-35"]
    with_spouse = [next(x["share_pct"] for x in marr if x["age_band"] == b and x["status"] == "目前有配偶") for b in bands4]
    without_spouse = [100 - v for v in with_spouse]
    image = Image.new("RGB", (1800, 850), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "114年底婚姻狀態｜衍生二分類（合計）", font=pil_font(42, True), fill="#173A5E")
    left, right = 170, 1700
    bar_y, bar_h = 260, 300
    gap = 55
    bar_w = int((right - left - gap * 3) / 4)
    for i, (band, v) in enumerate(zip(bands4, with_spouse)):
        x0 = left + i * (bar_w + gap)
        spouse_h = bar_h * v / 100
        draw.rectangle((x0, bar_y, x0 + bar_w, bar_y + bar_h), fill="#D9E3EA")
        draw.rectangle((x0, bar_y + bar_h - spouse_h, x0 + bar_w, bar_y + bar_h), fill="#168A8A")
        pct = f"{v:.1f}%"
        bbox = draw.textbbox((0, 0), pct, font=pil_font(24, True))
        draw.text((x0 + (bar_w-(bbox[2]-bbox[0]))/2, bar_y + bar_h - spouse_h - 38), pct, font=pil_font(24, True), fill="#168A8A")
        bbox2 = draw.textbbox((0, 0), band, font=pil_font(25, True))
        draw.text((x0 + (bar_w-(bbox2[2]-bbox2[0]))/2, bar_y + bar_h + 25), band, font=pil_font(25, True), fill="#384857")
    draw.rectangle((420, 690, 452, 716), fill="#168A8A")
    draw.text((468, 683), "目前有配偶", font=pil_font(22), fill="#384857")
    draw.rectangle((760, 690, 792, 716), fill="#D9E3EA")
    draw.text((808, 683), "目前無配偶", font=pil_font(22), fill="#384857")
    draw.text((1150, 683), "25–29為官方精確；其餘為模型估計", font=pil_font(21), fill="#A86B19")
    p = ASSET_DIR / "marriage_binary_114.png"
    image.save(p)
    figures["marriage"] = p

    edu = [x for x in data["education"]["official_5y"] if x["source_age_band"] == "25-29"]
    sex_order = ["合計", "男", "女"]
    edu_order = ["研究所", "大學", "專科", "高中職", "國中及以下"]
    colors = ["#173A5E", "#2E6F9E", "#65A9C7", "#E3A548", "#D9E3EA"]
    image = Image.new("RGB", (1800, 850), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "114年底25–29歲戶籍人口教育程度｜官方精確值", font=pil_font(42, True), fill="#173A5E")
    left, right, bar_y, bar_h = 250, 1600, 220, 340
    gap = 180
    bar_w = int((right - left - gap * 2) / 3)
    for i, sex in enumerate(sex_order):
        x0 = left + i * (bar_w + gap)
        y = bar_y + bar_h
        for category, color in reversed(list(zip(edu_order, colors))):
            value = next(x["share_pct"] for x in edu if x["sex"] == sex and x["education"] == category)
            h = bar_h * value / 100
            draw.rectangle((x0, y - h, x0 + bar_w, y), fill=color)
            y -= h
        bbox = draw.textbbox((0, 0), sex, font=pil_font(25, True))
        draw.text((x0 + (bar_w-(bbox[2]-bbox[0]))/2, bar_y + bar_h + 25), sex, font=pil_font(25, True), fill="#384857")
    legend_x = 180
    for idx, (category, color) in enumerate(zip(edu_order, colors)):
        x0 = legend_x + idx * 300
        draw.rectangle((x0, 700, x0 + 32, 726), fill=color)
        draw.text((x0 + 46, 693), category, font=pil_font(21), fill="#384857")
    p = ASSET_DIR / "education_25_29_114.png"
    image.save(p)
    figures["education"] = p

    wage = [x for x in data["wage"] if x["source_age_band"] == "25-29"]
    years = sorted({int(x["year"]) for x in wage})
    avg = [next(x["annual_salary_10k_ntd"] for x in wage if int(x["year"]) == y and x["measure"] == "平均數") for y in years]
    med = [next(x["annual_salary_10k_ntd"] for x in wage if int(x["year"]) == y and x["measure"] == "中位數") for y in years]
    image = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "新北工作地25–29歲全年總薪資｜官方原生年齡組", font=pil_font(42, True), fill="#173A5E")
    left, top, right, bottom = 150, 190, 1700, 700
    min_v, max_v = 44, 62
    def pos(year, value):
        x = left + (year - years[0]) / (years[-1] - years[0]) * (right - left)
        y = bottom - (value - min_v) / (max_v - min_v) * (bottom - top)
        return x, y
    for tick in range(44, 63, 3):
        y = pos(years[0], tick)[1]
        draw.line((left, y, right, y), fill="#E4E9ED", width=2)
        draw.text((65, y - 13), f"{tick}", font=pil_font(20), fill="#5A6773")
    for series, color in ((avg, "#2E6F9E"), (med, "#D9912B")):
        points = [pos(year, value) for year, value in zip(years, series)]
        draw.line(points, fill=color, width=7)
        for (x, y), value in zip(points, series):
            draw.ellipse((x-9, y-9, x+9, y+9), fill=color)
            label_y = y - 42 if color == "#2E6F9E" else y + 16
            draw.text((x - 22, label_y), f"{value:.1f}", font=pil_font(19, True), fill=color)
    for year in years:
        x, _ = pos(year, min_v)
        draw.text((x - 22, bottom + 25), str(year), font=pil_font(22, True), fill="#384857")
    draw.text((150, 760), "單位：萬元／年", font=pil_font(21), fill="#5A6773")
    draw.line((520, 772, 570, 772), fill="#2E6F9E", width=7)
    draw.text((585, 758), "平均數", font=pil_font(22), fill="#384857")
    draw.line((760, 772, 810, 772), fill="#D9912B", width=7)
    draw.text((825, 758), "中位數", font=pil_font(22), fill="#384857")
    p = ASSET_DIR / "wage_25_29_trend_108_113.png"
    image.save(p)
    figures["wage"] = p

    p = ASSET_DIR / "pclm_ipf_flow.png"
    draw_flow(
        p,
        "五歲年齡組轉為18–35歲：受約束的年齡拆分",
        [("官方五歲組", "年齡×分類×性別\n行政精確格"), ("PCLM", "估計非負、平滑\n單歲初始形狀"), ("IPF", "校準單歲人口與\n官方五歲組邊際"), ("目標年齡組", "18–24／25–29\n30–35／18–35")],
        [(0, 1), (1, 2), (2, 3)],
        "25–29完整落在原始五歲組，直接加總；只有跨18、19、35歲邊界的組別才使用模型。",
    )
    figures["pclm"] = p

    p = ASSET_DIR / "wage_decision_flow.png"
    draw_flow(
        p,
        "薪資估計的停止規則",
        [("同母體資料", "新北工作地×18–35\n全年總薪資"), ("平均數", "需要薪資總額或\n各組均數＋受僱權數"), ("中位數", "需要個體分布或\n薪資級距×人數"), ("發布決策", "資料足夠才計算\n否則HOLD")],
        [(0, 1), (0, 2), (1, 3), (2, 3)],
        "投保薪資、稅務薪資所得、每月主要工作收入的母體與概念不同，只能做輔助診斷或模型先驗。",
    )
    figures["wage_flow"] = p
    return figures


def page_break(doc: Document) -> None:
    doc.add_page_break()


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("研究方法草稿  /  GATE 4")
    set_run_font(r, size=9, bold=True, color=TEAL)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(30)
    title.paragraph_format.space_after = Pt(12)
    r = title.add_run("新北市青年18–35歲\n人口、婚姻、教育與薪資")
    set_run_font(r, size=28, bold=True, color=NAVY)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(22)
    r2 = sub.add_run("計算方法、資料可得性與儀表板視覺化規格")
    set_run_font(r2, size=15, bold=True, color=TEAL)

    add_callout(
        doc,
        "本版結論",
        "人口可直接精確加總；婚姻與教育的25–29歲可直接精確加總，18–24、30–35及18–35需以PCLM建立單歲形狀、再用IPF校準官方邊際；薪資目前僅25–29歲可用新北工作地官方原生值，18–35平均數與中位數均須維持HOLD，直到取得同母體受僱人數權數與薪資分布。",
        fill=PALE_TEAL,
        accent=TEAL,
    )

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    labels = ["文件版本", "資料截點", "地理範圍", "核心年齡", "證據狀態"]
    vals = ["V0.1（方法草稿）", "2026-08-25查核；人口／婚姻／教育以民國114年為主", "新北市；戶籍地與工作地分開標示", "18–24、25–29、30–35；合計18–35", "精確值／官方原生值／模型估計／停止估計四級"]
    for i, (label, val) in enumerate(zip(labels, vals)):
        set_cell_shading(table.cell(i, 0), NAVY)
        set_cell_text(table.cell(i, 0), label, bold=True, color=WHITE, size=8.5)
        set_cell_shading(table.cell(i, 1), "F5F8FA")
        set_cell_text(table.cell(i, 1), val, size=8.8)
    set_col_widths(table, [3.4, 13.4])
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("設計預設：standard_business_brief｜頁首：memo_masthead")
    set_run_font(r, size=7.5, color=GRAY)
    p2 = doc.add_paragraph()
    r2 = p2.add_run("用途：研究、系統設計與人工查核；未經裁示不得標示為新北市政府正式統計。")
    set_run_font(r2, size=8.2, bold=True, color=RED)


def add_executive_summary(doc: Document, data: dict) -> None:
    page_break(doc)
    add_heading(doc, "決策摘要", kicker="EXECUTIVE SUMMARY")
    add_callout(doc, "可立即進入儀表板原型的資料", "人口18–35可發布為官方行政精確值；婚姻與教育25–29可發布為官方行政精確值。婚姻邊界組已有模型結果但需人工覆核；教育邊界組已有完整方法與精確驗證組，尚待正式執行模型。薪資只發布官方原生25–29歲年值與來源年齡組，不生成18–35假精確值。", fill=PALE_GREEN, accent=GREEN)

    rows = [
        ["人口", "18–24／25–29／30–35／18–35", "MOI單一年齡×性別", "官方行政精確值", "可發布"],
        ["婚姻", "25–29", "MOI五歲年齡×婚姻", "官方行政精確值", "可發布"],
        ["婚姻", "18–24／30–35／18–35", "五歲婚姻＋單歲人口", "PCLM＋IPF模型估計", "覆核後發布"],
        ["教育", "25–29", "MOI五歲年齡×教育", "官方行政精確值", "可發布"],
        ["教育", "18–24／30–35／18–35", "五歲教育＋單歲人口", "PCLM＋IPF模型估計", "待執行／覆核"],
        ["薪資平均數", "25–29（108–113）", "DGBAS表6，新北工作地", "官方來源原生值", "可發布"],
        ["薪資平均數", "18–35", "缺同母體年齡薪資總額與人數", "HOLD", "不可估"],
        ["薪資中位數", "18–35", "缺個體分布或薪資級距×人數", "HOLD", "不可估"],
    ]
    add_table(doc, ["主題", "目標年齡", "目前資料", "值的身分", "Gate判定"], rows,
              widths=[2.4, 3.4, 4.7, 3.4, 2.7], font_size=7.8, status_col=4)
    add_body(doc, "114年雖已有全國全年總薪資平均數，但沒有新北工作地×目標年齡的同粒度交叉表，因此不得用全國值覆寫或補成新北青年值。", refs=["S05"])
    add_body(doc, "新北25–29歲官方全年總薪資在113年為平均59.9萬元、中位52.4萬元；這是目前最適合第一階段驗證的薪資基準。", refs=["S04"])

    add_heading(doc, "共同資料治理規則", level=2)
    for text in (
        "值來源必填：OFFICIAL_ADMIN_EXACT、OFFICIAL_SOURCE_NATIVE、MODEL_ESTIMATE_FROM_OFFICIAL_MARGINS或HOLD_NO_ESTIMATE。",
        "時間基準必填：戶籍主題為END_OF_YEAR_12_31；全年總薪資為FULL_CALENDAR_YEAR；兩者不可在同一公式中互當母數。",
        "地理角色必填：REGISTERED_RESIDENCE（戶籍地）與WORKPLACE（工作地）分開；只可在同一儀表板並列，不可直接相除。",
        "人口母數不可混用：戶籍人口P_REG不取代人力資源調查的民間人口P_LFS，也不取代薪資統計的受僱員工N_WAGE。",
        "模型估計必須伴隨方法碼、重聚合誤差、敏感度包絡與人工覆核狀態。",
    ):
        add_bullet(doc, text)


def add_population_section(doc: Document, data: dict, figures: dict[str, Path]) -> None:
    page_break(doc)
    add_heading(doc, "1｜人口年齡結構", kicker="POPULATION")
    add_body(doc, "定義：戶籍登記現住人口，依戶籍地統計；不是實際常住人口、工作地人口或人力資源調查民間人口。來源具村里、單一年齡與性別欄位，可直接產製18–35及三個子群，不需要Sprague、PCLM或IPF。", refs=["S01"])
    add_callout(doc, "母數代號", "P_REG(a,s,t)：時間t、單一年齡a、性別s的新北市戶籍登記現住人口。勞參率與失業率仍必須使用人力資源調查的P_LFS與LF，不以P_REG代入。", fill=SKY, accent=NAVY)

    add_heading(doc, "1.1 計算公式", level=2)
    add_formula(doc, "P-01｜目標年齡人口", "P(A,s,t) = Σ[a∈A] Σ[v∈NTPC] P(v,a,s,t)", "A∈{18–24, 25–29, 30–35, 18–35}；v為新北市村里。")
    add_formula(doc, "P-02｜性別合計", "P(A,合計,t) = P(A,男,t) + P(A,女,t)")
    add_formula(doc, "P-03｜年齡組占核心青年比", "Share(A,t) = P(A,t) / P(18–35,t) × 100%")
    add_formula(doc, "P-04｜年對年變動", "YoY(A,t) = [P(A,t) / P(A,t−1) − 1] × 100%")
    add_body(doc, "年度主圖採每年12月31日期末存量；月度圖可使用API的各月期末存量，做月對月或同月年對年比較，但不得把年度值平均拆成12個月。")

    pop = data["population"]["target_bands"]
    def pv(b, s="合計"):
        key = b.replace("–", "-")
        return next(x["population"] for x in pop if x["age_band"] == key and x["sex"] == s)
    add_heading(doc, "1.2 114年12月底驗算範例", level=2)
    rows = [[b, f"{pv(b,'男'):,}", f"{pv(b,'女'):,}", f"{pv(b):,}", "OFFICIAL_ADMIN_EXACT"] for b in ("18–24", "25–29", "30–35", "18–35")]
    add_table(doc, ["年齡", "男", "女", "合計", "值來源"], rows, widths=[2.7, 3.0, 3.0, 3.2, 4.5], status_col=4)
    add_formula(doc, "P-05｜總量一致性", f"261,472 + 250,128 + 334,338 = {pv('18-35'):,}", "三個互斥子群合計必須等於18–35歲總量。")
    add_picture(doc, figures["population"], 16.5, "114年12月底新北市18至24、25至29、30至35歲戶籍人口男女分組長條圖。", "圖1　人口圖表原型；來源：MOI-POP1Y，官方行政精確值。")

    add_heading(doc, "1.3 儀表板呈現規格", level=2)
    rows = [
        ["人口金額卡", "18–35總數、三子群人數", "年、性別、行政區", "顯示期末日期與ADMIN_EXACT"],
        ["分組長條圖", "三子群×性別", "年、行政區", "適合年度比較；零軸起點"],
        ["人口金字塔", "單一年齡×性別", "年、行政區", "使用單歲原始值，不插值"],
        ["月度折線", "每月月底人口", "性別、子群", "可MoM與YoY；不得由年值拆月"],
        ["區域熱圖", "行政區×青年占比", "年、性別", "分母為同區同年全體或18–35，必須標明"],
    ]
    add_table(doc, ["圖表", "指標", "篩選", "查核要求"], rows, widths=[2.5, 4.2, 3.1, 6.4], font_size=7.9)


def add_pclm_explanation(doc: Document, classification: str) -> None:
    add_formula(doc, "M-01｜PCLM觀測模型", "Y[g,c,s] ~ Poisson( μ[g,c,s] ),  μ = C · exp(Bθ)", f"g為官方五歲組；c為{classification}；C為單歲到五歲的聚合矩陣；B為B-spline基底。")
    add_formula(doc, "M-02｜平滑估計", "θ̂ = argmaxθ { ℓ(θ;Y) − (λ/2) ||D²θ||² }", "λ控制平滑程度；本計畫以內部留一組Poisson偏差選擇，並保留收斂與有效自由度。")
    add_formula(doc, "M-03｜IPF單歲人口校準", "x[a,c] ← x[a,c] × P1Y[a] / Σc x[a,c]")
    add_formula(doc, "M-04｜IPF官方分類邊際校準", "x[a,c] ← x[a,c] × Y[g,c] / Σ[a∈g] x[a,c]")
    add_formula(doc, "M-05｜目標年齡彙總", "N[A,c,s] = Σ[a∈A] x[a,c,s]")
    add_body(doc, "PCLM適合把粗年齡組還原為平滑單歲分布；IPF則反覆校準列、欄，使估計格同時符合已知單歲人口與官方五歲分類總量。", refs=["S09", "S10"])
    add_body(doc, "Sprague可作固定係數敏感度基準，但可能產生負值；本案不以它作主模型，只把它或『組內人口比例拆分』用來檢查PCLM結果是否過度依賴平滑假設。", refs=["S11"])


def add_marriage_section(doc: Document, data: dict, figures: dict[str, Path]) -> None:
    page_break(doc)
    add_heading(doc, "2｜婚姻狀態", kicker="MARITAL STATUS")
    add_body(doc, "定義：年底戶籍登記婚姻狀態存量，按戶籍地、性別與五歲年齡組統計；不是該年結婚／離婚件數。官方原始類別先將同性／異性有偶與離婚子類合併為未婚、有偶、離婚或終止結婚、喪偶四類。", refs=["S02"])
    add_formula(doc, "MAR-01｜衍生二分類", "目前有配偶 = 有偶；目前無配偶 = 未婚 + 離婚或終止結婚 + 喪偶")
    add_callout(doc, "精確與估計的界線", "25–29完整落在官方五歲組，直接加總；18–24需取15–19中的18、19歲，30–35需取35–39中的35歲，因此兩者與18–35都需模型拆分。", fill=PALE_AMBER, accent=AMBER)
    add_picture(doc, figures["pclm"], 16.5, "五歲年齡組依序經過PCLM、IPF後彙總為目標年齡組的流程圖。", "圖2　婚姻與教育共用的年齡拆分流程。")

    add_heading(doc, "2.1 公式與驗證", level=2)
    add_pclm_explanation(doc, "婚姻類別")
    add_formula(doc, "MAR-02｜婚姻占比", "Share[A,c,s] = N[A,c,s] / Σc N[A,c,s] × 100%")
    add_formula(doc, "MAR-03｜方法敏感度", "Envelope[A,c] = [ min(N_PCLM+IPF, N_PROP+IPF), max(·) ]", "此範圍是方法敏感度包絡，不是抽樣信賴區間。")

    rows = []
    for band in ("18-24", "25-29", "30-35", "18-35"):
        for status in ("目前有配偶", "目前無配偶"):
            x = next(v for v in data["marriage"] if v["age_band"] == band and v["status"] == status)
            rows.append([band, status, f"{x['population']:,.0f}", f"{x['share_pct']:.2f}%", x["value_origin_class"], x["method_code"]])
    add_table(doc, ["年齡", "分類", "人數", "占比", "值來源", "方法"], rows,
              widths=[2.1, 2.8, 2.6, 2.0, 4.0, 3.7], font_size=7.5, status_col=4)
    add_formula(doc, "MAR-04｜114年總量驗證", "134,360.896 + 711,577.104 = 845,938", "18–35兩類合計精確回到同年單一年齡戶籍人口；IPF最大邊際誤差約10⁻⁸。")
    add_picture(doc, figures["marriage"], 16.2, "114年底18至24、25至29、30至35與18至35歲目前有配偶與無配偶占比百分之百堆疊圖。", "圖3　婚姻狀態建議採100%堆疊長條；25–29為精確值，其他組為模型估計。")

    add_heading(doc, "2.2 儀表板呈現規格", level=2)
    add_table(doc, ["圖表", "用途", "設計規則", "不應做的事"], [
        ["100%堆疊長條", "比較各年齡／年度婚姻結構", "四分類為主，二分類可切換", "避免用多張圓餅比較年度"],
        ["性別啞鈴圖", "比較男女性有偶占比差", "顯示百分點差", "不可只顯示人數"],
        ["年度斜率圖", "113→114占比變化", "同一方法、同一年底存量", "不可與事件件數混線"],
        ["敏感度誤差棒", "模型組方法包絡", "標『非信賴區間』", "25–29不得畫模型誤差"],
    ], widths=[3.0, 4.0, 4.8, 4.6], font_size=7.8)


def add_education_section(doc: Document, data: dict, figures: dict[str, Path]) -> None:
    page_break(doc)
    add_heading(doc, "3｜教育程度", kicker="EDUCATION")
    add_body(doc, "正式來源為內政部『15歲以上現住人口按性別、年齡、婚姻狀況及教育程度分』，population是每一列『年度×戶籍地×性別×五歲年齡×婚姻×教育』交叉格人數。114年原始檔共13,674,529列，本次串流核對新北市1,820,448列。", refs=["S03"])
    add_callout(doc, "資料衝突處理", "主計總處表27／32在縣市列中把教育邊際與年齡邊際並排，並非年齡×教育交叉表；不可直接回答25–29教育分布。表50雖有年齡×教育，但地理是全國，只可作模型種子或外部比較。正式新北值以MOI交叉格為主。", fill=PALE_RED, accent=RED)

    add_heading(doc, "3.1 分類與公式", level=2)
    add_table(doc, ["原始教育類別", "儀表板分類", "理由"], [
        ["博畢、碩畢", "研究所", "避免稀疏格；仍保留原始碼可回溯"],
        ["大畢", "大學", "維持原生類別"],
        ["專畢", "專科", "維持原生類別"],
        ["高中畢", "高中職", "以來源標籤為準，不自行拆高中／高職"],
        ["國中畢、國小畢以下", "國中及以下", "政策呈現用；原始欄位仍保留"],
    ], widths=[4.3, 3.5, 8.5], font_size=8.0)
    add_formula(doc, "EDU-01｜新北五歲組精確格", "Y[g,e,s] = Σ[v∈NTPC] Σ[m] population[v,g,e,m,s]")
    add_formula(doc, "EDU-02｜25–29精確占比", "EduShare[25–29,e,s] = Y[25–29,e,s] / Σe Y[25–29,e,s] × 100%")
    add_body(doc, "教育儀表板只需年齡×教育時，先加總婚姻維度再建模，避免在稀疏的教育×婚姻小格上過度擬合。若未來政策問題要求婚姻×教育交互分析，再啟用多維IPF並提高人工查核層級。")
    add_pclm_explanation(doc, "教育類別（婚姻先加總）")

    edu25 = [x for x in data["education"]["official_5y"] if x["source_age_band"] == "25-29" and x["sex"] == "合計"]
    add_heading(doc, "3.2 25–29官方精確組驗算", level=2)
    rows = [[x["education"], f"{x['population']:,}", f"{x['share_pct']:.2f}%", "OFFICIAL_ADMIN_EXACT"] for x in edu25]
    add_table(doc, ["教育程度", "人數", "占比", "值來源"], rows, widths=[4.2, 3.4, 3.0, 5.9], status_col=3)
    edu_total = sum(x["population"] for x in edu25)
    add_formula(doc, "EDU-03｜跨來源總量核對", f"Σe Y[25–29,e] = {edu_total:,} = P_REG[25–29]", "MOI教育交叉格與MOI單一年齡人口在25–29組完全一致。")
    add_picture(doc, figures["education"], 16.2, "114年底新北市25至29歲教育程度構成，分合計、男性與女性的百分之百堆疊圖。", "圖4　教育圖表原型；目前僅25–29可標官方精確值。")

    add_heading(doc, "3.3 模型輸出與驗證門檻", level=2)
    add_table(doc, ["檢查", "計算", "通過標準", "失敗處理"], [
        ["25–29回算", "模型結果重新加總", "每類與官方精確格差異≤浮點容許值", "停止發布並檢查分類映射"],
        ["列邊際", "Σe x[a,e,s] vs P1Y[a,s]", "最大絕對誤差<10⁻⁶人", "重新執行IPF／檢查總量相容"],
        ["欄邊際", "Σa∈g x[a,e,s] vs Y[g,e,s]", "最大絕對誤差<10⁻⁶人", "檢查零格與結構零"],
        ["非負性", "min x[a,e,s]", "≥0", "不使用產生負值的Sprague主結果"],
        ["方法敏感度", "PCLM+IPF vs 比例+IPF／Sprague", "包絡與方向合理", "標高不確定性或HOLD"],
        ["年度穩定性", "113→114分類占比", "大變動有來源或制度解釋", "人工抽查原始列與版本"],
    ], widths=[2.6, 4.0, 4.6, 5.7], font_size=7.5)

    add_heading(doc, "3.4 儀表板呈現規格", level=2)
    add_table(doc, ["圖表", "主要問題", "發布規則"], [
        ["100%堆疊長條", "各年齡組教育構成如何不同？", "25–29實色；模型組加方法徽章／斜線紋理"],
        ["年齡×教育熱圖", "哪個年齡／教育格占比最高？", "顯示人數與列百分比切換"],
        ["性別啞鈴圖", "同年齡組男女性高等教育差距？", "高等教育定義固定並在tooltip列示"],
        ["年度小倍圖", "教育構成是否變化？", "只比較相同年底、相同分類映射"],
    ], widths=[3.2, 6.4, 7.2], font_size=7.9)


def add_wage_section(doc: Document, data: dict, figures: dict[str, Path]) -> None:
    page_break(doc)
    add_heading(doc, "4｜薪資平均數與中位數", kicker="WAGES")
    add_body(doc, "主要發布指標採主計總處表6：工業及服務業、本國籍全時受僱員工、按工作場所所在縣市與年齡別的全年總薪資。地理角色是WORKPLACE，不等於新北戶籍青年。", refs=["S04"])
    add_body(doc, "114年已發布全國工業及服務業全年平均薪資；它可做全國參考卡，但因缺少新北工作地×目標年齡交叉，不能補成114年新北18–35值。", refs=["S05"])

    add_heading(doc, "4.1 來源概念不可互換", level=2)
    add_table(doc, ["資料", "母體／地理", "薪資概念", "可扮演角色", "禁止用途"], [
        ["DGBAS表6", "本國籍全時受僱；工作地", "全年經常＋非經常性薪資", "正式主指標", "不可解讀為戶籍青年所得"],
        ["人力運用調查", "抽樣個人；調查定義", "主要工作月收入", "研究模型／微資料申請", "未調整不得替代全年總薪資"],
        ["MOF薪資所得", "所得人；全國年齡", "稅務薪資所得", "年齡形狀敏感度", "不可標成新北工作地薪資"],
        ["BLI投保人數", "投保制度涵蓋者", "只有人數，非實際薪資", "涵蓋率／權數診斷", "不可推估實際薪資"],
        ["投保薪資級距", "制度級距、上限限制", "申報基礎", "獨立制度指標", "不可當全年總薪資分布"],
    ], widths=[2.8, 3.7, 3.8, 3.2, 3.3], font_size=7.3)

    add_heading(doc, "4.2 平均數：何時可以算", level=2)
    add_formula(doc, "W-AVG-01｜個體／微資料加權平均", "ȳ[18–35] = Σ[i:18≤ageᵢ≤35] wᵢ yᵢ / Σ[i:18≤ageᵢ≤35] wᵢ", "yᵢ為同一母體全年總薪資；wᵢ為官方權數或擴大數。")
    add_formula(doc, "W-AVG-02｜已知組別均數與人數", "ȳ[18–35] = Σg N[g]·ȳ[g] / Σg N[g]", "N[g]必須是與ȳ[g]相同工作地、年度、僱用型態與國籍範圍的受僱人數。")
    add_formula(doc, "W-AVG-03｜只知寬組均數時的停止條件", "若只有 ȳ[30–39]，且未知 Σy[30–35] 與 N[30–35]，則 ȳ[30–35] 不可識別。")
    add_callout(doc, "目前判定：HOLD", "表6只有『未滿25、25–29、30–39』的平均數，未提供同母體受僱人數或年齡內薪資總額；因此不能用人口權數、勞保投保人數或等權平均硬合成18–35。", fill=PALE_RED, accent=RED)

    add_heading(doc, "4.3 中位數：何時可以算", level=2)
    add_formula(doc, "W-MED-01｜加權中位數", "Median = inf{ y : Σ[i:yᵢ≤y] wᵢ ≥ 0.5·Σi wᵢ }")
    add_formula(doc, "W-MED-02｜只有薪資級距時的組內插值", "Median ≈ L + [(0.5N − C_prev) / f_m] · h", "L為中位級距下界；C_prev為前一級距累積人數；f_m為中位級距人數；h為級距寬。僅能用同一薪資母體分布。")
    add_formula(doc, "W-MED-03｜不可使用的合併", "Weighted average of group medians ≠ median of pooled individuals")
    add_body(doc, "主計總處的中位數定義是把受僱員工按全年總薪資排序後取中間位置；公開寬組中位數與平均數不足以重建個體排序。", refs=["S06", "S07"])

    add_picture(doc, figures["wage_flow"], 16.4, "薪資資料先判斷同母體，再依平均數與中位數所需資料決定計算或停止的流程。", "圖5　薪資資料充分性與停止規則。")

    add_heading(doc, "4.4 目前可發布的薪資結果", level=2)
    wage113 = [x for x in data["wage"] if x["year"] == "113"]
    rows = []
    for age in ("未滿25", "25-29", "30-39"):
        avg = next(x["annual_salary_10k_ntd"] for x in wage113 if x["source_age_band"] == age and x["measure"] == "平均數")
        med = next(x["annual_salary_10k_ntd"] for x in wage113 if x["source_age_band"] == age and x["measure"] == "中位數")
        status = "可發布（來源原生）" if age == "25-29" else "僅來源年齡組參考"
        rows.append([age, f"{avg:.1f}", f"{med:.1f}", "萬元／年", status])
    add_table(doc, ["來源年齡組", "平均數", "中位數", "單位", "用途"], rows, widths=[3.2, 2.8, 2.8, 3.0, 5.1], status_col=4)
    add_picture(doc, figures["wage"], 16.3, "民國108至113年新北工作地25至29歲本國籍全時受僱員工全年總薪資平均數與中位數折線圖。", "圖6　薪資第一階段驗證圖；25–29為官方原生組，108–113均使用同一表6來源。")
    add_body(doc, "圖中平均數與中位數的差距不可直接稱為所得不均；若要討論分布，應取得四分位數、十分位數或個體／級距分布。")

    add_heading(doc, "4.5 可行的補資料路徑", level=2)
    add_table(doc, ["優先序", "資料路徑", "能解決的問題", "發布等級"], [
        ["A", "向主計總處申請第一類資料／客製表：年齡、工作地、薪資、權數", "18–35平均數與加權中位數", "通過揭露審查後可正式"],
        ["B", "取得同母體『單歲／細年齡×薪資級距×人數』", "平均數、CDF與中位數", "可發布官方衍生值"],
        ["C", "人力運用調查微資料＋校準權數", "主要工作月收入的研究估計", "調查估計；不得冒充全年總薪資"],
        ["D", "MOF年齡薪資所得＋BLI人數作小域模型先驗", "情境／敏感度", "EXPERIMENTAL；不進正式KPI"],
    ], widths=[1.5, 6.5, 5.0, 3.8], font_size=7.6, status_col=3)
    add_body(doc, "主計總處第一類資料申請目錄列有人力資源調查（4401）與人力運用調查（4402）；是否可取得足以形成新北18–35全年總薪資的欄位，仍需由提供單位審查與確認。", refs=["S08", "S12"])

    page_break(doc)
    add_heading(doc, "4.6 儀表板呈現規格", level=2)
    add_table(doc, ["圖表", "現在可做", "取得完整分布後", "標示"], [
        ["平均／中位折線", "25–29年度趨勢", "三子群與18–35", "工作地、母體、萬元／年"],
        ["平均－中位啞鈴", "25–29各年", "各年齡／性別比較", "差距不是不均指數"],
        ["箱型圖／小提琴圖", "不可", "有微資料或分位數後可", "樣本數、權數、極端值規則"],
        ["十等分位帶", "全國參考可另卡", "新北×年齡分布後可", "不可用全國覆寫新北"],
        ["114全國參考卡", "可", "—", "清楚標『全國，非新北』"],
    ], widths=[3.2, 3.7, 5.0, 4.6], font_size=7.7)

    add_heading(doc, "4.7 薪資卡片的發布狀態", level=2)
    add_table(doc, ["卡片", "目前狀態", "前端顯示規則"], [
        ["25–29平均數／中位數", "READY", "顯示官方原生值、113年、工作地與母體"],
        ["18–24、30–35、18–35平均數", "HOLD", "不補值；提示需同母體人數權數或微資料"],
        ["18–24、30–35、18–35中位數", "HOLD", "不以組別中位數相加或加權；提示需個體／級距分布"],
        ["114全國薪資", "REFERENCE", "獨立參考卡，禁止接入新北年齡趨勢線"],
    ], widths=[4.5, 2.7, 9.6], font_size=7.7, status_col=1)
    add_callout(doc, "發布防呆", "任何HOLD欄位在API與CSV中均保留NULL，並同步輸出hold_reason；前端不得以0取代。")


def add_dashboard_section(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "5｜跨主題儀表板規格", kicker="DASHBOARD SPEC")
    add_body(doc, "人口、婚姻、教育與薪資可以放在同一互動式儀表板，但只能在『共同篩選與敘事』層次介接；率、占比與加權平均的分母仍必須由各指標自己的母體提供。")

    add_heading(doc, "5.1 共同篩選器", level=2)
    add_table(doc, ["篩選器", "允許值", "介接規則"], [
        ["年度／月份", "年度；人口另有月份", "預設顯示共同年；不同資料期必須卡片級標示"],
        ["年齡", "18–24、25–29、30–35、18–35", "顯示source_age_band與target_age_band"],
        ["性別", "合計、男、女", "薪資表6目前無新北×年齡×性別交叉，不假造"],
        ["行政區", "新北市／區", "薪資只有新北工作地市級時停用區篩選"],
        ["證據等級", "精確、官方原生、模型、HOLD", "模型值預設帶方法徽章"],
    ], widths=[3.2, 5.4, 8.1], font_size=7.9)

    add_heading(doc, "5.2 每張圖必備的血緣欄位", level=2)
    fields = [
        ["source_alias", "MOI-POP1Y、MOI-MARITAL5Y、MOI-EDU5Y、DGBAS-WAGE-LOC"],
        ["source_url", "官方資料頁或API路徑"],
        ["source_period / period_basis", "11412／END_OF_YEAR_12_31；113／FULL_CALENDAR_YEAR"],
        ["geography / geography_role", "新北市；REGISTERED_RESIDENCE或WORKPLACE"],
        ["universe", "戶籍人口、民間人口、本國籍全時受僱員工等"],
        ["source_age_band / target_age_band", "原始年齡與換算目標分開"],
        ["value_origin_class", "精確／官方原生／模型／HOLD"],
        ["method_code / model_version", "M0、PCLM_IPF_BOUNDARY_ESTIMATE等"],
        ["uncertainty_low / high / type", "方法包絡或抽樣信賴區間，不混稱"],
        ["qa_status / reviewed_by / reviewed_at", "查核狀態與人工簽核"],
    ]
    add_table(doc, ["欄位", "內容"], fields, widths=[5.0, 11.8], font_size=7.9)

    add_heading(doc, "5.3 視覺標記", level=2)
    add_table(doc, ["值身分", "建議樣式", "Tooltip第一行"], [
        ["OFFICIAL_ADMIN_EXACT", "實色＋綠色精確徽章", "官方行政精確值（在登記範圍內）"],
        ["OFFICIAL_SOURCE_NATIVE", "實色＋藍色官方來源徽章", "官方來源原生統計；未重新估計"],
        ["MODEL_ESTIMATE_FROM_OFFICIAL_MARGINS", "斜線／琥珀徽章＋敏感度", "模型估計；PCLM＋IPF"],
        ["HOLD_NO_ESTIMATE", "空值卡＋紅色原因", "資料不足；未估計"],
    ], widths=[5.5, 5.4, 5.8], font_size=7.9, status_col=0)


def add_review_section(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "6｜人工查核表與下一Gate", kicker="HUMAN REVIEW")
    add_callout(doc, "建議Gate 4判定", "人口：通過。婚姻：模型結果通過機器檢核，待抽查後通過。教育：25–29精確組通過；邊界組待執行模型。薪資：25–29通過；18–35平均／中位數維持HOLD。", fill=PALE_AMBER, accent=AMBER)

    checks = [
        ["G4-HR-01", "人口定義", "確認『戶籍登記現住人口』用語與政策問題一致", "決策者", "□"],
        ["G4-HR-02", "人口API", "抽查11412三個單歲欄位與村里加總", "資料工程", "□"],
        ["G4-HR-03", "婚姻分類", "同／異性有偶與離婚子類映射是否符合政策用語", "業務＋法制", "□"],
        ["G4-HR-04", "婚姻模型", "核對25–29、IPF列欄誤差、敏感度包絡與113→114方向", "統計人員", "□"],
        ["G4-HR-05", "教育分類", "確認高中畢顯示為『高中職』是否可接受", "教育業務", "□"],
        ["G4-HR-06", "教育模型", "執行PCLM＋IPF後核對25–29精確組與所有邊際", "統計人員", "□"],
        ["G4-HR-07", "薪資母體", "確認政策要工作地青年，或改為戶籍／居住地青年", "決策者", "□"],
        ["G4-HR-08", "薪資補資料", "決定是否申請主計總處微資料／客製表", "計畫主持", "□"],
        ["G4-HR-09", "114全國薪資", "是否另建『全國參考』卡片，不與新北113主卡合併", "決策者", "□"],
        ["G4-HR-10", "發布標籤", "確認模型值顯示方法徽章、敏感度與版本", "產品＋統計", "□"],
    ]
    add_table(doc, ["查核ID", "主題", "人工動作", "建議角色", "確認"], checks,
              widths=[2.4, 2.6, 7.9, 3.0, 1.2], font_size=7.4)

    add_heading(doc, "6.1 建議執行順序", level=2)
    for idx, text in enumerate((
        "鎖定共同字典：年齡、年度、地理角色、母體、證據等級與方法碼。",
        "人口直接上線；婚姻匯入既有113／114模型成果並完成人工抽核。",
        "教育先發布25–29官方精確圖；再執行18–24、30–35、18–35的PCLM＋IPF。",
        "薪資保留25–29年度趨勢與113來源年齡組；18–35顯示資料缺口，不顯示推測值。",
        "提出薪資微資料／客製表申請，取得後再開啟平均數與中位數第二階段。",
        "所有資料管線輸出manifest、SHA-256、版本、處理時間與QA結果，供AWS排程與日後查核。",
    ), start=1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        set_run_font(r, size=9.2)

    add_heading(doc, "6.2 需裁示事項", level=2)
    add_table(doc, ["決策ID", "問題", "建議預設", "影響"], [
        ["G4-DEC-01", "薪資地理角色", "沿用工作地（表6）", "若改戶籍地需另找／申請資料"],
        ["G4-DEC-02", "教育五大類映射", "研究所／大學／專科／高中職／國中及以下", "影響圖表與跨年一致性"],
        ["G4-DEC-03", "模型值發布", "通過人工抽核後可發布，必帶模型標籤", "影響婚姻與教育邊界組"],
        ["G4-DEC-04", "薪資補資料路徑", "先申請微資料／官方客製表，不做跨概念小域估計", "決定18–35薪資時程與可信度"],
        ["G4-DEC-05", "114全國薪資卡", "另卡顯示，不納入新北青年年度線", "避免年度與地理混淆"],
    ], widths=[2.6, 4.2, 6.2, 3.8], font_size=7.8)


def add_appendix(doc: Document, data: dict) -> None:
    page_break(doc)
    add_heading(doc, "附錄A｜公式與發布矩陣", kicker="APPENDIX")
    rows = [
        ["P_REG_18_35", "人口", "Σ單歲行政數", "845,938（11412）", "官方行政精確值", "可"],
        ["MAR_WITH_SPOUSE_18_35", "婚姻", "PCLM＋IPF後加總", "134,360.896（114）", "模型估計", "覆核後"],
        ["EDU_25_29", "教育", "官方五歲組直接加總", "250,128（114）", "官方行政精確值", "可"],
        ["EDU_18_35", "教育", "PCLM＋IPF後加總", "待執行", "模型估計", "待覆核"],
        ["WAGE_MEAN_25_29", "薪資", "DGBAS表6原生值", "59.9萬元／年（113）", "官方來源原生", "可"],
        ["WAGE_MEDIAN_25_29", "薪資", "DGBAS表6原生值", "52.4萬元／年（113）", "官方來源原生", "可"],
        ["WAGE_MEAN_18_35", "薪資", "同母體加權平均", "NULL", "HOLD", "不可"],
        ["WAGE_MEDIAN_18_35", "薪資", "加權個體中位數／級距CDF", "NULL", "HOLD", "不可"],
    ]
    add_table(doc, ["欄位代號", "主題", "公式／方法", "目前值", "值身分", "發布"], rows,
              widths=[4.1, 2.0, 4.4, 3.1, 3.0, 1.5], font_size=7.2, status_col=4)

    add_heading(doc, "附錄B｜引用來源", level=1)
    for code, (label, url) in SOURCES.items():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"[{code}] {label}：")
        set_run_font(r, size=8.2, bold=True, color=NAVY)
        add_hyperlink(p, "開啟官方／方法來源", url)

    page_break(doc)
    add_heading(doc, "附錄C｜查核註記", level=1)
    add_bullet(doc, "行政精確值是對該登記制度與統計範圍的精確彙總，不代表無漏登、遷徙落差或實際居住偏差。")
    add_bullet(doc, "主計總處薪資與人力調查結果仍有各自母體、定義與統計方法；應沿用官方名稱與精度。")
    add_bullet(doc, "PCLM、IPF與Sprague是資料拆分／校準工具，不會創造原資料不存在的可識別資訊；模型結果需與精確邊際共同發布。")
    add_bullet(doc, "本文件中的婚姻模型數值沿用V2.3既有處理成果；教育邊界組仍屬方法草稿，尚未產出正式估計值。")
    add_bullet(doc, "所有公式使用的年齡端點均為含端點：18–24含18與24；30–35含30與35。")


def add_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "新北市青年18–35歲人口、婚姻、教育與薪資計算及視覺化方法草稿"
    props.subject = "政府公開資料、年齡拆分、PCLM、IPF、薪資資料充分性與儀表板規格"
    props.author = "Codex research method draft"
    props.keywords = "新北市, 青年, 人口, 婚姻, 教育, 薪資, PCLM, IPF, 資料治理"
    props.comments = "Design preset: standard_business_brief; header template: memo_masthead"


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    figures = make_figures(data)
    doc = Document()
    setup_document(doc)
    add_core_properties(doc)
    add_cover(doc)
    add_executive_summary(doc, data)
    add_population_section(doc, data, figures)
    add_marriage_section(doc, data, figures)
    add_education_section(doc, data, figures)
    add_wage_section(doc, data, figures)
    add_dashboard_section(doc)
    add_review_section(doc)
    add_appendix(doc, data)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
