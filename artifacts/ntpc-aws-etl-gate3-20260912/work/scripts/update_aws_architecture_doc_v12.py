from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
PACKAGE = PROJECT / "deliverables" / "NTPC_Youth_Complete_Data_Package_G5_V3.2_20260826"
SOURCE = PACKAGE / "07_AWS與Kiro" / "AWS_Kiro_介接計算操作與每日更新_V1.1.docx"
OUTPUT = PACKAGE / "07_AWS與Kiro" / "AWS_Kiro_介接計算操作與每日更新_V1.2.docx"

NAVY = "0F172A"
BLUE = "1F4E78"
LIGHT_BLUE = "DCE6F1"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "E2F0D9"
PALE_AMBER = "FFF2CC"
WHITE = "FFFFFF"


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def format_cell(cell, *, bold=False, color=NAVY, size=8.6) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "Microsoft JhengHei"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft JhengHei")
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        table.rows[0].cells[index].text = text
        shade(table.rows[0].cells[index], BLUE)
        format_cell(table.rows[0].cells[index], bold=True, color=WHITE, size=font_size)
    set_repeat_header(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        for index, text in enumerate(row_values):
            row.cells[index].text = text
            if len(table.rows) % 2 == 1:
                shade(row.cells[index], LIGHT_GRAY)
            format_cell(row.cells[index], size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, label: str, text: str, fill=LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = f"{label}｜{text}"
    shade(cell, fill)
    format_cell(cell, bold=False, size=9.2)
    if cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].font.bold = True
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_together = True


def create_numbering_sequence(doc: Document) -> int:
    """Create a fresh decimal list so this appended section restarts at 1."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.append(level_justification)
    paragraph_properties = OxmlElement("w:pPr")
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "720")
    indentation.set(qn("w:hanging"), "360")
    paragraph_properties.append(indentation)
    level.append(paragraph_properties)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    num.append(abstract_reference)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_properties.append(level)
    num_properties.append(number)
    paragraph_properties.append(num_properties)


def add_detail(doc, title: str, data_text: str, architecture_text: str, acceptance_text: str) -> None:
    doc.add_heading(title, level=2)
    add_bullets(doc, [
        f"資料呈現：{data_text}",
        f"架構呈現：{architecture_text}",
        f"驗收證據：{acceptance_text}",
    ])


def append_existing_score_rows(doc: Document) -> None:
    reliability_table = doc.tables[13]
    extra = [
        ["卓越營運", "Kiro規格、IaC、版本化重跑、dev／staging／production Gate與回復SOP", "乾淨環境重跑、sam validate、NO_CHANGE與rollback演練"],
        ["永續性", "來源未變不重算；Lambda／Glue／Fargate用完即停；Athena只掃必要Parquet", "追蹤跳過重算次數、DPU／vCPU時間與Athena掃描量"],
    ]
    for values in extra:
        row = reliability_table.add_row()
        for index, value in enumerate(values):
            row.cells[index].text = value
            shade(row.cells[index], LIGHT_GRAY if len(reliability_table.rows) % 2 else WHITE)
            format_cell(row.cells[index], size=8.5)
    set_table_geometry(reliability_table, [1300, 4760, 3300])


def build() -> Path:
    doc = Document(SOURCE)
    doc.tables[0].rows[1].cells[1].text = "G5 V3.2／文件修訂 V1.2"
    doc.tables[0].rows[1].cells[3].text = "AWS Well-Architected部署前強化版；未執行正式部署"
    for cell in doc.tables[0].rows[1].cells:
        format_cell(cell, size=8.8)
    doc.paragraphs[2].text = "每日審視｜三運算分流｜JWT白名單｜受控Agent工具｜發布閘門｜雲端QA"
    append_existing_score_rows(doc)

    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("10. AWS架構評分六大關鍵修正", level=1)
    add_callout(
        doc,
        "修正結論",
        "核心仍為三母體、三份長格式CSV、S3分層、Step Functions協調及AgentCore問答。本次補強不改研究方法或儀表板母體邏輯，而是把可重跑、安全、沿襲、發布與監控變成可驗收元件。",
        PALE_GREEN,
    )
    add_table(
        doc,
        ["關鍵", "資料中的呈現", "AWS／Kiro架構中的呈現", "目前狀態與部署Gate"],
        [
            ["1 可重跑版本與路徑", "三CSV與方法腳本鎖定G5 V3.2；重跑後record_id、來源與QA須一致。", "腳本由檔案位置向上尋找專案根目錄；Kiro／CI明確傳入package路徑。", "程式已修正；須在乾淨工作目錄重跑並比對SHA-256。"],
            ["2 AWS IaC資源定義", "raw、curated、published、quarantine、manifest及三母體目錄分離。", "新增KMS、S3、DynamoDB、Scheduler、Step Functions、SNS、DLQ、Glue Database、Athena及CloudWatch藍圖。", "IaC藍圖已建立；Transform workflow與正式參數仍須在dev驗證。"],
            ["3 JWT登入與白名單", "查詢資料本身為公開彙總，但登入Email與問答紀錄屬受控營運資料。", "API Gateway驗證OIDC JWT；Lambda再核對兩個核准Email，runtimeUserId只取JWT sub。", "程式與4項測試通過；正式issuer／audience待部署時填入。"],
            ["4 ChatBot受控工具", "答案只能取三CSV、方法、來源主檔及資料列來源關聯；不得混用三母體。", "AgentCore Gateway只註冊六個唯讀工具；Policy Engine default-deny，不開放public-web-browser取指標。", "OpenAPI與Cedar範本已建立；Gateway target及ENFORCE尚待部署。"],
            ["5 發布閘門", "publish_status決定可查詢性；HOLD_DATA_GAP保留空列，語意是無資料而非0。", "Quality Gate先檢查整體PASS及三發布檔，再以S3條件式Put切換data_catalog；失敗維持舊版。", "政策、QA及catalog switch程式已建立；須演練未核准、競爭寫入與rollback。"],
            ["6 雲端QA與監控", "37,412筆指標、11個來源、70,060筆來源關聯；14,720項本地不變性QA全部通過。", "Glue DQDL檢查欄位與狀態；Python檢查跨列公式；CloudWatch、SNS、DLQ記錄失敗與通知。", "本地QA與語法檢查通過；Owner、通知信箱、預算與門檻待部署裁示。"],
        ],
        [1320, 2440, 3200, 2400],
        font_size=8.1,
    )

    add_detail(
        doc,
        "10.1 可重跑版本與路徑",
        "build、validate、source manifest與dashboard export都以同一資料包根目錄為基準；發布CSV不因執行位置不同而改讀舊資料。",
        "CI／Kiro從設定注入package路徑；程式找不到專案根目錄即fail closed，不默默回退舊版。",
        "從資料包內reproducible_scripts與專案scripts各執行一次，輸出列數、record_id集合、來源雜湊及14,720項QA一致。",
    )
    add_detail(
        doc,
        "10.2 AWS IaC資源定義",
        "S3 raw保留官方原始檔，curated保留Parquet，published保留三CSV與catalog，quarantine保留失敗候選。",
        "AWS資料平台IaC採SSE-KMS、S3版本化、DynamoDB PITR、Scheduler臺北時區、Step Functions Standard與Athena掃描上限。",
        "sam validate／cfn-lint、dev部署、來源未變、QA失敗、重試、DLQ、發布及回復逐項演練。",
    )
    add_detail(
        doc,
        "10.3 JWT登入與白名單",
        "統計值為公開彙總；Email白名單只用於入口授權，不能寫入三份公開CSV或來源主檔。",
        "JWT驗證issuer、audience、簽章與期限；Lambda檢查Email並從sub建立使用者身分，API Key只可另作流量配額。",
        "允許帳號回傳200；非白名單回傳403；缺少、過期或錯誤audience的Token回傳401；日誌不保存完整問題文字。",
    )
    add_detail(
        doc,
        "10.4 ChatBot受控工具",
        "list dimensions先確認可用年度、性別與地理；query工具依universe分流；method與source工具提供公式及證據。",
        "AgentCore Gateway把六個API轉成工具；Policy Engine採default-deny與forbid-wins，正式Agent不具有S3寫入或發布權限。",
        "以固定繁中題庫逐題比對CSV數值、母體、公式、方法、不確定性、來源URL及缺資料拒答。",
    )
    add_detail(
        doc,
        "10.5 發布閘門",
        "可發布狀態與HOLD狀態以publication-policy.json集中定義；薪資12筆中位數缺口不得被0補值。",
        "Step Functions只有在Machine QA PASS且必要時取得人工核准後才調用CatalogSwitch；S3 If-Match／If-None-Match避免競爭覆寫。",
        "故意製造schema漂移、QA失敗、未核准、舊ETag及重複run_id，確認舊版仍可服務且失敗候選進quarantine。",
    )
    add_detail(
        doc,
        "10.6 雲端QA與監控",
        "DQDL檢查完整性、唯一性、年度、年齡、發布狀態與QA狀態；Python持續檢查29區、男女、百分比、LF=E+U及UR+ESLF=100%。",
        "CloudWatch記錄run_id、狀態、延遲、錯誤、掃描量及Agent用量；SNS通知，Scheduler失敗送SQS DLQ。",
        "以資料新鮮度、來源失敗率、QA失敗數、發布延遲、Athena掃描量、Agent錯誤率及成本設定告警。",
    )

    doc.add_heading("11. AWS Well-Architected六支柱對應", level=1)
    add_table(
        doc,
        ["支柱", "本案設計", "可查核證據", "本次或未來"],
        [
            ["卓越營運", "Kiro規格驅動、IaC、環境Gate、版本化作業與回復SOP", "requirements／design／tasks、SAM／CDK、run manifest", "本次必要"],
            ["安全性", "KMS、封鎖S3公開、JWT白名單、最小IAM、Policy Engine", "template.yaml、policy.cedar、登入測試與CloudTrail", "本次必要"],
            ["可靠性", "冪等、重試、DLQ、PITR、quarantine、原子切換與舊版續服", "DynamoDB state、S3版本、Step Functions歷程", "本次必要"],
            ["效能效率", "依工作負載選Lambda／Glue／Fargate；Parquet及Athena分區", "基準測試、處理時間、查詢P95與掃描量", "本次必要"],
            ["成本最佳化", "NO_CHANGE不重算、按需服務、Athena掃描上限與Budget", "Cost Explorer標籤、Budget告警、每次run成本", "Budget值待裁示"],
            ["永續性", "事件驅動、來源未變即停止、工作完成釋放資源", "跳過重算次數、DPU／vCPU時間與資料掃描量", "持續優化"],
        ],
        [1300, 3500, 3000, 1560],
        font_size=8.3,
    )
    add_callout(
        doc,
        "評分邊界",
        "Lake Formation、跨帳號／跨區域DR、Iceberg與AgentCore線上評估列為未來規劃；它們不改變三母體及目前MVP主架構，也不阻擋本階段。",
        PALE_AMBER,
    )

    doc.add_heading("12. 部署前最小驗收順序", level=1)
    acceptance_num_id = create_numbering_sequence(doc)
    for text in [
        "在乾淨環境重跑三CSV、來源主檔、資料列來源關聯及14,720項QA。",
        "對兩份SAM／CloudFormation執行validate、lint與最小權限檢查。",
        "在dev填入真實JWT issuer、audience、TransformWorkflowArn、Owner、通知信箱與預算。",
        "執行NO_CHANGE、來源變更、QA失敗、人工未核准、競爭切換與rollback演練。",
        "執行ChatBot固定題庫、來源引用、三母體隔離、HOLD拒答與提示注入測試。",
        "staging驗收通過並取得人工部署核准後，才可production deploy。",
    ]:
        paragraph = doc.add_paragraph(text)
        apply_numbering(paragraph, acceptance_num_id)
        paragraph.paragraph_format.space_after = Pt(3)

    doc.add_heading("12.1 本次修正檔案索引", level=2)
    add_table(
        doc,
        ["功能", "檔案位置", "用途"],
        [
            ["資料沿襲", "09_發布資料包/04_來源主檔.csv；05_資料列來源關聯.csv", "讓每筆指標連回原子來源與角色"],
            ["發布政策", "11_介接設定與追溯紀錄/config/publication-policy.json", "集中定義可查詢、HOLD與模型值規則"],
            ["資料平台", "12_AWS_Kiro_實作檔案/AWS資料平台IaC", "資料湖、排程、QA、發布、Athena與告警"],
            ["JWT橋接", "12_AWS_Kiro_實作檔案/AWS橋接", "The Weekly Blend登入身分與白名單驗證"],
            ["受控工具", "12_AWS_Kiro_實作檔案/AgentCore/governed-tools", "六個唯讀工具與Policy Engine範本"],
            ["Kiro清單", "KIRO規格/specs/ntpc-youth-g5/AWS架構評分修正清單.md", "交接狀態、證據與部署剩餘Gate"],
        ],
        [1450, 4400, 3510],
        font_size=8.3,
    )

    doc.add_heading("12.2 AWS官方依據補充", level=2)
    add_bullets(doc, [
        "AWS Well-Architected Data Analytics Lens：https://docs.aws.amazon.com/wellarchitected/latest/analytics-lens/well-architected-design-principles.html",
        "AWS Glue Data Quality與DQDL：https://docs.aws.amazon.com/glue/latest/dg/glue-data-quality.html",
        "API Gateway JWT authorizer：https://docs.aws.amazon.com/apigateway/latest/developerguide/http-api-jwt-authorizer.html",
        "AgentCore Policy Engine：https://docs.aws.amazon.com/bedrock-agentcore/latest/devguide/policy-getting-started.html",
        "AgentCore Observability：https://docs.aws.amazon.com/bedrock-agentcore/latest/devguide/observability.html",
    ])

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
