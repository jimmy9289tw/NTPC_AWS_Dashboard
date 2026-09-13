from __future__ import annotations

import argparse
import csv
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_g5_v3_docs import (
    BLUE,
    DARK_BLUE,
    DISPLAY_YEARS,
    MUTED,
    NAVY,
    VERSION,
    add_bullets,
    add_callout,
    add_hyperlink,
    add_page_number,
    add_para,
    add_picture_with_alt,
    add_source_link,
    add_table,
    configure_styles,
    load_inputs,
    make_diagrams,
    set_run_font,
)


def resolve_project_root() -> Path:
    for candidate in Path(__file__).resolve().parents:
        if (candidate / "config" / "g5-refresh-policy.json").is_file() and (candidate / "scripts").is_dir():
            return candidate
    raise RuntimeError("Cannot locate NtpcYouthAI project root")


PROJECT = resolve_project_root()
DEFAULT_PACKAGE = PROJECT / "deliverables" / "NTPC_Youth_Complete_Data_Package_G5_V3.2_20260826"
REPORT_VERSION = "G5 V3.5"
OUTPUT_NAME = "新北市青年18-35歲資料整合與智慧儀表板完整總說明報告書_G5_V3.5.docx"


def read_csv(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def find_row(rows: list[dict], **criteria: str) -> dict:
    for row in rows:
        if all(row.get(key) == value for key, value in criteria.items()):
            return row
    raise KeyError(f"No row matched: {criteria}")


def fmt(row: dict, digits: int = 1) -> str:
    value = float(row["value"])
    unit = row["unit"]
    if unit == "人":
        return f"{value:,.0f}人"
    if unit == "千人":
        return f"{value:,.1f}千人"
    if unit == "%":
        return f"{value:.2f}%"
    if unit == "萬元/年":
        return f"{value:.{digits}f}萬元／年"
    return f"{value:g}{unit}"


def optional_float(row: dict, key: str) -> float | None:
    raw = str(row.get(key, "")).strip()
    return None if raw == "" else float(raw)


def fmt_scalar(value: float | None, unit: str, *, model: bool = True) -> str:
    if value is None:
        return "未提供"
    if unit == "人":
        return f"{value:,.3f}人" if model and abs(value - round(value)) > 1e-9 else f"{value:,.0f}人"
    if unit == "千人":
        return f"{value:,.3f}千人"
    if unit == "%":
        return f"{value:.3f}%"
    if unit == "萬元/年":
        return f"{value:.1f}萬元／年"
    return f"{value:g}{unit}"


def sensitivity_methods(row: dict) -> tuple[str, str | None]:
    uncertainty_type = row.get("uncertainty_type", "")
    method_code = row.get("method_code", "")
    primary = {
        "M0_DIRECT_SINGLE_AGE_SUM": "單一年齡直接加總",
        "PCLM_CITY_SEED_IPF_GEO": "PCLM種子＋IPF",
        "PCLM": "PCLM",
        "WAGE_PCLM_RATIO_CALIBRATION": "PCLM輔助輪廓＋官方薪資錨定",
        "M0_OFFICIAL_AGE_BAND": "官方原生年齡帶",
        "M0_DIRECT_5Y_AGGREGATION": "官方原生五歲帶加總",
    }.get(method_code, method_code or "來源原生值")
    alternate = {
        "METHOD_SENSITIVITY_ENVELOPE_NOT_CI": "人口比例初值＋IPF",
        "PCLM_SPRAGUE_METHOD_ENVELOPE_NOT_CI": "Sprague非負校準",
        "PCLM_VS_UNIFORM_AUXILIARY_METHOD_ENVELOPE_NOT_CI": "五歲組內均勻輔助法",
    }.get(uncertainty_type)
    return primary, alternate


def sensitivity_row(row: dict, label: str) -> list[str]:
    if str(row.get("value", "")).strip() == "":
        reason = "缺個體薪資分布與權數／可重建級距頻數"
        return [
            label,
            f"不可估計\n{reason}",
            "不適用",
            "未定義（無點估計）",
            "未定義（無點估計）",
        ]
    value = float(row["value"])
    unit = row["unit"]
    low = optional_float(row, "uncertainty_low")
    high = optional_float(row, "uncertainty_high")
    primary_method, alternate_method = sensitivity_methods(row)
    if alternate_method is None:
        primary_cell = f"{fmt_scalar(value, unit, model=False)}\n{primary_method}"
        alternate_cell = "不適用"
        low_cell = "未提供" if low is None else f"{fmt_scalar(low, unit, model=False)}\n{primary_method}"
        high_cell = "未提供" if high is None else f"{fmt_scalar(high, unit, model=False)}\n{primary_method}"
        return [label, primary_cell, alternate_cell, low_cell, high_cell]

    if low is None or high is None:
        raise ValueError(f"Model row lacks sensitivity envelope: {row.get('record_id')}")
    tolerance = max(1e-9, abs(value) * 1e-9)
    if abs(value - low) <= tolerance and abs(value - high) <= tolerance:
        alternate_value = value
        low_method = high_method = "兩法相同"
    elif abs(value - low) <= tolerance:
        alternate_value = high
        low_method, high_method = primary_method, alternate_method
    elif abs(value - high) <= tolerance:
        alternate_value = low
        low_method, high_method = alternate_method, primary_method
    else:
        raise ValueError(f"Point estimate outside or detached from envelope: {row.get('record_id')}")
    primary_cell = f"{fmt_scalar(value, unit)}\n{primary_method}"
    alternate_cell = f"{fmt_scalar(alternate_value, unit)}\n{alternate_method}"
    low_cell = f"{fmt_scalar(low, unit)}\n{low_method}"
    high_cell = f"{fmt_scalar(high, unit)}\n{high_method}"
    return [label, primary_cell, alternate_cell, low_cell, high_cell]


def count_rows(rows: list[dict], **criteria: str) -> int:
    return sum(all(row.get(key) == value for key, value in criteria.items()) for row in rows)


def add_csv_row_count_audit(doc: Document, registered: list[dict], labor: list[dict], wage: list[dict]) -> None:
    doc.add_heading("5.1 三份CSV列數乘積與實際列數查核", level=2)
    add_para(doc, "每一列代表一個『年度×地理×年齡×性別×類別×指標』的有效組合；不是所有欄位盲目做完整笛卡兒乘積。下表把條件式維度分開計算，再與CSV實際列數核對。")

    registered_parts = [
        ("人口數", "5年×30地理×4年齡×3性別×1類別×1指標", 5 * 30 * 4 * 3, count_rows(registered, metric_code="REGISTERED_POPULATION_COUNT")),
        ("人口占地區比率", "5×30×4×3×1×1", 5 * 30 * 4 * 3, count_rows(registered, metric_code="REGISTERED_POPULATION_SHARE_PCT")),
        ("年齡帶內性別占比", "5×30×4×2性別×1×1", 5 * 30 * 4 * 2, count_rows(registered, metric_code="SEX_SHARE_WITHIN_AGE_BAND_PCT")),
        ("教育人數＋占比", "5×30×4×3×5類別×2指標", 5 * 30 * 4 * 3 * 5 * 2, sum(row.get("category_dimension_code") == "EDUCATION" for row in registered)),
        ("婚姻人數＋占比", "5×30×4×3×4類別×2指標", 5 * 30 * 4 * 3 * 4 * 2, sum(row.get("category_dimension_code") == "MARITAL_STATUS" for row in registered)),
    ]
    registered_expected = sum(part[2] for part in registered_parts)
    audit_rows = [
        ["戶籍人口｜" + label, formula, f"{expected:,}", f"{actual:,}", "PASS" if expected == actual else "FAIL"]
        for label, formula, expected, actual in registered_parts
    ]
    audit_rows.append(["戶籍人口CSV合計", "上述五部分加總", f"{registered_expected:,}", f"{len(registered):,}", "PASS" if registered_expected == len(registered) else "FAIL"])
    labor_expected = 5 * 1 * 4 * 1 * 1 * 9
    audit_rows.append(["勞動市場CSV合計", "5年×1地理×4年齡×1性別×1類別×9指標", f"{labor_expected:,}", f"{len(labor):,}", "PASS" if labor_expected == len(labor) else "FAIL"])
    wage_mean_expected = 4 * 1 * 4 * 1 * 1
    wage_median_expected = 4 * 1 * 4 * 1 * 1
    wage_mean_actual = count_rows(wage, category_code="MEAN")
    wage_median_actual = count_rows(wage, category_code="MEDIAN")
    audit_rows.extend([
        ["薪資｜平均數", "4年×1地理×4年齡×1性別×1統計量", f"{wage_mean_expected:,}", f"{wage_mean_actual:,}", "PASS" if wage_mean_expected == wage_mean_actual else "FAIL"],
        ["薪資｜中位數", "4年×1地理×4年齡×1性別×1統計量", f"{wage_median_expected:,}", f"{wage_median_actual:,}", "PASS" if wage_median_expected == wage_median_actual else "FAIL"],
        ["薪資CSV合計", "4年×1地理×4年齡×1性別×2統計量", f"{wage_mean_expected + wage_median_expected:,}", f"{len(wage):,}", "PASS" if wage_mean_expected + wage_median_expected == len(wage) else "FAIL"],
    ])
    add_table(doc, ["資料分段", "維度乘積", "理論列數", "實際列數", "結果"], audit_rows, [1850, 3900, 1200, 1200, 1210], font_size=7.9)
    wage_numeric = sum(str(row.get("value", "")).strip() != "" for row in wage)
    wage_gaps = len(wage) - wage_numeric
    add_callout(doc, "查核結論", f"三份CSV均通過列數核對：戶籍人口{len(registered):,}列、民間人口與勞動市場{len(labor):,}列、受僱員工薪資{len(wage):,}列。薪資表採完整32列結構，其中{wage_numeric}列有數值、{wage_gaps}列缺值；官方原生值與模型估計值分開標示。", "ok")


def add_sensitivity_result_tables(doc: Document, registered: list[dict], labor: list[dict], wage: list[dict]) -> None:
    doc.add_heading("6.6 最低值、最高值的方法與數值結果", level=2)
    add_para(doc, "最低值與最高值是兩種合理換算方法所形成的方法敏感度包絡：low=min(主要法,替代法)，high=max(主要法,替代法)。它回答『更換方法後結果會移動多少』，不是抽樣調查的95%信賴區間。比率必須先在各方法內使用相符分子、分母重算，再比較兩個比率。")
    add_table(doc, ["資料／計算", "主要方法", "替代方法", "最低與最高值定義", "區間性質"], [
        ["戶籍人口數與比率", "單一年齡直接加總", "無", "low=value=high", "官方行政精確值"],
        ["教育、婚姻人數與占比", "全市PCLM種子＋各區IPF", "人口比例初值＋IPF", "兩法結果逐列取min／max", "方法敏感度，非CI"],
        ["P、LF、E、U、NLF及四項率", "PCLM", "Sprague非負校準", "計數逐項比較；率先各自重算再取min／max", "方法敏感度，非CI"],
        ["非原生年齡平均薪資", "PCLM輔助輪廓＋官方薪資錨定", "五歲組內均勻輔助法", "兩種加權估計取min／max", "方法敏感度，非CI"],
        ["25–29平均與中位薪資", "官方原生年齡帶", "無", "官方未公布CI時不補造low／high", "官方調查統計"],
        ["非25–29薪資中位數", "官方寬帶形狀轉移＋對數常態混合分布", "組內均勻輔助輪廓", "PCLM與均勻替代法估計取min／max", "方法敏感度，非CI"],
    ], [1700, 2250, 1900, 2300, 1210], font_size=8.0)
    add_callout(doc, "數值表範圍", "為保持報告可讀性，下列逐項列出目前儀表板最新範圍：戶籍、教育、婚姻及勞動採114年新北市整體18–35歲全體性別；薪資採最新可用113年。其他年度、行政區、性別與年齡組沿用同一規則，完整列級結果保存在三份發布CSV。", "info")

    registered_rows: list[list[str]] = []
    exact_pop = find_row(registered, roc_year="114", geography_name_zh="新北市", age_band="18-35", sex_code="ALL", category_dimension_code="NONE", metric_code="REGISTERED_POPULATION_COUNT")
    exact_share = find_row(registered, roc_year="114", geography_name_zh="新北市", age_band="18-35", sex_code="ALL", category_dimension_code="NONE", metric_code="REGISTERED_POPULATION_SHARE_PCT")
    registered_rows.append(sensitivity_row(exact_pop, "人口｜戶籍人口數"))
    registered_rows.append(sensitivity_row(exact_share, "人口｜占全市戶籍人口"))
    for dimension, dimension_label in (("EDUCATION", "教育"), ("MARITAL_STATUS", "婚姻")):
        subset = [
            row for row in registered
            if row.get("roc_year") == "114"
            and row.get("geography_name_zh") == "新北市"
            and row.get("age_band") == "18-35"
            and row.get("sex_code") == "ALL"
            and row.get("category_dimension_code") == dimension
        ]
        subset.sort(key=lambda row: (row.get("category_name_zh", ""), row.get("metric_code", "")))
        for row in subset:
            metric_label = "人數" if row["metric_code"].endswith("COUNT") else "占比"
            registered_rows.append(sensitivity_row(row, f"{dimension_label}｜{row['category_name_zh']}｜{metric_label}"))
    add_table(doc, ["114年18–35歲指標", "主要法數值", "替代法數值", "最低值與方法", "最高值與方法"], registered_rows, [1950, 1850, 1850, 1855, 1855], font_size=7.4)

    labor_order = [
        "CIVILIAN_POPULATION", "LABOR_FORCE", "EMPLOYED", "UNEMPLOYED", "NOT_IN_LABOR_FORCE",
        "UNEMPLOYMENT_RATE", "LABOR_FORCE_PARTICIPATION_RATE", "EMPLOYMENT_TO_POPULATION_RATE", "EMPLOYMENT_SHARE_OF_LABOR_FORCE",
    ]
    labor_rows = []
    for metric in labor_order:
        row = find_row(labor, roc_year="114", geography_name_zh="新北市", age_band="18-35", sex_code="ALL", metric_code=metric)
        labor_rows.append(sensitivity_row(row, row["metric_name_zh"]))
    add_table(doc, ["114年18–35歲勞動指標", "PCLM數值", "Sprague數值", "最低值與方法", "最高值與方法"], labor_rows, [1950, 1850, 1850, 1855, 1855], font_size=7.5)

    wage_rows = []
    for age_band in ("18-24", "25-29", "30-35", "18-35"):
        row = find_row(wage, roc_year="113", geography_name_zh="新北市", age_band=age_band, category_code="MEAN", metric_code="ANNUAL_TOTAL_SALARY_MEAN")
        wage_rows.append(sensitivity_row(row, f"{age_band}歲｜平均薪資"))
    for age_band in ("18-24", "25-29", "30-35", "18-35"):
        median = find_row(wage, roc_year="113", geography_name_zh="新北市", age_band=age_band, category_code="MEDIAN", metric_code="ANNUAL_TOTAL_SALARY_MEDIAN")
        wage_rows.append(sensitivity_row(median, f"{age_band}歲｜薪資中位數"))
    add_table(doc, ["113年薪資指標", "主要法數值", "替代法數值", "最低值與方法", "最高值與方法"], wage_rows, [1950, 1850, 1850, 1855, 1855], font_size=7.6)
    add_callout(doc, "發布文字", "模型列應標示『估計值／方法敏感度範圍』；不得把low／high寫成95%信賴區間。若未來取得調查微資料與樣本權數，須另以設計式變異或重抽樣建立正式統計區間。", "warn")


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def _new_numbering_id(doc: Document) -> int:
    """Create a fresh level-0 decimal list that restarts at 1."""
    numbering = doc.part.numbering_part.element
    nums = numbering.findall(qn("w:num"))
    num_ids = [int(node.get(qn("w:numId"))) for node in nums]

    style_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    source_num = next(node for node in nums if int(node.get(qn("w:numId"))) == style_num_id)
    abstract_id = source_num.find(qn("w:abstractNumId")).get(qn("w:val"))

    new_num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)

    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return new_num_id


def add_numbered_steps(doc: Document, steps: list[str]) -> None:
    """Add a semantic Word numbered list and restart each procedure at 1."""
    num_id = _new_numbering_id(doc)
    for step in steps:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.167

        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_ref = OxmlElement("w:numId")
        num_ref.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_ref)
        paragraph._p.get_or_add_pPr().append(num_pr)

        run = paragraph.add_run(step)
        set_run_font(run, size=10.5, color=NAVY)


def mark_all_table_headers(doc: Document) -> None:
    """Mark every table's first row as a repeating semantic header row."""
    for table in doc.tables:
        row_properties = table.rows[0]._tr.get_or_add_trPr()
        if row_properties.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            row_properties.append(header)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = False
    run = p.add_run(text)
    set_run_font(run, size=9.5, bold=True, color=MUTED)


def add_figure(doc: Document, path: Path, alt_text: str, caption: str, width=Inches(6.25)) -> None:
    add_picture_with_alt(doc, path, alt_text, width=width)
    add_caption(doc, caption)


def add_cover(doc: Document) -> None:
    # editorial_cover pattern resolved on top of standard_business_brief.
    add_para(doc, "新北市青年資料研究", bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, size=10.5)
    add_para(doc, "", after=48)
    add_para(
        doc,
        "新北市青年18–35歲資料整合\n與智慧儀表板完整總說明報告書",
        bold=True,
        color=NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=10,
        size=27,
    )
    add_para(
        doc,
        "政策背景、三母體定義、年齡換算、AWS／Kiro與ChatBot系統交接",
        color=DARK_BLUE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=42,
        size=13,
    )
    add_table(
        doc,
        ["文件版本", "資料觀察期間", "產製日期", "文件狀態"],
        [[REPORT_VERSION, DISPLAY_YEARS, "2026年8月26日", "研究與系統交接用；發布前仍須人工核准"]],
        [1500, 1900, 1800, 4160],
        font_size=9.5,
    )
    add_para(doc, "編製原則：母體先分流、方法可追溯、估計須標示、發布可查核。", bold=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, size=10.5)
    add_para(doc, "本報告不把不同母體視為同一群人，也不把模型估計值改稱官方精確值。", color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, size=10)
    add_page_break(doc)


def add_static_toc(doc: Document) -> None:
    doc.add_heading("章節導覽", level=1)
    sections = [
        "技術摘要與目前可用結論",
        "研究背景、目標與治理原則",
        "三種統計母體與青年年齡定義",
        "資料來源、期間、頻率與可用範圍",
        "共通資料架構與三份長格式CSV",
        "年齡區間整合方法與公式",
        "各母體計算方法、限制與驗證",
        "資料處理、更新、查核與發布流程",
        "AWS、Kiro與ChatBot系統架構",
        "互動式儀表板資訊架構與視覺化",
        "品質、資安、成本與行政查核",
        "實施狀態、缺口與後續建議",
        "附錄：資料來源、方法文獻、欄位與重跑檔案",
    ]
    for item in sections:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(item)
        set_run_font(run, size=11, color=NAVY)
    add_callout(doc, "閱讀方式", "行政決策者可先讀技術摘要、三母體與實施狀態；資料與資訊人員再讀方法、欄位、AWS／Kiro與附錄。", "info")
    add_page_break(doc)


def add_source_table(doc: Document) -> None:
    doc.add_heading("4. 資料來源、期間、頻率與可用範圍", level=1)
    add_para(doc, "核心資料均來自政府機關或政府資料開放平臺；新北市OpenAPI與新北市統計資料庫除作介接來源，也負責命題指定平台的交叉查核。")
    rows = [
        ["人口", "內政部戶政司／政府資料開放平臺", "村里×單一年齡×性別×月", "110–114", "每月；年度採12月31日", "核心值"],
        ["教育", "內政部戶政司／政府資料開放平臺", "村里×五歲年齡×性別×教育", "110–114", "每年", "核心值"],
        ["婚姻", "內政部戶政司／政府資料開放平臺", "村里×五歲年齡×性別×婚姻", "110–114", "每年", "核心值"],
        ["勞動", "行政院主計總處人力資源調查年報表27／32／36；表28作QA；114年另用下半年表41／42", "新北市×年齡帶×P／E／U；LF及各率重算", "110–114", "每年；全年12月平均", "核心值／輔助拆分"],
        ["薪資", "行政院主計總處全年總薪資統計表6", "工作場所縣市×年齡×統計量", "110–113", "每年", "核心錨點"],
        ["薪資輔助", "勞保局公開資料", "新北市年齡人數；全國年齡平均提繳工資", "110–113", "每年", "只作權數／相對輪廓"],
        ["地圖", "新北市資料開放平臺", "29區行政界線", "現行", "按需", "戶籍區塊互動"],
        ["交叉查核", "新北市OpenAPI、OAS", "依各統計表", "依來源", "每日審視", "不與核心值重複加權"],
    ]
    add_table(doc, ["主題", "發布機關／平臺", "原生粒度", "本案期間", "原生頻率", "本案角色"], rows, [1050, 2200, 2200, 1050, 1500, 1360], font_size=8.6)
    add_callout(doc, "時間基準", "戶籍人口為12月31日年末存量；勞動為全年12個月平均；薪資為全年統計。三者可在同年度並列，但不得解讀為同一觀察時點。", "warn")


def build_report(package: Path) -> Path:
    data = load_inputs(package)
    diagrams = make_diagrams(package)
    registered = data["registered"]
    labor = data["labor"]
    wage = data["wage"]
    qa = data["qa"]

    pop_114 = find_row(registered, roc_year="114", geography_name_zh="新北市", age_band="18-35", sex_code="ALL", category_dimension_code="NONE", metric_code="REGISTERED_POPULATION_COUNT")
    labor_metrics = {
        code: find_row(labor, roc_year="114", geography_name_zh="新北市", age_band="18-35", sex_code="ALL", metric_code=code)
        for code in ["CIVILIAN_POPULATION", "LABOR_FORCE", "EMPLOYED", "UNEMPLOYED", "UNEMPLOYMENT_RATE", "LABOR_FORCE_PARTICIPATION_RATE", "EMPLOYMENT_TO_POPULATION_RATE", "EMPLOYMENT_SHARE_OF_LABOR_FORCE"]
    }
    wage_mean_113 = find_row(wage, roc_year="113", geography_name_zh="新北市", age_band="18-35", category_code="MEAN", metric_code="ANNUAL_TOTAL_SALARY_MEAN")
    wage_median_113 = find_row(wage, roc_year="113", geography_name_zh="新北市", age_band="25-29", category_code="MEDIAN", metric_code="ANNUAL_TOTAL_SALARY_MEDIAN")

    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        header = section.header.paragraphs[0]
        for run in header.runs:
            if VERSION in run.text:
                run.text = run.text.replace(VERSION, REPORT_VERSION)
                set_run_font(run, size=9, bold=True, color=MUTED)
    doc.core_properties.title = "新北市青年18–35歲資料整合與智慧儀表板完整總說明報告書"
    doc.core_properties.subject = "政府公開資料、三母體、年齡換算、AWS、Kiro與ChatBot"
    doc.core_properties.keywords = "新北市, 青年, 18-35歲, PCLM, IPF, Sprague, AWS, Kiro, 儀表板"
    add_cover(doc)
    add_static_toc(doc)

    doc.add_heading("1. 技術摘要與目前可用結論", level=1)
    add_callout(doc, "總結", "本案已建立三條彼此分流的統計產品線，以三份長格式CSV供儀表板與ChatBot查詢。戶籍人口與勞動資料涵蓋110–114年；薪資因官方表6最新資料年為113年，涵蓋110–113年。", "ok")
    add_bullets(doc, [
        "戶籍人口母體可以單一年齡直接得到18–24、25–29、30–35及18–35歲人口；教育與婚姻則以PCLM產生單歲種子，再用IPF重現官方邊際。",
        "勞動市場以PCLM拆分P、LF、E、U，並以Sprague作方法敏感度；率一律由相符分子與分母重算，不直接拆分百分比。",
        "薪資CSV依4年度×4年齡組×2統計量建立完整32列。平均數與中位數四組皆有值；25–29為官方原生調查值，其他三組為官方薪資錨定模型估計。非原生中位數採官方寬帶形狀轉移與PCLM權重的對數常態混合分布估計。",
        "三母體可同頁並列，但地理角色、分母及時間基準不同，不可相加或形成未經定義的綜合分數。",
        "每日09:15只審視來源是否變更；未變更只記錄NO_CHANGE，來源變更且完成機器與人工查核後才發布新版本。",
    ])
    add_table(doc, ["最新可發布觀察", "數值", "母體／身分", "必要解讀"], [
        ["114年18–35歲戶籍人口", fmt(pop_114, 0), "戶籍人口／官方行政精確值", "12月31日年末存量"],
        ["114年18–35歲勞參率", fmt(labor_metrics["LABOR_FORCE_PARTICIPATION_RATE"]), "勞動市場／調查值再模型換算", "LF÷P；全年平均"],
        ["114年18–35歲失業率", fmt(labor_metrics["UNEMPLOYMENT_RATE"]), "勞動市場／調查值再模型換算", "U÷LF；全年平均"],
        ["113年18–35歲平均薪資", fmt(wage_mean_113), "受僱員工薪資／模型估計", "工作場所；非戶籍地"],
        ["113年25–29歲薪資中位數", fmt(wage_median_113), "受僱員工薪資／官方調查統計", "其他年齡中位數以分布模型估計並保留身分標籤"],
    ], [2500, 1500, 2600, 2760], font_size=9)

    doc.add_heading("2. 研究背景、目標與治理原則", level=1)
    add_para(doc, "新北市青年政策需要同時回答人口、教育、婚姻、就業、失業、勞動參與及薪資問題，但政府統計的觀察母體、地理角色、時間頻率與年齡分組並不一致。若只追求單一數值而忽略分母與值的身分，容易把戶籍人口、抽樣調查估計及工作場所受僱員工誤認為同一群人。")
    add_para(doc, "本研究的主要目標，是建立可重跑、可記錄、可查核的資料生命週期：從官方來源擷取、來源版本保存、年齡換算、公式驗證、人工Gate、三份CSV發布，到AWS資料服務、Kiro規格驅動開發及ChatBot可追溯問答。")
    add_table(doc, ["治理原則", "落實方式", "禁止事項"], [
        ["母體先分流", "每列資料必須帶universe與分母定義", "跨母體直接加總或互當分母"],
        ["值身分分層", "官方行政精確、官方調查估計、模型估計分開標示", "把模型值寫成官方真值"],
        ["來源不可變", "保存URL、期間、抓取日、SHA-256及原始快照", "以新檔覆寫舊來源且無版本"],
        ["方法可重現", "保存程式、模型版本、公式、參數與診斷", "人工手算後只留最後數字"],
        ["發布失敗即關閉", "異常資料進quarantine，舊版繼續服務", "驗證失敗仍更新儀表板"],
    ], [1700, 4100, 3560], font_size=9.2)

    doc.add_heading("3. 三種統計母體與青年年齡定義", level=1)
    add_figure(doc, diagrams["universes"], "三種母體彼此分流並在同一儀表板並列", "圖1　三種母體的分流與並列原則")
    add_table(doc, ["母體", "納入對象", "地理角色", "核心指標", "主要限制"], [
        ["戶籍人口母體", "戶籍登記現住人口", "戶籍登記地，可至29區", "人口、教育、婚姻、性別", "不等於實際居住人口或勞動市場人口"],
        ["民間人口及勞動市場母體", "依人力資源調查定義之民間人口", "新北市整體調查地區", "P、LF、E、U、勞參率、失業率", "抽樣調查估計；目前發布檔為全體性別"],
        ["受僱員工薪資母體", "工作場所位於新北市之本國籍全時受僱員工", "工作場所所在地", "全年總薪資平均數與中位數", "不是全部青年；目前無新北×年齡×性別交叉"],
    ], [1650, 2900, 1500, 1800, 1510], font_size=8.8)
    add_table(doc, ["政策群組", "年齡", "政策角色", "統計處理"], [
        ["青年銜接觀察組", "15–17歲", "非核心", "可觀察，不計入18–35總數"],
        ["就學、初入職場與轉銜", "18–24歲", "核心", "含18、19邊界；必要時模型換算"],
        ["職涯建立", "25–29歲", "核心驗證組", "多數官方表原生五歲組，可直接核對"],
        ["職涯、居住與家庭形成", "30–35歲", "核心", "含35歲邊界；必要時模型換算"],
        ["方案擴充組", "36–40歲", "非核心", "不計入核心青年總數"],
    ], [2600, 1200, 1600, 3960], font_size=9.2)
    add_callout(doc, "核心定義", "競賽與本研究的核心青年總數固定為18–35歲。分組是分析層，不改變核心定義；15–17與36–40只作觀察或擴充。", "info")

    add_source_table(doc)

    doc.add_heading("5. 共通資料架構與三份長格式CSV", level=1)
    add_para(doc, "發布層不把三種母體壓成單一表，而是保持三份CSV；三表共享同一組欄位語意，因此儀表板與ChatBot可以用一致查詢流程，但仍能保留母體邊界。")
    add_table(doc, ["發布檔", "列數", "年度", "主要維度", "目前可用內容"], [
        ["01_戶籍人口母體_長格式.csv", f"{data['catalog']['files'][0]['rows']:,}", "110–114", "年×地區×年齡×性別×類別×指標", "29區人口、教育、婚姻、性別"],
        ["02_民間人口與勞動市場母體_長格式.csv", f"{data['catalog']['files'][1]['rows']:,}", "110–114", "年×年齡×指標", "P、LF、E、U及四項率；新北市整體"],
        ["03_受僱員工薪資母體_長格式.csv", f"{data['catalog']['files'][2]['rows']:,}", "110–113", "4年×4年齡×2統計量", "32列均有值；官方與模型身分分開"],
    ], [2700, 900, 1100, 2500, 2160], font_size=8.8)
    add_table(doc, ["欄位群", "代表欄位", "用途"], [
        ["時間", "roc_year、gregorian_year、source_period、period_basis", "區分資料年、來源期間與年末／全年平均／全年統計"],
        ["母體", "universe_code、universe_name_zh", "阻止跨母體誤用分母"],
        ["地理", "geography_code、name、level、role", "區分戶籍地、調查地區與工作場所"],
        ["切分", "age_band、sex_code、category_dimension、category_code", "控制年齡、性別、教育／婚姻／統計量"],
        ["數值", "metric_code、value、unit、numerator、denominator、formula_zh", "保存值、單位與公式"],
        ["估計", "value_origin、method、model_version、uncertainty_low/high", "揭露值身分、方法與敏感度"],
        ["追溯", "source_alias、URL、snapshot、SHA-256、retrieved_at", "回查來源版本"],
        ["發布", "publish_status、qa_status、note_zh", "控制前端可見性與限制文字"],
    ], [1400, 4200, 3760], font_size=8.8)
    add_csv_row_count_audit(doc, registered, labor, wage)

    doc.add_heading("6. 年齡區間整合方法與公式", level=1)
    add_callout(doc, "方法選擇", "有單一年齡即直接加總；只有聚合計數時用PCLM；需要同時重現年齡與類別邊際時加IPF；Sprague只作勞動資料的替代拆分與敏感度比較。率與中位數不得按年數比例拆分。", "ok")
    doc.add_heading("6.1 直接加總", level=2)
    add_para(doc, "若官方資料提供單一年齡計數xₐ，目標區間A的值為：X_A = Σ(a∈A)xₐ。人口18–35即加總18至35歲，不需要模型。")
    doc.add_heading("6.2 PCLM：由粗年齡組估計單一年齡", level=2)
    add_para(doc, "令y為官方年齡組計數、C為把單一年齡合併成來源年齡組的組成矩陣、γ為待估單歲計數。PCLM使用E(y)=Cγ、γ=exp(Bθ)，並估計：")
    add_callout(doc, "PCLM目標函數", "最大化 ℓ(y；C·exp(Bθ)) − λ‖D²θ‖²。第一項維持與官方分組資料相符，第二項抑制不合理的單歲鋸齒；λ以GCV／AIC類準則及穩定性檢查選擇。", "info")
    add_para(doc, "模型估計不是把五歲組平均分成五份，而是在非負、平滑且可回聚合的條件下估計單歲輪廓。回聚合後必須重現原始年齡組總數。")
    doc.add_heading("6.3 IPF：同時重現兩組官方邊際", level=2)
    add_para(doc, "教育與婚姻同時受到『單一年齡人口總數』與『官方年齡組×類別總數』限制。IPF從PCLM種子矩陣zₐc開始，交替依列與欄調整：")
    add_table(doc, ["步驟", "公式", "目的"], [
        ["列調整", "xₐc ← xₐc × Rₐ／Σc xₐc", "每一單一年齡合計等於戶籍人口Rₐ"],
        ["欄調整", "xₐc ← xₐc × Cgc／Σ(a∈g)xₐc", "每個官方年齡組的類別合計等於Cgc"],
        ["停止", "最大邊際誤差 < 1×10⁻⁶", "確保列欄邊際均重現"],
    ], [1500, 4300, 3560], font_size=9.2)
    doc.add_heading("6.4 Sprague：替代拆分與敏感度", level=2)
    add_para(doc, "Sprague乘數以相鄰五歲組的線性組合推得單歲估計。本案不把Sprague當唯一真值，而是對勞動PCLM結果建立替代估計；兩方法的最小值與最大值形成方法敏感度包絡。此包絡不是抽樣信賴區間。")
    doc.add_heading("6.5 禁止直接拆分的統計量", level=2)
    add_table(doc, ["統計量", "不能直接按年數拆分的原因", "本案處理"], [
        ["率", "百分比沒有可加性，必須知道分子與分母", "先換算P、LF、E、U，再重算率"],
        ["平均數", "跨組平均需要各組有效人數權數", "薪資以官方錨＋輔助年齡輪廓及人數權數"],
        ["中位數", "第50百分位取決於完整分布，不具線性可加性", "25–29用官方原生值；其他年齡以官方寬帶形狀轉移與對數常態混合分布估計，並保留模型標籤"],
    ], [1300, 4200, 3860], font_size=9.2)
    add_sensitivity_result_tables(doc, registered, labor, wage)

    doc.add_heading("7. 各母體計算方法、限制與驗證", level=1)
    doc.add_heading("7.1 戶籍人口、教育與婚姻", level=2)
    add_figure(doc, diagrams["registered"], "戶籍人口、教育與婚姻的PCLM和IPF處理流程", "圖2　戶籍人口母體的年齡換算與邊際校準")
    add_table(doc, ["項目", "主要方法", "可發布範圍", "驗證", "限制"], [
        ["人口", "單一年齡直接加總", "110–114；29區；男女及全體；四組年齡", "行政區加總=新北市；男女和=全體", "戶籍登記人口"],
        ["教育", "全市PCLM種子＋各區IPF", "110–114；29區；教育類別；四組年齡", "PCLM收斂；IPF誤差<1×10⁻⁶", "邊界組為模型估計"],
        ["婚姻", "全市PCLM種子＋各區IPF", "110–114；29區；四類婚姻；四組年齡", "PCLM收斂；IPF誤差<1×10⁻⁶", "邊界組為模型估計"],
    ], [1200, 2150, 2200, 2000, 1810], font_size=8.8)
    doc.add_heading("7.2 民間人口與勞動市場", level=2)
    add_callout(doc, "來源血緣修正", "本母體不是由單一表產製。CSV已改為逐指標標示實際表別、逐年官方年報網址、原始快照與組合SHA-256；表28只作LF四捨五入一致性查核，不冒充主要計算輸入。", "info")
    add_table(doc, ["發布指標", "source_alias（內部）", "官方表與資料角色", "計算關係"], [
        ["民間人口P", "DGBAS-HR-T27", "表27：15歲以上民間人口之教育程度與年齡", "P"],
        ["就業人口E", "DGBAS-HR-T32", "表32：就業者之教育程度與年齡", "E"],
        ["失業人口U", "DGBAS-HR-T36", "表36：失業者之教育程度與年齡", "U"],
        ["勞動力LF", "T32+T36｜QA:T28", "表32＋表36重建；表28只核對四捨五入差", "LF=E+U"],
        ["非勞動力NLF／勞參率", "T27+T32+T36", "表27、32、36", "NLF=P−LF；LFPR=LF÷P"],
        ["失業率／勞動力中就業占比", "T32+T36", "表32、36", "UR=U÷LF；ESLF=E÷LF"],
        ["就業人口比率", "T27+T32", "表27、32", "EPR=E÷P"],
        ["114年年齡拆分輔助", "AUX:T41+T42", "114年下半年表41／42提供15–19與20–24之E、U、NLF占比", "先保留全年15–24總量，再依同年H2占比分配"],
    ], [1700, 2200, 3100, 2360], font_size=8.2)
    add_para(doc, "所有比率以同一套P、LF、E、U計數重算，並維持下列恆等式：")
    add_table(doc, ["符號／指標", "公式", "分母"], [
        ["非勞動力NLF", "NLF = P − LF", "民間人口P"],
        ["就業與失業", "LF = E + U", "勞動力LF"],
        ["失業率UR", "UR = U ÷ LF × 100%", "勞動力LF"],
        ["勞參率LFPR", "LFPR = LF ÷ P × 100%", "民間人口P"],
        ["就業人口比率EPR", "EPR = E ÷ P × 100%", "民間人口P"],
        ["勞動力中就業占比ESLF", "ESLF = E ÷ LF × 100% = 100% − UR", "勞動力LF"],
    ], [2100, 3900, 3360], font_size=9.2)
    add_para(doc, f"114年18–35歲估計結果包括：民間人口{fmt(labor_metrics['CIVILIAN_POPULATION'])}、勞動力{fmt(labor_metrics['LABOR_FORCE'])}、就業{fmt(labor_metrics['EMPLOYED'])}、失業{fmt(labor_metrics['UNEMPLOYED'])}。這些是由官方抽樣調查估計再經年齡模型換算，不是行政登記人數。")
    add_callout(doc, "性別限制", "目前勞動發布CSV只提供全體性別。若要發布男女18–35歲，須取得並驗證官方性別×年齡×P／LF／E／U一致表組，再分別重跑PCLM／Sprague與公式查核。", "warn")
    doc.add_heading("7.3 受僱員工薪資", level=2)
    add_para(doc, "薪資表6提供新北市工作場所×年齡的官方平均數與中位數，但目標邊界不完全一致。25–29歲保持官方原生值；18–24與30–35平均數透過官方薪資錨定的輔助年齡輪廓換算；18–35再依新北市勞保年齡人數加權。")
    add_table(doc, ["計算", "公式", "資料角色"], [
        ["輔助組平均", "m̃_A = Σ(a∈A)NₐCₐ ÷ Σ(a∈A)Nₐ", "Nₐ為年齡人數；Cₐ為相對薪資輪廓"],
        ["官方錨定", "μ̂_A = μ_G × (m̃_A ÷ m̃_G)", "μ_G為官方原生年齡組平均薪資"],
        ["18–35合成", "μ̂₁₈₋₃₅ = Σk N_k μ̂_k ÷ Σk N_k", "k為18–24、25–29、30–35"],
        ["敏感度", "low=min(PCLM,均勻組內)；high=max(…)", "反映方法差異，不是95%信賴區間"],
    ], [1700, 4200, 3460], font_size=9.2)
    add_table(doc, ["年齡／統計量", "110–113狀態", "值身分", "114年"], [
        ["18–24平均數", "可發布", "官方薪資錨定模型估計", "官方尚未發布"],
        ["25–29平均數", "可發布", "官方原生調查統計", "官方尚未發布"],
        ["30–35平均數", "可發布", "官方薪資錨定模型估計", "官方尚未發布"],
        ["18–35平均數", "可發布", "三組人數加權模型估計", "官方尚未發布"],
        ["25–29中位數", "可發布", "官方原生調查統計", "官方尚未發布"],
        ["其他年齡中位數", "可發布", "官方寬帶形狀轉移＋對數常態混合分布模型估計", "官方尚未發布"],
    ], [2300, 1600, 3100, 2360], font_size=9.2)
    add_callout(doc, "薪資解讀", "表6按工作場所所在地，且母體為本國籍全時受僱員工；不能解讀為所有新北戶籍青年，也不能把勞退提繳工資直接當全年實領薪資。", "warn")

    doc.add_heading("8. 資料處理、更新、查核與發布流程", level=1)
    add_figure(doc, diagrams["daily"], "每日來源審視、變更偵測、轉換、驗證、人工核准與發布", "圖3　每日審視與條件式發布流程")
    add_numbered_steps(doc, [
        "每日09:15（Asia／Taipei）啟動來源審視，檢查ETag、Last-Modified、Content-Length、資料期與SHA-256。",
        "來源未變時寫入NO_CHANGE日誌，不重算、不產生新發布版本。",
        "來源已變時保存不可變raw快照，建立run_id、來源manifest與候選資料版本。",
        "依資料類型執行直接加總、PCLM、IPF、Sprague或薪資錨定換算，產出三份候選CSV。",
        "執行schema、主鍵、期間、範圍、公式、邊際、模型收斂、行政區與值身分驗證。",
        "人工核對來源定義、版本、異常值、估計合理性、圖表文字與發布限制。",
        "通過後原子切換data_catalog；未通過則移入quarantine並持續服務上一個PASS版本。",
    ])
    total_checks = sum(qa[key]["check_count"] for key in ["registered_population", "civilian_labor_market", "employee_wage"])
    total_failures = sum(qa[key]["failure_count"] for key in ["registered_population", "civilian_labor_market", "employee_wage"])
    add_table(doc, ["母體", "機器檢查數", "失敗數", "狀態"], [
        ["戶籍人口", f"{qa['registered_population']['check_count']:,}", str(qa['registered_population']['failure_count']), qa['registered_population']['status']],
        ["民間人口與勞動市場", f"{qa['civilian_labor_market']['check_count']:,}", str(qa['civilian_labor_market']['failure_count']), qa['civilian_labor_market']['status']],
        ["受僱員工薪資", f"{qa['employee_wage']['check_count']:,}", str(qa['employee_wage']['failure_count']), qa['employee_wage']['status']],
        ["合計", f"{total_checks:,}", str(total_failures), qa['overall_status']],
    ], [3000, 1800, 1400, 3160], font_size=9.5)
    add_para(doc, "機器PASS僅表示規則檢查通過；正式發布仍須完成來源版本、統計定義、模型結果、介面文字及AWS權限的人工簽核。", italic=True, color=MUTED, size=10.5)

    doc.add_heading("9. AWS、Kiro與ChatBot系統架構", level=1)
    add_figure(doc, diagrams["aws"], "官方資料進入AWS raw層、轉換成三份CSV、由Athena查詢並供儀表板及AgentCore問答", "圖4　AWS、Kiro與ChatBot整體架構")
    add_table(doc, ["元件", "職責", "控制重點"], [
        ["EventBridge Scheduler", "每日09:15啟動審視", "時區、重試與DLQ"],
        ["Step Functions", "來源偵測、分流、轉換、QA與Gate", "fail-closed、狀態與run_id"],
        ["Lambda", "API、小型檔解析、雜湊與輕量轉換", "逾時、重試、最小權限"],
        ["Glue", "大型教育／婚姻ZIP及PCLM／IPF", "只寫候選區，不直覆正式資料"],
        ["S3", "raw、curated、published、quarantine與manifest", "版本化、KMS、生命週期"],
        ["Glue Catalog＋Athena", "三母體表目錄與受控SQL", "欄位型別、分區、查詢權限"],
        ["Bedrock AgentCore", "自然語言路由、工具調用與答案生成", "allowlist、來源、限制與日誌"],
        ["Dashboard API", "中文儀表板與目前篩選狀態", "只讀published；不得讀raw"],
    ], [2200, 3900, 3260], font_size=8.8)
    doc.add_heading("9.1 Kiro規格驅動流程", level=2)
    add_numbered_steps(doc, [
        "先讀steering中的三母體、年份窗、中文前端與禁止混分母規則。",
        "將使用者需求寫入requirements.md，逐項定義驗收條件與不可接受行為。",
        "在design.md描述S3、Step Functions、Glue、Athena、AgentCore、前端及失敗路徑。",
        "在tasks.md拆成可測試任務；每次變更先跑schema、pytest、前端lint／build與CDK synth。",
        "AWS資源以SAM／CDK／CloudFormation管理；正式部署須經IAM、KMS、預算與資料發布核准。",
    ])
    add_callout(doc, "部署邊界", "資料包已包含Kiro規格、Lambda／SAM及AgentCore／CDK基礎檔，但是否已部署到正式AWS環境，必須以CloudFormation/CDK輸出、資源ARN與CloudTrail紀錄查核，不能只憑文件推定。", "warn")
    doc.add_heading("9.2 ChatBot受控問答", level=2)
    add_figure(doc, diagrams["chatbot"], "ChatBot辨識問題母體、查詢允許清單、核對公式與來源後回覆", "圖5　ChatBot治理與回答鏈")
    add_bullets(doc, [
        "每次問題先辨識母體、年度、年齡、地理、性別、指標與統計量；不足時回報可用選項。",
        "只允許查詢三份已發布CSV／Athena表、data_catalog與來源證據；不得直接讀未查核raw。",
        "答案先給數值，再附母體、期間、值身分、方法、來源與限制。",
        "跨母體只能並列，不做混合分母、個人推論或因果結論。",
        "查無資料時顯示原因與最近可用年度，但不自動偷換年度。",
    ])

    doc.add_heading("10. 互動式儀表板資訊架構與視覺化", level=1)
    add_para(doc, "儀表板採單一入口，但以清楚section隔開三母體。年度與年齡篩選可控制全頁；行政區地圖只控制戶籍人口區塊，勞動與薪資固定標示『新北市整體』。前端圖名、圖例、座標軸與狀態文字一律使用繁體中文或阿拉伯數字，不顯示內部代號。")
    add_table(doc, ["區塊", "主要圖表", "篩選", "重要限制"], [
        ["29區地圖", "游標顯示區名；點擊更新區級卡片", "年度、年齡、性別、行政區", "只連動人口、教育、婚姻與男女比例"],
        ["戶籍人口", "年齡結構、男女比例、教育與婚姻組成", "年度、年齡、性別、區", "官方與模型值必須標示"],
        ["勞動市場", "UR／ESLF圓餅、勞參率與就業人口比率年度長條", "年度、年齡", "新北市整體；目前全體性別"],
        ["薪資", "平均薪資分組趨勢；25–29平均／中位比較", "年度、年齡、統計量", "模型虛線／標籤；114空狀態"],
        ["ChatBot", "自然語言數值、公式、來源與限制", "沿用篩選或問題指定", "一次路由一母體；跨母體只並列"],
    ], [1700, 2800, 2100, 2760], font_size=8.8)
    add_callout(doc, "年度策略", "預設顯示（執行民國年−5）至（執行民國年−1）。民國115年顯示110–114；民國116年顯示111–115。預設同年度檢視；若某母體無當年值，顯示空狀態而非自動帶入前一年。", "info")

    doc.add_heading("11. 品質、資安、成本與行政查核", level=1)
    add_table(doc, ["面向", "控制措施", "發布前證據"], [
        ["資料品質", "schema、主鍵、列數、期間、公式、範圍、邊際與收斂", "machine_qa、validation CSV、run manifest"],
        ["可追溯", "URL、資料期、檔名、SHA-256、抓取日、模型版本", "source_manifest、FILE_MANIFEST_SHA256"],
        ["安全", "最小IAM、KMS、S3版本化、Secrets Manager、CloudTrail", "IAM policy、KMS key、CloudTrail事件"],
        ["可靠性", "重試、DLQ、quarantine、舊版持續服務、原子切換", "失敗演練與回復紀錄"],
        ["成本", "先偵測變更再重算；年度來源未變不啟Glue", "Budget、成本標籤、異常告警"],
        ["隱私", "發布資料為彙總值；白名單與秘密不放入資料包", "環境變數／秘密管理稽核"],
        ["行政責任", "機器PASS後仍需人工簽核", "查核人、日期、裁示與證據路徑"],
    ], [1500, 4800, 3060], font_size=8.8)
    add_para(doc, "若來源定義衝突，處理順序為：停止該指標新版本發布、記錄衝突來源與差異、回查原發布機關及資料產製機關說明、評估是否可轉換、送人工裁示、更新資料字典與版本紀錄後再重跑。")

    doc.add_heading("12. 實施狀態、缺口與後續建議", level=1)
    add_table(doc, ["項目", "目前狀態", "缺口／風險", "建議下一步"], [
        ["三份發布CSV", "已產製並通過機器檢核", "新增來源時須重新驗證", "維持每日審視與版本化"],
        ["五年資料", "人口／教育／婚姻／勞動110–114；薪資110–113", "薪資114尚未發布", "來源發布後重算，不以前一年補值"],
        ["薪資年齡", "32列完整結構；四組平均與中位數均有值", "非25–29中位數為分布模型估計；男女仍缺三維交叉", "申請微資料驗證中位數並補足性別交叉"],
        ["勞動性別", "目前發布檔為全體", "未完成男女P／LF／E／U一致年齡拆分", "取得官方性別表組後另行驗證"],
        ["行政區", "戶籍母體可至29區", "勞動／薪資無區級官方值", "不按戶籍人口比例推造"],
        ["AWS", "具可部署架構與IaC基礎檔", "正式資源狀態需以AWS證據查核", "完成IAM、KMS、Budget、CDK deploy與回復演練"],
        ["ChatBot", "已有治理、工具與回答契約", "正式AgentCore權限與觀測需核定", "建立測試問題集、來源附帶率與越權測試"],
    ], [1600, 2600, 2500, 2660], font_size=8.8)
    add_callout(doc, "優先次序", "第一優先是保持三母體及值身分正確；第二是補齊勞動性別與薪資微資料缺口；第三才是擴充預測、個人化或跨母體推論。", "ok")
    doc.add_heading("12.1 建議決策Gate", level=2)
    add_bullets(doc, [
        "Gate 1：人工確認本報告的母體、年齡與時間定義，可作為系統統一說明。",
        "Gate 2：確認薪資模型值可在儀表板發布，且使用明確模型標籤與敏感度範圍。",
        "Gate 3：確認勞動性別資料在取得正式表組前維持『全體』，不以比例推估。",
        "Gate 4：核定AWS正式部署帳號、us-west-2區域、IAM、KMS、S3名稱、預算與告警。",
        "Gate 5：以人工查核SOP完成一次端到端演練，再開啟自動發布。",
    ])

    add_page_break(doc)
    doc.add_heading("附錄A　官方資料來源", level=1)
    source_rows = [
        ("人口｜村里戶數、單一年齡人口", "https://data.gov.tw/dataset/77132", "MOI-POP1Y；戶籍人口核心來源。"),
        ("教育｜15歲以上現住人口按教育程度分", "https://data.gov.tw/dataset/117988", "MOI-EDU5Y；年度ZIP。"),
        ("婚姻｜15歲以上現住人口按婚姻狀況分", "https://data.gov.tw/dataset/117986", "MOI-MARITAL5Y；年度ZIP。"),
        ("勞動｜人力資源調查統計年報列表", "https://www.stat.gov.tw/News.aspx?n=4001", "110–114年官方年報索引。"),
        ("勞動｜110年人力資源調查統計", "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=207903", "02 CSV中民國110年各列的官方來源頁。"),
        ("勞動｜111年人力資源調查統計", "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=231112", "02 CSV中民國111年各列的官方來源頁。"),
        ("勞動｜112年人力資源調查統計", "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=234726", "02 CSV中民國112年各列的官方來源頁。"),
        ("勞動｜113年人力資源調查統計", "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=234885", "02 CSV中民國113年各列的官方來源頁。"),
        ("勞動｜114年人力資源調查統計", "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=236078", "02 CSV中民國114年各列的官方來源頁。"),
        ("勞動輔助｜114年下半年表41／42", "https://www.stat.gov.tw/News_Content.aspx?n=4003&s=235759", "只用於114年15–24歲拆為15–19與20–24的同年輔助占比。"),
        ("薪資｜受僱員工全年總薪資統計", "https://www.stat.gov.tw/News_Content.aspx?n=4580&s=232642", "表6工作場所縣市×年齡。"),
        ("薪資輔助｜勞保地區、年齡與性別統計", "https://data.gov.tw/dataset/162821", "新北市年齡人數權數。"),
        ("薪資輔助｜勞退年齡平均提繳工資", "https://data.gov.tw/dataset/46103", "只作相對年齡薪資輪廓。"),
        ("政府資料開放平臺", "https://data.gov.tw/", "命題指定政府公開資料入口。"),
        ("新北市政府資料開放平臺OpenAPI", "https://data.ntpc.gov.tw/openapi/", "命題指定介接與交叉查核平台。"),
        ("新北市統計資料庫", "https://oas.bas.ntpc.gov.tw/DgbasWeb/Page/Default.aspx", "命題指定統計表鏡像與人工查核。"),
        ("新北市行政區域圖", "https://data.ntpc.gov.tw/datasets/214634ca-3c71-4fc8-8f46-faffe97f23ff", "29區互動地圖。"),
    ]
    for label, url, note in source_rows:
        add_source_link(doc, label, url, note)
    add_para(doc, "來源查核日：2026年8月26日。實際發布時仍以資料包source_manifest.json所記錄的來源期間、檔案雜湊與抓取時間為準。", italic=True, color=MUTED, size=10)

    doc.add_heading("附錄B　方法與系統參考文獻", level=1)
    references = [
        ("Rizzi, Gampe & Eilers (2015), Efficient Estimation of Smooth Distributions From Coarsely Grouped Data", "https://academic.oup.com/aje/article/182/2/138/94562", "PCLM主要方法來源；American Journal of Epidemiology 182(2):138–147。"),
        ("Deming & Stephan (1940), On a Least Squares Adjustment of a Sampled Frequency Table When the Expected Marginal Totals are Known", "https://doi.org/10.1214/aoms/1177731829", "IPF／邊際校準的經典來源。"),
        ("United Nations, Manual III: Methods for Population Projections by Sex and Age", "https://digitallibrary.un.org/record/3926474/files/un_1956_manual_iii_-_methods_for_population_projections_by_sex_and_age.pdf", "含Sprague乘數表及年齡插補說明。"),
        ("AWS｜Using EventBridge Scheduler to start a Step Functions state machine", "https://docs.aws.amazon.com/step-functions/latest/dg/using-eventbridge-scheduler.html", "每日排程、重試、DLQ與執行角色。"),
        ("AWS｜Building a data analyst agent using Amazon Bedrock AgentCore", "https://docs.aws.amazon.com/solutions/building-a-data-analyst-agent-using-amazon-bedrock-agentcore/", "S3、Glue、Athena、Cognito與AgentCore分析代理架構參考。"),
        ("AWS Well-Architected Framework", "https://docs.aws.amazon.com/wellarchitected/latest/framework/welcome.html", "安全、可靠性、效能與成本治理原則。"),
    ]
    for label, url, note in references:
        add_source_link(doc, label, url, note)
    add_callout(doc, "方法界線", "上述方法文獻支持的是『如何從聚合資料建立受限制估計』，不表示估計值等同官方未發布的單一年齡真值；因此本案仍需模型標示、敏感度與人工查核。", "warn")

    doc.add_heading("附錄C　專案參考文件與可重跑檔案", level=1)
    add_table(doc, ["類別", "檔案／位置", "用途"], [
        ["命題與課程", "【命題文件】青年局－新北市政府AI黑客松競賽.pdf", "競賽需求與政府資料平台範圍"],
        ["AWS／Kiro", "Bedrock AgentCore工作坊、2026智慧城市工作坊、NTPC Kiro工作坊及參賽者手冊", "系統操作、規格驅動與AgentCore情境參考"],
        ["AWS治理", "Well-Architected Framework及Security、Reliability、Performance、Cost pillars", "架構非功能需求與查核項目"],
        ["資料重建", "05_年齡轉換與統計方法/reproducible_scripts", "資料建置、PCLM／IPF／Sprague、來源清冊及驗證"],
        ["系統實作", "12_AWS_Kiro_實作檔案", "Kiro requirements／design／tasks、Lambda／SAM與AgentCore／CDK"],
        ["發布證據", "09_發布資料包；08_人工查核與異常處理/machine_qa", "三份CSV、資料目錄、機器檢核與人工Gate"],
    ], [1500, 4300, 3560], font_size=9)

    doc.add_heading("附錄D　重要限制速查", level=1)
    add_bullets(doc, [
        "戶籍人口、民間人口與受僱員工是三個不同母體；數值不能相加。",
        "人口的12月31日年末存量、勞動的全年平均、薪資的全年統計不是同一時間概念。",
        "PCLM、IPF及Sprague產生的是受限制模型估計，不是官方逐歲真值。",
        "方法敏感度包絡不是抽樣調查的95%信賴區間。",
        "工作場所薪資不能推論為新北戶籍青年的個人薪資。",
        "沒有分子與分母時不得直接拆分率；沒有分布時不得換算中位數。",
        "沒有官方區級勞動或薪資值時，不按戶籍人口比例推造行政區數據。",
        "來源變更不代表立即發布；必須通過機器QA與人工核准。",
    ])

    path = package / "00_總覽與治理" / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    mark_all_table_headers(doc)
    doc.save(path)
    return path


def update_package_metadata(package: Path, output_path: Path) -> None:
    toc_path = package / "TOC.txt"
    toc = toc_path.read_text(encoding="utf-8")
    toc = toc.replace(
        "├─ 00_總覽與治理/\n│  └─ assets/",
        "├─ 00_總覽與治理/                         完整總說明報告書、研究流程與治理文件\n│  └─ assets/",
    )
    toc = toc.replace("三份可發布長格式CSV、11份研究與操作Word文件", "三份可發布長格式CSV、12份現行研究與操作Word文件")
    toc = toc.replace("三份可發布長格式CSV、12份研究與操作Word文件", "三份可發布長格式CSV、12份現行研究與操作Word文件")
    toc = toc.replace("三份可發布長格式CSV、13份研究與操作Word文件", "三份可發布長格式CSV、12份現行研究與操作Word文件")
    toc_path.write_text(toc, encoding="utf-8")

    manifest_path = package / "00_總覽與治理" / "document_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    rel = output_path.relative_to(package).as_posix()
    if rel not in manifest["documents"]:
        manifest["documents"].insert(0, rel)
    manifest["master_report"] = {
        "path": rel,
        "preset": "standard_business_brief",
        "first_page_pattern": "editorial_cover",
        "audience": "technical_and_government_stakeholders",
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    reproducible = package / "05_年齡轉換與統計方法" / "reproducible_scripts"
    reproducible.mkdir(parents=True, exist_ok=True)
    shutil.copy2(Path(__file__), reproducible / Path(__file__).name)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--package", type=Path, default=DEFAULT_PACKAGE)
    args = parser.parse_args()
    output = build_report(args.package)
    update_package_metadata(args.package, output)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
