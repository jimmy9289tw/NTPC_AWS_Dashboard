from __future__ import annotations

import csv
import math
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
PACKAGE = (
    PROJECT
    / "deliverables"
    / "NTPC_Labor_18-35_Sprague_PCLM_Pilot_V1.1_20260825"
)
ANNUAL = PACKAGE / "annual_110_114"
HALF_YEAR = PACKAGE / "half_year_114H2"
ASSET_DIR = PACKAGE / "report_assets"
OUTPUT = PACKAGE / "新北市青年18-35歲就業失業年齡轉換與驗證報告_V1.3.docx"

DOC_SKILL = Path(
    r"C:\Users\LOCAL_USER\.codex\plugins\cache\openai-primary-runtime\documents\26.819.11345\skills\documents"
)
sys.path.insert(0, str(DOC_SKILL / "scripts"))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CAUTION = "FFF4CE"
RISK = "FDE7E9"
PASS_FILL = "E8F5E9"
CJK_FONT = "Microsoft JhengHei"
LATIN_FONT = "Calibri"


def load_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def set_run_font(run, size=None, bold=None, italic=None, color=None, latin=LATIN_FONT):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, *, bold=False, color=None, size=9.2, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, weights, *, header_fill=LIGHT_GRAY, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        set_cell_text(
            table.rows[0].cells[idx],
            header,
            bold=True,
            color=INK,
            size=font_size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(table.rows[0].cells[idx], header_fill)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            try:
                float(str(value).replace("%", "").replace(",", ""))
                numeric = True
            except ValueError:
                numeric = False
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or numeric else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[idx], value, size=font_size, align=align)
    widths = column_widths_from_weights(weights, 9360)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=9360,
        indent_dxa=120,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120},
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_para(doc, text="", *, bold=False, italic=False, size=11, color=None, after=6, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_label_value(doc, label, value, *, fill=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Pt(7)
    p.paragraph_format.right_indent = Pt(7)
    p.paragraph_format.line_spacing = 1.10
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill or LIGHT_GRAY)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "4")
        border.set(qn("w:color"), "D7DBE2")
        borders.append(border)
    p_pr.append(borders)
    label_run = p.add_run(label + "｜")
    set_run_font(label_run, size=10.4, bold=True, color=INK)
    value_run = p.add_run(value)
    set_run_font(value_run, size=10.4)
    return p


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), LATIN_FONT)
    r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
    r_fonts.set(qn("w:eastAsia"), CJK_FONT)
    r_pr.append(r_fonts)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    end = paragraph.add_run(" 頁")
    set_run_font(end, size=9, color=MUTED)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def set_running_furniture(doc: Document):
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    for header in (section.header, section.even_page_header):
        p = header.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run("新北市青年 AI 儀表板｜研究方法與驗證紀錄")
        set_run_font(run, size=9, color=MUTED)
    for footer in (section.footer, section.even_page_footer):
        p = footer.paragraphs[0]
        p.text = ""
        add_page_number(p)


def heading(doc, level, text):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def start_new_page(doc):
    """Use a non-empty page-break-before paragraph for stable Word/LibreOffice pagination."""
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 0.1
    r = p.add_run("\u00a0")
    set_run_font(r, size=1, color=WHITE)
    return p


def make_workflow(path: Path):
    width, height = 1500, 560
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = r"C:\Windows\Fonts\msjh.ttc"
    font = ImageFont.truetype(font_path, 28)
    small = ImageFont.truetype(font_path, 22)
    title = ImageFont.truetype(font_path, 34)
    draw.text((60, 35), "全年就業／失業年齡轉換流程", font=title, fill=(31, 77, 120))
    boxes = [
        ("官方年報", ["表27／28", "表32／36"]),
        ("恆等式調和", ["LF = E + U", "NLF = P - LF"]),
        ("114輔助拆分", ["同年H2", "E／U／NLF占比"]),
        ("年齡展開", ["PCLM主估", "Sprague對照"]),
        ("目標與查核", ["三組＋18–35", "回加總／非負"]),
    ]
    x0, y0, bw, bh, gap = 65, 165, 250, 220, 36
    for idx, (top, bottom) in enumerate(boxes):
        x = x0 + idx * (bw + gap)
        fill = (232, 238, 245) if idx not in (2, 5) else ((255, 244, 206) if idx == 2 else (232, 245, 233))
        draw.rounded_rectangle((x, y0, x + bw, y0 + bh), radius=18, fill=fill, outline=(46, 116, 181), width=3)
        bbox = draw.textbbox((0, 0), top, font=font)
        draw.text((x + (bw - (bbox[2] - bbox[0])) / 2, y0 + 40), top, font=font, fill=(11, 37, 69))
        yy = y0 + 105
        for line in bottom:
            bbox = draw.textbbox((0, 0), line, font=small)
            draw.text((x + (bw - (bbox[2] - bbox[0])) / 2, yy), line, font=small, fill=(70, 70, 70))
            yy += 37
        if idx < len(boxes) - 1:
            ax = x + bw + 7
            ay = y0 + bh // 2
            draw.line((ax, ay, ax + gap - 14, ay), fill=(46, 116, 181), width=5)
            draw.polygon([(ax + gap - 14, ay), (ax + gap - 28, ay - 10), (ax + gap - 28, ay + 10)], fill=(46, 116, 181))
    draw.text((60, 455), "守恆原則：所有拆分後單一年齡值必須回加總至官方發布年齡帶；模型值不得改標為官方直接發布值。", font=small, fill=(122, 90, 0))
    img.save(path)


def make_grouped_bar(path: Path, results: list[dict]):
    selected = [row for row in results if row["method"] == "PCLM" and row["age_group"] in ("18-24", "25-29", "30-35")]
    data = {(int(row["roc_year"]), row["age_group"]): float(row["unemployment_rate_pct"]) for row in selected}
    years = [110, 111, 112, 113, 114]
    groups = ["18-24", "25-29", "30-35"]
    colors = [(46, 116, 181), (91, 155, 213), (165, 165, 165)]
    width, height = 1450, 800
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = r"C:\Windows\Fonts\msjh.ttc"
    title = ImageFont.truetype(font_path, 36)
    axis = ImageFont.truetype(font_path, 24)
    small = ImageFont.truetype(font_path, 21)
    draw.text((80, 35), "110–114年新北市青年年齡組失業率（PCLM主估）", font=title, fill=(31, 77, 120))
    left, top, right, bottom = 105, 120, 1370, 680
    ymax = 14
    for tick in range(0, ymax + 1, 2):
        y = bottom - (bottom - top) * tick / ymax
        draw.line((left, y, right, y), fill=(225, 225, 225), width=2)
        label = f"{tick}%"
        draw.text((55, y - 14), label, font=small, fill=(90, 90, 90))
    cluster = (right - left) / len(years)
    bar_w = 48
    bar_gap = 12
    for yi, year in enumerate(years):
        center = left + cluster * (yi + 0.5)
        total = len(groups) * bar_w + (len(groups) - 1) * bar_gap
        x_start = center - total / 2
        for gi, group in enumerate(groups):
            value = data[(year, group)]
            x1 = x_start + gi * (bar_w + bar_gap)
            y1 = bottom - (bottom - top) * value / ymax
            draw.rectangle((x1, y1, x1 + bar_w, bottom), fill=colors[gi])
            text = f"{value:.1f}"
            bbox = draw.textbbox((0, 0), text, font=small)
            draw.text((x1 + (bar_w - (bbox[2] - bbox[0])) / 2, y1 - 27), text, font=small, fill=(50, 50, 50))
        ytext = f"{year}年"
        bbox = draw.textbbox((0, 0), ytext, font=axis)
        draw.text((center - (bbox[2] - bbox[0]) / 2, bottom + 20), ytext, font=axis, fill=(50, 50, 50))
    lx = 790
    for gi, group in enumerate(groups):
        x = lx + gi * 180
        draw.rectangle((x, 735, x + 25, 760), fill=colors[gi])
        draw.text((x + 36, 731), group, font=axis, fill=(50, 50, 50))
    img.save(path)


def make_population_bridge(path: Path):
    width, height = 1500, 790
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = r"C:\Windows\Fonts\msjh.ttc"
    title = ImageFont.truetype(font_path, 34)
    head = ImageFont.truetype(font_path, 27)
    body = ImageFont.truetype(font_path, 22)
    small = ImageFont.truetype(font_path, 19)

    draw.text((55, 35), "P的資料橋接：共用維度，但母體與指標分開", font=title, fill=(31, 77, 120))
    lanes = [
        (55, 135, 430, 560, (232, 238, 245), "人力資源調查", ["P_LFS：15歲以上民間人口", "LF／E／U／NLF", "全年為12個月平均", "LFPR、EPR只能用P_LFS"]),
        (535, 135, 910, 560, (232, 245, 233), "戶籍人口API", ["P_REG：戶籍登記現住人口", "可到單一年齡、男／女、月份", "年度可另算12個月平均", "適合人口結構與人口金字塔"]),
        (1015, 135, 1445, 560, (255, 244, 206), "共享語意層", ["geography／year／month", "age_group／sex／time_grain", "source_id／population_universe", "value_origin／method／status"]),
    ]
    for x1, y1, x2, y2, fill, lane_title, lines in lanes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=fill, outline=(46, 116, 181), width=3)
        bbox = draw.textbbox((0, 0), lane_title, font=head)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y1 + 35), lane_title, font=head, fill=(11, 37, 69))
        yy = y1 + 115
        for line in lines:
            draw.ellipse((x1 + 30, yy + 8, x1 + 41, yy + 19), fill=(46, 116, 181))
            draw.text((x1 + 55, yy), line, font=body, fill=(55, 55, 55))
            yy += 68

    for x1, x2 in ((430, 535), (910, 1015)):
        y = 348
        draw.line((x1 + 10, y, x2 - 20, y), fill=(46, 116, 181), width=5)
        draw.polygon([(x2 - 20, y), (x2 - 38, y - 12), (x2 - 38, y + 12)], fill=(46, 116, 181))

    draw.rounded_rectangle((55, 625, 1445, 735), radius=16, fill=(253, 231, 233), outline=(180, 70, 80), width=3)
    draw.text((85, 650), "禁止：用P_REG取代P_LFS計算LFPR／EPR，或把年末戶籍人口與全年平均勞動資料直接當同一母體。", font=head, fill=(130, 35, 45))
    draw.text((85, 697), "允許：以共享維度並列、篩選、做人口結構背景；所有圖表保留來源與母體標籤。", font=small, fill=(70, 70, 70))
    img.save(path)


def fmt(value, digits=2):
    return f"{float(value):,.{digits}f}"


def build():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    annual_results = load_csv(ANNUAL / "annual_target_age_results.csv")
    annual_validation = load_csv(ANNUAL / "annual_validation_checks.csv")
    sensitivity = load_csv(ANNUAL / "annual_method_sensitivity.csv")
    aux_validation = load_csv(ANNUAL / "annual_114_h2_auxiliary_share_validation.csv")
    aux_split = load_csv(ANNUAL / "annual_114_h2_auxiliary_split.csv")
    backtest = load_csv(ANNUAL / "annual_coarsening_backtest.csv")
    h2_results = load_csv(HALF_YEAR / "target_age_results.csv")
    h2_reconciliation = load_csv(HALF_YEAR / "source_rounding_reconciliation.csv")

    workflow = ASSET_DIR / "labor_age_workflow.png"
    trend = ASSET_DIR / "unemployment_rate_annual.png"
    population_bridge = ASSET_DIR / "population_data_bridge.png"
    make_workflow(workflow)
    make_grouped_bar(trend, annual_results)
    make_population_bridge(population_bridge)

    doc = Document()
    configure_styles(doc)
    set_running_furniture(doc)

    # Memo masthead opening.
    add_para(doc, "研究方法與查核報告", bold=True, size=11, color=BLUE, after=4)
    add_para(doc, "新北市青年（18–35歲）就業與失業\n年齡區間轉換試算與驗證", bold=True, size=24, color=INK, after=7)
    add_para(doc, "Sprague／PCLM｜110–114年全年序列＋114年下半年方法驗證", size=13, color=MUTED, after=16)
    add_table(
        doc,
        ["項目", "內容"],
        [
            ["版本", "V1.3｜2026-08-25"],
            ["研究範圍", "新北市；18–24、25–29、30–35、18–35歲"],
            ["時間範圍", "110–114年全年平均；114年下半年另作方法驗證"],
            ["正式主估", "PCLM；Sprague作敏感度與合理性對照"],
            ["發布定位", "官方人力資源調查量與其衍生模型估計；非戶籍行政精確值"],
        ],
        [1.3, 5.2],
        font_size=10.0,
    )

    heading(doc, 1, "一、結論先行")
    add_label_value(doc, "為何不只114年下半年？", "110–114年全年官方表均已納入；114年下半年只用來補足該年全年表未拆出的15–19／20–24細節，並進行Sprague與PCLM交叉驗證。", fill=BLUE_GRAY)
    add_label_value(doc, "可否形成年度比較？", "可以。年度值必須使用人力資源調查的12個月平均，不得誤標為12月底期末值，也不得與114年下半年平均混成同一條趨勢。", fill=PASS_FILL)
    add_label_value(doc, "目前最可靠的讀法", "25–29歲為官方原生年齡帶；18–24、30–35及18–35含邊界拆分，均需標為模型估計。114年18–24失業率與勞動力中就業占比可列暫估；勞參率及就業人口比率（EPR）仍待微資料或全年細年齡表複核。", fill=CAUTION)
    add_label_value(doc, "本版修正", "更正P的統計身分，新增ESLF／UR、LFPR、EPR儀表板規格、戶籍人口資料橋接及男女性別可用性邊界。", fill=BLUE_GRAY)

    heading(doc, 1, "二、研究問題、範圍與定義")
    add_para(doc, "本試算回答兩個行政與資料工程問題：第一，如何把政府發布但年齡分組不同的就業／失業資料，一致轉為競賽所需的18–35歲；第二，如何讓每一個轉換結果可追溯、可重算、可查核，並清楚區分官方調查估計與模型估計。")
    add_table(
        doc,
        ["代號", "定義", "本案計算"],
        [
            ["P／P_LFS", "15歲以上民間人口（人力資源調查母體）", "表27控制人口量；P = LF + NLF"],
            ["LF", "勞動力人口（調查估計）", "LF = E + U"],
            ["E", "就業人口（調查估計）", "官方表32；千人"],
            ["U", "失業人口（調查估計）", "官方表36；千人"],
            ["NLF", "非勞動力人口（調查估計／調和值）", "本案以NLF = P - LF調和"],
            ["P_REG", "戶籍登記現住人口（行政統計）", "人口結構用；不得替代P_LFS"],
            ["UR", "失業率", "UR = U / LF × 100%"],
            ["LFPR", "勞動力參與率", "LFPR = LF / P × 100%"],
            ["EPR", "就業人口比率", "EPR = E / P × 100%"],
            ["ESLF", "勞動力中就業占比", "ESLF = E / LF × 100% = 100% - UR"],
        ],
        [0.75, 2.5, 3.25],
    )
    add_para(doc, "重要更正：LF、E、U及調查表中的NLF屬人力資源調查樣本估計；P則是人力資源調查定義下的15歲以上民間人口控制總數，估計程序以同時期性別×年齡戶籍人口校準。P不是單純的未校準樣本估計，也不等於戶政的戶籍登記現住人口。來源人數均以千人發布，表間可能有1千人取整差；本案保留E、U、P邊際，再以LF=E+U及NLF=P−LF調和。", italic=True, color=DARK_BLUE, after=8)

    heading(doc, 2, "2.1 目標年齡帶與數值身分")
    add_table(
        doc,
        ["年齡帶", "處理", "發布標記"],
        [
            ["18–24", "需估計18、19歲在15–19（或114年全年15–24）中的分布", "模型估計"],
            ["25–29", "官方原生五歲年齡帶；不需拆邊界", "官方調查估計（年齡帶直接）"],
            ["30–35", "30–34官方原生值＋估計35歲", "模型估計"],
            ["18–35", "18–24＋25–29＋30–35", "模型估計"],
        ],
        [1.0, 3.7, 1.8],
    )

    heading(doc, 2, "2.2 公式總覽、分母與計算順序")
    add_para(doc, "所有比率一律先把目標年齡內的人數加總，再以加總後的分子除以加總後的分母；不得平均既有年齡組比率。令X代表P、LF、E、U或NLF，a代表單一年齡。", color=DARK_BLUE)
    formula_rows = [
        ["人口恆等式", "P = LF + NLF", "民間人口由勞動力與非勞動力構成"],
        ["勞動力恆等式", "LF = E + U", "勞動力由就業者與失業者構成"],
        ["非勞動力", "NLF = P - LF", "由同一人口母體反推"],
        ["失業率UR", "UR = U / LF × 100%", "分母為勞動力，不含NLF"],
        ["勞參率LFPR", "LFPR = LF / P × 100%", "分母為15歲以上民間人口"],
        ["就業人口比率EPR", "EPR = E / P × 100%", "民間人口中目前就業的比例"],
        ["勞動力中就業占比ESLF", "ESLF = E / LF × 100%", "亦等於100% - UR"],
    ]
    add_table(doc, ["項目", "公式", "分母／意義"], formula_rows, [1.65, 2.4, 2.45], font_size=8.8)
    add_para(doc, "年齡加總公式（X ∈ {P, LF, E, U, NLF}）：", bold=True, color=INK, after=3)
    add_para(doc, "X₁₈₋₂₄ = Σₐ₌₁₈²⁴ Xₐ；X₂₅₋₂₉ = Σₐ₌₂₅²⁹ Xₐ；X₃₀₋₃₅ = Σₐ₌₃₀³⁵ Xₐ；X₁₈₋₃₅ = Σₐ₌₁₈³⁵ Xₐ", bold=True, size=11.5, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_para(doc, "率值重算公式（以18–35為例）：UR₁₈₋₃₅ = ΣU₁₈₋₃₅ / ΣLF₁₈₋₃₅；LFPR₁₈₋₃₅ = ΣLF₁₈₋₃₅ / ΣP₁₈₋₃₅；EPR₁₈₋₃₅ = ΣE₁₈₋₃₅ / ΣP₁₈₋₃₅；ESLF₁₈₋₃₅ = ΣE₁₈₋₃₅ / ΣLF₁₈₋₃₅。", size=10.2, color=DARK_BLUE)
    add_label_value(doc, "公式範例｜110年25–29歲", "P=268、LF=252、E=236、U=16、NLF=16（千人）。UR=16÷252=6.35%；LFPR=252÷268=94.03%；EPR=236÷268=88.06%；ESLF=236÷252=93.65%=100%-6.35%。", fill=BLUE_GRAY)

    heading(doc, 1, "三、官方資料與年度範圍")
    add_para(doc, "每個年度使用同一組主計總處人力資源調查年報表，地區列固定為New Taipei City。表29、37提供官方直接率值；核心人數由表27、28、32、36讀取與調和。")
    add_table(
        doc,
        ["表次", "內容", "本案角色"],
        [
            ["表27", "15歲以上民間人口之教育程度與年齡", "P_LFS控制人口邊際"],
            ["表28", "勞動力之教育程度與年齡", "LF四捨五入對照"],
            ["表29", "年齡組別勞動力參與率", "25–29官方LFPR交叉檢視"],
            ["表32", "就業者之教育程度與年齡", "E邊際"],
            ["表36", "失業者之教育程度與年齡", "U邊際"],
            ["表37", "年齡組別失業率", "官方發布率值交叉檢視"],
            ["表41", "新北市歷年年齡別失業率", "114H2 15–24 UR=12.4%"],
            ["表42", "新北市勞動力狀況與年齡", "114H2直接列兩組人數"],
        ],
        [0.8, 3.8, 1.9],
    )
    source_pages = {
        110: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=207903",
        111: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=231112",
        112: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=234726",
        113: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=234885",
        114: "https://www.stat.gov.tw/News_Content.aspx?n=4001&s=236078",
    }
    for year, url in source_pages.items():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"民國{year}年年報：")
        set_run_font(r, size=10.5, bold=True)
        add_hyperlink(p, "開啟官方年報頁面", url)

    add_para(doc, "資料欄位變更：110–113年全年表提供15–19與20–24；114年全年表只提供15–24合計。因此114年需加入同年下半年表42作輔助拆分，但最終仍回加總至114年全年官方15–24總數。", italic=True, color=DARK_BLUE, after=6)
    heading(doc, 2, "3.1 表42欄位確認：H2已直接提供15–19與20–24")
    add_para(doc, "已逐格確認DGBAS-NTPC-LABOR-AGE-T42-114H2.xlsx之工作表「42」：第19列為15–19歲、第20列為20–24歲；B至F欄依序為P、LF、E、U、NLF。這兩個H2年齡組是官方表格直接發布，不是PCLM或Sprague拆分結果。男性年齡資料位於第51–64列，女性位於第73–86列，因此114H2具備按性別各自進行年齡轉換的必要人數。", bold=True, color=DARK_BLUE)
    add_table(
        doc,
        ["T42列", "年齡", "P", "LF", "E", "U", "NLF"],
        [
            ["18", "15–24合計", "354", "138", "121", "17", "216"],
            ["19", "15–19", "165", "17", "14", "2", "148"],
            ["20", "20–24", "189", "122", "107", "15", "67"],
        ],
        [0.65, 1.35, 0.9, 0.9, 0.9, 0.9, 0.9],
        font_size=8.6,
    )
    add_para(doc, "原表單位為千人且四捨五入：子列LF為17+122=139，與15–24合計列138相差1千人；子列NLF為148+67=215，與合計列216亦相差1千人。故本案不能把所有顯示整數直接當成未取整真值，而須在±0.5千人的發布取整區間內調和。", italic=True, color=MUTED, size=9.4)

    heading(doc, 2, "3.2 P與勞動指標的官方估計性質")
    add_para(doc, "主計總處方法說明載明：人力資源調查採分層二段抽樣，樣本資料先作不偏估計，再以同時期性別×年齡戶籍人口校準；除人口總數等以同時期人口統計為基準者外，其餘調查結果存在抽樣估計誤差。因此本報告不再以一句『P、LF、E、U皆為抽樣調查估計』概括所有欄位。")
    add_table(
        doc,
        ["欄位", "正確身分", "儀表板標示"],
        [
            ["P_LFS", "調查人口母體下、經戶籍年齡×性別控制校準的人口量；非P_REG", "官方人力資源調查人口量"],
            ["LF／E／U", "官方抽樣調查估計；正式率以未取整估計值計算", "官方調查估計"],
            ["NLF", "官方調查估計；本案為守恆另以P−LF調和", "調查估計／調和值"],
            ["模型年齡帶", "由上述官方量經PCLM或Sprague轉換", "模型估計；顯示method"],
        ],
        [1.0, 3.7, 1.8],
        font_size=8.7,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("官方估計方法：")
    set_run_font(r, size=10.2, bold=True)
    add_hyperlink(p, "人力資源調查統計之編製方法與估計誤差", "https://www.stat.gov.tw/public/Data/11122153012VTN8S5UB.pdf")

    heading(doc, 1, "四、處理流程與計算方法")
    picture = doc.add_picture(str(workflow), width=Inches(6.25))
    picture._inline.docPr.set("descr", "新北市青年就業與失業年齡轉換流程圖")
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p = doc.add_paragraph("圖1　官方來源、年齡轉換與查核流程")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    for run in p.runs:
        set_run_font(run, size=9.5, italic=True, color=MUTED)

    heading(doc, 2, "4.1 PCLM（正式主估）")
    add_para(doc, "PCLM將每個發布年齡帶視為單一年齡潛在分布的加總。令 y 為官方分組人數、C為分組加總矩陣、B為樣條基底、θ為待估參數，單一年齡期望值為 γ = exp(Bθ)，分組期望值為 μ = Cγ。模型最大化：")
    add_para(doc, "ℓ(θ) − (λ/2)‖D²θ‖²", bold=True, size=13, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(doc, "第一項是Poisson概似，第二項懲罰相鄰年齡曲線的二階差分，使結果平滑且保持非負。本案在λ = 0.1、1、10、100、1,000、10,000中，以內部年齡帶逐一留一交叉驗證之Poisson deviance選取λ；最後按每一官方年齡帶比例校準，使單一年齡回加總完全等於來源邊際。")

    heading(doc, 2, "4.2 Sprague（敏感度基準）")
    add_para(doc, "Sprague第五差分公式以固定係數矩陣S將五歲組y拆為單一年齡x：x = Sy。此法在官方人口統計與聯合國人口估計中長期使用，優點是保留原五歲組總數；限制是輸入需為等寬五歲組，且小量事件（例如失業人數）可能產生負值。本案遇負值即標警示、截為零，並重新按官方五歲組校準；因此Sprague不作正式主估，只作PCLM敏感度對照。")

    heading(doc, 2, "4.3 114年全年15–24的輔助拆分")
    add_para(doc, "114年全年缺15–19／20–24，但同年下半年表42已直接發布兩組。H2不需要先做年齡拆分；只因T42人數取整至千人，需先用T41的114H2官方15–24失業率12.4%作取整調和，再計算年齡組成比例。")
    add_para(doc, "取整區間：若官方顯示人數為X̃千人，則潛在未取整值X*滿足 X* ∈ [X̃−0.5, X̃+0.5)；若官方顯示率為r̃（小數一位），則r* ∈ [r̃−0.05, r̃+0.05)。", bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_para(doc, "勞動力調和限制：LF* = E* + U*；r* = U* / LF* × 100%。在所有取整區間與恆等式均成立的條件下，選擇與官方顯示值距離最小的可行解：", size=10.5)
    add_para(doc, "min Σ(X*−X̃)² + [(r*−r̃)/0.1]²，X ∈ {LF,E,U}", bold=True, size=11.5, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    add_para(doc, "兩個H2子組另滿足 X*₁₅₋₁₉ + X*₂₀₋₂₄ = X*₁₅₋₂₄，且各自仍位於T42發布整數的±0.5千人區間內。調和只修正取整不一致，沒有改變官方直接提供的15–19／20–24年齡分組。", italic=True, color=DARK_BLUE)
    h2_lookup = {row["age_group"]: row for row in h2_reconciliation}
    add_table(
        doc,
        ["年齡", "P*", "LF*", "E*", "U*", "NLF*"],
        [
            [
                age,
                fmt(h2_lookup[age]["reconciled_P_thousand"], 3),
                fmt(h2_lookup[age]["reconciled_LF_thousand"], 3),
                fmt(h2_lookup[age]["reconciled_E_thousand"], 3),
                fmt(h2_lookup[age]["reconciled_U_thousand"], 3),
                fmt(h2_lookup[age]["reconciled_NLF_thousand"], 3),
            ]
            for age in ("15-19", "20-24")
        ],
        [1.0, 1.1, 1.1, 1.1, 1.1, 1.1],
        font_size=8.5,
    )
    add_para(doc, "以上星號值為取整調和後的中間計算值，仍屬官方調查估計的衍生值，不是官方新增發布精度。對E、U、NLF三個互斥成分c分別計算H2之15–19占比：", italic=True, color=MUTED, size=9.4)
    add_para(doc, "s₍c₎ = N₁₅₋₁₉,₍c,H2₎ / (N₁₅₋₁₉,₍c,H2₎ + N₂₀₋₂₄,₍c,H2₎)", bold=True, size=12.5, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=7)
    add_para(doc, "再以 N₁₅₋₁₉,₍c,annual₎ = s₍c₎ × N₁₅₋₂₄,₍c,annual₎ 分配全年值；N₂₀₋₂₄,₍c,annual₎ = N₁₅₋₂₄,₍c,annual₎ − N₁₅₋₁₉,₍c,annual₎。最後令LF = E + U、P = LF + NLF，使每個成分的全年15–24總數與會計恆等式維持精確。")
    add_table(
        doc,
        ["成分", "114H2之15–19占比", "全年15–24", "配置15–19", "配置20–24"],
        [
            [
                row["component"],
                f"{100*float(row['auxiliary_15_19_share']):.2f}%",
                fmt(row["annual_15_24_total_thousand"]),
                fmt(row["allocated_15_19_thousand"]),
                fmt(row["allocated_20_24_thousand"]),
            ]
            for row in aux_split
        ],
        [0.75, 1.5, 1.15, 1.55, 1.55],
    )

    heading(doc, 1, "五、驗證步驟與合理性判準")
    validation_steps = [
        ["V1", "來源定位", "逐年確認官方頁面、表次、New Taipei City列與單位", "錯表／全國值立即停止"],
        ["V2", "欄位結構", "比對110–113與114表頭位移及15–24分組變更", "欄位對應必須留存"],
        ["V3", "四捨五入調和", "保留E、U、P邊際；計算LF=E+U、NLF=P-LF", "表28差異≤1千人"],
        ["V4", "非負檢查", "逐單一年齡檢查E、U、NLF≥0", "負值不得發布"],
        ["V5", "回加總", "單一年齡重聚合至來源年齡帶", "誤差≤1e−8千人"],
        ["V6", "25–29錨點", "模型結果必須完全復原官方25–29", "誤差≤1e−8千人"],
        ["V7", "方法敏感度", "比較PCLM與Sprague的E、U及率值", "差異列入方法不確定性"],
        ["V8", "粗分組回測", "將110–113的15–19／20–24先合為15–24再反推", "辨認114粗分組風險"],
        ["V9", "輔助比例檢核", "114H2占比與110–113全年歷史範圍比較", "超界則標WARN"],
        ["V10", "時間口徑", "全年與半年分開，全年定義為12個月平均", "不得混成同一趨勢"],
    ]
    add_table(doc, ["編號", "檢核", "作法", "判準"], validation_steps, [0.65, 1.2, 3.4, 1.25], font_size=8.6)
    add_para(doc, "回加總公式：對任一官方年齡帶g與成分X，Σₐ∈g X̂ₐ = Xg,source；回加總誤差 = |Σₐ∈g X̂ₐ − Xg,source|。本案判準為≤1e−8千人。", bold=True, color=DARK_BLUE, size=10.0)

    fail_count = sum(row["status"] == "FAIL" for row in annual_validation)
    warn_count = sum(row["status"] == "WARN" for row in annual_validation)
    add_label_value(doc, "自動驗證結果", f"40筆年度目標結果；FAIL = {fail_count}，WARN = {warn_count}。警示均為Sprague校準前對小量E／U產生負值；PCLM未出現結構驗證失敗。", fill=PASS_FILL if fail_count == 0 else RISK)

    heading(doc, 2, "5.1 為什麼方法合理")
    add_para(doc, "合理性不是以『模型名稱知名』直接成立，而是同時滿足：方法與資料結構相容、官方邊際守恆、結果非負、官方原生年齡帶可完全復原、不同方法差異可量化、輔助假設可被歷史資料檢查。PCLM能處理不等寬分組並透過懲罰概似平滑，適合作為主估；Sprague具有常見且可查的第五差分係數，適合作為等寬五歲資料的敏感度基準。")

    heading(doc, 1, "六、年度結果（PCLM主估）")
    picture = doc.add_picture(str(trend), width=Inches(5.95))
    picture._inline.docPr.set("descr", "110至114年新北市18至24、25至29、30至35歲失業率群組長條圖")
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p = doc.add_paragraph("圖2　110–114年三個青年年齡組失業率；114年18–24含同年H2輔助拆分")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        set_run_font(run, size=9.3, italic=True, color=MUTED)

    start_new_page(doc)
    heading(doc, 2, "6.1 人數結果")
    count_rows = []
    for row in annual_results:
        if row["method"] != "PCLM":
            continue
        count_rows.append(
            [
                row["roc_year"],
                row["age_group"],
                fmt(row["civilian_population_thousand"]),
                fmt(row["labor_force_thousand"]),
                fmt(row["employed_thousand"]),
                fmt(row["unemployed_thousand"]),
                fmt(row["not_in_labor_force_thousand"]),
            ]
        )
    add_table(doc, ["年度", "年齡", "P", "LF", "E", "U", "NLF"], count_rows[:10], [0.65, 0.8, 1.0, 1.0, 1.0, 1.0, 1.05], font_size=8.0)
    start_new_page(doc)
    add_para(doc, "表6-1　年度人數結果（續）", italic=True, color=MUTED, size=9.2, after=4)
    add_table(doc, ["年度", "年齡", "P", "LF", "E", "U", "NLF"], count_rows[10:], [0.65, 0.8, 1.0, 1.0, 1.0, 1.0, 1.05], font_size=8.0)
    add_para(doc, "註：人數單位均為千人。25–29為官方原生年齡帶；其餘含模型拆分。兩位小數僅供檢視，不代表官方發布精度提升。", italic=True, color=MUTED, size=9.2, after=6)

    start_new_page(doc)
    heading(doc, 2, "6.2 比率結果（由同一組人數重新計算）")
    rate_rows = []
    for row in annual_results:
        if row["method"] != "PCLM":
            continue
        rate_rows.append(
            [
                row["roc_year"],
                row["age_group"],
                fmt(row["unemployment_rate_pct"]),
                fmt(row["labor_force_participation_rate_pct"]),
                fmt(row["employment_to_population_rate_pct"]),
                fmt(row["employment_share_of_labor_force_pct"]),
            ]
        )
    add_table(doc, ["年度", "年齡", "UR(%)", "LFPR(%)", "EPR(%)", "ESLF(%)"], rate_rows[:10], [0.65, 0.8, 1.25, 1.25, 1.25, 1.3], font_size=8.1)
    start_new_page(doc)
    add_para(doc, "表6-2　年度比率結果（續）", italic=True, color=MUTED, size=9.2, after=4)
    add_table(doc, ["年度", "年齡", "UR(%)", "LFPR(%)", "EPR(%)", "ESLF(%)"], rate_rows[10:], [0.65, 0.8, 1.25, 1.25, 1.25, 1.3], font_size=8.1)
    add_para(doc, "公式：UR=U/LF；LFPR=LF/P；EPR=E/P；ESLF=E/LF=100%-UR。25–29之表內率值是以官方千人人數重算的QA值；正式儀表板應另顯示表29／37官方直接發布率值。", italic=True, color=MUTED, size=9.2, after=6)

    official_rates = {
        "110": (6.5, 94.1),
        "111": (5.6, 93.2),
        "112": (5.6, 92.4),
        "113": (5.6, 92.5),
        "114": (6.4, 93.6),
    }
    official_rows = []
    for row in annual_results:
        if row["method"] != "PCLM" or row["age_group"] != "25-29":
            continue
        official_ur, official_lfpr = official_rates[row["roc_year"]]
        official_rows.append([
            row["roc_year"],
            fmt(row["unemployment_rate_pct"]),
            f"{official_ur:.1f}",
            f"{float(row['unemployment_rate_pct'])-official_ur:.2f}",
            fmt(row["labor_force_participation_rate_pct"]),
            f"{official_lfpr:.1f}",
            f"{float(row['labor_force_participation_rate_pct'])-official_lfpr:.2f}",
        ])
    heading(doc, 3, "6.2.1 25–29官方率值交叉檢視")
    add_table(doc, ["年度", "UR重算", "UR官方", "差", "LFPR重算", "LFPR官方", "差"], official_rows, [0.65, 1.05, 1.0, 0.8, 1.15, 1.1, 0.85], font_size=8.0)
    add_para(doc, "25–29人數與官方年齡帶一致；率值差異來自官方千人人數取整。官方直接率使用未公開的未取整調查估計值計算，故正式展示以官方率為主、重算率作公式查核。", italic=True, color=DARK_BLUE, size=9.3)

    start_new_page(doc)
    heading(doc, 2, "6.3 方法敏感度")
    ur_sens = [row for row in sensitivity if row["metric"] == "unemployment_rate_pct" and row["age_group"] != "25-29"]
    max_row = max(ur_sens, key=lambda row: abs(float(row["Sprague_minus_PCLM"])))
    add_label_value(doc, "PCLM vs. Sprague", f"110–114年18–24、30–35與18–35失業率的最大絕對差為{abs(float(max_row['Sprague_minus_PCLM'])):.2f}個百分點（{max_row['roc_year']}年、{max_row['age_group']}）。此差異可作方法敏感度範圍，但不是抽樣95%信賴區間。", fill=BLUE_GRAY)
    sens_rows = []
    for row in ur_sens:
        sens_rows.append([row["roc_year"], row["age_group"], fmt(row["PCLM_value"]), fmt(row["Sprague_value"]), fmt(row["Sprague_minus_PCLM"])])
    add_para(doc, "方法差公式：Δmethod = RateSprague − RatePCLM；絕對敏感度 = |Δmethod|。此值衡量方法選擇差異，不是抽樣信賴區間。", bold=True, color=DARK_BLUE, size=9.8)
    add_table(doc, ["年度", "年齡", "PCLM(%)", "Sprague(%)", "差(百分點)"], sens_rows[:8], [0.75, 0.9, 1.45, 1.45, 1.45], font_size=8.5)
    start_new_page(doc)
    add_para(doc, "表6-3　方法敏感度（續）", italic=True, color=MUTED, size=9.2, after=4)
    add_table(doc, ["年度", "年齡", "PCLM(%)", "Sprague(%)", "差(百分點)"], sens_rows[8:], [0.75, 0.9, 1.45, 1.45, 1.45], font_size=8.5)

    start_new_page(doc)
    heading(doc, 1, "七、114年輔助假設與誤差警示")
    aux_rows = []
    for row in aux_validation:
        aux_rows.append(
            [
                row["component"],
                f"{100*float(row['historical_15_19_share_min']):.2f}–{100*float(row['historical_15_19_share_max']):.2f}%",
                f"{100*float(row['auxiliary_114H2_15_19_share']):.2f}%",
                fmt(row["distance_outside_historical_range_percentage_point"]),
                row["status"],
            ]
        )
    table = add_table(doc, ["成分", "110–113歷史範圍", "114H2占比", "超界(百分點)", "狀態"], aux_rows, [0.8, 1.75, 1.25, 1.4, 1.3], font_size=8.8)
    for row_idx, data in enumerate(aux_rows, start=1):
        set_cell_shading(table.rows[row_idx].cells[-1], PASS_FILL if data[-1] == "PASS" else CAUTION)
    add_para(doc, "超界距離公式：d = max(L−s, 0, s−H)，其中[L,H]為110–113歷史範圍、s為114H2占比。NLF之d=68.76%−64.56%=4.21個百分點。", bold=True, color=DARK_BLUE, size=9.8)
    add_para(doc, "E與U的114H2占比均位於110–113年全年歷史範圍內；NLF之15–19占比高出歷史上緣4.21個百分點。因此114年18–24的失業率與ESLF（由E、U決定）相對可用，但LFPR與EPR涉及NLF，必須標為暫估並列入人工查核。", bold=True, color=DARK_BLUE, after=8)

    # Coarsening diagnostic summary.
    nlf_errors = [float(row["absolute_error_thousand"]) for row in backtest if row["component"] == "NLF"]
    add_label_value(doc, "為何不能只把15–24直接交給PCLM？", f"110–113回測顯示，先合併15–19與20–24、再只靠粗分組反推時，NLF兩組的平均絕對誤差約{sum(nlf_errors)/len(nlf_errors):.2f}千人、最大{max(nlf_errors):.2f}千人。這證明114年需加入同年H2細分輔助，而非直接接受粗分組模型結果。", fill=CAUTION)

    start_new_page(doc)
    heading(doc, 1, "八、114年下半年方法驗證（不得混入年度趨勢）")
    h2_count_rows = []
    h2_rate_rows = []
    for row in h2_results:
        if row["age_group"] not in ("18-24", "25-29", "30-35", "18-35"):
            continue
        method_label = "PCLM" if row["method"] == "PCLM" else "Sprague（非負校準）"
        h2_count_rows.append([row["age_group"], method_label, fmt(row["employed_thousand"]), fmt(row["unemployed_thousand"])])
        h2_rate_rows.append([row["age_group"], method_label, fmt(row["unemployment_rate_pct"]), fmt(row["labor_force_participation_rate_pct"]), fmt(row["employment_to_population_rate_pct"]), fmt(row["employment_share_of_labor_force_pct"])])
    add_table(doc, ["年齡", "方法", "E(千人)", "U(千人)"], h2_count_rows, [1.0, 2.5, 1.5, 1.5], font_size=8.6)
    add_table(doc, ["年齡", "方法", "UR(%)", "LFPR(%)", "EPR(%)", "ESLF(%)"], h2_rate_rows, [0.8, 1.9, 1.05, 1.1, 1.1, 1.15], font_size=8.2)
    add_para(doc, "114年下半年表42直接具有15–19與20–24，故可在同一來源上完整比較PCLM與Sprague。此處用途是驗證方法與提供114全年輔助比例；期間口徑為2025年7–12月平均，不能當作114年全年值。", italic=True, color=DARK_BLUE)

    start_new_page(doc)
    heading(doc, 1, "九、可發布程度與人工查核表")
    publish_rows = [
        ["25–29就業／失業／失業率", "可用", "官方原生年齡帶；仍屬官方調查估計", "確認表次、年度與單位"],
        ["18–24、30–35、18–35失業率（110–113）", "可用但須標模型", "PCLM主估；Sprague最大差0.24百分點", "儀表板顯示方法與來源"],
        ["114年18–24失業率", "暫估可用", "E/U輔助比例通過歷史範圍檢核", "取得微資料後複核"],
        ["114年18–24 LFPR／EPR", "暫緩正式發布", "NLF輔助比例超出歷史範圍4.21百分點", "需微資料或官方全年細分"],
        ["114年18–24 ESLF", "暫估可用", "ESLF=E/LF=100%-UR，不含NLF", "與UR同步複核"],
        ["114H2數值", "僅方法驗證", "半年平均，不是全年", "不得混入年度趨勢"],
    ]
    add_table(doc, ["項目", "建議", "理由", "人工查核"], publish_rows, [2.0, 1.2, 2.1, 1.2], font_size=8.7)

    start_new_page(doc)
    heading(doc, 1, "十、限制與後續改善")
    limitations = [
        ("抽樣不確定性", "目前方法差異不是正式95%信賴區間；若要正式誤差界線，需人力資源調查微資料、樣本設計與權數，或官方變異數／標準誤表。"),
        ("官方四捨五入", "來源人數只到千人，任何兩位小數皆為模型計算結果，不得宣稱增加官方精度。"),
        ("114輔助期間", "用下半年組成比例拆全年，假設H2的年齡組成可代表全年；NLF已觸發警示。"),
        ("平滑假設", "PCLM假設單一年齡分布平滑；Sprague假設鄰近五歲組可用固定第五差分係數展開。青年就學轉職轉銜可能造成局部突變。"),
        ("年度口徑", "人力資源調查年報值為12個月平均，不是12月底；跨資料庫比較時必須另建time_grain與reference_period欄位。"),
    ]
    for title, body in limitations:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.10
        r = p.add_run(f"{title}：")
        set_run_font(r, size=11, bold=True, color=DARK_BLUE)
        r2 = p.add_run(body)
        set_run_font(r2, size=11)

    heading(doc, 2, "10.1 建議的下一個Gate")
    next_gate = [
        ["G4-1", "向主計總處申請或確認可得之人力資源調查微資料、權數與變異數計算方式", "正式抽樣誤差"],
        ["G4-2", "以微資料直接加權估計18–24、25–29、30–35及18–35", "消除年齡邊界模型"],
        ["G4-3", "將PCLM／Sprague／輔助比例、λ、版本、輸入雜湊寫入AWS處理紀錄", "可追溯與重算"],
        ["G4-4", "儀表板同時顯示value_origin、method、time_grain、provisional_flag", "避免誤讀"],
    ]
    add_table(doc, ["Gate", "動作", "目的"], next_gate, [0.8, 4.4, 1.3], font_size=9.0)

    start_new_page(doc)
    heading(doc, 1, "十一、儀表板規格、人口橋接與性別維度")
    add_label_value(doc, "設計原則", "圖表先回答行政問題，再揭露來源、母體、時間口徑與估計身分。任何可切換的年齡、年度與性別篩選，都必須保證分子、分母來自同一資料母體與同一期間。", fill=BLUE_GRAY)

    heading(doc, 2, "11.1 指定圖表與目前可用程度")
    dashboard_rows = [
        ["ESLF／UR年度圓餅圖", "每個年度一個雙切片甜甜圈；ESLF+UR=100%", "110–114全體可做；114年18–24須暫估標記", "同頁另提供100%堆疊長條切換，較易跨年比較"],
        ["LFPR年度長條圖", "年度×年齡組分組長條；全體／男／女篩選", "全體可做；114年18–24暫緩正式發布", "公開版Y軸固定0–100%，柱上顯示一位小數"],
        ["EPR年度長條圖", "年度×年齡組分組長條；E÷P_LFS", "全體可做；114年18–24暫緩正式發布", "不得以P_REG作分母"],
        ["人口年齡結構", "單一年齡人口金字塔、年齡組占比、年度變化", "可接P_REG；與P_LFS分開", "並列背景，不混成同一人口數"],
        ["男／女分組", "全域sex篩選及男女性別差距圖", "114H2人數可試算；年度目標年齡組尚未全數就緒", "性別各自估計後再加總查核"],
    ]
    add_table(doc, ["圖表", "呈現方式", "資料就緒度", "必要限制"], dashboard_rows, [1.55, 2.1, 1.7, 1.6], font_size=8.2)

    heading(doc, 2, "11.2 圖表計算公式與顯示規則")
    dashboard_formula_rows = [
        ["圓餅UR切片", "Pie_UR(g,s,t) = 100 × U(g,s,t) / [E(g,s,t)+U(g,s,t)]", "官方原生帶優先用表37直接UR"],
        ["圓餅ESLF切片", "Pie_ESLF(g,s,t) = 100 − Pie_UR(g,s,t)", "由同一未取整UR取補數，保證合計100%"],
        ["勞參率", "LFPR(g,s,t) = 100 × [E(g,s,t)+U(g,s,t)] / P_LFS(g,s,t)", "原生帶優先用表29直接LFPR"],
        ["就業人口比率", "EPR(g,s,t) = 100 × E(g,s,t) / P_LFS(g,s,t)", "P_LFS與E須同年、同期間、同年齡、同性別"],
        ["戶籍年度平均", "P_REG_ANNAVG(a,s,y) = (1/n) × Σₘ P_REG(a,s,y,m)", "n應為12；不足12月即標不完整"],
        ["性別差距", "Gap_metric(g,t) = Metric_male(g,t) − Metric_female(g,t)", "單位為百分點；不得把兩個率直接相加"],
    ]
    add_table(doc, ["項目", "公式", "規則"], dashboard_formula_rows, [1.4, 3.6, 1.5], font_size=8.3)
    add_para(doc, "g為年齡組、s為性別、t為期間。圖表計算使用未取整值，畫面標籤再四捨五入至小數一位；不得先將分組率四捨五入後再平均。圓餅圖是同一年度、同一年齡組、同性別下的勞動力組成，而不是人口組成。", italic=True, color=DARK_BLUE, size=9.4)

    heading(doc, 2, "11.3 P能否接到年齡人口結構？")
    add_para(doc, "可以在系統語意層連接，但不能把兩者當成同一個P。人力資源調查P_LFS與戶籍人口P_REG的母體、時間粒度與統計性質不同；正確方式是共用地理、年齡、性別與時間維度，保留兩套人口量與來源標籤。")
    picture = doc.add_picture(str(population_bridge), width=Inches(6.25))
    picture._inline.docPr.set("descr", "人力資源調查人口與戶籍人口透過共享維度橋接但不混用母數的示意圖")
    p = doc.add_paragraph("圖3　P_LFS與P_REG的資料橋接與禁止混用規則")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        set_run_font(run, size=9.3, italic=True, color=MUTED)
    add_table(
        doc,
        ["比較面向", "P_LFS", "P_REG"],
        [
            ["人口定義", "15歲以上民間人口（調查母體）", "戶籍登記現住人口（行政統計）"],
            ["常見時間", "全年12個月平均／半年平均", "月末；可另算12個月平均"],
            ["主要用途", "LFPR、EPR與勞動市場人數", "人口結構、單一年齡、人口金字塔"],
            ["能否互代", "否", "否"],
        ],
        [1.25, 2.7, 2.55],
        font_size=8.7,
    )
    add_para(doc, "若同頁比較全年勞動指標與戶籍人口，優先把12個月月末戶籍人口計算為年度平均；若僅有12月底，欄位必須標記P_REG_DEC_END，不得標成年度平均。P_REG可以作背景分布與篩選維度，但不得用來重算LFPR或EPR。", bold=True, color=DARK_BLUE, size=9.7)

    heading(doc, 2, "11.4 男女性別呈現：可做範圍與缺口")
    sex_rows = [
        ["114H2表42", "P／LF／E／U／NLF均有男、女與細年齡帶", "可各自跑PCLM／Sprague，建立18–24、25–29、30–35、18–35半年試算", "尚未寫入目前target_age_results.csv"],
        ["110–114表29／37", "LFPR與UR有Total／Male／Female及官方原生年齡帶", "25–29可直接按性別年度比較；其他原生帶亦可另頁呈現", "率值不足以反推所有性別人數"],
        ["110–114表27／28／32／36", "目前資料包提取的是全體人數", "可支援全體目標年齡組", "不足以發布完整年度18–35男女性別EPR與人數"],
        ["正式性別模型", "需性別×年齡的P／E／U／NLF人數或微資料權數", "每一性別獨立轉換、重算率值", "不得先算全體後按人口性別比拆分"],
    ]
    add_table(doc, ["來源", "性別資訊", "可做", "限制"], sex_rows, [1.3, 2.15, 2.25, 1.3], font_size=8.2)
    add_label_value(doc, "114年官方直接例", "表29之25–29歲LFPR：全體93.6%、男94.0%、女93.1%；表37之25–29歲UR：全體6.4%、男6.4%、女6.5%。這些可直接畫年度性別圖，但不得與模型年齡組混標。", fill=PASS_FILL)
    add_para(doc, "性別模型查核公式：對每一成分X∈{P,LF,E,U,NLF}，|X_male + X_female − X_total|應落在來源取整可解釋範圍；並分別檢查LF_s=E_s+U_s、P_s=LF_s+NLF_s、回加總與非負。", bold=True, color=DARK_BLUE, size=9.7)

    start_new_page(doc)
    heading(doc, 2, "11.5 建議頁面配置與其他呈現")
    layout_rows = [
        ["全域篩選", "年度／期間、年齡組、性別、方法、資料身分", "篩選後所有分子分母同步"],
        ["KPI列", "P_LFS、LF、E、U、NLF、UR、LFPR、EPR", "每張卡顯示來源、官方／模型、暫估狀態"],
        ["勞動組成", "ESLF／UR年度甜甜圈＋100%堆疊長條切換", "甜甜圈看單年組成，長條看跨年差異"],
        ["年度比較", "LFPR、EPR分組長條；UR趨勢折線", "固定色彩與年齡組順序"],
        ["人口結構", "P_REG人口金字塔、單一年齡曲線、青年占比", "獨立標示戶籍人口母體"],
        ["性別差距", "男／女並列長條或差距點圖", "差值用百分點；小樣本需警示"],
        ["方法與品質", "PCLM與Sprague敏感度、來源表、更新日期、查核狀態", "讓估計可追溯、可紀錄、可查核"],
    ]
    add_table(doc, ["區塊", "建議視覺", "行政／查核目的"], layout_rows, [1.25, 3.3, 1.95], font_size=8.5)
    add_para(doc, "另建議加入年齡組×年度熱圖，以快速辨認UR、LFPR、EPR高低；以及『資料品質抽屜』顯示source_id、reference_period、time_grain、value_origin_class、method、provisional_flag與人工查核結果。若未取得抽樣變異數，不得顯示虛構的信賴區間。", italic=True, color=DARK_BLUE, size=9.7)

    heading(doc, 2, "11.6 上線前查核表")
    dashboard_checks = [
        ["DB-QA-01", "ESLF+UR以未取整值計算等於100%", "自動", "阻擋發布"],
        ["DB-QA-02", "LFPR、EPR分母必須為同一筆P_LFS", "自動", "阻擋發布"],
        ["DB-QA-03", "P_REG與P_LFS不可合併成單一人口欄", "資料模型", "阻擋發布"],
        ["DB-QA-04", "全年、半年、月末與年度平均不得混線", "自動", "阻擋發布"],
        ["DB-QA-05", "性別結果須通過男+女≈全體及各恆等式", "自動＋人工", "不通過則隱藏性別切換"],
        ["DB-QA-06", "每個圖表可展開來源URL、表次、方法與版本", "人工", "未完成不得對外"],
        ["DB-QA-07", "114年18–24 LFPR／EPR保留暫估警示", "人工", "不得顯示為正式值"],
    ]
    add_table(doc, ["編號", "查核事項", "方式", "處置"], dashboard_checks, [0.95, 3.8, 1.1, 1.15], font_size=8.4)

    start_new_page(doc)
    heading(doc, 1, "十二、參考文獻與方法來源")
    sources = [
        ("Rizzi, Gampe & Eilers (2015), Efficient Estimation of Smooth Distributions From Coarsely Grouped Data, DOI", "https://doi.org/10.1093/aje/kwv020"),
        ("DemoTools graduation documentation（Sprague、PCLM等）", "https://timriffe.github.io/DemoTools/articles/graduation_with_demotools.html"),
        ("United Nations, World Population Prospects 2019 Methodology（Sprague第五差分應用與限制）", "https://population.un.org/wpp/assets/Files/WPP2019_Methodology.pdf"),
        ("主計總處，人力資源調查統計之編製方法與估計誤差", "https://www.stat.gov.tw/public/Data/11122153012VTN8S5UB.pdf"),
        ("主計總處，人力資源調查資料品質說明", "https://www.stat.gov.tw/News_NoticeCalendar_Content_temp.aspx?MetaI_D=149&n=3717&year=2026"),
        ("主計總處人力資源調查年報總覽", "https://www.stat.gov.tw/News.aspx?n=4001"),
    ]
    for title, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(title + "：")
        set_run_font(r, size=10.3, bold=True)
        add_hyperlink(p, "開啟來源", url)

    heading(doc, 1, "附錄A｜可重算檔案")
    files = [
        ("全年主結果", "annual_110_114/annual_target_age_results.csv"),
        ("單一年齡結果", "annual_110_114/annual_single_age_estimates.csv"),
        ("全年驗證資料包", "annual_110_114/（來源調和、驗證、敏感度、回測、114輔助比例）"),
        ("執行筆記", "annual_110_114/annual_labor_age_graduation_audit.ipynb"),
        ("114H2試算", "half_year_114H2/target_age_results.csv"),
    ]
    add_table(doc, ["用途", "相對路徑"], files, [1.8, 4.7], font_size=9.0)
    add_para(doc, "本報告不把任何估計值宣稱為官方行政精確值。正式對外發布前，應由統計／業務承辦人依第九節查核表簽認。", bold=True, color=DARK_BLUE, after=0)

    # Preset audit facts are encoded above: Letter, 1-inch margins, 9360 DXA,
    # standard_business_brief style tokens, memo masthead, fixed table geometry.
    doc.core_properties.title = "新北市青年18-35歲就業失業年齡轉換與驗證報告"
    doc.core_properties.subject = "PCLM與Sprague年度試算、完整公式、驗證、人口橋接與儀表板規格"
    doc.core_properties.author = "新北市青年AI研究方法工作檔"
    doc.core_properties.keywords = "新北市, 青年, 18-35, 失業率, 就業, PCLM, Sprague, 儀表板, 性別, 戶籍人口"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
