from __future__ import annotations

import csv
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def emit(label: str, value: object) -> None:
    print(f"\n===== {label} =====")
    if isinstance(value, str):
        print(value)
    else:
        print(json.dumps(value, ensure_ascii=False, indent=2, default=str))


def inspect_docx(path: Path) -> None:
    doc = Document(path)
    lines: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            lines.append(text)
    for idx, table in enumerate(doc.tables, start=1):
        lines.append(f"[TABLE {idx}]")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
    emit(str(path), "\n".join(lines))


def inspect_csv(path: Path, limit: int = 15) -> None:
    for encoding in ("utf-8-sig", "cp950", "utf-8"):
        try:
            with path.open("r", encoding=encoding, newline="") as fh:
                reader = csv.reader(fh)
                rows = []
                for _, row in zip(range(limit), reader):
                    rows.append(row)
            emit(str(path), {"encoding": encoding, "rows": rows})
            return
        except UnicodeDecodeError:
            continue
    emit(str(path), "Unable to decode CSV")


def inspect_json(path: Path, limit: int = 5) -> None:
    raw = path.read_text(encoding="utf-8-sig")
    obj = json.loads(raw)
    summary: object = obj
    if isinstance(obj, list):
        summary = {"length": len(obj), "sample": obj[:limit]}
    elif isinstance(obj, dict):
        compact = {}
        for key, value in obj.items():
            if isinstance(value, list):
                compact[key] = {"length": len(value), "sample": value[:limit]}
            else:
                compact[key] = value
        summary = compact
    emit(str(path), summary)


def inspect_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as zf:
        items = [{"name": i.filename, "size": i.file_size} for i in zf.infolist()]
        previews: dict[str, object] = {}
        for info in zf.infolist()[:3]:
            if not info.filename.lower().endswith((".csv", ".txt")):
                continue
            with zf.open(info) as fh:
                prefix = fh.read(65536)
            encoding = "utf-8-sig" if prefix.startswith(b"\xef\xbb\xbf") else "utf-8"
            text = prefix.decode(encoding, errors="replace")
            previews[info.filename] = {
                "encoding": f"{encoding} (replacement-safe preview)",
                "lines": text.splitlines()[:12],
            }
        emit(str(path), {"items": items, "previews": previews})


def inspect_xlsx(path: Path, row_limit: int = 18, col_limit: int = 18) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict[str, object] = {}
    for ws in wb.worksheets:
        rows = []
        for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if ridx > row_limit:
                break
            rows.append(list(row[:col_limit]))
        out[ws.title] = {"max_row": ws.max_row, "max_column": ws.max_column, "rows": rows}
    emit(str(path), out)


def inspect_pdf(path: Path, page_limit: int = 8) -> None:
    try:
        import pdfplumber
    except ImportError:
        emit(str(path), "pdfplumber unavailable")
        return
    with pdfplumber.open(path) as pdf:
        pages = []
        for page in pdf.pages[:page_limit]:
            pages.append((page.extract_text() or "")[:7000])
    emit(str(path), {"page_count": len(pdf.pages), "sample_pages": pages})


def main() -> None:
    for arg in sys.argv[1:]:
        path = Path(arg)
        suffix = path.suffix.lower()
        if suffix == ".docx":
            inspect_docx(path)
        elif suffix == ".csv":
            inspect_csv(path)
        elif suffix == ".json":
            inspect_json(path)
        elif suffix == ".zip":
            inspect_zip(path)
        elif suffix == ".xlsx":
            inspect_xlsx(path)
        elif suffix == ".pdf":
            inspect_pdf(path)
        else:
            emit(str(path), f"Unsupported suffix: {suffix}")


if __name__ == "__main__":
    main()
