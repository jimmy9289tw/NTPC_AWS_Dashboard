from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "deliverables" / "NTPC_Youth_18-35_Government_Open_Data_Package_V2.1_20260824"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def docx_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def population_total(path: Path, start_age: int, end_age: int) -> int:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    rows = [row for row in payload["responseData"] if str(row["district_code"]).startswith("650")]
    return sum(
        int(row[f"people_age_{age:03d}_m"]) + int(row[f"people_age_{age:03d}_f"])
        for row in rows
        for age in range(start_age, end_age + 1)
    )


def check_hash_manifest() -> tuple[int, list[str]]:
    manifest_path = PACKAGE / "00_總覽與治理" / "file_hashes_sha256.csv"
    with manifest_path.open("r", encoding="utf-8-sig", newline="") as handle:
        records = list(csv.DictReader(handle))
    failures: list[str] = []
    for record in records:
        path = PACKAGE / record["relative_path"]
        if not path.exists():
            failures.append(f"missing:{record['relative_path']}")
            continue
        if path.stat().st_size != int(record["bytes"]):
            failures.append(f"size:{record['relative_path']}")
        if sha256(path) != record["sha256"]:
            failures.append(f"sha256:{record['relative_path']}")
    expected = {
        path.relative_to(PACKAGE).as_posix()
        for path in PACKAGE.rglob("*")
        if path.is_file()
        and path != manifest_path
        and not any(part.startswith("_qa") for part in path.parts)
    }
    listed = {record["relative_path"] for record in records}
    failures.extend(f"unlisted:{item}" for item in sorted(expected - listed))
    failures.extend(f"stale:{item}" for item in sorted(listed - expected))
    return len(records), failures


def main() -> None:
    failures: list[str] = []
    docs = sorted(PACKAGE.rglob("*V2.1.docx"))
    if len(docs) != 8:
        failures.append(f"docx_count={len(docs)}")
    texts = {path.name: docx_text(path) for path in docs}
    corpus = "\n".join(texts.values())
    for forbidden in ("V2.0", "3.76", "MOI-POP1Y-current.csv"):
        if forbidden in corpus:
            failures.append(f"forbidden_text:{forbidden}")
    required = {
        "01_年齡人口結構_處理方法與驗證範例_V2.1.docx": ("OpenAPI", "7,752", "7,781", "845,938"),
        "02_失業率_處理方法與驗證範例_V2.1.docx": ("6.4%", "6.52%", "LF", "PCLM"),
        "04_教育程度_處理方法與驗證範例_V2.1.docx": ("臺灣地區", "只有兩組邊際", "IPF"),
        "07_整合查核與AWS_Kiro流程_V2.1.docx": ("source_period", "retrieved_at", "EventBridge", "fail closed"),
    }
    for filename, terms in required.items():
        text = texts.get(filename, "")
        for term in terms:
            if term not in text:
                failures.append(f"missing_text:{filename}:{term}")

    periods = {
        "11412": (7752, 4, 1032, 845938),
        "11507": (7781, 4, 1039, 832214),
    }
    for period, (row_count, page_count, ntpc_rows, total) in periods.items():
        path = PACKAGE / "01_年齡人口結構" / "raw" / f"MOI-POP1Y-{period}.json"
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        rows = payload["responseData"]
        if payload["period"] != period or payload["totalDataSize"] != row_count or len(rows) != row_count:
            failures.append(f"population_meta:{period}")
        if payload["totalPage"] != page_count:
            failures.append(f"population_pages:{period}")
        if sum(str(row["district_code"]).startswith("650") for row in rows) != ntpc_rows:
            failures.append(f"population_ntpc_rows:{period}")
        if population_total(path, 18, 35) != total:
            failures.append(f"population_total:{period}")

    integration_path = PACKAGE / "07_整合與查核" / "reference" / "青年18-35歲資料整合表_V1.5.csv"
    with integration_path.open("r", encoding="utf-8-sig", newline="") as handle:
        integration = list(csv.DictReader(handle))
    ids = [row["record_id"] for row in integration]
    if len(integration) != 72 or len(set(ids)) != 72:
        failures.append(f"integration_rows_or_ids:{len(integration)}/{len(set(ids))}")
    t50 = [row for row in integration if row["source_alias"] == "DGBAS-HR-T50"]
    if not t50 or any(row["source_geography"] != "臺灣地區（無縣市維度）" for row in t50):
        failures.append("t50_geography")

    hash_count, hash_failures = check_hash_manifest()
    failures.extend(hash_failures)
    print(f"DOCX={len(docs)} INTEGRATION_ROWS={len(integration)} HASH_ROWS={hash_count}")
    print("POP_11412_18_35=845938 POP_11507_18_35=832214")
    if failures:
        print("QA_FAILED")
        for failure in failures:
            print(failure)
        raise SystemExit(1)
    print("QA_PASSED")


if __name__ == "__main__":
    main()
