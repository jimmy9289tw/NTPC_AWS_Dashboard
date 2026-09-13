from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "deliverables" / "NTPC_Youth_Complete_Data_Package_G5_V3.2_20260826"
OUTPUT = PACKAGE / "07_AWS與Kiro" / "Kiro本地實作交接流程與驗證SOP_V1.0.docx"


# Design tokens: compact_reference_guide + editorial_cover, CJK font override.
NAVY = "173B57"
TEAL = "167D87"
TEAL_LIGHT = "E5F3F4"
BLUE_LIGHT = "EDF4F8"
WARM = "F7F5F0"
AMBER = "E8A33D"
AMBER_LIGHT = "FFF3DE"
RED = "B84942"
RED_LIGHT = "FBEAE7"
INK = "1E2F3B"
MID = "536775"
GRID = "CBD6DD"
WHITE = "FFFFFF"
ASCII_FONT = "Calibri"
CJK_FONT = "Microsoft JhengHei"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def set_keep_together(paragraph) -> None:
    paragraph.paragraph_format.keep_together = True


def set_east_asia(run, font=CJK_FONT) -> None:
    run.font.name = ASCII_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    set_east_asia(run)


def add_bottom_border(paragraph, color=GRID, size="6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color, before, after in (
        ("Title", 31, NAVY, 0, 12),
        ("Subtitle", 14, TEAL, 0, 20),
        ("Heading 1", 21, NAVY, 16, 7),
        ("Heading 2", 15, TEAL, 12, 4),
        ("Heading 3", 12, NAVY, 8, 3),
    ):
        style = styles[style_name]
        style.font.name = ASCII_FONT
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["List Bullet"].font.name = ASCII_FONT
    styles["List Bullet"]._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    styles["List Number"].font.name = ASCII_FONT
    styles["List Number"]._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    code = styles.add_style("CodeBlock", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string(INK)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(3)
    code.paragraph_format.keep_together = True

    label = styles.add_style("SmallLabel", 1)
    label.font.name = ASCII_FONT
    label.font.size = Pt(8.5)
    label.font.bold = True
    label.font.color.rgb = RGBColor.from_string(TEAL)
    label._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    label.paragraph_format.space_after = Pt(2)


def add_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run("新北市青年資料證據台  ·  Kiro本地實作交接SOP")
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(NAVY)
        set_east_asia(run)
        add_bottom_border(p)

        footer = section.footer
        p = footer.paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("V1.0  ·  2026-08-31  ·  ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID)
        set_east_asia(run)
        add_page_number(p)


def add_text(doc, text, style=None, bold=False, color=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    set_east_asia(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_east_asia(r)
    return p


def add_number(doc, text, level=0):
    p = doc.add_paragraph(style="List Number" if level == 0 else "List Number 2")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_east_asia(r)
    return p


def add_callout(doc, title, body, fill=TEAL_LIGHT, accent=TEAL):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.line_spacing = Pt(1)
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [0.10, 6.40])
    set_cant_split(table.rows[0])
    table.cell(0, 0).text = ""
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    set_east_asia(r)
    p2 = table.cell(0, 1).add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.color.rgb = RGBColor.from_string(INK)
    set_east_asia(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    set_cell_shading(table.cell(0, 0), WARM)
    set_cell_margins(table.cell(0, 0), top=90, start=120, bottom=90, end=120)
    p = table.cell(0, 0).paragraphs[0]
    p.style = doc.styles["CodeBlock"]
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_standard_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(WHITE)
        set_east_asia(r)
    for row_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        set_cant_split(table.rows[-1])
        for idx, value in enumerate(row_data):
            if row_idx % 2:
                set_cell_shading(cells[idx], BLUE_LIGHT)
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor.from_string(INK)
            set_east_asia(r)
    return table


def add_step(doc, number, title, purpose, inputs, action, evidence, stop_rule):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [0.72, 5.78])
    set_cant_split(table.rows[0])
    badge, body = table.rows[0].cells
    set_cell_shading(badge, NAVY)
    set_cell_shading(body, BLUE_LIGHT)
    badge_p = badge.paragraphs[0]
    badge_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    badge_p.paragraph_format.space_after = Pt(0)
    r = badge_p.add_run(str(number))
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    set_east_asia(r)
    p = body.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    set_east_asia(r)
    for label, value in (("目的", purpose), ("輸入", inputs), ("執行", action), ("證據", evidence), ("停止條件", stop_rule)):
        p = body.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"{label}：")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(TEAL)
        set_east_asia(r)
        r = p.add_run(value)
        set_east_asia(r)
    doc.add_paragraph("↓", style=None).alignment = WD_ALIGN_PARAGRAPH.CENTER


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    configure_styles(doc)
    add_header_footer(doc)

    # Cover
    doc.add_paragraph().paragraph_format.space_after = Pt(34)
    p = add_text(doc, "KIRO 實作交接", style="SmallLabel")
    p.runs[0].font.size = Pt(10)
    add_text(doc, "Kiro本地實作交接\n流程與驗證SOP", style="Title")
    add_text(doc, "新北市青年資料證據台 · G8 Gate · 不含AWS部署授權", style="Subtitle")
    add_callout(
        doc,
        "本版結論",
        "現在只做流程確認、本地建置、驗證與可查核交付；不需要AWS帳號、CostCenter、預算、憑證或雲端寫入。未來可將本SOP直接交給Kiro作為實作入口。",
    )
    meta = add_standard_table(
        doc,
        ["版本", "日期", "Gate狀態", "目前邊界"],
        [["V1.0", "2026-08-31", "READY_FOR_KIRO", "僅本地／不部署AWS"]],
        [0.85, 1.2, 1.65, 2.8],
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_text(doc, "文件控制", style="Heading 2")
    add_bullet(doc, "適用資料包：NTPC_Youth_Complete_Data_Package_G5_V3.2_20260826。")
    add_bullet(doc, "適用角色：資料工程、前端、AI Agent、治理查核與Kiro實作人員。")
    add_bullet(doc, "本文件是執行SOP，不取代官方統計定義、Code Book或來源證據。")
    p = add_text(doc, "一頁執行摘要", style="Heading 1")
    p.paragraph_format.page_break_before = True
    add_callout(
        doc,
        "現在交給Kiro的任務",
        "從既有資料包讀取四份CSV、治理規格、Code Book、lineage與八個唯讀工具契約；在本地完成建置、測試、差異報告與可重跑證據。",
    )
    add_text(doc, "執行流程", style="Heading 2")
    flow = [
        (1, "驗證資料包", "清單、哈希、目錄完整性"),
        (2, "鎖定契約", "三母體、第四CSV、缺值與地理邊界"),
        (3, "本地重跑", "資料轉換、驗證、lineage、dashboard JSON"),
        (4, "Kiro實作", "前端、AI問答、匯出與證據抽屜"),
        (5, "自動驗收", "資料、契約、儀表板與治理測試"),
        (6, "人工核對", "繁中、母體、來源、缺值與政策越界"),
    ]
    for number, title, desc in flow:
        t = doc.add_table(rows=1, cols=2)
        set_table_geometry(t, [0.65, 5.85])
        set_cell_shading(t.cell(0, 0), NAVY if number % 2 else TEAL)
        set_cell_shading(t.cell(0, 1), BLUE_LIGHT if number % 2 else TEAL_LIGHT)
        p = t.cell(0, 0).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(number))
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor.from_string(WHITE)
        set_east_asia(r)
        p = t.cell(0, 1).paragraphs[0]
        r = p.add_run(f"{title}｜")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        set_east_asia(r)
        r = p.add_run(desc)
        set_east_asia(r)
        if number < len(flow):
            arrow = doc.add_paragraph("↓")
            arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
            arrow.paragraph_format.space_before = Pt(0)
            arrow.paragraph_format.space_after = Pt(0)

    add_text(doc, "通過條件", style="Heading 2")
    for item in (
        "四份發布CSV保留正確母體與輔助層身分，不得合併成單一分母。",
        "重跑程式後，欄位、主鍵、數值身分、單位、來源URL及缺值語意不變。",
        "本地測試結果達既有基線，並產生可查核的log或JSON摘要。",
        "Kiro交付包含修改清單、測試結果、未解決差異與不執行AWS的聲明。",
    ):
        add_bullet(doc, item)

    p = add_text(doc, "1. 目的、範圍與排除項", style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(36)
    add_text(doc, "1.1 目的", style="Heading 2")
    add_text(doc, "本SOP將目前已完成的數據、契約與驗證轉成Kiro可執行的本地實作流程，使實作結果可重跑、可記錄、可查核，並能在未來取得授權後再穩定對接AWS。")
    add_text(doc, "1.2 本階段納入", style="Heading 2")
    for item in (
        "資料包完整性、哈希與來源追溯確認。",
        "四份發布CSV、Code Book、catalog、lineage與參數契約讀取。",
        "資料轉換、資料品質、前端、匯出與AI問答的本地建置。",
        "機器驗證、人工查核、差異回報與交付證據。",
    ):
        add_bullet(doc, item)
    add_text(doc, "1.3 本階段明確排除", style="Heading 2")
    add_callout(
        doc,
        "不得執行",
        "agentcore deploy、cdk deploy、sam deploy、AWS資源建立、正式S3寫入、密鑰或secret上傳、生產網域發布、Google Drive上傳。",
        fill=RED_LIGHT,
        accent=RED,
    )

    add_text(doc, "2. 四個資料平面與分母邊界", style="Heading 1")
    add_standard_table(
        doc,
        ["資料平面", "母體／分母", "地理粒度", "可使用", "禁止"],
        [
            ["戶籍人口", "戶籍登記現住人口", "新北市／29區", "人口、教育、婚姻、性別、密度", "不得當勞動或薪資分母"],
            ["勞動市場", "官方人力資源調查民間人口／勞動力", "新北市全市", "P、LF、E、U、UR、LFPR、EPR、ESLF", "不得切成行政區"],
            ["受僱薪資", "官方薪資統計涵蓋的受僱員工", "新北市全市／工作場所口徑", "平均薪資、中位數、發布缺口", "不得稱設籍青年所得"],
            ["政策服務輔助", "活動、據點與服務行政記錄；非人口母體", "全市／據點所在區", "參與人次、場次、據點、容量、曝光強度", "不得稱曝光率、觸達率或不重複人數"],
        ],
        [1.05, 1.45, 1.05, 1.55, 1.40],
    )
    add_callout(doc, "關鍵原則", "四份CSV可以被同一個儀表板讀取，但不能因為「放在同一頁」就變成同一母體或同一分母。")

    add_text(doc, "3. Kiro輸入資料包地圖", style="Heading 1")
    add_standard_table(
        doc,
        ["讀取順序", "資料夾／檔案", "Kiro應取得的資訊"],
        [
            ["1", "TOC.txt／README.txt", "層級、建議閱讀路徑、發布與排除原則"],
            ["2", "12_AWS_Kiro_實作檔案/KIRO規格/steering/g5-data-governance.md", "三母體、地理邊界、缺值、繁中與Agent禁止事項"],
            ["3", "KIRO規格/specs/ntpc-youth-g5/", "requirements、design、tasks及驗收條件"],
            ["4", "11_介接設定與追溯紀錄/config/", "來源、參數、發布規則、八工具契約"],
            ["5", "09_發布資料包/", "四份CSV、Code Book、catalog、lineage與介接規則"],
            ["6", "scripts/／dashboard/", "可重跑資料管線、匯出與前端測試"],
        ],
        [0.75, 3.20, 2.55],
    )

    add_text(doc, "4. Kiro本地實作SOP", style="Heading 1")
    add_step(doc, 0, "建立工作副本", "避免覆寫官方原始快照與已發布檔。", "原始repository與完整資料包。", "在獨立branch或可回復的工作副本開始；不修改official_original。", "branch名稱、起始commit、工作目錄。", "來源資料夾不完整、有未辨識密鑰或無法追溯起始版本。")
    add_step(doc, 1, "驗證清單與哈希", "確認輸入檔未缺失、未被靜默改寫。", "FILE_MANIFEST_SHA256.csv、TOC.txt、README.txt。", "依manifest比對檔名、大小與SHA-256；識別本次允許差異。", "manifest驗證摘要與差異清單。", "不在計畫內的檔案缺失或哈希變更。")
    add_step(doc, 2, "讀取治理與實作契約", "在編碼前鎖定不可破壞邏輯。", "steering、requirements、design、tasks、agent-tool-contracts.json與schema。", "建立「契約→實作→測試」對照表；任何新欄位先寫契約再改程式。", "對照表、受影響測試、納入／排除清單。", "母體、地理粒度、單位、數值身分或缺值語意有衝突。")
    add_step(doc, 3, "執行本地預檢", "先確認專案可以被重跑。", "本地Python、Node.js與repository。", "執行preflight與G7 readiness；僅讀本地檔案。", "預檢終端輸出及machine_qa摘要。", "任一必要檢查FAIL，或檢查要求AWS寫入。")
    add_code(doc, ["python scripts/preflight.py", "python scripts/validate_g7_aws_kiro_readiness.py"])
    add_step(doc, 4, "重跑資料與追溯輸出", "確認程式可以由原始快照重建發布層。", "原始快照、處理程式、config、schemas。", "按順序建置外部政策資料、三母體、lineage與dashboard JSON。", "四份CSV、Code Book、lineage、JSON與QA結果。", "主鍵重複、公式不守恆、來源不足、年度或單位默默變更。")
    add_code(doc, [
        "python scripts/build_external_policy_data.py",
        "python scripts/validate_external_policy_data.py",
        "python scripts/build_g5_v3_data.py",
        "python scripts/validate_g5_v3_data.py",
        "python scripts/build_g5_normalized_lineage.py",
        "python scripts/export_g5_dashboard_json.py",
    ])
    add_step(doc, 5, "建置儀表板與AI問答", "將契約落實為繁中、可探索、可匯出且可追溯的使用者體驗。", "dashboard、四份CSV／精簡JSON、八工具契約。", "完成篩選、資料卡、缺值、政策訊號、CSV匯出、Code Book、AI回答與來源抽屜。", "前端建置結果、畫面快照、工具呼叫log與測試。", "簡體中文、內部代號顯示在主UI、母體混用、缺值補造或Agent越權。")
    add_code(doc, ["cd dashboard", "npm ci", "npm test", "npm run build"])
    add_step(doc, 6, "自動驗收與人工查核", "確認功能正確，也確認政策敘述沒有超過證據。", "機器QA、測試報告、儀表板與政策卡。", "核對數值、空值、圖表七項證據、地圖單位、全市／行政區邊界、AI拒答與不可得結論。", "測試計數、人工查核表、已知限制與殘餘風險。", "任一P0邏輯失敗、無來源就給數值、或從人口數直接推論服務不足。")
    add_step(doc, 7, "交付Kiro實作結果", "使下一位人可知道改了什麼、如何驗證、還缺什麼。", "程式差異、測試、文件與已知限制。", "產生交付摘要、檔案清單、測試證據、待裁示項目、未執行AWS聲明與可回復commit。", "Kiro交付報告與對應commit。", "無法重現、差異不明、人工裁示尚未紀錄或不得宣稱的AWS成果。")

    add_text(doc, "5. AI Agent八個唯讀工具契約", style="Heading 1")
    add_standard_table(
        doc,
        ["工具", "作用", "關鍵限制"],
        [
            ["list_available_dimensions", "列出可用年度、年齡、性別、地理與指標", "只回契約白名單"],
            ["query_registered_metrics", "查戶籍人口、教育、婚姻與密度", "行政區限於戶籍與有據點空間資料"],
            ["query_labor_metrics", "查P、LF、E、U及比率", "僅新北市全市；必須回官方調查或模型身分"],
            ["query_wage_metrics", "查平均薪資、中位數與最新可用值", "114年未發布不得偽裝成114年"],
            ["query_policy_service_evidence", "查活動、參與人次與曝光強度", "曝光強度＝人次÷場次；非比率"],
            ["list_service_facilities", "列據點、位置、容量與可用狀態", "未有服務覆蓋資料時必須明示缺口"],
            ["explain_method", "說明公式、年齡轉換、缺值與限制", "不以自由文字改寫正式定義"],
            ["get_source_evidence", "回來源URL、擷取日期、檔案哈希與lineage", "不列raw S3、secret或任意本機路徑"],
        ],
        [2.05, 2.55, 1.90],
    )
    add_callout(doc, "Agent的答案契約", "回答必須同時揭露母體、資料年度、數值身分、計算方法、更新日期、資料來源與可用性狀態；無證據時明確拒答或說明缺口。")

    add_text(doc, "6. 資料重跑與機器驗證順序", style="Heading 1")
    add_text(doc, "建議保持下列順序，不要在前一階段FAIL時繼續將候選資料發布給儀表板。")
    validation_rows = [
        ["A", "資料包完整性", "manifest、檔案存在、SHA-256", "不預期差異為0"],
        ["B", "Schema與契約", "欄位、型別、enum、工具白名單", "全部PASS"],
        ["C", "資料公式", "比率、分母、男女加總、曝光強度", "重算容許差內"],
        ["D", "粒度與身分", "年度、年齡、性別、地理、數值身分", "不越界、不重複主鍵"],
        ["E", "前端與匯出", "繁中、空狀態、篩選、CSV、Code Book", "35／35基線或更新後同等PASS"],
        ["F", "AI治理", "拒答、來源、母體、缺值、不得任意SQL", "全部治理測試PASS"],
    ]
    add_standard_table(doc, ["關卡", "檢查層", "檢查內容", "通過準則"], validation_rows, [0.65, 1.45, 2.80, 1.60])

    add_text(doc, "7. 人工查核表", style="Heading 1")
    add_standard_table(
        doc,
        ["編號", "查核項目", "通過準則", "不通過時"],
        [
            ["HR-01", "三母體", "圖表、匯出、AI回答均明示正確母體", "停止發布；回到契約與查詢路由"],
            ["HR-02", "地理邊界", "勞動與薪資僅全市；行政區只顯示有證據資料", "移除偽切割並增加拒答測試"],
            ["HR-03", "缺值與尚未發布", "不補零、不偽裝當期；顯示最新可用年與落後年數", "改正資料及UI邏輯"],
            ["HR-04", "曝光強度", "數值＝人次÷場次；單位人次／場；不稱曝光率", "停止該指標發布並更正Code Book"],
            ["HR-05", "繁體中文與代號", "主UI只顯示繁中與阿拉伯數字；代號留在查核層", "全文字審視與snapshot回歸"],
            ["HR-06", "政策推論", "說明可得與不可得結論；不從人口多直推服務不足", "政策卡降級為待證據問題"],
            ["HR-07", "來源可追溯", "每組數值可回到URL、擷取日期、檔案哈希、處理方法", "隔離該批資料並補齊lineage"],
            ["HR-08", "AWS邊界", "無AWS建立、寫入、部署或密鑰更動", "立即停止；記錄外部變更並要求授權"],
        ],
        [0.65, 1.60, 2.85, 1.40],
    )

    add_text(doc, "8. Kiro可直接使用的任務提示", style="Heading 1")
    add_callout(doc, "使用方式", "將下列區塊作為Kiro的起始任務，再依當次要實作的頁面或功能補上明確範圍。")
    add_code(doc, [
        "請依《Kiro本地實作交接流程與驗證SOP_V1.0》執行。",
        "本次僅進行本地實作，不得建立、修改或部署任何AWS資源。",
        "請依順序讀取TOC、README、steering、requirements、design、tasks、",
        "agent-tool-contracts.json、四份發布CSV、Code Book、catalog與lineage。",
        "不得混用戶籍、勞動、薪資三母體；政策服務CSV只是輔助證據。",
        "全市勞動與薪資不得偽切成行政區；曝光強度不得命名為曝光率或觸達率。",
        "保留官方行政精確值、官方調查估計值、模型估計值及缺值身分。",
        "修改完成後，執行本地資料、契約、儀表板與治理測試。",
        "交付時提供：修改檔案清單、測試結果、差異、限制、待人工查核項與未執行AWS聲明。",
    ])

    p = add_text(doc, "9. 異常處理與停止規則", style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(36)
    add_standard_table(
        doc,
        ["異常", "系統處置", "人工處置", "不得作法"],
        [
            ["來源哈希改變", "建新raw snapshot，不覆寫舊版", "確認是官方更新或異常", "靜默換檔"],
            ["Schema drift", "隔離候選資料，標示HOLD", "裁示新欄位定義與版本", "自動忽略欄位"],
            ["年度缺值", "回傳not_available與latest_available_year", "確認官方是否已發布", "補0、延用前期或假造月值"],
            ["模型QA超界", "隔離模型輸出，保留前一PASS版", "檢查輸入、約束、敏感度與文獻", "將超界估計值發布"],
            ["Agent無證據", "拒答並說明缺少的資料", "評估是否新增合法來源", "幻覺、推測或無來源政策結論"],
            ["本地測試FAIL", "停止交付，保留log", "判斷是預期基線變更或缺陷", "只改測試來變PASS"],
        ],
        [1.35, 2.10, 1.85, 1.20],
    )

    add_text(doc, "10. 未來AWS部署附錄（本次不執行）", style="Heading 1")
    add_callout(
        doc,
        "附錄狀態：DEFERRED",
        "本節只保留未來所需的資訊，不是目前Kiro實作的前置條件。未來只有在使用者明確授權後，才進入dev部署。",
        fill=AMBER_LIGHT,
        accent=AMBER,
    )
    for item in (
        "確認目標AWS帳號、區域與只允許dev的環境邊界。",
        "填入Owner、CostCenter、每月預算、告警收件人與原型到期日。",
        "確認KMS、日誌保存期、API key／secret輪替及最小權限。",
        "先執行validate／synth，diff與成本估算；取得人工核准後才部署。",
        "dev完成NO_CHANGE、QA失敗、回復、權限、效能與成本演練後，才能提出下一階段建議。",
    ):
        add_bullet(doc, item)
    add_text(doc, "人工裁示原始清單位於 08_人工查核與異常處理/G8_dev部署前人工裁示清單.md，已明確標示「未來使用／目前暫緩」。")

    add_callout(
        doc,
        "11. 交付收尾",
        "Kiro實作者交付修改檔案、功能與契約對照、測試指令與PASS／FAIL證據，並明列未執行AWS部署、正式發布或雲端寫入。資料／政策查核者再完成HR-01～HR-08、日期、證據、殘餘風險及真正需要裁示的外部行為。",
    )
    doc.core_properties.title = "Kiro本地實作交接流程與驗證SOP V1.0"
    doc.core_properties.subject = "新北市青年資料證據台 Kiro 本地實作交接"
    doc.core_properties.author = "NTPC Youth Data Evidence Platform"
    doc.core_properties.keywords = "Kiro, SOP, data governance, local implementation, validation"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
