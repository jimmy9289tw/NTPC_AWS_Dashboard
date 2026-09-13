"""唯讀競賽附件、封存目錄及套件能力，輸出本 Gate 紀錄。"""
from pathlib import Path
import hashlib, importlib.util, json, zipfile
from pypdf import PdfReader
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
BASE=ROOT.parent/'ntpc-lsf-v1-20260912/NTPC_Youth_LSF_V1_20260912'

def main():
    pdf=Path('C:/Users/jimmy/Downloads/黑客松競賽環境規範與限制_20260722.pdf')
    xlsx=Path('C:/Users/jimmy/Downloads/Supported AWS Services List 20260722.xlsx')
    pages=[{'page':i+1,'text':p.extract_text()} for i,p in enumerate(PdfReader(pdf).pages)]
    book=load_workbook(xlsx,read_only=True,data_only=True)
    sheets={s.title:[list(row) for row in s.iter_rows(values_only=True)] for s in book.worksheets};book.close()
    services=[]
    targets={'lambda','scheduler','events','states','s3','ecs','iam','logs','cloudwatch','cloudformation','ecr','codebuild','glue'}
    for sheet,rows in sheets.items():
        for i,row in enumerate(rows):
            if row and str(row[0]).split(' ')[0] in targets:
                services.append({'sheet':sheet,'row':i+1,'cells':row})
    archives=[]
    for path in BASE.rglob('*.zip'):
        with zipfile.ZipFile(path) as z:
            names=z.namelist()
            hits=[{'name':i.filename,'bytes':i.file_size} for i in z.infolist() if not i.is_dir() and any(term in i.filename.replace('\\','/') for term in ['MOI-EDU-SINGLEAGE/114','MOI-MARITAL5Y/114','WAGE-AUX/20260826','NTPC-LAND-AREA/20260831','NTPC-YOUTH-SERVICE/20260831','annual_target_age_results.csv','source_rounding_reconciliation.csv'])]
            archives.append({'path':str(path),'members':len(names),'matched_inputs':hits})
    result={'rules_pdf':{'path':str(pdf),'sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'pages':pages},'service_list':{'path':str(xlsx),'sha256':hashlib.sha256(xlsx.read_bytes()).hexdigest(),'sheets':list(sheets),'matches':services},'archives':archives,'dependencies':{name:importlib.util.find_spec(name) is not None for name in ['boto3','numpy','scipy','openpyxl','xlrd','pypdf','pytest','pandas']}}
    (ROOT/'preflight.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(json.dumps({'pdf_pages':len(pages),'services':[{'sheet':s['sheet'],'row':s['row'],'name':s['cells'][0]} for s in services],'archives':[{'path':a['path'],'members':a['members'],'matched_inputs':len(a['matched_inputs'])} for a in archives],'dependencies':result['dependencies']},ensure_ascii=False,default=str))

if __name__=='__main__':main()
